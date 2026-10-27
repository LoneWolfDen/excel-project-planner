"""Assignments sheet (Excel Table "AssignmentTable"): more people on a task, one row per extra person.

A task's first person stays on Tasks (Resource, Units), so a one-person task is still one row to type. Each row here
adds another person to the task with that UID, like a row in MS Project's Assign Resources dialog. The row works
out its own work and cost (the task's Duration x HoursPerDay x Units) and its time-phasing, from the task's forecast
working days, so every per-resource total (Resources, Resource Report, Sprints) is two SUMIFS: one over Tasks and one
over this table. Tasks adds the rows' work and cost into the task's Work and Cost, and lists the names in Resource
Names. Nothing here is read back by the scheduling columns, so an assignment never moves a date.
"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from .core import Col, Tab, add_table, col_range, fill, name, title

TABLE = "AssignmentTable"
A = Tab(TABLE)
MAX_LISTED = 4          # extra people named in Tasks > Resource Names (the rest are counted: "+2 more")
KEY_STEP = 1000         # _Key = Task UID x KEY_STEP + n for a task's n-th valid row (n < KEY_STEP)

# key -> column name; defined names asg_<...> for the columns other sheets read
N = {"UID": "Task UID", "Res": "Resource", "Units": "Units", "Task": "Task", "Work": "Work", "Cost": "Cost",
     "Check": "Check", "Pos": "_TaskPos", "Valid": "_Valid", "Key": "_Key", "Text": "_Text",
     "SWD": "_StartWD", "FWD": "_FinishWD", "WPD": "_WorkPerWD", "WBefore": "_WorkBeforeStart",
     "Done": "_WorkDone", "AC": "_ActualCost", "Late": "_Late", "Crit": "_Critical", "Over": "_OverAllocated"}
NAMES = {"UID": "asg_TaskUID", "Res": "asg_Resource", "Work": "asg_Work", "Cost": "asg_Cost", "Valid": "asg_Valid",
         "Key": "asg_Key", "Text": "asg_Text", "SWD": "asg_StartWD", "FWD": "asg_FinishWD", "WPD": "asg_WorkPerWD",
         "WBefore": "asg_WorkBeforeStart", "Done": "asg_WorkDone", "AC": "asg_ActualCost", "Late": "asg_Late",
         "Crit": "asg_Critical", "Over": "asg_OverAllocated"}


def R(key):
    return A.row(N[key])


def task(col):
    """The assigned task's value in a Tasks column (by its tsk_ name); the position is looked up once, in _TaskPos."""
    return f"INDEX({col},{R('Pos')})"


def over_allocated(resource, start, finish):
    """Yes when resource is over 100% (Resource Report utilisation, after leave) in a week from start to finish.
    Only the weeks in the Resource Report window and the resources in its grid are known; others count as not over."""
    week_of = lambda d: f"(INT(({d}-ResourceWeekStart)/7)+1)"
    first, last = f"MAX(1,{week_of(start)})", f"MIN(COLUMNS(rw_Util),{week_of(finish)})"
    row = f"MATCH({resource},res_Name,0)"
    return (f"IFERROR(IF({first}>{last},FALSE(),"
            f"COUNTIF(INDEX(rw_Util,{row},{first}):INDEX(rw_Util,{row},{last}),\">1\")>0),FALSE())")


def formulas():
    uid, res, pos, ok = R("UID"), R("Res"), R("Pos"), f'{R("Valid")}=1'
    units = f'IF({R("Units")}="",DefaultUnits,{R("Units")})'
    f = {}
    f["Pos"] = f'IF(N({uid})=0,0,IFERROR(MATCH({uid},tsk_UID,0),-1))'
    f["Task"] = f'IF({pos}<=0,"",{task("tsk_Name")}&"")'
    f["Check"] = (f'IF(AND(N({uid})=0,{res}=""),"",IF(N({uid})=0,"Enter the Task UID",IF({pos}<0,"UID "&{uid}&" not found",'
                  f'IF({task("tsk_Name")}="","UID "&{uid}&" is an empty row",'
                  f'IF({task("tsk_IsSummary")},"UID "&{uid}&" is a summary: assign one of its tasks",'
                  f'IF({res}="","Enter a resource",IF(ISNA(MATCH({res},res_Name,0)),{res}&" is not on Resources",'
                  f'IF({res}={task("tsk_Resource")},"Already the task\'s Resource (Tasks)",'
                  f'IF(COUNTIFS({A.upto(N["UID"])},{uid},{A.upto(N["Res"])},{res})>1,"Listed twice","")))))))))')
    f["Valid"] = f'IF(AND({pos}>0,{res}<>"",{R("Check")}=""),1,0)'
    f["Key"] = f'IF({ok},{uid}*{KEY_STEP}+COUNTIFS({A.upto(N["UID"])},{uid},{A.upto(N["Valid"])},1),"")'
    f["Text"] = f'IF({ok},{res}&IF(OR({R("Units")}="",{R("Units")}=1),"","["&ROUND({R("Units")}*100,0)&"%]"),"")'
    f["Work"] = f'IF({ok},N({task("tsk_Duration")})*HoursPerDay*{units},0)'
    f["Cost"] = f'IF({R("Work")}=0,0,{R("Work")}*IFERROR(INDEX(res_Rate,MATCH({res},res_Name,0)),0))'
    # time-phasing, as on Tasks: the task's forecast working days (counted from the project start), work spread evenly
    f["SWD"] = f'IF({ok},N({task("tsk_StartWD")}),0)'
    f["FWD"] = f'IF({ok},N({task("tsk_FinishWD")}),0)'
    f["WPD"] = f'IF(OR({R("Work")}=0,{R("FWD")}=0),0,{R("Work")}/({R("FWD")}-{R("SWD")}+1))'
    f["WBefore"] = f'{R("WPD")}*MAX(0,{R("SWD")}-1)'
    f["Done"] = f'IF({ok},{R("Work")}*N({task("tsk_LeafProgress")}),0)'
    # the task's actual cost shared by cost: with no Actual Cost typed this is Cost x Progress, as on Tasks
    total = task("tsk_LeafTotalCost")
    # (nested IFs, never AND / OR: those evaluate every argument, and INDEX at an invalid position is an error)
    f["AC"] = f'IF(NOT({ok}),0,IF(N({total})=0,0,N({task("tsk_LeafAC")})*{R("Cost")}/{total}))'
    f["Late"] = f'IF(NOT({ok}),0,IF(N({task("tsk_LateKey")})>0,1,0))'
    f["Crit"] = f'IF(NOT({ok}),0,IF({task("tsk_Critical")}="Yes",1,0))'
    f["Over"] = (f'IF(NOT({ok}),"",IF(N({task("tsk_LeafProgress")})>=1,"",IF('
                 f'{over_allocated(res, task("tsk_FcStart"), task("tsk_FcFinish"))},"Yes","")))')
    return f


# key, kind, width, number format, centre, note
LAYOUT = [
    ("UID", "in", 9, None, True, "UID of a detail task on Tasks."),
    ("Task", "calc", 34, None, False, "The task's name (looked up)."),
    ("Res", "in", 22, None, False, "Another person on the task (dropdown). The first person goes on Tasks > Resource."),
    ("Units", "in", 8, "0%", True, "Share of the person's time on this task (blank = DefaultUnits)."),
    ("Work", "calc", 10, "#,##0.0", False, "Hours = the task's Duration x HoursPerDay x Units."),
    ("Cost", "calc", 11, "#,##0.00", False, "Work x the resource's Std Rate."),
    ("Check", "calc", 30, None, False, "Explains a row that is ignored."),
] + [(k, "help", 10, None, False, None) for k in ("Pos", "Valid", "Key", "Text", "SWD", "FWD", "WPD", "WBefore",
                                                   "Done", "AC", "Late", "Crit", "Over")]


def build_assignments(wb, n_rows, rows=()):
    """rows: (task UID, resource, units) for the sample or an import."""
    ws = wb.create_sheet("Assignments")
    title(ws, "Assignments", "More than one person on a task: the first goes on Tasks (Resource, Units); add a row "
                             "here for each other person. Work, cost, the Resource Report and Sprints include them. "
                             "Add rows to the table for more.")
    ws.column_dimensions["A"].width = 3
    f = formulas()
    values = {"UID": [r[0] for r in rows], "Res": [r[1] for r in rows], "Units": [r[2] for r in rows]}
    cols = [Col(N[k], kind, width, fmt, None if kind == "in" else f[k], center, note, k, values.get(k, []))
            for k, kind, width, fmt, center, note in LAYOUT]
    L = add_table(ws, TABLE, 5, 2, cols, n_rows)
    for key, defined in NAMES.items():
        name(wb, defined, A.col(N[key]))
    first, last = 6, 5 + n_rows
    for rule, key in ((DataValidation(type="list", formula1="=res_Name"), "Res"),
                      (DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1"), "UID"),
                      (DataValidation(type="decimal", operator="between", formula1="0", formula2="10"), "Units")):
        ws.add_data_validation(rule)
        rule.add(col_range(L[key], first, last))
    check = f"${L['Check']}{first}"
    ws.conditional_formatting.add(col_range(L["Check"], first, last), FormulaRule(
        formula=[f'{check}<>""'], fill=fill("FFC7CE"), font=Font(color="9C0006")))
    over = f"${L['Over']}{first}"
    ws.conditional_formatting.add(col_range(L["Res"], first, last), FormulaRule(
        formula=[f'{over}="Yes"'], fill=fill("FFC7CE"), font=Font(color="9C0006", bold=True)))
    ws.column_dimensions.group(L["Pos"], L["Over"], hidden=True, outline_level=1)
    ws.freeze_panes = "C6"
    return ws
