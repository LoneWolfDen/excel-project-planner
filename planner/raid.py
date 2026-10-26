"""RAID sheet: risks, assumptions, issues and dependencies (Excel Table "RAIDLog").

Everything a PM types is an input; the sheet only scores, rates and flags. Score = probability x impact (an issue
has already happened, so it counts as probability 5); the rating thresholds are settings. A linked task UID shows
the task's name and forecast finish, so a risk sits next to the work it threatens. The Status Report lists the
top open items (any type) and counts high-rated ones in its RAG.
"""
import datetime as dt

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from .core import DATE, Col, Tab, add_table, col_range, fill, linked_task, name, title

D = dt.date
TYPES = ["Risk", "Assumption", "Issue", "Dependency"]
STATUSES = ["Open", "In Progress", "Closed"]
# Ref, Type, Title, Owner, Raised, Due, Status, Probability, Impact, Task UID, Response, Update
SAMPLE = [
    (1, "Risk", "Payment provider certification takes longer than planned", "Jordan Lee", D(2026, 10, 1),
     D(2026, 10, 30), "Open", 4, 4, 14, "Start the certification paperwork now; book the provider's test slot.",
     "Paperwork submitted 20-Oct."),
    (2, "Risk", "UX designer on leave for two weeks in November", "Casey Kim", D(2026, 10, 5), D(2026, 11, 2), "Open",
     3, 3, 18, "Agree cover with the design agency.", ""),
    (3, "Issue", "Core API endpoints running 6 days behind plan", "Priya Sharma", D(2026, 10, 20), D(2026, 10, 23),
     "In Progress", None, 4, 13, "Add a second developer for two weeks (change request CR 2).", "Escalated to sponsor."),
    (4, "Assumption", "Client UAT team available from 7 December", "Morgan Patel", D(2026, 9, 15), D(2026, 11, 20),
     "Open", 2, 4, 22, "Confirm names and dates at the November steering committee.", ""),
    (5, "Dependency", "Test environment delivered by IT operations by 9 November", "Alex Chen", D(2026, 9, 30),
     D(2026, 11, 9), "Open", 3, 5, 21, "Weekly check-in with IT operations.", "Hardware ordered."),
    (6, "Risk", "Scope creep on the reporting screens", "Priya Sharma", D(2026, 9, 10), D(2026, 10, 9), "Closed", 2, 3,
     17, "Reporting screens moved to phase 2 (decision D 2).", "Closed 09-Oct."),
]
NAMES = {"Ref": "raid_Ref", "Type": "raid_Type", "Title": "raid_Title", "Owner": "raid_Owner", "Due": "raid_Due",
         "Status": "raid_Status", "Score": "raid_Score", "Rating": "raid_Rating", "Overdue": "raid_Overdue",
         "Task": "raid_Task", "Response": "raid_Response", "_Label": "raid_Label", "_TopKey": "raid_TopKey"}


def build_raid(wb, blank, n_rows=100):
    ws = wb.create_sheet("RAID")
    title(ws, "RAID Log: Risks, Assumptions, Issues, Dependencies",
          "Type in the blue columns. Score = Probability x Impact (1-5 each; an issue counts as probability 5). "
          "Rating thresholds: Settings > RAIDHighScore / RAIDMediumScore. Link a task by its UID.")
    ws.column_dimensions["A"].width = 3
    data = [] if blank else SAMPLE
    t = Tab("RAIDLog")
    r = t.row
    blank_row = f'{r("Title")}=""'
    open_ = f'{r("Status")}<>"Closed"'
    values = lambda i: [row[i] for row in data]
    letters = add_table(ws, "RAIDLog", 4, 2, [
        Col("Ref", "in", 6, center=True, values=list(range(1, n_rows + 1)),
            note="Permanent number, quoted in reports and other logs (e.g. RAID 3)."),
        Col("Type", "in", 12, values=values(1)),
        Col("Title", "in", 44, values=values(2)),
        Col("Owner", "in", 16, values=values(3)),
        Col("Raised", "in", 12, DATE, values=values(4)),
        Col("Due", "in", 12, DATE, values=values(5), note="Date the next action or the mitigation is due."),
        Col("Status", "in", 11, center=True, values=values(6)),
        Col("Probability", "in", 9, center=True, values=values(7), note="1 (rare) to 5 (almost certain). Ignored for issues."),
        Col("Impact", "in", 8, center=True, values=values(8), note="1 (minor) to 5 (severe)."),
        Col("Score", width=7, center=True, note="Probability x Impact; an issue counts as probability 5.",
            formula=f'IF({blank_row},"",IF({r("Type")}="Issue",5,N({r("Probability")}))*N({r("Impact")}))'),
        Col("Rating", width=8, center=True,
            formula=(f'IF({blank_row},"",IF({r("Score")}>=RAIDHighScore,"High",'
                     f'IF({r("Score")}>=RAIDMediumScore,"Medium","Low")))')),
        Col("Overdue", width=9, center=True, note="Not closed and the due date is before the status date.",
            formula=f'IF(OR({blank_row},N({r("Due")})=0),"",IF(AND({open_},{r("Due")}<StatusDate),"Overdue",""))'),
        Col("Task UID", "in", 8, center=True, values=values(9), note="UID of the task this item affects (optional)."),
        Col("Task", width=28, formula=linked_task(r("Task UID"))),
        Col("Task Finish", width=12, fmt=DATE, formula=(
            f'IF(N({r("Task UID")})=0,"",IFERROR(INDEX(tsk_FcFinish,MATCH({r("Task UID")},tsk_UID,0)),""))'),
            note="The linked task's forecast finish."),
        Col("Response", "in", 44, values=values(10), note="Mitigation, contingency or action."),
        Col("Update", "in", 30, values=values(11), note="Latest progress note."),
        Col("_Label", "help", 8, formula=f'IF({blank_row},"",{r("Type")}&" "&{r("Ref")}&": "&{r("Title")})',
            note="Type, Ref and Title in one cell, for the Status Report."),
        Col("_TopKey", "help", 8, note="Open items ranked by score (the Status Report's top list).",
            formula=f'IF(AND(NOT({blank_row}),{open_}),{r("Score")}+1-(ROW()-ROW({t.head("Ref")}))/100000,"")'),
    ], n_rows)
    for column, defined in NAMES.items():
        name(wb, defined, f"RAIDLog[{column}]")
    ws.column_dimensions.group(letters["_Label"], letters["_TopKey"], hidden=True, outline_level=1)
    first, last = 5, 4 + n_rows
    rng = lambda k: col_range(letters[k], first, last)
    for rule, keys in ((DataValidation(type="list", formula1="=L_RAIDType"), ("Type",)),
                       (DataValidation(type="list", formula1="=L_RAIDStatus"), ("Status",)),
                       (DataValidation(type="whole", operator="between", formula1="1", formula2="5"), ("Probability", "Impact")),
                       (DataValidation(type="date", operator="greaterThan", formula1="1"), ("Raised", "Due")),
                       (DataValidation(type="list", formula1="=res_Name", showErrorMessage=False), ("Owner",))):
        ws.add_data_validation(rule)
        for k in keys:
            rule.add(rng(k))
    cf = ws.conditional_formatting
    rating = letters["Rating"]
    for word, (bg, fg) in {"High": ("FFC7CE", "9C0006"), "Medium": ("FFEB9C", "9C5700"), "Low": ("C6EFCE", "006100")}.items():
        cf.add(rng("Rating"), FormulaRule(formula=[f'{rating}{first}="{word}"'], fill=fill(bg), font=Font(color=fg, bold=True)))
    cf.add(rng("Overdue"), FormulaRule(formula=[f'{letters["Overdue"]}{first}="Overdue"'], font=Font(color="C00000", bold=True)))
    cf.add(f"B{first}:{letters['Update']}{last}", FormulaRule(formula=[f'${letters["Status"]}{first}="Closed"'],
                                                               font=Font(color="A6A6A6")))
    ws.freeze_panes = f"E{first}"
    return ws
