"""Status Report: overall RAG, RAG by area, earned value and the period commentary, for steering meetings.

Each RAG is calculated from the plan against the Settings thresholds. The project manager can override it (PM RAG)
and explain why (Comment); the RAG column shows the override when there is one, which is how most PMOs report.
Earned-value cells have defined names (evm_BAC, evm_SPI, ...) so the Dashboard and other reports reuse them
instead of repeating the sums.
"""
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation

from . import rag
from .core import DATE, Col, Tab, add_table, font, name, title

SHEET = "Status Report"
WIDTHS = {"A": 3, "B": 36, "C": 56, "D": 16, "E": 12, "F": 10, "G": 10, "H": 50}
MONEY = "#,##0"

# area, measure, value, severity of the value (0 green, 1 amber, 2 red), number format
AREAS = [
    ("Schedule", "Forecast finish vs plan (working days)", "FinishVarianceDays", rag.slip("{v}"), "+0;-0;0"),
    ("Schedule performance", "SPI = earned value / planned value", "evm_SPI", rag.index("{v}"), "0.00"),
    ("Cost performance", "CPI = earned value / actual cost", "evm_CPI", rag.index("{v}"), "0.00"),
    ("Deadlines", "Tasks forecast to miss their deadline", 'COUNTIF(tsk_DeadlineCheck,"Late*")', rag.count("{v}", 2), "0"),
    ("Critical path", "Critical tasks forecast to finish late",
     'COUNTIFS(tsk_Critical,"Yes",tsk_IsLeaf,1,tsk_Variance,">0")', rag.count("{v}", 1), "0"),
    ("Plan quality", "Ignored links and outline errors (see Tasks)", 'COUNTIF(tsk_LinkCheck,"?*")+COUNTIF(tsk_WBS,"~?")',
     rag.count("{v}", 1), "0"),
]

# defined name, label, formula, meaning, format
EVM = [
    ("evm_BAC", "Budget at completion (BAC)", "SUM(tsk_LeafBudget)", "Baseline Cost once a baseline is set, otherwise Cost.", MONEY),
    ("evm_PV", "Planned value (PV)", "SUM(tsk_LeafPV)", "Budget of the work planned before the status date.", MONEY),
    ("evm_EV", "Earned value (EV)", "SUM(tsk_LeafEV)", "Budget of the work done (Budget x Progress).", MONEY),
    ("evm_AC", "Actual cost (AC)", "SUM(tsk_LeafAC)", "Actual Cost on Tasks; Cost x Progress where blank.", MONEY),
    ("evm_SV", "Schedule variance (SV)", "evm_EV-evm_PV", "EV - PV. Negative = behind plan.", MONEY),
    ("evm_CV", "Cost variance (CV)", "evm_EV-evm_AC", "EV - AC. Negative = over cost.", MONEY),
    ("evm_SPI", "Schedule performance index (SPI)", 'IF(evm_PV>0,evm_EV/evm_PV,"")', "EV / PV. Below 1 = behind.", "0.00"),
    ("evm_CPI", "Cost performance index (CPI)", 'IF(evm_AC>0,evm_EV/evm_AC,"")', "EV / AC. Below 1 = over cost.", "0.00"),
    ("evm_EAC", "Estimate at completion (EAC)", "IF(N(evm_CPI)>0,evm_AC+(evm_BAC-evm_EV)/evm_CPI,evm_AC+evm_BAC-evm_EV)",
     "AC + (BAC - EV) / CPI: forecast cost at the current CPI.", MONEY),
    ("evm_ETC", "Estimate to complete (ETC)", "evm_EAC-evm_AC", "EAC - AC: forecast cost of the remaining work.", MONEY),
    ("evm_VAC", "Variance at completion (VAC)", "evm_BAC-evm_EAC", "BAC - EAC. Negative = forecast overrun.", MONEY),
    ("evm_TCPI", "To-complete performance index (TCPI)", 'IF(evm_BAC>evm_AC,(evm_BAC-evm_EV)/(evm_BAC-evm_AC),"")',
     "CPI the remaining work needs to finish on budget.", "0.00"),
    ("evm_PctEarned", "% complete (earned)", "IF(evm_BAC>0,evm_EV/evm_BAC,0)", "EV / BAC.", "0%"),
    ("evm_PctPlanned", "% complete (planned)", "IF(evm_BAC>0,evm_PV/evm_BAC,0)", "PV / BAC.", "0%"),
]

NOTES = ["Summary", "Achievements this period", "Planned for next period", "Issues and escalations",
         "Decisions needed", "Risks to watch"]


def _rag_columns(table, measures, values, severities):
    t = Tab(table)
    return [
        Col("Area", width=WIDTHS["B"], values=measures[0]),
        Col("Measure", width=WIDTHS["C"], values=measures[1]),
        Col("Value", width=WIDTHS["D"], values=values),
        Col("Calculated", width=WIDTHS["E"], values=severities, center=True,
            note="RAG from the thresholds on Settings (RAGAmberSlipDays, RAGRedSlipDays, RAGAmberIndex, RAGRedIndex)."),
        Col("PM RAG", "in", WIDTHS["F"], center=True, note="Optional override by the project manager: Green, Amber or Red."),
        Col("RAG", width=WIDTHS["G"], center=True,
            formula=f'IF({t.row("PM RAG")}<>"",{t.row("PM RAG")},{t.row("Calculated")})',
            note="Reported RAG: the PM override if there is one, otherwise the calculated RAG."),
        Col("Comment", "in", WIDTHS["H"], note="Why: especially when the PM RAG differs from the calculated one."),
    ]


def build_status_report(wb):
    ws = wb.create_sheet(SHEET)
    title(ws, None)
    ws["A1"] = '=ProjectName&"  -  Status Report"'
    ws["A2"] = ('="Status date "&TEXT(StatusDate,"ddd dd-mmm-yyyy")&"   Project manager: "&ProjectManager&'
                '"   Forecast finish "&TEXT(ProjectFinish,"ddd dd-mmm-yyyy")&" ("&TEXT(FinishVarianceDays,"+0;-0;0")&'
                '" working days vs plan)"')
    for col, width in WIDTHS.items():
        ws.column_dimensions[col].width = width

    areas = Tab("StatusRAG")
    worst = f'MAX(IF(COUNTIF({areas.col("RAG")},"Red")>0,2,0),IF(COUNTIF({areas.col("RAG")},"Amber")>0,1,0))'
    overall_value = f'=COUNTIF({areas.col("RAG")},"Red")&" red, "&COUNTIF({areas.col("RAG")},"Amber")&" amber"'
    add_table(ws, "OverallStatus", 4, 2, _rag_columns(
        "OverallStatus", (["Overall"], ["Worst of the areas below"]), [overall_value], [f"={rag.rag_text(worst)}"]), 1, autofilter=False)
    name(wb, "ProjectRAG", "OverallStatus[RAG]")

    top = 7
    add_table(ws, "StatusRAG", top, 2, _rag_columns(
        "StatusRAG", ([a[0] for a in AREAS], [a[1] for a in AREAS]), [f"={a[2]}" for a in AREAS],
        [f"={rag.rag_text(a[3].format(v=areas.row('Value')))}" for a in AREAS]), len(AREAS), autofilter=False)
    for i, a in enumerate(AREAS):
        ws.cell(top + 1 + i, 4).number_format = a[4]
        ws.cell(top + 1 + i, 4).alignment = Alignment(horizontal="right")
    for r in range(5, top + 1 + len(AREAS)):
        ws.cell(r, 2).font = font(bold=True)
    rag_range = f"E5:G{top + len(AREAS)}"
    rag.colour(ws, rag_range, "E5")
    pick = DataValidation(type="list", formula1="=L_RAG")
    ws.add_data_validation(pick)
    pick.add("F5")
    pick.add(f"F{top + 1}:F{top + len(AREAS)}")

    top += len(AREAS) + 3
    ws.cell(top - 1, 2, "EARNED VALUE at the status date").font = font(bold=True, color="1F4E78")
    add_table(ws, "EarnedValue", top, 2, [
        Col("Measure", width=WIDTHS["B"], values=[e[1] for e in EVM]),
        Col("Meaning", width=WIDTHS["C"], values=[e[3] for e in EVM]),
        Col("Value", width=WIDTHS["D"], values=[f"={e[2]}" for e in EVM]),
    ], len(EVM), autofilter=False)
    for i, (defined, _, _, _, fmt) in enumerate(EVM):
        r = top + 1 + i
        ws.cell(r, 4).number_format = fmt
        ws.cell(r, 4).font = font(bold=True)
        name(wb, defined, f"'{SHEET}'!$D${r}")
    evm_top = top

    top += len(EVM) + 3
    ws.cell(top - 1, 2, "COMMENTARY for this period (type in the Notes column)").font = font(bold=True, color="1F4E78")
    add_table(ws, "StatusNotes", top, 2, [
        Col("Section", width=WIDTHS["B"], values=NOTES),
        Col("Notes", "in", WIDTHS["C"]),
    ], len(NOTES), autofilter=False)
    for i in range(len(NOTES)):
        ws.cell(top + 1 + i, 3).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(top + 1 + i, 2).alignment = Alignment(vertical="top")
        ws.row_dimensions[top + 1 + i].height = 36

    chart = BarChart()
    chart.type, chart.style, chart.title, chart.legend = "bar", 10, "Earned value (BAC, PV, EV, AC)", None
    chart.add_data(Reference(ws, min_col=4, min_row=evm_top + 1, max_row=evm_top + 4))
    chart.set_categories(Reference(ws, min_col=2, min_row=evm_top + 1, max_row=evm_top + 4))
    chart.height, chart.width = 7, 16
    ws.add_chart(chart, f"E{evm_top}")
    ws.page_setup.orientation = "landscape"     # one page when printed or saved as PDF
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.freeze_panes = "A4"
    return ws
