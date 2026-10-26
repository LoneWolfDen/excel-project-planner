"""Dashboard: KPI table, status counts, next milestones and two charts. Formulas use defined names only."""
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font

from .core import DATE, Col, Tab, add_table, fill, font
from .tasks import STATUS_COLOURS

WD = "WorkWeek,Holidays"
KPIS = [
    ("Planned start", "MIN(tsk_PlanStart)", DATE),
    ("Planned finish", "MAX(tsk_PlanFinish)", DATE),
    ("Forecast finish", "MAX(tsk_FcFinish)", DATE),
    ("Finish variance (work days)",
     f"IF(MAX(tsk_FcFinish)>=MAX(tsk_PlanFinish),NETWORKDAYS.INTL(MAX(tsk_PlanFinish),MAX(tsk_FcFinish),{WD})-1,"
     f"1-NETWORKDAYS.INTL(MAX(tsk_FcFinish),MAX(tsk_PlanFinish),{WD}))", "+0;-0;0"),
    ("Overall progress (duration-weighted)", "IFERROR(SUM(tsk_WtDone)/SUM(tsk_WtDur),0)", "0%"),
    ("Planned progress by status date", "IFERROR(SUMPRODUCT(tsk_WtDur*(tsk_PlanFinish<StatusDate))/SUM(tsk_WtDur),0)", "0%"),
    ("Detail tasks", "SUM(tsk_IsLeaf)", "0"),
    ("Milestones", 'SUMPRODUCT(tsk_IsLeaf*(tsk_Duration=0)*(tsk_Name<>""))', "0"),
    ("Total work (hrs)", "SUMPRODUCT(tsk_Work,tsk_IsLeaf)", "#,##0"),
    ("Total cost", "SUMPRODUCT(tsk_Cost,tsk_IsLeaf)", "#,##0.00"),
    ("Current sprint", 'IF(OR(StatusDate<Sprint1Start,StatusDate>SprintLastEnd),"(outside sprint range)",'
                       'INDEX(spr_Name,MATCH(StatusDate,spr_Start,1)))', "@"),
    ("Baseline", 'IF(COUNT(tsk_BaseStart)=0,"Not set (planned = calculated)","Set on "&COUNT(tsk_BaseStart)&" tasks")', "@"),
    ("Open critical tasks", 'COUNTIFS(tsk_Critical,"Yes",tsk_IsLeaf,1)', "0"),
    ("Missed deadlines", 'COUNTIF(tsk_DeadlineCheck,"Late*")', "0"),
    ("Link problems (see Tasks > Link Check)", 'COUNTIF(tsk_LinkCheck,"?*")', "0"),
]
ALERT_IF_POSITIVE = {"Finish variance (work days)", "Missed deadlines", "Link problems (see Tasks > Link Check)"}


def build_dashboard(wb, sprints_ws, sprint_letters, n_chart_sprints=12):
    ws = wb.create_sheet("Dashboard")
    ws["A1"] = '=ProjectName&"  -  Dashboard"'
    ws["A1"].font = font(size=18, bold=True, color="1F4E78")
    ws["A2"] = '="Project manager: "&ProjectManager&"     Status date: "&TEXT(StatusDate,"ddd dd-mmm-yyyy")'
    ws["A2"].font = font(italic=True, color="595959")
    for col, width in zip("ABCDEF", (3, 38, 22, 4, 18, 12)):
        ws.column_dimensions[col].width = width

    add_table(ws, "KPIs", 4, 2, [
        Col("Key indicator", width=38, values=[k[0] for k in KPIS]),
        Col("Value", width=22, values=[f"={k[1]}" for k in KPIS]),
    ], len(KPIS), autofilter=False)
    for i, (label, _, fmt) in enumerate(KPIS):
        r = 5 + i
        ws.cell(r, 2).font = font(bold=True)
        value = ws.cell(r, 3)
        value.number_format = fmt
        value.font = font(size=11, bold=True)
        value.alignment = Alignment(horizontal="right")
        if label in ALERT_IF_POSITIVE:
            ws.conditional_formatting.add(f"C{r}", FormulaRule(formula=[f"C{r}>0"], font=Font(color="C00000", bold=True),
                                                               fill=fill("FFC7CE")))
            ws.conditional_formatting.add(f"C{r}", FormulaRule(formula=[f"C{r}<=0"], font=Font(color="548235", bold=True)))

    status = Tab("StatusSummary")
    add_table(ws, "StatusSummary", 4, 5, [
        Col("Status", width=18, values=list(STATUS_COLOURS)),
        Col("Count", width=12, formula=f"COUNTIFS(tsk_Status,{status.row('Status')},tsk_IsLeaf,1)"),
    ], len(STATUS_COLOURS), autofilter=False)
    for i, (s, colour) in enumerate(STATUS_COLOURS.items()):
        ws.cell(5 + i, 5).font = font(bold=True, color=colour)
        ws.cell(5 + i, 6).font = font(bold=True)

    top = 5 + len(KPIS) + 2
    nxt = Tab("NextMilestones")
    k = f"(ROW()-ROW({nxt.head('Next milestones (forecast)')}))"
    nth = f"SMALL(tsk_MilestoneKey,{k})"
    add_table(ws, "NextMilestones", top, 2, [
        Col("Next milestones (forecast)", width=38, formula=f'IFERROR(INDEX(tsk_Name,MATCH({nth},tsk_MilestoneKey,0)),"")'),
        Col("Forecast date", width=22, fmt=DATE, formula=f'IFERROR(INDEX(tsk_FcFinish,MATCH({nth},tsk_MilestoneKey,0)),"")'),
    ], 5, autofilter=False)

    chart = BarChart()
    chart.type, chart.style, chart.title, chart.legend = "bar", 10, "Detail tasks by status", None
    chart.add_data(Reference(ws, min_col=6, min_row=4, max_row=4 + len(STATUS_COLOURS)), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=5, min_row=5, max_row=4 + len(STATUS_COLOURS)))
    chart.height, chart.width = 7, 14
    ws.add_chart(chart, "E12")

    work = sprints_ws[f"{sprint_letters['Planned Work']}5"].column
    cap = sprints_ws[f"{sprint_letters['Capacity']}5"].column
    nm = sprints_ws[f"{sprint_letters['Sprint']}5"].column
    chart2 = BarChart()
    chart2.title = "Sprint planned work vs capacity (hrs)"
    chart2.add_data(Reference(sprints_ws, min_col=work, min_row=5, max_row=5 + n_chart_sprints), titles_from_data=True)
    chart2.add_data(Reference(sprints_ws, min_col=cap, min_row=5, max_row=5 + n_chart_sprints), titles_from_data=True)
    chart2.set_categories(Reference(sprints_ws, min_col=nm, min_row=6, max_row=5 + n_chart_sprints))
    chart2.height, chart2.width = 8, 22
    ws.add_chart(chart2, f"B{top + 8}")
    return ws
