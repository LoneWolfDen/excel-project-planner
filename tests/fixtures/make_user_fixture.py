"""Write tests/fixtures/planner_v1.8_user.xlsx: the v1.8 sample as a project manager would leave it after some weeks
of use, for TypedDataTest (B-1012). No real user workbook is available, so this stands in for one: it types into
every kind of input the workbook has, including the ones that hold a formula by default (stages, sprint length).

    python tests/fixtures/make_user_fixture.py

The file is committed; rerun this only when the template's inputs change, and say so in the PR.
"""
import datetime as dt
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from openpyxl.utils.cell import range_boundaries  # noqa: E402

from planner.build import make_workbook  # noqa: E402

D = dt.date
OUT = Path(__file__).resolve().parent / "planner_v1.8_user.xlsx"


def table(wb, name):
    """(sheet, {header: column}, first data row) of an Excel Table."""
    for ws in wb.worksheets:
        if name in ws.tables:
            c1, r1, c2, _ = range_boundaries(ws.tables[name].ref)
            return ws, {ws.cell(r1, c).value: c for c in range(c1, c2 + 1)}, r1 + 1
    raise KeyError(name)


def put(wb, name, row, values):
    """Type values (header -> value) into data row `row` (0 = first) of table `name`."""
    ws, col, first = table(wb, name)
    for header, value in values.items():
        ws.cell(first + row, col[header]).value = value


def named(wb, n, value):
    sheet, ref = wb.defined_names[n].attr_text.split("!")
    wb[sheet.strip("'")][ref.replace("$", "")].value = value


def main():
    wb = make_workbook(baselined=True)
    # settings: a typed sprint length and thresholds over the project type's defaults; reschedule off; a status date
    named(wb, "SprintLength", 10)
    named(wb, "RAGRedSlipDays", 8)
    named(wb, "RescheduleAfterStatus", "No")
    named(wb, "StatusDateInput", D(2026, 11, 1))
    named(wb, "ProjectManager", "R. Okonkwo")
    # the team renamed two stages and changed a %; added a project type and a resource type to the lists
    put(wb, "StageTable", 2, {"Stage": "Building", "% Complete": 0.45})
    put(wb, "StageTable", 3, {"Stage": "Ready for Test"})
    put(wb, "ProjectPresets", 3, {"Project Type": "Data migration", "Sprint Length": 7, "Amber Slip": 2,
                                  "Red Slip": 4, "Stage 1": "Not Started", "% 1": 0, "Stage 2": "Mapped",
                                  "% 2": 0.3, "Stage 3": "Loaded", "% 3": 0.7, "Stage 4": "Reconciled", "% 4": 0.9,
                                  "Stage 5": "Signed Off", "% 5": 0.95, "Stage 6": "Complete", "% 6": 1})
    put(wb, "ListResType", 3, {"Resource Type": "Contractor"})
    put(wb, "HolidayTable", 4, {"Date": D(2026, 11, 26), "Description": "Thanksgiving (US team)"})
    # tasks: progress, notes, an owner, a pasted Predecessors text with a lag, a typed % override
    tasks = {13: {"Stage": "Building", "Notes": "API contract agreed with the payments team 22-Oct."},
             14: {"Actual Start": D(2026, 10, 28), "% Complete": 0.2},
             17: {"Predecessors": "16, 13SS+3d", "Notes": "Waits for the design system v2."},
             20: {"Owner": "Taylor Morgan"}, 21: {"Story Points": 13}}
    for uid, values in tasks.items():
        put(wb, "Tasks", uid - 1, values)
    # people: leave, a later start, an allocation change, a personal focus
    put(wb, "ResourceTable", 2, {"Allocation %": 0.6, "Focus %": 0.7, "Start": D(2026, 9, 21), "Nov-26": "9, 10, 11"})
    put(wb, "ResourceTable", 5, {"Dec-26": "21, 22, 23, 24(0.5)", "Notes": "Covers hypercare in January."})
    put(wb, "RateCard", 1, {"Day Rate": 590})
    # sprints: goals and one moved end date
    put(wb, "SprintTable", 3, {"Goal": "Payments end to end in staging", "Set End": D(2026, 11, 6)})
    put(wb, "SprintTable", 4, {"Goal": "Accessibility fixes and SIT entry"})
    # status: PM RAG overrides with reasons, this period's notes
    put(wb, "StatusRAG", 2, {"PM RAG": "Amber", "Comment": "Build complete deadline at risk; recovery plan agreed."})
    put(wb, "StatusRAG", 4, {"PM RAG": "Green", "Comment": "Two risks closed this week."})
    put(wb, "StatusNotes", 0, {"Notes": "Payments integration started; API 60% done."})
    put(wb, "StatusNotes", 4, {"Notes": "Approve the UAT window move to 7-Dec."})
    # a pasted history row and a cost snapshot row
    put(wb, "StatusHistory", 4, {"Status Date": D(2026, 10, 25), "Forecast Finish": D(2027, 1, 6),
                                 "Planned Finish": D(2026, 12, 31), "Finish Var": 4, "% Done": 0.5,
                                 "% Planned": 0.58, "SPI": 0.86, "Points Done": 38, "Total Points": 84,
                                 "Open Risks": 2, "Open Issues": 1, "Pending Changes": 1, "RAG": "Amber"})
    put(wb, "CostHistory", 4, {"Status Date": D(2026, 10, 25), "BAC": 55133, "PV": 33400, "EV": 29100,
                               "AC": 31040, "CPI": 0.94, "EAC": 58700})
    # the log: a new risk and a decision
    ws, col, first = table(wb, "Log")
    free = next(r for r in range(first, first + 150) if not ws.cell(r, col["Title"]).value)
    put(wb, "Log", free - first, {"Type": "Risk", "Title": "Payments sandbox outages", "Owner": "Jordan Lee",
                                  "Raised": D(2026, 10, 29), "Status": "Open", "Probability": 3, "Impact": 4,
                                  "Response": "Mock the gateway for SIT; escalate to the vendor.", "Task UID": 14})
    put(wb, "Log", free - first + 1, {"Type": "Decision", "Title": "Freeze scope for release 1",
                                      "Owner": "Priya Sharma", "Raised": D(2026, 10, 30), "Status": "Approved",
                                      "Options and Reasons": "Defer reporting module to release 2."})
    wb.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
