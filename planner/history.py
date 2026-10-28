"""History sheet: status snapshots for trend charts (Excel Table "StatusHistory").

A formula can't remember last week's value, so history is kept the way the baseline is: SnapshotNow, at the
bottom of the Status sheet, shows this status date's numbers (formulas), and on each reporting date the PM copies
that row and pastes it as values onto the first empty row of StatusHistory (inputs). The trend charts read
StatusHistory: forecast vs planned finish, progress done vs planned, the schedule SPI, and a story-point burn-up that
shows scope changes. Money is not here: the Cost sheet keeps its own history, so deleting it leaves no cost behind.
Nothing here needs a macro; a macro or script would only automate the same copy and paste.
"""
import datetime as dt

from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter

from .core import DATE, Col, add_table, font, missing, title, v11

D = dt.date
# header, formula (names only), number format
FIELDS = [
    ("Status Date", "StatusDate", DATE),
    ("Forecast Finish", "ProjectFinish", DATE),
    ("Planned Finish", "PlannedFinish", DATE),
    ("Finish Var", "FinishVarianceDays", "+0;-0;0"),
    ("% Done", "sch_Progress", "0%"),
    ("% Planned", "sch_Planned", "0%"),
    ("SPI", "sch_SPI", "0.00"),
    ("Points Done", "agl_Done", "0.#"),
    ("Total Points", "TotalPoints", "0.#"),
    ("Open Risks", 'COUNTIFS(log_Type,"Risk",log_Open,1)', "0"),
    ("Open Issues", 'COUNTIFS(log_Type,"Issue",log_Open,1)', "0"),
    ("Pending Changes", 'COUNTIFS(log_Type,"Change",log_Open,1)', "0"),
    ("RAG", "ProjectRAG", "@"),
]
HISTORY_ROWS = 60
# template 1.1 (F-1106): risk exposure = the open RAID items' scores added up (probability x impact), so its trend is
# a burndown of how much risk the project carries
V11_FIELDS = [("Risk Exposure", "SUMIFS(log_Score,log_Open,1)", "0")]
EXPOSURE = [44, 50, 47, 52]                       # the sample's four snapshots


def fields():
    """(position in FIELDS, field) for the fields this build can calculate (B-913: no points without Agile, no
    RAID counts without the Log; template 1.1 adds V11_FIELDS after them)."""
    every = FIELDS + (V11_FIELDS if v11() else [])
    return [(i, f) for i, f in enumerate(every) if not missing(f[1])]


def sample_history():
    """Four weekly snapshots before the sample's status date (illustrative; the PM pastes the real ones)."""
    rows = []
    for when, finish, slip, done_pct, planned, done, total, risks, issues, pending, rag in (
            (D(2026, 9, 27), D(2026, 12, 28), 0, 0.17, 0.19, 5, 76, 2, 0, 0, "Green"),
            (D(2026, 10, 4), D(2026, 12, 29), 1, 0.27, 0.31, 18, 78, 3, 0, 1, "Green"),
            (D(2026, 10, 11), D(2026, 12, 31), 3, 0.36, 0.43, 23, 82, 3, 0, 0, "Amber"),
            (D(2026, 10, 18), D(2027, 1, 4), 5, 0.44, 0.54, 33, 82, 4, 1, 0, "Amber")):
        rows.append([when, finish, D(2026, 12, 28), slip, done_pct, planned, round(done_pct / planned, 2), done, total,
                     risks, issues, pending, rag])
    if v11():
        rows = [row + [exposure] for row, exposure in zip(rows, EXPOSURE)]
    return rows


def build_history(wb, blank):
    ws = wb.create_sheet("History")
    title(ws, "Status History",
          "On each reporting date: on the Status sheet, select the SnapshotNow row at the bottom, Copy, then click "
          "the first empty row of the table below and Paste Special > Values. The trend charts update from it.")
    ws.column_dimensions["A"].width = 3
    data = [] if blank else sample_history()
    top = 5
    ws.cell(top - 1, 2, "StatusHistory: one row per reporting date (values pasted from Status > SnapshotNow)").font = \
        font(bold=True, color="1F4E78")
    add_table(ws, "StatusHistory", top, 2, [Col(h, "in", 11 if fmt != DATE else 13, fmt, center=True,
                                                values=[row[i] for row in data])
                                            for i, (h, _, fmt) in fields()], HISTORY_ROWS)
    _charts(ws, top)
    ws.freeze_panes = f"C{top + 1}"
    return ws


def _charts(ws, top):
    """Trend charts over the StatusHistory rows (charts need A1 ranges; empty rows leave no points)."""
    col = {h: 2 + k for k, (_, (h, _, _)) in enumerate(fields())}
    first, last = top, top + HISTORY_ROWS
    dates = Reference(ws, min_col=col["Status Date"], min_row=first + 1, max_row=last)

    def chart(title_text, headers, anchor, y_format=None):
        if any(h not in col for h in headers):
            return
        c = LineChart()
        c.title = title_text
        for h in headers:
            c.add_data(Reference(ws, min_col=col[h], min_row=first, max_row=last), titles_from_data=True)
        for s in c.series:
            s.smooth = False
        c.set_categories(dates)
        c.x_axis.number_format = "dd-mmm"
        if y_format:
            c.y_axis.number_format = y_format
        c.height, c.width = 7, 14
        ws.add_chart(c, anchor)

    right = 2 + len(fields()) + 1
    at = lambda row: f"{get_column_letter(right)}{row}"
    chart("Forecast vs planned finish", ["Forecast Finish", "Planned Finish"], at(4), "dd-mmm-yy")
    chart("Progress: done vs planned", ["% Done", "% Planned"], at(19), "0%")
    chart("SPI (schedule)", ["SPI"], at(34), "0.00")
    chart("Story-point burn-up (done vs scope)", ["Points Done", "Total Points"], at(49))
    chart("Risk exposure (open RAID score)", ["Risk Exposure"], at(64), "0")
