"""Performance and structure checks for planner workbooks. Run on every PR (see .github/workflows/ci.yml).

    python scripts/perf_check.py Project_Planner.xlsx --profile sample
    python scripts/perf_check.py stress.xlsx --profile stress --timing

Static checks (openpyxl only):
  * formula count per sheet
  * volatile functions (OFFSET, INDIRECT, TODAY, NOW, RAND, ...): every use must be on the allow-list in
    perf_budget.json, because a volatile cell makes Excel recalculate it, and everything that depends on it,
    after every edit
  * A1 references inside formulas: there must be none; formulas use structured references and defined names
  * Excel's hard limits: balanced parentheses, at most 8,192 characters and 64 nested functions per formula
    (LibreOffice silently repairs some of these; Excel refuses the formula or repairs the file)
  * conditional-formatting rules per sheet (B-1016): a budget on the total, and no sheet may carry many rules
    that each cover one row (a rule per task row: Excel re-evaluates every rule on screen at each redraw, and
    the file grows with the plan); one rule over a whole column range does the same job
Timing (LibreOffice via python3-uno, with --timing):
  * full recalculation and a single-edit recalculation, median of three runs

Writes a Markdown report to stdout (and to $GITHUB_STEP_SUMMARY in CI) and exits 1 when a budget is exceeded.
"""
import argparse
import json
import os
import re
import statistics
import subprocess
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.formula import Tokenizer

ROOT = Path(__file__).resolve().parents[1]
SINGLE_ROW_RANGES_MAX = 20      # per sheet: more looks like a rule written for each row
VOLATILE = ("OFFSET", "INDIRECT", "TODAY", "NOW", "RAND", "RANDBETWEEN", "CELL", "INFO")
A1 = re.compile(r"^(?:'?[^!']+'?!)?\$?(?:[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?|[A-Z]{1,3}:\$?[A-Z]{1,3}|\d+:\$?\d+)$")


def formulas(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    yield ws.title, c, c.value


MAX_LENGTH, MAX_NESTING = 8192, 64


def excel_limit_problems(formula):
    """Problems Excel would reject: unbalanced parentheses, too long, nested too deeply."""
    depth = deepest = 0
    in_string = False
    brackets = 0
    for ch in formula:
        if ch == '"':
            in_string = not in_string
        elif in_string:
            continue
        elif ch == "[":
            brackets += 1
        elif ch == "]":
            brackets -= 1
        elif brackets:
            continue                                   # inside a structured reference
        elif ch == "(":
            depth += 1
            deepest = max(deepest, depth)
        elif ch == ")":
            depth -= 1
            if depth < 0:
                return ["unbalanced parentheses"]
    problems = []
    if depth or in_string or brackets:
        problems.append("unbalanced parentheses, quotes or brackets")
    if len(formula) > MAX_LENGTH:
        problems.append(f"{len(formula)} characters (Excel allows {MAX_LENGTH})")
    if deepest > MAX_NESTING:
        problems.append(f"nesting depth {deepest} (Excel allows {MAX_NESTING})")
    return problems


def a1_references(formula, names):
    """A1-style cell or range operands in a formula (structured references and defined names are fine)."""
    found = []
    for tok in Tokenizer(formula).items:
        if tok.type == "OPERAND" and tok.subtype == "RANGE":
            v = tok.value
            if "[" in v or v in names or v.split("!")[-1] in names:
                continue
            if A1.match(v):
                found.append(v)
    return found


def missing_references(wb):
    """Operands that name a sheet, table or defined name the workbook doesn't have, in cell formulas, data
    validation, conditional formatting and defined names: what a build that leaves a module out must never keep
    (it would show #NAME? or #REF! in Excel). A1 addresses (formats, validation) are not names, so they pass."""
    names, sheets = set(wb.defined_names), set(wb.sheetnames)
    tables = {t for ws in wb.worksheets for t in ws.tables}
    found = []

    def check(where, formula):
        for tok in Tokenizer(formula).items:
            if tok.type != "OPERAND" or tok.subtype != "RANGE":
                continue
            v = tok.value.lstrip(":")                  # INDEX(...):Table[...] arrives as ":Table[...]"
            if "!" in v:
                ok = v.split("!")[0].strip("'") in sheets
            elif "[" in v:
                ok = v.split("[")[0] in tables
            else:
                ok = v in names or bool(A1.match(v)) or v.upper() in ("TRUE", "FALSE")
            if not ok:
                found.append(f"{where}: {v}")

    for sheet, cell, f in formulas(wb):
        check(f"{sheet}!{cell.coordinate}", f)
    for ws in wb.worksheets:
        for dv in ws.data_validations.dataValidation:
            for f in (dv.formula1, dv.formula2):
                if f:
                    check(f"{ws.title} validation {dv.sqref}", "=" + f.lstrip("="))
        for rng in ws.conditional_formatting:
            for rule in rng.rules:
                for f in rule.formula or []:
                    check(f"{ws.title} format {rng.sqref}", "=" + f)
    for n, d in wb.defined_names.items():
        if d.attr_text and not d.attr_text.startswith('"'):
            check(f"name {n}", "=" + d.attr_text)
    return found


def cf_report(wb):
    """{sheet: {"ranges", "rules", "single_row"}}: conditional-formatting ranges, the rules on them, and how many
    ranges cover a single row (the one-rule-per-row pattern)."""
    from openpyxl.utils.cell import range_boundaries
    out = {}
    for ws in wb.worksheets:
        ranges = rules = single = 0
        for cf in ws.conditional_formatting:
            ranges += 1
            rules += len(cf.rules)
            single += sum(1 for r in cf.sqref.ranges if range_boundaries(r.coord)[1] == range_boundaries(r.coord)[3])
        if ranges:
            out[ws.title] = {"ranges": ranges, "rules": rules, "single_row": single}
    return out


def static_report(path):
    wb = openpyxl.load_workbook(path)
    names = set(wb.defined_names)
    per_sheet = Counter()
    volatile = defaultdict(Counter)          # sheet -> function -> cells
    a1 = []
    limits = []
    longest = 0
    for sheet, cell, f in formulas(wb):
        coord = cell.coordinate
        longest = max(longest, len(f))
        for problem in excel_limit_problems(f):
            limits.append(f"{sheet}!{coord}: {problem}")
        per_sheet[sheet] += 1
        upper = f.upper()
        for fn in VOLATILE:
            if re.search(rf"(?<![A-Z_.]){fn}\(", upper):
                volatile[sheet][fn] += 1
        refs = a1_references(f, names)
        if refs:
            a1.append(f"{sheet}!{coord}: {', '.join(refs)}")
    return {"formulas": dict(per_sheet), "formulas_total": sum(per_sheet.values()),
            "volatile": {s: dict(c) for s, c in volatile.items()},
            "volatile_cells": sum(sum(c.values()) for c in volatile.values()), "a1_references": a1,
            "excel_limit_problems": limits, "longest_formula": longest, "missing_references": missing_references(wb),
            "cf": cf_report(wb)}


def time_recalc(path, runs=3):
    """Median seconds for a full recalculation and for single edits followed by an incremental recalculation.

    Edits: a duration near the top of the plan (worst case: every task below may move), the last task's duration,
    a note (should recalculate almost nothing), and the duration again with Settings > Live resource grids = No
    (B-1011: bulk entry without the per-person grids)."""
    import uno                                            # python3-uno, part of LibreOffice
    from com.sun.star.beans import PropertyValue

    def prop(n, v):
        p = PropertyValue()
        p.Name, p.Value = n, v
        return p

    wb = openpyxl.load_workbook(path, read_only=True)
    grids = wb.defined_names["LiveGrids"].attr_text if "LiveGrids" in wb.defined_names else None
    ws = wb["Tasks"]
    header_row = next(r for r in range(1, 20) if any(c.value == "Duration" for c in ws[r]))
    head = {c.value: c.column_letter for c in ws[header_row] if c.value}
    rows = [r for r, (v,) in enumerate(ws.iter_rows(min_row=header_row + 1, min_col=1, max_col=1, values_only=True),
                                       start=header_row + 1) if v not in (None, "")]
    wb.close()
    named = [r for r in rows]
    edits = {"edit_recalc_s": f"{head['Duration']}{header_row + 2}",        # second task: a detail task near the top
             "edit_last_task_s": f"{head['Duration']}{named[-1] if named else header_row + 2}",
             "edit_note_s": f"{head['Notes']}{header_row + 2}"}

    port = 2084
    proc = subprocess.Popen(["soffice", "--headless", "--invisible", "--norestore", "--nologo",
                             f"--accept=socket,host=localhost,port={port};urp;"],
                            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    out = {}
    try:
        ctx = uno.getComponentContext()
        resolver = ctx.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", ctx)
        for _ in range(120):
            try:
                remote = resolver.resolve(f"uno:socket,host=localhost,port={port};urp;StarOffice.ComponentContext")
                break
            except Exception:
                time.sleep(0.5)
        else:
            raise RuntimeError("LibreOffice did not start")
        desktop = remote.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", remote)
        doc = desktop.loadComponentFromURL(uno.systemPathToFileUrl(str(Path(path).resolve())), "_blank", 0,
                                           (prop("Hidden", True),))
        full = []
        for _ in range(runs):
            t = time.perf_counter()
            doc.calculateAll()
            full.append(time.perf_counter() - t)
        out["full_recalc_s"] = statistics.median(full)
        sheet = doc.Sheets.getByName("Tasks")
        for key, cell in edits.items():
            target = sheet.getCellRangeByName(cell)
            numeric = key != "edit_note_s"
            original = target.getValue() if numeric else target.getString()
            times = []
            for i in range(runs):
                if numeric:
                    target.setValue(original + 1 + i)
                else:
                    target.setString(f"{original} edit {i}")
                t = time.perf_counter()
                doc.calculate()
                times.append(time.perf_counter() - t)
            out[key] = statistics.median(times)
        if grids:
            sheet_name, ref = grids.split("!")
            switch = doc.Sheets.getByName(sheet_name.strip("'")).getCellRangeByName(ref.replace("$", ""))
            switch.setString("No")
            doc.calculate()
            target = sheet.getCellRangeByName(edits["edit_recalc_s"])
            original, times = target.getValue(), []
            for i in range(runs):
                target.setValue(original + 1 + i)
                t = time.perf_counter()
                doc.calculate()
                times.append(time.perf_counter() - t)
            out["edit_grids_off_s"] = statistics.median(times)
        doc.close(True)
        try:
            desktop.terminate()
        except Exception:
            pass
    finally:
        try:
            proc.wait(timeout=30)
        except subprocess.TimeoutExpired:
            proc.kill()
    return out


def check(report, budget):
    problems = []
    if report["formulas_total"] > budget["formulas_max"]:
        problems.append(f"{report['formulas_total']} formulas > budget {budget['formulas_max']}")
    allowed = budget["volatile_allowed"]
    for sheet, fns in report["volatile"].items():
        for fn, n in fns.items():
            limit = allowed.get(sheet, {}).get(fn, 0)
            if n > limit:
                problems.append(f"{sheet}: {n} cells use volatile {fn} (allowed {limit})")
    if report["excel_limit_problems"]:
        problems.append(f"{len(report['excel_limit_problems'])} formulas break Excel limits, e.g. "
                        + "; ".join(report["excel_limit_problems"][:5]))
    if report["a1_references"]:
        problems.append(f"{len(report['a1_references'])} formulas contain A1 references, e.g. "
                        + "; ".join(report["a1_references"][:5]))
    if report["missing_references"]:
        problems.append(f"{len(report['missing_references'])} references to a missing sheet, table or name, e.g. "
                        + "; ".join(report["missing_references"][:5]))
    cf = report.get("cf", {})
    total = sum(v["rules"] for v in cf.values())
    if "cf_rules_max" in budget and total > budget["cf_rules_max"]:
        problems.append(f"{total} conditional-formatting rules > budget {budget['cf_rules_max']}")
    for sheet, v in cf.items():
        if v["single_row"] > SINGLE_ROW_RANGES_MAX:
            problems.append(f"{sheet}: {v['single_row']} conditional-formatting ranges cover one row each "
                            f"(at most {SINGLE_ROW_RANGES_MAX}): use one rule over the whole column range")
    for k in ("full_recalc_s", "edit_recalc_s"):
        if k in report and report[k] > budget[k + "_max"]:
            problems.append(f"{k} {report[k]:.3f}s > budget {budget[k + '_max']}s")
    return problems


def markdown(path, profile, report, budget, problems):
    lines = [f"### Performance: `{Path(path).name}` ({profile})", "",
             "| Metric | Value | Budget |", "|---|---|---|",
             f"| Formulas | {report['formulas_total']:,} | {budget['formulas_max']:,} |",
             f"| Volatile formula cells | {report['volatile_cells']:,} | allow-listed per sheet |",
             f"| Formulas with A1 references | {len(report['a1_references'])} | 0 |",
             f"| References to a missing sheet, table or name | {len(report['missing_references'])} | 0 |",
             f"| Formulas breaking Excel limits (parentheses, length, nesting) | {len(report['excel_limit_problems'])} | 0 |",
             f"| Longest formula | {report['longest_formula']:,} chars | {MAX_LENGTH:,} |",
             f"| Conditional-formatting rules | {sum(v['rules'] for v in report.get('cf', {}).values()):,} | "
             f"{budget.get('cf_rules_max', '-')} |"]
    for k, label in (("full_recalc_s", "Full recalculation (LibreOffice)"),
                     ("edit_recalc_s", "Edit a duration near the top, then recalculate"),
                     ("edit_last_task_s", "Edit the last task's duration"), ("edit_note_s", "Edit a note"),
                     ("edit_grids_off_s", "Edit a duration with Live resource grids off")):
        if k in report:
            limit = budget.get(k + "_max")
            lines.append(f"| {label} | {report[k]:.3f} s | {f'{limit} s' if limit else '-'} |")
    lines += ["", "Formulas by sheet: " + ", ".join(f"{s} {n:,}" for s, n in sorted(report["formulas"].items())), ""]
    lines += ["Conditional-formatting rules by sheet: " + ", ".join(
        f"{s} {v['rules']}" for s, v in sorted(report.get("cf", {}).items())), ""]
    vol = "; ".join(f"{s}: " + ", ".join(f"{fn} {n}" for fn, n in fns.items()) for s, fns in report["volatile"].items())
    lines += [f"Volatile functions: {vol or 'none'}", ""]
    lines += (["**Budget exceeded:**", ""] + [f"- {p}" for p in problems]) if problems else ["All budgets met."]
    return "\n".join(lines) + "\n"


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("workbook")
    p.add_argument("--profile", default="sample", help="budget profile in perf_budget.json")
    p.add_argument("--budget", default=str(ROOT / "perf_budget.json"))
    p.add_argument("--timing", action="store_true", help="also time recalculation in LibreOffice (needs python3-uno)")
    p.add_argument("--json", help="write the raw report here")
    a = p.parse_args(argv)
    budget = json.loads(Path(a.budget).read_text())[a.profile]
    report = static_report(a.workbook)
    if a.timing:
        report.update(time_recalc(a.workbook))
    problems = check(report, budget)
    md = markdown(a.workbook, a.profile, report, budget, problems)
    print(md)
    if os.environ.get("GITHUB_STEP_SUMMARY"):
        with open(os.environ["GITHUB_STEP_SUMMARY"], "a") as fh:
            fh.write(md)
    if a.json:
        Path(a.json).write_text(json.dumps(report, indent=2))
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
