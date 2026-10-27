"""Gantt sheet: a task panel (Excel Table "GanttView") and a timeline drawn purely with conditional formatting.

The timeline cells hold no formulas at all: bars, the planned outline, milestones, deadlines, the status date and
non-working days are conditional-format rules, which Excel only evaluates for the cells on screen. That keeps the
calculation chain small (see docs/ARCHITECTURE.md).
"""
from openpyxl.formatting.rule import FormulaRule, Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from .core import (HEADER_FILLS, NOTE_FONT, SHORT_DATE, Col, Tab, add_table, fill, font, name)
from .tasks import group_columns, group_outline

TABLE = "GanttView"
HEADER_ROW = 5
WD = "WorkWeek,Holidays"
# planned (baseline) = border, current schedule = fill (docs/REAL_WORLD_FIT.md): red is kept for slip
COLOURS = {"critical": "F4B183", "summary": "404040", "done": "70AD47", "remaining": "9DC3E6", "slip": "E06666",
           "milestone": "7030A0", "deadline": "C00000", "planned": "000000", "planned_gap": "C9D5E8", "status": "FF0000",
           "nonworking": "EDEDED"}


def period_end(d):
    """The day after the timeline column that starts on d: GanttPeriod 1 / 7 are days, 30 / 91 calendar months /
    quarters (EDATE, so months keep their real lengths)."""
    return f"IF(GanttPeriod<28,{d}+GanttPeriod,EDATE({d},IF(GanttPeriod>=90,3,1)))"


def column_start(k):
    """The first day of timeline column k (1 = first). Month and quarter scales start on the 1st of GanttStart's
    month / calendar quarter."""
    first_month = ("DATE(YEAR(GanttStart),MONTH(GanttStart)-IF(GanttPeriod>=90,MOD(MONTH(GanttStart)-1,3),0),1)")
    return (f"IF(GanttPeriod<28,GanttStart+({k}-1)*GanttPeriod,"
            f"EDATE({first_month},({k}-1)*IF(GanttPeriod>=90,3,1)))")


def build_gantt(wb, n_rows, n_days, levels=()):
    """levels: the Level of each task the workbook is built with, for the Excel outline grouping."""
    ws = wb.create_sheet("Gantt")
    ws["A1"] = '=ProjectName&"  -  Gantt Chart"'
    ws["A1"].font = font(size=16, bold=True, color="1F4E78")
    ws["A2"] = ('="Status date: "&TEXT(StatusDate,"ddd dd-mmm-yyyy")&"   |   Forecast finish: "&TEXT(MAX(tsk_FcFinish),'
                '"dd-mmm-yyyy")&"   |   Planned finish: "&TEXT(MAX(tsk_PlanFinish),"dd-mmm-yyyy")&IF(COUNTA(tsk_Name)>ROWS('
                'GanttView[ID]),"   |   Showing the first "&ROWS(GanttView[ID])&" tasks","")')
    ws["A2"].font = font(bold=True, color="404040")

    g = Tab(TABLE)
    n = g.row("_Row")
    blank = f'IFERROR(INDEX(tsk_Name,{n})="",TRUE())'          # also true past the end of the Tasks table
    task = lambda col: f"INDEX({col},{n})"
    show = lambda col: f'IF({blank},"",{task(col)})'
    pct = g.row("%")
    fs, ff = g.row("Start"), g.row("Finish")
    # the panel's essentials, then the rest in a collapsed column group (+ above it), so more timeline fits
    cols = [
        Col("ID", width=5, formula=show("tsk_ID"), center=True),
        Col("Task", width=34, formula=f'IF({blank},"",REPT("    ",{task("tsk_Level")}-1)&{task("tsk_Name")})'),
        Col("Start", width=10, fmt=SHORT_DATE, formula=show("tsk_FcStart")),
        Col("Finish", width=10, fmt=SHORT_DATE, formula=show("tsk_FcFinish")),
        Col("%", width=6, fmt="0%", center=True, formula=show("tsk_Progress")),
        Col("Resources", width=13, formula=f'IF({blank},"",{task("tsk_ResourceNames")}&"")'),
        Col("WBS", width=7, formula=show("tsk_WBS"), center=True),
        Col("Dur", width=5, center=True, formula=(
            f'IF({blank},"",IF({task("tsk_IsSummary")},NETWORKDAYS.INTL({task("tsk_FcStart")},{task("tsk_FcFinish")},{WD}),'
            f'N({task("tsk_Duration")})))')),
        Col("Planned Start", width=10, fmt=SHORT_DATE, formula=show("tsk_PlanStart")),
        Col("Planned Finish", width=10, fmt=SHORT_DATE, formula=show("tsk_PlanFinish")),
        Col("Status", width=11, center=True, formula=show("tsk_Status")),
        Col("Show", width=6, center=True, formula=f'IF({blank},"",{task("tsk_Show")}&"")',
            note="Yes down to Settings > OutlineShowLevel: filter on Yes to collapse the outline."),
        # helpers read by the timeline's conditional formatting
        Col("_Row", "help", 5, formula=f"ROW()-ROW({g.head('ID')})"),
        Col("_IsSummary", "help", 5, formula=f"IF({blank},FALSE(),{task('tsk_IsSummary')})"),
        Col("_IsMilestone", "help", 5, formula=f"IF({blank},FALSE(),AND({task('tsk_IsLeaf')}=1,N({task('tsk_Duration')})=0))"),
        Col("_DoneFrom", "help", 8, "dd-mmm;;-", formula=(
            f"IF({blank},0,IF(AND({task('tsk_IsLeaf')}=1,NOT({g.row('_IsMilestone')}),N({pct})>0),{fs},0))")),
        Col("_DoneTo", "help", 8, "dd-mmm;;-", formula=(
            f"IF({g.row('_DoneFrom')}=0,0,IF({pct}>=1,{ff},MAX({g.row('_DoneFrom')},{fs}+ROUND(({ff}-{fs}+1)*{pct},0)-1)))")),
        Col("_RemFrom", "help", 8, "dd-mmm;;-", formula=(
            f"IF({blank},0,IF(OR({task('tsk_IsLeaf')}<>1,{g.row('_IsMilestone')},N({pct})>=1),0,"
            f"IF({g.row('_DoneTo')}>0,{g.row('_DoneTo')}+1,{fs})))")),
        Col("_RemTo", "help", 8, "dd-mmm;;-", formula=f"IF({g.row('_RemFrom')}=0,0,{ff})"),
        Col("_ShowPlan", "help", 5, formula=f"IF({blank},FALSE(),OR(NOT(BaselineSet),N({task('tsk_BaseStart')})>0))",
            note="Draw the planned outline: every task until a baseline is set, then only tasks that have one (a task "
                 "added after the baseline shows its bar without an outline)."),
        Col("_Deadline", "help", 8, "dd-mmm;;-", formula=f"IF({blank},0,N({task('tsk_Deadline')}))"),
        Col("_Critical", "help", 5, formula=f'IF({blank},FALSE(),{task("tsk_Critical")}="Yes")'),
    ]
    L = add_table(ws, TABLE, HEADER_ROW, 1, cols, n_rows)
    group_outline(ws, levels, HEADER_ROW + 1)
    for c in cols:
        for r in range(HEADER_ROW + 1, HEADER_ROW + n_rows + 1):
            ws[f"{L[c.name]}{r}"].font = font(size=9)
    ws.row_dimensions[HEADER_ROW].height = 30
    ws.column_dimensions.group(L["_Row"], L["_Critical"], hidden=True, outline_level=1)
    ws.sheet_properties.outlinePr.summaryRight = False             # the + sits above Resources, left of the group
    group_columns(ws, L["WBS"], L["Status"], 1)                    # Show stays visible between the two groups

    origin_col = len(cols) + 1                       # narrow spacer column; the timeline starts after it
    ws.column_dimensions[get_column_letter(origin_col)].width = 1
    name(wb, "gantt_Origin", f"Gantt!${get_column_letter(origin_col)}${HEADER_ROW}")
    first, last = origin_col + 1, origin_col + n_days
    k = "(COLUMN()-COLUMN(gantt_Origin))"
    for col in range(first, last + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = 2.9
        day = column_start(k)
        ws.cell(HEADER_ROW, col, f"={day}").number_format = "d"
        # row 4: month (days / weeks) or year (months / quarters) where it changes; row 3: weekday, month or quarter
        ws.cell(4, col, f'=IF(GanttPeriod<28,IF(OR({k}=1,DAY({day})<=GanttPeriod),TEXT({day},"mmm yy"),""),'
                        f'IF(OR({k}=1,MONTH({day})<=IF(GanttPeriod>=90,3,1)),TEXT({day},"yyyy"),""))')
        ws.cell(3, col, f'=IF(GanttPeriod=1,LEFT(TEXT({day},"ddd"),1),IF(GanttPeriod=30,LEFT(TEXT({day},"mmm"),1),'
                        f'IF(GanttPeriod>=90,"Q"&INT((MONTH({day})+2)/3),"")))')
        for r, size, colour in ((3, 8, "404040"), (4, 8, "404040"), (5, 8, "FFFFFF")):
            c = ws.cell(r, col)
            c.font = font(size=size, bold=(r == 4), color=colour)
            c.alignment = Alignment(horizontal="left" if r == 4 else "center", vertical="center")
        ws.cell(HEADER_ROW, col).fill = HEADER_FILLS["calc"]

    _timeline_rules(ws, L, first, last, n_rows)
    _legend(ws, first)
    ws.freeze_panes = f"{get_column_letter(first)}{HEADER_ROW + 1}"
    ws.sheet_view.zoomScale = 90
    return ws


def _timeline_rules(ws, L, first, last, n_rows):
    top, bottom = HEADER_ROW + 1, HEADER_ROW + n_rows
    g1 = get_column_letter(first)
    grid = f"{g1}{top}:{get_column_letter(last)}{bottom}"
    d = f"{g1}${HEADER_ROW}"                                   # this column's date (header row)
    at = lambda col: f"${L[col]}{top}"                         # this row's helper cell
    end = period_end(d)
    in_period = lambda x: f"AND({x}>0,{x}>={d},{x}<{end})"
    overlaps = lambda a, b: f"AND({a}>0,{d}<={b},{end}-1>={a})"
    thin = lambda colour: Side(style="thin", color=colour)
    ps, pf = at("Planned Start"), at("Planned Finish")
    planned = f"AND({at('_ShowPlan')},ISNUMBER({ps}),{overlaps(ps, pf)})"
    first_day, last_day = in_period(ps), f"AND({pf}>0,{pf}>={d},{pf}<{end})"
    edge = lambda *sides: Border(**{s: thin(COLOURS["planned"]) for s in sides})
    # Excel applies every rule that is true and, for each property, the first one listed wins
    rules = [
        FormulaRule(formula=[in_period(at("_Deadline"))], fill=fill(COLOURS["deadline"])),
        FormulaRule(formula=[f"AND({at('_IsSummary')},{overlaps(at('Start'), at('Finish'))})"],
                    fill=fill(COLOURS["summary"])),
        FormulaRule(formula=[f"AND({at('_IsMilestone')},{in_period(at('Start'))})"], fill=fill(COLOURS["milestone"])),
        FormulaRule(formula=[overlaps(at("_DoneFrom"), at("_DoneTo"))], fill=fill(COLOURS["done"])),
        # slip: remaining work after the planned finish (only where the outline is drawn)
        FormulaRule(formula=[f"AND({at('_ShowPlan')},ISNUMBER({pf}),{d}>{pf},{overlaps(at('_RemFrom'), at('_RemTo'))})"],
                    fill=fill(COLOURS["slip"])),
        FormulaRule(formula=[f"AND({at('_Critical')},{overlaps(at('_RemFrom'), at('_RemTo'))})"], fill=fill(COLOURS["critical"])),
        FormulaRule(formula=[overlaps(at("_RemFrom"), at("_RemTo"))], fill=fill(COLOURS["remaining"])),
        # planned days the current bar has left (work moved later or earlier): a pale gap inside the outline
        FormulaRule(formula=[planned], fill=fill(COLOURS["planned_gap"])),
        # the planned outline: top and bottom on every planned day, closed on its first and last day
        FormulaRule(formula=[f"AND({planned},{first_day},{last_day})"], border=edge("top", "bottom", "left", "right")),
        FormulaRule(formula=[f"AND({planned},{first_day})"], border=edge("top", "bottom", "left")),
        FormulaRule(formula=[f"AND({planned},{last_day})"], border=edge("top", "bottom", "right")),
        FormulaRule(formula=[planned], border=edge("top", "bottom")),
        FormulaRule(formula=[in_period("StatusDate")],
                    border=Border(left=thin(COLOURS["status"]), right=thin(COLOURS["status"]))),
        FormulaRule(formula=[f"AND(GanttPeriod=1,NETWORKDAYS.INTL({d},{d},{WD})=0)"], fill=fill(COLOURS["nonworking"])),
    ]
    for rule in rules:
        ws.conditional_formatting.add(grid, rule)
    header = f"{g1}{HEADER_ROW}:{get_column_letter(last)}{HEADER_ROW}"
    ws.conditional_formatting.add(header, FormulaRule(formula=[in_period("StatusDate")], fill=fill("C00000")))
    # month / quarter scales: rows 3-4 name the month and year, so the header's day number (always 1) is hidden
    hide = DifferentialStyle(numFmt=NumberFormat(numFmtId=200, formatCode=";;;"))
    ws.conditional_formatting.add(header, Rule(type="expression", formula=["GanttPeriod>=28"], dxf=hide))
    panel = f"A{top}:{L['Status']}{bottom}"
    ws.conditional_formatting.add(panel, FormulaRule(formula=[at("_IsSummary")], font=Font(bold=True), fill=fill("DDEBF7")))
    from .tasks import STATUS_COLOURS
    for status, colour in STATUS_COLOURS.items():
        ws.conditional_formatting.add(f"{L['Status']}{top}:{L['Status']}{bottom}",
                                      FormulaRule(formula=[f'{at("Status")}="{status}"'], font=Font(color=colour, bold=True)))


def _legend(ws, first):
    planned = Side(style="thin", color=COLOURS["planned"])
    items = [("Planned (outline)", None, Border(top=planned, bottom=planned, left=planned, right=planned)),
             ("Done (% complete)", COLOURS["done"], None), ("Remaining - on track", COLOURS["remaining"], None),
             ("Slip: past the planned finish", COLOURS["slip"], None), ("Remaining - critical", COLOURS["critical"], None),
             ("Planned, no work now", COLOURS["planned_gap"], None), ("Summary (forecast)", COLOURS["summary"], None),
             ("Milestone (forecast)", COLOURS["milestone"], None), ("Deadline", COLOURS["deadline"], None),
             ("Status date", None, Border(left=Side(style="medium", color=COLOURS["status"]))),
             ("Non-working day", COLOURS["nonworking"], None)]
    for i, (text, colour, border) in enumerate(items):
        row, col = 1 + i // 6, first + (i % 6) * 11      # rows 1-2 only: rows 3-4 carry the month / year scale
        swatch = ws.cell(row, col)
        if colour:
            swatch.fill = fill(colour)
        if border:
            swatch.border = border
        ws.cell(row, col + 1, text).font = font(size=9)
    ws.cell(3, 1, "Timeline: conditional formatting only (no formulas), so it costs no recalculation.").font = NOTE_FONT
