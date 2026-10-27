"""Resource Report: a summary per resource, weekly hours and utilisation grids with over-allocation in red, the
leave by week, and the same hours and utilisation by sprint (the Agile sheet's sprint calendar).

Every grid cell finds its resource and week (or sprint) by ROW()/COLUMN() minus a named anchor, so each grid holds
one formula. Hours come from a hidden grid of cumulative work at each week boundary
(sprints.work_upto: SUMIFS over the time-phased helper columns on Tasks, working days only); a week's hours are
the difference of two neighbours. Working each boundary out once halves the SUMIFS, and boundaries before the
project or after it finishes need none. Weeks start on Settings > ResourceWeekStart, so the window scrolls by
changing one date. The summary reads totals from ResourceTable and the grids; it adds no per-task SUMPRODUCT.
"""
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .core import BOX, CALC_FILL, CALC_FONT, Col, Tab, add_table, fill, font, header_cell, name, title
from .leave import leave_days, leave_row, month_column, month_start, team_leave_hours
from .resources import by_resource, count_by_resource
from .agile import LEAVE_MONTHS
from .workload import work_between, work_upto

SHEET = "Resource Report"
DEFAULT_WEEKS = 26
WD = "WorkWeek,Holidays"


def build_resource_report(wb, n_resources, n_weeks=DEFAULT_WEEKS, n_sprints=26):
    ws = wb.create_sheet(SHEET)
    title(ws, None)
    ws["A1"] = '=ProjectName&"  -  Resource Report"'
    ws["A2"] = ('="Weeks from "&TEXT(ResourceWeekStart,"ddd dd-mmm-yyyy")&" (Settings > ResourceWeekStart). Hours are '
                'each task\'s work spread over its forecast working days; red = more than Allocation % x HoursPerDay x '
                'working days in the week, less leave (Resources sheet, leave months)."')
    ws.column_dimensions["A"].width = 3

    summary_top = 5
    _summary(ws, summary_top, n_resources)
    load_top = summary_top + n_resources + 4
    util_top = load_top + n_resources + 7
    leave_top = util_top + n_resources + 5
    _grid(wb, ws, "load", load_top, n_resources, n_weeks, "Weekly hours by resource", util_top - load_top)
    _grid(wb, ws, "util", util_top, n_resources, n_weeks, "Weekly utilisation by resource (hours / capacity after leave)")
    _leave_grid(wb, ws, leave_top, n_resources, n_weeks)
    sprint_top = leave_top + n_resources + 7
    sprint_util_top = sprint_top + n_resources + 5
    _sprint_grid(wb, ws, "load", sprint_top, n_sprints, n_resources,
                 "Hours by resource and sprint (Agile sheet)  -  red = over capacity after leave",
                 sprint_util_top - sprint_top)
    _sprint_grid(wb, ws, "util", sprint_util_top, n_sprints, n_resources,
                 "Utilisation by resource and sprint (% of capacity after leave)")
    sprint_leave_top = sprint_util_top + n_resources + 4
    _sprint_leave_grid(wb, ws, sprint_leave_top, n_sprints, n_resources)
    _cumulative(wb, ws, sprint_leave_top + n_resources + 3, n_resources, n_weeks)
    _team_chart(ws, load_top, n_resources, n_weeks)
    ws.freeze_panes = "D4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def _summary(ws, top, n_resources):
    t = Tab("ResourceSummary")
    k = f"ROW()-ROW({t.head('Resource')})"
    res = t.row("Resource")
    blank = f'{res}=""'
    from_table = lambda col: f'IF({blank},"",INDEX(ResourceTable[{col}],{k}))'
    week_row = lambda grid: f"INDEX({grid},{k},0)"
    ws.cell(top - 1, 2, "RESOURCES").font = font(bold=True, color="1F4E78")
    L = add_table(ws, "ResourceSummary", top, 2, [
        Col("Resource", width=22, formula=f'IFERROR(INDEX(res_Name,{k})&"","")'),
        Col("Allocation", width=10, fmt="0%", formula=from_table("Allocation %"), center=True),
        Col("Tasks", width=10, fmt="0", formula=from_table("Tasks"), center=True),
        Col("Work", width=10, fmt="#,##0", formula=from_table("Work"), note="Assigned hours."),
        Col("Work Done", width=10, fmt="#,##0", formula=f'IF({blank},"",{by_resource("Done", res)})',
            note="Work x Progress."),
        Col("Remaining", width=10, fmt="#,##0", formula=f'IF({blank},"",{t.row("Work")}-{t.row("Work Done")})'),
        Col("Late Tasks", width=10, fmt="0", center=True,
            formula=f'IF({blank},"",{count_by_resource(res, ',tsk_LateKey,">0"')})',
            note="Unfinished tasks forecast to finish after plan, or not started on time (Status > late tasks)."),
        Col("Open Critical", width=10, fmt="0", center=True,
            formula=f'IF({blank},"",{count_by_resource(res, ',tsk_Critical,"Yes",tsk_IsLeaf,1')})'),
        Col("Peak Week (h)", width=10, fmt="#,##0", center=True, formula=f'IF({blank},"",MAX({week_row("rw_Load")}))',
            note="Most hours in any week shown in the grid below."),
        Col("Peak Use", width=10, fmt="0%", center=True, formula=f'IF({blank},"",MAX({week_row("rw_Util")}))',
            note="Highest weekly utilisation in the grid."),
        Col("Weeks Over", width=10, fmt="0", center=True,
            formula=f'IF({blank},"",COUNTIF({week_row("rw_Util")},">1"))',
            note="Weeks in the grid above 100% of capacity (after leave)."),
        Col("Leave Days", width=10, fmt="0.#", center=True, formula=f'IF({blank},"",SUM({week_row("rw_Leave")}))',
            note="Working days of leave in the weeks shown below (Resources sheet, leave months)."),
    ], n_resources, autofilter=False)
    first, last = top + 1, top + n_resources
    w, u = L["Weeks Over"], L["Peak Use"]
    ws.conditional_formatting.add(f"{w}{first}:{w}{last}", FormulaRule(
        formula=[f"N({w}{first})>0"], font=Font(color="9C0006", bold=True), fill=fill("FFC7CE")))
    ws.conditional_formatting.add(f"{u}{first}:{u}{last}", FormulaRule(
        formula=[f"AND(ISNUMBER({u}{first}),{u}{first}>1)"], font=Font(color="9C0006", bold=True)))


def _grid(wb, ws, kind, top, n_resources, n_weeks, heading, util_offset=0):
    """Rows = resources, columns = weeks. The origin cell (the "Resource" header) anchors every position.
    kind: "load" (hours) or "util" (hours / capacity after leave, from the leave grid)."""
    ws.cell(top - 2, 2, heading).font = font(bold=True, size=12, color="1F4E78")
    origin = f"rw_{kind.capitalize()}Origin"
    name(wb, origin, f"'{SHEET}'!$B${top}")
    header_cell(ws.cell(top, 2), "Resource")
    header_cell(ws.cell(top, 3), "Allocation")
    i = f"(ROW()-ROW({origin})-1)"               # resource number (row 1 below the working-days row)
    j = f"(COLUMN()-COLUMN({origin})-1)"          # week number
    week = f"(ResourceWeekStart+7*({j}-1))"
    first_col, last_col = 4, 3 + n_weeks
    for c in range(first_col, last_col + 1):
        header_cell(ws.cell(top, c), None)
        ws.cell(top, c).value = f"={week}"
        ws.cell(top, c).number_format = "dd-mmm"
        days = ws.cell(top + 1, c, f"=NETWORKDAYS.INTL({week},{week}+6,{WD})")
        days.font = font(italic=True, size=9)
        days.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = 9
    ws.cell(top + 1, 2, "Working days").font = font(italic=True, size=9)
    L = get_column_letter
    if kind == "load":
        name(wb, "rw_WeekDays", f"'{SHEET}'!${L(first_col)}${top + 1}:${L(last_col)}${top + 1}")
    resource, units = f"INDEX(res_Name,{i})", f"INDEX(res_Allocation,{i})"
    first_row, last_row = top + 2, top + 1 + n_resources
    for r in range(first_row, last_row + 1):
        ws.cell(r, 2, f'=IFERROR({resource}&"","")').font = font(bold=True)
        u = ws.cell(r, 3, f'=IFERROR(IF({resource}="","",{units}),"")')
        u.number_format = "0%"
        for c in range(first_col, last_col + 1):
            if kind == "load":
                v = f'=IFERROR(IF({resource}="","",INDEX(rw_Cum,{i},{j}+1)-INDEX(rw_Cum,{i},{j})),"")'
                fmt = "#,##0;;-"
            else:
                # capacity after leave; a week with work but no capacity left shows 999%
                capacity = f"({units}*HoursPerDay*(INDEX(rw_WeekDays,{j})-INDEX(rw_Leave,{i},{j})))"
                load = f"INDEX(rw_Load,{i},{j})"
                v = (f'=IFERROR(IF(OR({resource}="",N({units})=0,INDEX(rw_WeekDays,{j})=0),"",'
                     f'IF({capacity}<=0,IF(N({load})>0,9.99,""),{load}/{capacity})),"")')
                fmt = "0%;;-"
            cell = ws.cell(r, c, v)
            cell.number_format, cell.border, cell.font = fmt, BOX, CALC_FONT
    grid = f"{L(first_col)}{first_row}:{L(last_col)}{last_row}"
    name(wb, f"rw_{kind.capitalize()}", f"'{SHEET}'!${L(first_col)}${first_row}:${L(last_col)}${last_row}")
    tl = f"{L(first_col)}{first_row}"
    if kind == "load":
        # red when the matching utilisation cell (same column, util_offset rows down) is over 100%: leave counts
        util_cell = f"{L(first_col)}{first_row + util_offset}"
        ws.conditional_formatting.add(grid, FormulaRule(
            formula=[f"AND(ISNUMBER({util_cell}),{util_cell}>1)"], fill=fill("FFC7CE"), font=Font(color="9C0006", bold=True)))
        # team totals: the chart below plots them
        totals = [("Team hours", f"SUM(INDEX(rw_Load,0,{j}))"),
                  ("Team capacity", f'SUMIFS(res_Allocation,res_Type,"Work")*HoursPerDay*INDEX(rw_WeekDays,{j})'
                                    f'-{team_leave_hours("rw_Leave", j)}')]
        for n, (label, formula) in enumerate(totals):
            r = last_row + 1 + n
            ws.cell(r, 2, label).font = font(bold=True)
            for c in range(first_col, last_col + 1):
                cell = ws.cell(r, c, f"={formula}")
                cell.number_format, cell.border, cell.font, cell.fill = "#,##0;;-", BOX, font(bold=True), CALC_FILL
    elif kind == "util":
        ws.conditional_formatting.add(grid, ColorScaleRule(
            start_type="num", start_value=0, start_color="FFFFFF", mid_type="num", mid_value=0.8, mid_color="C6EFCE",
            end_type="num", end_value=1.2, end_color="F8696B"))


def _week_months(week, j):
    """The two calendar months a week can touch, with their LeaveTable columns (helper rows of the leave grid)."""
    return [(month_start(week, k), f"INDEX(rw_LeaveCol{k},{j})") for k in (0, 1)]


def _leave_grid(wb, ws, top, n_resources, n_weeks):
    """Working days of leave per person per week (Resources sheet, leave months). Two helper rows hold the leave column of the
    week's month and the next month, and a helper column holds each person's LeaveTable row, so every cell reuses
    them instead of repeating the look-ups."""
    ws.cell(top - 2, 2, "Leave by resource (working days, from the Resources sheet)").font = font(bold=True, size=12,
                                                                                               color="1F4E78")
    origin = "rw_LeaveOrigin"
    name(wb, origin, f"'{SHEET}'!$B${top}")
    header_cell(ws.cell(top, 2), "Resource")
    header_cell(ws.cell(top, 3), "_Leave Row", "help")
    i = f"(ROW()-ROW({origin})-2)"
    j = f"(COLUMN()-COLUMN({origin})-1)"
    week = f"(ResourceWeekStart+7*({j}-1))"
    first_col, last_col = 4, 3 + n_weeks
    L = get_column_letter
    for k in (0, 1):
        r = top + 1 + k
        ws.cell(r, 2, f"_Leave column: {'week' if k == 0 else 'next'} month").font = font(italic=True, size=8,
                                                                                         color="7F7F7F")
        for c in range(first_col, last_col + 1):
            ws.cell(r, c, f"={month_column(month_start(week, k))}").font = font(size=8, color="7F7F7F")
        name(wb, f"rw_LeaveCol{k}", f"'{SHEET}'!${L(first_col)}${r}:${L(last_col)}${r}")
    for c in range(first_col, last_col + 1):
        header_cell(ws.cell(top, c), None)
        ws.cell(top, c).value = f"={week}"
        ws.cell(top, c).number_format = "dd-mmm"
    resource = f"INDEX(res_Name,{i})"
    first_row, last_row = top + 3, top + 2 + n_resources
    days = leave_days(week, f"{week}+6", f"INDEX(rw_LeaveRow,{i})", _week_months(week, j))
    for r in range(first_row, last_row + 1):
        ws.cell(r, 2, f'=IFERROR({resource}&"","")').font = font(bold=True)
        ws.cell(r, 3, f'=IFERROR(IF({resource}="",0,{leave_row(resource)}),0)').font = font(size=8, color="7F7F7F")
        for c in range(first_col, last_col + 1):
            cell = ws.cell(r, c, f'=IFERROR(IF({resource}="",0,{days}),0)')        # 0, not "": summed for the team
            cell.number_format, cell.border, cell.font = "0.#;;-", BOX, CALC_FONT
    name(wb, "rw_LeaveRow", f"'{SHEET}'!$C${first_row}:$C${last_row}")
    name(wb, "rw_Leave", f"'{SHEET}'!${L(first_col)}${first_row}:${L(last_col)}${last_row}")
    ws.conditional_formatting.add(f"{L(first_col)}{first_row}:{L(last_col)}{last_row}", FormulaRule(
        formula=[f"N({L(first_col)}{first_row})>0"], fill=fill("E4DFEC"), font=Font(color="7030A0", bold=True)))


def _team_chart(ws, load_top, n_resources, n_weeks):
    totals_row = load_top + 2 + n_resources
    chart = BarChart()
    chart.title, chart.y_axis.title = "Team hours per week vs capacity", "Hours"
    weeks = Reference(ws, min_col=4, max_col=3 + n_weeks, min_row=load_top)
    chart.series.append(Series(Reference(ws, min_col=4, max_col=3 + n_weeks, min_row=totals_row), title="Team hours"))
    chart.set_categories(weeks)
    capacity = LineChart()
    capacity.series.append(Series(Reference(ws, min_col=4, max_col=3 + n_weeks, min_row=totals_row + 1),
                                  title="Team capacity"))
    capacity.set_categories(weeks)
    capacity.series[0].smooth = False
    chart.gapWidth = 40
    chart += capacity
    chart.height, chart.width = 8, 22
    ws.add_chart(chart, "P4")


def _cumulative(wb, ws, top, n_resources, n_weeks):
    """Hidden grid: each resource's cumulative work up to the end of the day before week b (b = 0 .. n_weeks).

    Week j's hours are column j+1 minus column j. Before the project starts the total is 0; from the forecast
    finish on it is the resource's whole work (one SUMIFS); only boundaries inside the project need work_upto."""
    origin = "rw_CumOrigin"
    name(wb, origin, f"'{SHEET}'!$B${top}")
    ws.cell(top, 2, "Cumulative hours at each week boundary (helper for the weekly grids)").font = font(italic=True)
    i = f"(ROW()-ROW({origin}))"
    b = f"(COLUMN()-COLUMN({origin})-2)"
    day = f"(ResourceWeekStart+7*{b}-1)"
    resource = f"INDEX(res_Name,{i})"
    total = by_resource("Work", resource)
    upto = work_upto(f"INDEX(rw_CumWD,{b}+1)", resource)
    formula = (f'=IFERROR(IF({resource}="","",IF({day}<ProjectStart,0,IF({day}>=ProjectFinish,{total},{upto}))),"")')
    first_col, last_col = 4, 4 + n_weeks
    L = get_column_letter
    # each boundary's working-day number, once per column instead of in every cell's SUMIFS
    for c in range(first_col, last_col + 1):
        ws.cell(top, c, f"=MAX(0,NETWORKDAYS.INTL(ProjectStart,{day},{WD}))").font = CALC_FONT
    name(wb, "rw_CumWD", f"'{SHEET}'!${L(first_col)}${top}:${L(last_col)}${top}")
    for r in range(top + 1, top + 1 + n_resources):
        for c in range(first_col, last_col + 1):
            ws.cell(r, c, formula).font = CALC_FONT
    name(wb, "rw_Cum", f"'{SHEET}'!${L(first_col)}${top + 1}:${L(last_col)}${top + n_resources}")
    for r in range(top, top + 1 + n_resources):
        ws.row_dimensions[r].hidden = True


def _sprint_grid(wb, ws, kind, top, n_sprints, n_resources, heading, util_offset=0):
    """Rows = resources, columns = sprints. The origin cell (the "Resource" header) anchors every position."""
    ws.cell(top - 2, 2, heading).font = font(bold=True, size=12, color="1F4E78")
    origin = f"{kind}_Origin"
    name(wb, origin, f"'{SHEET}'!$B${top}")
    header_cell(ws.cell(top, 2), "Resource")
    header_cell(ws.cell(top, 3), "Allocation")
    i = f"(ROW()-ROW({origin})-1)"                 # resource number (row 1 below the sprint-start row)
    j = f"(COLUMN()-COLUMN({origin})-1)"            # sprint number
    sprint_name = f"INDEX(spr_Name,{j})"
    first_col, last_col = 4, 3 + n_sprints
    L = get_column_letter
    for c in range(first_col, last_col + 1):
        header_cell(ws.cell(top, c), None)
        ws.cell(top, c).value = f'=IFERROR({sprint_name}&"","")'
        d = ws.cell(top + 1, c, f'=IFERROR(IF({sprint_name}="","",INDEX(spr_Start,{j})),"")')
        d.number_format = "dd-mmm"
        d.font = font(italic=True, size=9)
        d.alignment = Alignment(horizontal="center")
        ws.column_dimensions[L(c)].width = 9                      # as the weekly grids' columns
    ws.cell(top + 1, 2, "Sprint start").font = font(italic=True, size=9)
    resource, units = f"INDEX(res_Name,{i})", f"INDEX(res_Allocation,{i})"
    first_row, last_row = top + 2, top + 1 + n_resources
    for r in range(first_row, last_row + 1):
        ws.cell(r, 2, f'=IFERROR({resource}&"","")').font = font(bold=True)
        u = ws.cell(r, 3, f'=IFERROR(IF({resource}="","",{units}),"")')
        u.number_format = "0%"
        for c in range(first_col, last_col + 1):
            if kind == "load":
                hours = work_between(f"INDEX(spr_FirstWD,{j})", f"INDEX(spr_LastWD,{j})", resource)
                v = f'=IFERROR(IF(OR({resource}="",{sprint_name}=""),"",{hours}),"")'
                fmt = "#,##0;;-"
            else:
                load = f"INDEX(load_Grid,{i},{j})"
                capacity = f"(N(INDEX(cap_Days,{i},{j}))*HoursPerDay)"      # Resources: dates, leave, allocation
                v = (f'=IFERROR(IF(OR({resource}="",{sprint_name}="",N({units})=0),"",'
                     f'IF({capacity}<=0,IF(N({load})>0,9.99,""),{load}/{capacity})),"")')
                fmt = "0%;;-"
            cell = ws.cell(r, c, v)
            cell.number_format, cell.border, cell.font = fmt, BOX, CALC_FONT
    grid = f"{L(first_col)}{first_row}:{L(last_col)}{last_row}"
    if kind == "load":
        name(wb, "load_Grid", f"'{SHEET}'!${L(first_col)}${first_row}:${L(last_col)}${last_row}")
        util_cell = f"{L(first_col)}{first_row + util_offset}"     # the matching utilisation cell: leave counts
        ws.conditional_formatting.add(grid, FormulaRule(
            formula=[f"AND(ISNUMBER({util_cell}),{util_cell}>1)"], fill=fill("FFC7CE"), font=Font(color="9C0006", bold=True)))
    else:
        ws.conditional_formatting.add(grid, ColorScaleRule(
            start_type="num", start_value=0, start_color="FFFFFF", mid_type="num", mid_value=0.8, mid_color="C6EFCE",
            end_type="num", end_value=1.2, end_color="F8696B"))


def _sprint_leave_grid(wb, ws, top, n_sprints, n_resources):
    """Hidden grid: working days of leave per person per sprint (Resources sheet, leave months), read by the sprint
    utilisation grid and the Agile sheet's capacity. Each person's Resources row is looked up once, in column C;
    each sprint's month columns come from SprintTable."""
    origin = "spr_LeaveOrigin"
    name(wb, origin, f"'{SHEET}'!$B${top}")
    ws.cell(top, 2, "Leave days by sprint (helper for the sprint utilisation grid)").font = font(italic=True)
    i = f"(ROW()-ROW({origin}))"
    j = f"(COLUMN()-COLUMN({origin})-1)"
    resource = f"INDEX(res_Name,{i})"
    start, end = f"INDEX(spr_Start,{j})", f"INDEX(spr_End,{j})"
    months = [(month_start(start, k), f"INDEX(spr_LeaveCol{k},{j})") for k in range(LEAVE_MONTHS)]
    days = leave_days(start, end, f"INDEX(spr_LeaveRow,{i})", months)
    L = get_column_letter
    last_col = L(3 + n_sprints)
    for r in range(top + 1, top + 1 + n_resources):
        ws.cell(r, 3, f'=IFERROR(IF({resource}="",0,{leave_row(resource)}),0)').font = CALC_FONT
        for c in range(4, 4 + n_sprints):
            ws.cell(r, c, f'=IFERROR(IF(OR({resource}="",INDEX(spr_Name,{j})=""),0,{days}),0)').font = CALC_FONT
    name(wb, "spr_LeaveRow", f"'{SHEET}'!$C${top + 1}:$C${top + n_resources}")
    name(wb, "spr_LeaveGrid", f"'{SHEET}'!$D${top + 1}:${last_col}${top + n_resources}")
    for r in range(top, top + 1 + n_resources):
        ws.row_dimensions[r].hidden = True
