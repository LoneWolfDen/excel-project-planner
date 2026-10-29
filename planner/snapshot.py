"""TakeSnapshot sheet (Release 1.0, optional module "snapshot"; not in Simple): the two copy-and-paste jobs a
formula can't do, laid out as steps with links. A formula can't remember last week's value or freeze a plan, so
both are a copy and a Paste Special > Values, as in the classic workbook; this sheet puts the row to copy and the
places to paste one click away.

  * Snapshot: this status date's numbers (the same formulas as Status > SnapshotNow, read by the snap_ names) to
    paste under the last row of History.
  * Baseline: the links select Tasks' four BL Source columns and the four Baseline columns (defined names over the
    Tasks table, so they follow inserted rows and columns). Settings > TakeBaseline says whether BL Source
    re-baselines every task or only the ones without a baseline.
"""
from openpyxl.styles import Alignment
from openpyxl.worksheet.hyperlink import Hyperlink

from .core import NOTE_FONT, TITLE_FONT, Col, add_table, font, has, header_cell, name, v12
from .history import fields

SHEET = "TakeSnapshot"


def _step(ws, r, label, text, link=None):
    ws.cell(r, 2, label).font = font(bold=True)
    c = ws.cell(r, 3, text)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if link:                                          # a hyperlink to a sheet or a defined name, not a formula
        c.hyperlink = Hyperlink(ref=c.coordinate, location=link)
        c.font = font(color="1F4E78", underline="single")
    ws.row_dimensions[r].height = 30


def build_snapshot(wb):
    ws = wb.create_sheet(SHEET)
    ws["B1"] = '=ProjectName&"  -  Snapshot and baseline"'
    ws["B1"].font = TITLE_FONT
    ws["B2"] = '="Status date "&TEXT(StatusDate,"ddd dd-mmm-yyyy")' + ('&".  "&chk_Verdict' if has("check") else "")
    ws["B2"].font = font(bold=True, color="404040")
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 100
    header_cell(ws.cell(4, 2), "SNAPSHOT", "in")
    header_cell(ws.cell(4, 3), "Each reporting date, after the status date and progress are in", "in")
    _step(ws, 5, "1. Check", "Fix what the Check sheet lists (it says All clear when there is nothing).",
          "'Check'!A1" if has("check") else None)
    _step(ws, 6, "2. Copy", "Select the row of numbers in the SNAPSHOT ROW table below (click this link) and press "
                            "Ctrl+C.", "snap_Row")
    _step(ws, 7, "3. Paste", "Go to History (this link), click the first empty row of its table and use Paste "
                             "Special > Values (Ctrl+Alt+V, V). The trend charts update." +
          ("  Cost: paste its own snapshot row on the Cost sheet onto its CostHistory the same way." if has("cost")
           else ""), "'History'!A1")
    header_cell(ws.cell(9, 2), "BASELINE", "in")
    header_cell(ws.cell(9, 3), "When the plan is agreed (and after an approved change, to re-baseline)", "in")
    _step(ws, 10, "1. Which tasks", "Settings > TakeBaseline: No = only tasks without a baseline get one (the rest "
                                    "keep theirs); Yes = every task is re-baselined.", "TakeBaseline")
    _step(ws, 11, "2. Copy", "Click this link: it selects Tasks' four BL Source columns (open the BASELINE "
                             "section with + to see them). Press Ctrl+C.", "tsk_BLSource")
    _step(ws, 12, "3. Paste", "Click this link: it selects the four Baseline columns. Paste Special > Values "
                              "(Ctrl+Alt+V, V). The Gantt outlines and the variance columns follow.", "tsk_Baseline")
    _step(ws, 13, "4. Reset", "Set Settings > TakeBaseline back to No.", "TakeBaseline")
    top = 16
    if v12():
        # B-201: before re-baselining, keep the agreed dates as baseline 2 (or 3) to compare against later
        header_cell(ws.cell(15, 2), "KEEP A BASELINE", "in")
        header_cell(ws.cell(15, 3), "Before re-baselining: keep the current baseline's dates as Baseline 2 or 3", "in")
        _step(ws, 16, "1. Copy", "Click this link: it selects Tasks' Baseline Start and Baseline Finish. Ctrl+C.",
              "tsk_BaseDates")
        _step(ws, 17, "2. Paste", "Click this link (Baseline 2) and Paste Special > Values. Use Baseline 3 for a third "
                                  "set.", "tsk_Baseline2")
        _step(ws, 18, "3. Compare", "Settings > CompareBaseline: 1, 2 or 3 picks the baseline Planned Start / Finish, "
                                    "the variances and the Gantt compare against.", "CompareBaseline")
        top = 21
    ws.cell(top - 1, 2, "SNAPSHOT ROW: this status date's numbers (formulas: copy them, paste values on History)"
            ).font = font(bold=True, color="1F4E78")
    letters = add_table(ws, "SnapshotRow", top, 2, [Col(h, width=13, fmt=fmt, formula=f, center=True)
                                                     for _, (h, f, fmt) in fields()], 1, autofilter=False)
    first, last = letters[fields()[0][1][0]], letters[fields()[-1][1][0]]
    name(wb, "snap_Row", f"{SHEET}!${first}${top + 1}:${last}${top + 1}")
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 100
    ws.cell(top + 3, 2, "Nothing here needs a macro; a macro would only automate the same copy and paste.").font = \
        NOTE_FONT
    return ws
