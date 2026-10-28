"""Move a plan into the latest template (B-901): read what the user typed into an older planner workbook and build
the current workbook around it.

Only typed values are read (formula cells are skipped), and always by header name, never by position, so a column
that moved or a sheet that was reordered still lines up. The plan's core (tasks, people, assignments, holidays,
leave, settings) goes through planner.build.ProjectData, as an MS Project import does, so the new workbook's rows,
UIDs and sizes are worked out by the builder. Every other input table (the logs, status history, sprint goals,
PM RAG and notes, resource notes) is then copied onto the new workbook column by column, matched by table and
header. Dropdown lists keep the new template's values and add any the user had added.

Inputs that a version renamed or moved are mapped in RENAMED, RENAMED_TABLES and LOGS and in the steps below;
workbooks carry their version in the PlannerVersion name (absent before v1.7). The same code runs on the command line
(scripts/upgrade_planner.py) and in the browser page docs/export.html.
"""
import io
import re

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from .core import BOX, HEADER_FILLS, INPUT_FONT

# Tasks headers an older version used, and the column they are now
RENAMED = {"Tasks": {"Start No Earlier Than": "Constraint Date"}}
# tables a version renamed: old name -> (new name, {old header: new header})
RENAMED_TABLES = {"ListRAIDType": ("ListLogType", {"RAID Type": "Log Type"})}
# before v1.7 the Log was three tables: old table -> (Type of its rows, or None to keep theirs, {old header: Log's})
LOGS = {"RAIDLog": (None, {}),
        "DecisionLog": ("Decision", {"Date": "Raised", "Decision": "Title", "RAID Ref": "Related Ref"}),
        "ChangeLog": ("Change", {"Requested By": "Owner", "Reason": "Options and Reasons"})}
# tables the builder fills from ProjectData; everything else with inputs is copied by header
BUILT = {"Tasks", "HolidayTable", "LeaveTable", "AssignmentTable"} | set(LOGS)
INPUT_FILL = HEADER_FILLS["in"].start_color.rgb[-6:]
MONTH = re.compile(r"^[A-Z][a-z]{2}-\d\d$")          # a leave month header: Nov-26


class UpgradeError(Exception):
    """A file that isn't a planner workbook (shown as-is by the command line and the browser page)."""


def _tables(wb):
    """{table name: (sheet, [header], [[typed value or None per column] per row])}, formulas left out."""
    out = {}
    for ws in wb.worksheets:
        for name, table in ws.tables.items():
            c1, r1, c2, r2 = range_boundaries(table.ref if hasattr(table, "ref") else table)
            head = [ws.cell(r1, c).value for c in range(c1, c2 + 1)]
            rows = []
            for r in range(r1 + 1, r2 + 1):
                row = []
                for c in range(c1, c2 + 1):
                    v = ws.cell(r, c).value
                    row.append(None if isinstance(v, str) and v.startswith("=") else v)
                rows.append(row)
            out[name] = (ws, head, rows)
    return out


def _records(tables, name, renamed=None):
    """A table's rows as {header: typed value}, trailing empty rows dropped."""
    if name not in tables:
        return []
    _, head, rows = tables[name]
    head = [(renamed or {}).get(h, h) for h in head]
    records = [{h: v for h, v in zip(head, row) if h is not None} for row in rows]
    while records and all(v in (None, "") for v in records[-1].values()):
        records.pop()
    return records


def _named_value(wb, name):
    if name not in wb.defined_names:
        return None
    ref = wb.defined_names[name].attr_text
    if "!" not in ref:                                   # a constant name, e.g. PlannerVersion = "1.7"
        return ref.strip('"') if ref.startswith('"') else None
    sheet, cell = ref.split("!")
    sheet = sheet.strip("'")
    if sheet not in wb.sheetnames:
        return None
    v = wb[sheet][cell.replace("$", "")].value
    return None if isinstance(v, str) and v.startswith("=") else v


def read_planner(data):
    """(ProjectData, copies, Sizes, warnings, version) from an older planner workbook's bytes. copies: the other
    input tables, {table: (headers, rows)}, for copy_inputs."""
    from .build import ProjectData, Sizes
    from .settings import settings_rows, week_preset
    from .settings import RELEASE_ROWS, V11_ROWS
    from .tasks import N, input_keys, take_links
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data))
    except Exception as e:                                   # not a zip / not a workbook
        raise UpgradeError(f"This isn't an Excel workbook ({e.__class__.__name__}).")
    tables = _tables(wb)
    if "Tasks" not in tables:
        raise UpgradeError("This workbook has no Tasks table: is it a planner workbook?")
    warnings = []
    version = _named_value(wb, "PlannerVersion") or "1.6 or earlier"

    settings = {}
    for n, _label, _default, _fmt, is_input, _note in settings_rows(True, 26) + RELEASE_ROWS + V11_ROWS:
        v = _named_value(wb, n)
        if is_input and v is not None:
            settings[n] = v

    if "WorkingWeek" not in settings:                # before v1.7 the week was only the seven Yes / No days
        settings["WorkingWeek"] = week_preset(settings)

    key_of = {N[k]: k for k in input_keys()}
    # before v1.7 each link was three inputs (Pred1 UID / Type / Lag ...): they become the Predecessors text
    key_of.update({f"Pred{k} {label}": f"{x}{k}" for k in range(1, 10) for x, label in (("P", "UID"), ("T", "Type"),
                                                                                       ("L", "Lag"))})
    key_of.update({"Resource": "Res", "Units": "Units"})      # before v1.7: one person + the Assignments sheet
    tasks, costed = [], 0
    for rec in _records(tables, "Tasks", RENAMED["Tasks"]):
        row = {key_of[h]: v for h, v in rec.items() if h in key_of and v not in (None, "")}
        if isinstance(rec.get("Actual Cost"), (int, float)):       # typed on Tasks before v1.8: now calculated
            row["Notes"] = f'{row.get("Notes") or ""} [Actual cost typed before v1.8: {rec["Actual Cost"]:g}]'.strip()
            costed += 1
        tasks.append(take_links(row))
    if costed:
        warnings.append(f"{costed} tasks had a typed Actual Cost: cost is now worked out on the Cost sheet (work done "
                        "x day rates), so each value is kept in the task's Notes.")
    if any(isinstance(r.get("Baseline Cost"), (int, float)) for r in _records(tables, "Tasks")):
        warnings.append("Baseline Cost is no longer typed: the Cost sheet's budget is the baseline's work at today's "
                        "day rates.")
    while tasks and not tasks[-1].get("Name"):
        tasks.pop()

    # before v1.8: Max Units and an hourly Std Rate on Resources; since: Allocation %, and day rates on the Cost sheet
    hours = settings.get("HoursPerDay") or 8
    day_rates = {r.get("Resource"): r.get("Day Rate") for r in _records(tables, "RateCard") if r.get("Resource")}
    resources = [(r.get("Resource Name"), r.get("Initials"), r.get("Group"), r.get("Allocation %", r.get("Max Units")),
                  day_rates.get(r["Resource Name"], None if r.get("Std Rate") is None else r["Std Rate"] * hours),
                  r.get("Email"), r.get("Type") or "Work", r.get("Start"), r.get("End"), r.get("Focus %"))
                 for r in _records(tables, "ResourceTable") if r.get("Resource Name")]
    holidays = [(r["Date"], r.get("Description")) for r in _records(tables, "HolidayTable") if r.get("Date")]
    leave, months = {}, None
    # leave: a Leave sheet before v1.7, month columns on Resources (Nov-26 ...) since
    source, key = ("LeaveTable", "Resource") if "LeaveTable" in tables else ("ResourceTable", "Resource Name")
    if source in tables:
        months = [h for h in tables[source][1] if isinstance(h, str) and MONTH.match(h)] or None
        for r in _records(tables, source):
            entries = {m: str(r[m]) for m in months or [] if r.get(m) not in (None, "")}
            if r.get(key) and entries:
                leave[r[key]] = entries
    people = {r[0] for r in resources}
    for person in sorted(set(leave) - people):
        warnings.append(f"Leave for {person} was not carried: {person} is not on the Resources sheet.")
        del leave[person]
    assignments = [(r["Task UID"], r.get("Resource"), r.get("Units"))
                   for r in _records(tables, "AssignmentTable") if r.get("Task UID")]

    copies = {name: (head, rows) for name, (_, head, rows) in tables.items() if name not in BUILT}
    for old, (new, headers) in RENAMED_TABLES.items():
        if old in copies:
            head, rows = copies.pop(old)
            copies[new] = ([headers.get(h, h) for h in head], rows)
    if any(t in tables for t in LOGS):
        copies["Log"] = _merge_logs(tables, warnings)
    if "StatusHistory" in copies and "BAC" in copies["StatusHistory"][0]:   # before v1.8: cost on History too
        head, rows = copies["StatusHistory"]
        cost = ["Status Date", "BAC", "PV", "EV", "AC", "CPI", "EAC"]
        copies["CostHistory"] = (cost, [[row[head.index(h)] if h in head else None for h in cost] for row in rows])
        copies["StatusHistory"] = (["% Done" if h == "% Earned" else h for h in head], rows)
    if "StatusRAG" in copies and "Area" in copies["StatusRAG"][0]:     # rows are copied by position: drop the gone row
        head, rows = copies["StatusRAG"]
        area = head.index("Area")
        for row in rows:
            if row[area] == "Cost performance" and any(row[head.index(h)] not in (None, "")
                                                       for h in ("PM RAG", "Comment") if h in head):
                warnings.append("The Cost performance RAG is now on the Cost sheet (from CPI), so its PM RAG and "
                                "comment were not carried.")
        copies["StatusRAG"] = (head, [row for row in rows if row[area] != "Cost performance"])
    if any(isinstance(r.get("Points Capacity"), (int, float)) for r in _records(tables, "SprintTable")):
        warnings.append("Points Capacity per sprint is now calculated from people (Resources: allocation, leave, "
                        "Focus %), so the typed values were not carried.")
    default = Sizes()
    sizes = Sizes(tasks=max(default.tasks, len(tasks)), resources=max(default.resources, len(resources)),
                  sprints=max(default.sprints, len(tables.get("SprintTable", (None, None, []))[2])))
    project = ProjectData(settings=settings, tasks=tasks, resources=resources, holidays=holidays, leave=leave,
                          leave_months=months, assignments=assignments)
    from .build import PLANNER_VERSIONS
    by_planner = {v: t for t, v in PLANNER_VERSIONS.items()}
    project.release = by_planner.get(str(version), False)    # a release-line workbook stays in its edition
    modules = _named_value(wb, "PlannerModules")          # since v1.8: a Simple workbook upgrades to a Simple one
    if modules is not None:
        project.modules = [m for m in str(modules).split(",") if m]
    return project, copies, sizes, warnings, version


def _merge_logs(tables, warnings):
    """The RAID, decision and change logs of a v1.6 or older workbook as one Log table (headers, rows). RAID items
    keep their Ref; decisions and change requests get the next free ones (the Log has one Ref sequence), and a
    warning lists them so references typed in notes can be updated. The rows after them carry on the numbering."""
    from .log import ROWS
    items, warned = [], {}
    for table, (kind, headers) in LOGS.items():
        for rec in _records(tables, table, headers):
            if rec.get("Title") in (None, ""):
                continue
            if kind:
                rec["Type"] = kind
            items.append((table, rec))
    refs = [rec["Ref"] for table, rec in items if table == "RAIDLog" and isinstance(rec.get("Ref"), (int, float))]
    ref = int(max(refs, default=0))
    for table, rec in items:
        if table != "RAIDLog" or not isinstance(rec.get("Ref"), (int, float)):
            ref += 1
            if rec.get("Ref") not in (None, ""):
                warned.setdefault(rec["Type"], []).append(f"{rec['Type']} {rec['Ref']} is now {rec['Type']} {ref}")
            rec["Ref"] = ref
    for kind, moves in warned.items():
        warnings.append(f"{kind} refs now follow the Log's numbering ({', '.join(moves)}): update any notes that "
                        "quote the old numbers.")
    head = ["Ref", "Type"] + sorted({h for _, rec in items for h in rec} - {"Ref", "Type"})
    rows = [[rec.get(h) for h in head] for _, rec in items]
    rows += [[ref + k] + [None] * (len(head) - 1) for k in range(1, ROWS - len(rows) + 1)]   # the free Refs
    return head, rows


def copy_inputs(wb, copies, warnings):
    """Write the old tables' typed values into the new workbook's input columns with the same table and header.
    Dropdown lists (List... tables) keep the template's values and add the user's extra ones."""
    new = _tables(wb)
    for name, (old_head, old_rows) in copies.items():
        if name not in new:
            if any(v not in (None, "") for row in old_rows for v in row):
                warnings.append(f"The {name} table is not in the new template; its contents were not carried.")
            continue
        ws, head, _ = new[name]
        table = ws.tables[name]
        c1, r1, c2, r2 = range_boundaries(table.ref)
        for i, h in enumerate(head):
            fill = (ws.cell(r1, c1 + i).fill.start_color.rgb or "")[-6:]
            if h not in old_head or fill != INPUT_FILL:
                continue
            values = [row[old_head.index(h)] for row in old_rows]
            if name.startswith("List"):
                current = [ws.cell(r, c1 + i).value for r in range(r1 + 1, r2 + 1)]
                values = [v for v in current if v is not None] + [v for v in values if v not in current and v is not None]
                if len(values) > r2 - r1:                # grow the list's table (lists sit side by side)
                    r2 = r1 + len(values)
                    table.ref = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"
                    for r in range(r1 + 1, r2 + 1):
                        ws.cell(r, c1 + i).font, ws.cell(r, c1 + i).border = INPUT_FONT, BOX
            room = r2 - r1
            if any(v not in (None, "") for v in values[room:]):
                warnings.append(f"{name}: only the first {room} rows of '{h}' fit; add rows to the table and copy "
                                "the rest from the old workbook.")
            for k, v in enumerate(values[:room]):
                ws.cell(r1 + 1 + k, c1 + i).value = v


def rebuild(data, release=None):
    """(the latest template's workbook holding the old workbook's inputs, warnings, the old version). release: True
    moves a classic workbook into the Release 1.0 one, "1.1" into template 1.1, False back to classic; None keeps the
    workbook's own edition."""
    from .build import make_workbook
    from .core import MODULES
    project, copies, sizes, warnings, version = read_planner(data)
    if release is not None:
        project.release = release
    if project.modules is not None:
        if project.release and set(MODULES) <= set(project.modules):
            project.modules = None                       # a full workbook: Release 1.0 adds Check and TakeSnapshot
        elif not project.release:
            project.modules = [m for m in project.modules if m in MODULES]
    wb = make_workbook(sizes=sizes, data=project, release_edition=project.release)
    copy_inputs(wb, copies, warnings)
    return wb, warnings, version


def convert(data, release=None):
    """(xlsx bytes of the latest template holding the old workbook's inputs, warnings)."""
    wb, warnings, _version = rebuild(data, release)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), warnings
