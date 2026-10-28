"""Build the planner workbook.

    python scripts/build_planner.py Project_Planner.xlsx --recalc
    python scripts/build_planner.py Project_Planner_Blank.xlsx --blank --recalc
    python scripts/build_planner.py stress.xlsx --tasks 1000 --stress 33      # performance testing
    python scripts/build_planner.py big.xlsx --tasks 2000 --resources 100 --grid-resources 60 --sprints 52 \
        --gantt-columns 365 --weeks 52                                         # a bigger template
"""
import argparse
from dataclasses import dataclass, field
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

import openpyxl

from . import sample
from .agile import build_agile
from .gantt import build_timeline
from .instructions import build_instructions
from .msp import build_msp_resources, build_msp_tasks
from .core import MODULES, all_modules, has, name, protect_view, release, set_modules, set_release
from .cost import build_cost
from .history import build_history
from .log import build_log
from .report_resources import build_resource_report
from .roadmap import build_roadmap
from .resources import RESOURCES, build_resources
from .settings import PRESETS, build_holidays, build_lists, build_settings, build_stages, stages_of
from .status import build_status
from .tasks import build_tasks, configure, take_people, take_stage

VIEWS = ("Status", "Roadmap", "Resource Report")    # protected without a password; inputs stay unlocked
# Release 1.0 protects every sheet whose rows are fixed (formula cells locked, blue inputs open); Tasks, Resources,
# Settings and History stay open because rows are added there (Excel can't add table rows on a protected sheet)
RELEASE_VIEWS = ("Start", "Status", "Roadmap", "Resource Report", "Agile", "Log", "Cost", "Check", "TakeSnapshot")
ORDER = ["Start", "Status", "Check", "Roadmap", "Resource Report", "Agile", "Log", "History", "TakeSnapshot",
         "Settings", "Tasks", "Resources", "Cost", "MSP_Tasks", "MSP_Resources"]
TAB_COLOURS = {"Start": "1F4E78", "Status": "548235", "Tasks": "2F75B5", "Roadmap": "548235",
               "Resources": "7030A0", "Log": "C55A11", "Agile": "548235",
               "Resource Report": "548235", "History": "C55A11", "Settings": "BF8F00", "Cost": "806000",
               "MSP_Tasks": "595959", "MSP_Resources": "595959", "Check": "C00000", "TakeSnapshot": "C55A11"}
@dataclass(frozen=True)
class Sizes:
    """How many rows / columns each table and grid starts with. Tables grow when users add rows; the fixed grids
    (Gantt timeline, per-resource grids, weeks) are sized here, so a bigger plan is a rebuild, not a formula edit."""
    tasks: int = 200
    resources: int = 50          # ResourceTable and MSP_Resources rows
    grid_resources: int = 25     # rows in the per-resource grids (Resource Report, by week and by sprint)
    sprints: int = 26            # SprintTable rows (SprintCount can be at most this)
    gantt_columns: int = 150     # Gantt timeline columns (days or weeks, Settings > GanttPeriod)
    weeks: int = 26              # Resource Report weekly grids

    def check(self):
        limits = {"tasks": (1, 20000), "resources": (1, 2000), "grid_resources": (1, 500), "sprints": (1, 520),
                  "gantt_columns": (7, 1000), "weeks": (1, 260)}
        for field, (low, high) in limits.items():
            value = getattr(self, field)
            if not low <= value <= high:
                raise SystemExit(f"--{field.replace('_', '-')} must be between {low} and {high} (got {value})")
        if self.grid_resources > self.resources:
            raise SystemExit("--grid-resources can't be more than --resources")


@dataclass
class ProjectData:
    """A project to load into a new workbook instead of the sample (planner.mspdi_import builds one from MS Project).
    Tasks are dicts keyed by the Tasks input keys (UID, Level, Name, Dur, P1, T1, L1, ..., BS, BF, BD, BW, BC)."""
    settings: dict = field(default_factory=dict)       # {setting name: value}, e.g. {"ProjectStartDate": date}
    tasks: list = field(default_factory=list)
    resources: list = field(default_factory=list)      # (name, initials, group, allocation, day rate, email, type)
    holidays: list = field(default_factory=list)       # (date, description)
    leave: dict = field(default_factory=dict)          # {person: {"Nov-26": "16, 17(0.5)"}}
    leave_months: list = None                          # leave month headers (Resources sheet columns)
    assignments: list = field(default_factory=list)    # (task UID, resource, units): people beyond the first
    stages: list = None                                # (stage, %) rows of the Stages table; None = the default
    modules: list = None                               # optional sheets (core.MODULES); None = all (B-913)
    release: bool = False                              # a Release 1.0 workbook (upgrade / reorder keep the edition)


VERSION = "1.8"          # PlannerVersion in every workbook: planner/upgrade.py reads it to know what to convert
RELEASE_VERSION = "2.0"  # PlannerVersion of the Release 1.0 workbooks (the planner code that builds them)
TEMPLATE_VERSION = "1.0"  # TemplateVersion: the release users see (Start, every printed footer)


def make_workbook(blank=False, n_tasks=200, stress=0, sizes=None, data=None, import_wizard=False, baselined=False,
                  project_type=None, modules=None, release_edition=False):
    """import_wizard: also build MSP_Tasks / MSP_Resources, formula copies of Tasks and Resources under MS Project's
    field names for its Import Wizard. Off by default: the XML converter (docs/export.html) carries more."""
    """The planner workbook: the sample project, a blank template (blank=True), or a given project (data).
    project_type: one of settings.PRESETS (B-1018) for the stages, sprint length and RAG thresholds.
    modules: the optional sheets to include (core.MODULES; None = all, () = the Simple workbook, B-913).
    release_edition: the Release 1.0 workbook (pickers, the leave day grid, Check, TakeSnapshot ...); the classic
    workbook is unchanged by it."""
    set_release(release_edition)
    set_modules(modules if modules is not None else data.modules if data is not None else None)
    configure()
    sizes = sizes or Sizes(tasks=n_tasks)
    sizes.check()
    n_tasks = sizes.tasks
    if data is not None:
        blank = True                                   # no sample RAID, decisions, history ...
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    simple = release_edition and data is None and not blank and not stress and modules is not None and not modules
    rows = data.tasks if data is not None else sample.simple_tasks(baselined) if simple else \
        sample.tasks(blank, stress, baselined)
    if len(rows) > n_tasks:
        raise SystemExit(f"{len(rows)} tasks do not fit in {n_tasks} rows; raise --tasks")
    values = dict(data.settings) if data and data.settings else {}
    if project_type:
        values["ProjectType"] = project_type
    build_settings(wb, blank, sizes.sprints, values or None)
    build_lists(wb)
    build_stages(wb, data.stages if data else None)
    build_holidays(wb, blank, holidays=data.holidays if data else None)
    people = data.resources if data else [] if blank else list(RESOURCES)
    build_resources(wb, blank, sizes.resources, data.resources if data else None, data.leave if data else None,
                    data.leave_months if data else None, sizes.sprints, sizes.grid_resources)
    more = {}                                          # people beyond the first: into each task's Resources text
    for uid, resource, units in (data.assignments if data is not None else sample.assignments(blank, stress)):
        more.setdefault(uid, []).append((resource, units))
    # a typed % becomes the nearest Stage of the plan's own stages, or of its Project type's (B-1018)
    stages = (data.stages if data and data.stages else
              stages_of(values.get("ProjectType")))
    rows = [take_stage(take_people(dict(r), more.get(r.get("UID"), [])), stages) for r in rows]
    if release_edition and data is None:              # the sample shows the pickers (imports keep the text cells)
        rows = [sample.to_pickers(r) for r in rows]
    tasks_ws, letters = build_tasks(wb, rows, n_tasks)
    build_timeline(wb, tasks_ws, letters, n_tasks, sizes.gantt_columns)
    charted = min(12, sizes.sprints)                 # sprints shown on the Status and Agile charts
    agile_ws, sprint_letters = build_agile(wb, blank, sizes.sprints, charted) if has("agile") else (None, None)
    if import_wizard:
        build_msp_tasks(wb, n_tasks)
        build_msp_resources(wb, sizes.resources)
    if has("log"):
        build_log(wb, blank)
    if has("history"):
        build_history(wb, blank)
    build_status(wb, agile_ws, sprint_letters, charted)
    build_roadmap(wb)
    if has("resources"):
        build_resource_report(wb, sizes.grid_resources, sizes.weeks, sizes.sprints)
    if has("cost"):
        build_cost(wb, blank, n_tasks, people, sizes.resources)
    if has("check"):
        from .check import build_check
        build_check(wb)
    if has("snapshot") and has("history"):            # the snapshot is pasted onto History
        from .snapshot import build_snapshot
        build_snapshot(wb)
    build_instructions(wb, sizes, import_wizard)
    if release_edition:
        from .release import finish_release
        finish_release(wb, TEMPLATE_VERSION)
    for view in RELEASE_VIEWS if release_edition else VIEWS:     # B-1019: formulas can't be typed over by accident
        if view in wb.sheetnames:
            protect_view(wb[view])
    name(wb, "PlannerVersion", f'"{RELEASE_VERSION if release_edition else VERSION}"')
    name(wb, "PlannerModules", '"' + ",".join(m for m in all_modules() if has(m)) + '"')   # an upgrade keeps the same set
    wb._sheets = [wb[n] for n in ORDER if n in wb.sheetnames]
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = TAB_COLOURS[ws.title]
    wb.active = wb.sheetnames.index("Tasks")
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.calcMode = "auto"                   # B-1001: views follow every edit
    return wb


def _tables(z):
    """{displayName: (path, xml)} for every table part in an .xlsx zip."""
    out = {}
    for n in z.namelist():
        if n.startswith("xl/tables/") and n.endswith(".xml"):
            xml = z.read(n).decode("utf-8")
            out[re.search(r'displayName="([^"]+)"', xml).group(1)] = (n, xml)
    return out


CALC_PR = '<calcPr calcId="191029" calcMode="auto" fullCalcOnLoad="1"/>'


def restore_tables(original, recalculated):
    """LibreOffice drops each table's calculated-column formulas when it saves. Excel needs them to fill formulas
    into newly inserted rows, so put the original table definitions back (keeping LibreOffice's part ids).
    It also drops fullCalcOnLoad, so put automatic calculation and a recalculation on open back (B-1001): without
    them a file opened while Excel is in manual mode shows stale views until F9."""
    with zipfile.ZipFile(original) as z:
        source = _tables(z)
    tmp = Path(str(recalculated) + ".tmp")
    with zipfile.ZipFile(recalculated) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        target = {path: name for name, (path, _) in _tables(zin).items()}
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in target and target[item.filename] in source:
                lo_id = re.search(r'<table[^>]*\sid="(\d+)"', data.decode("utf-8")).group(1)
                xml = source[target[item.filename]][1]
                data = re.sub(r'(<table[^>]*\s)id="\d+"', rf'\g<1>id="{lo_id}"', xml, count=1).encode("utf-8")
            elif item.filename == "xl/workbook.xml":
                data = re.sub(r"<calcPr[^>]*/>", CALC_PR, data.decode("utf-8"), count=1).encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(recalculated)


def save(wb, out, recalc):
    if not recalc:
        wb.save(out)
        return
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise SystemExit("--recalc needs LibreOffice (soffice) on PATH")
    with tempfile.TemporaryDirectory() as tmp:
        raw = Path(tmp) / Path(out).name
        wb.save(raw)
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(Path(tmp) / "out"),
                        str(raw)], check=True, capture_output=True, timeout=1800)
        done = Path(tmp) / "out" / raw.name
        restore_tables(raw, done)
        shutil.copyfile(done, out)


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("out", help="output .xlsx path")
    p.add_argument("--blank", action="store_true", help="empty template: no sample tasks, resources, holidays")
    p.add_argument("--recalc", action="store_true",
                   help="recalculate with LibreOffice so the saved file carries values (for previews)")
    d = Sizes()
    p.add_argument("--tasks", type=int, default=d.tasks, help="rows in the Tasks table (and MSP_Tasks)")
    p.add_argument("--resources", type=int, default=d.resources, help="rows in the Resources table")
    p.add_argument("--grid-resources", type=int, default=d.grid_resources,
                   help="resources shown in the per-resource grids (Resource Report)")
    p.add_argument("--sprints", type=int, default=d.sprints, help="rows in the sprint table (Agile sheet)")
    p.add_argument("--gantt-columns", type=int, default=d.gantt_columns, help="columns in the Gantt timeline (right of the Tasks table)")
    p.add_argument("--weeks", type=int, default=d.weeks, help="weeks in the Resource Report grids")
    p.add_argument("--stress", type=int, default=0, help="repeat the sample project N times (performance tests)")
    p.add_argument("--baselined", action="store_true",
                   help="the sample with its baseline set (the committed Project_Planner.xlsx), so the Gantt shows "
                        "planned against current")
    p.add_argument("--project-type", choices=list(PRESETS),
                   help="Start > Project type for a --blank template: its stages, sprint length and RAG thresholds")
    p.add_argument("--import-wizard-sheets", action="store_true",
                   help="also build MSP_Tasks / MSP_Resources for MS Project's Import Wizard (the XML converter "
                        "is the better route)")
    p.add_argument("--modules", help="optional sheets to include, comma-separated, from " + ", ".join(MODULES)
                   + " (default: all; '' = none)")
    p.add_argument("--simple", action="store_true",
                   help="the Simple workbook: Start, Status, Roadmap, Tasks (with its Gantt), Resources, Settings")
    p.add_argument("--release", action="store_true",
                   help="the Release 1.0 workbook (release-1.0/): predecessor and resource pickers, a leave day grid, "
                        "Duration as 3 / 3d / 2w, the Check and TakeSnapshot sheets, more protection, print set-up")
    a = p.parse_args(argv)
    modules = () if a.simple else None if a.modules is None else [m.strip() for m in a.modules.split(",") if m.strip()]
    sizes = Sizes(a.tasks, a.resources, a.grid_resources, a.sprints, a.gantt_columns, a.weeks)
    save(make_workbook(a.blank, stress=a.stress, sizes=sizes, import_wizard=a.import_wizard_sheets,
                       baselined=a.baselined, project_type=a.project_type, modules=modules,
                       release_edition=a.release), a.out, a.recalc)
    print("saved", a.out)


if __name__ == "__main__":
    main()
