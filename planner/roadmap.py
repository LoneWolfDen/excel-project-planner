"""Roadmap sheet: one landscape page of phases and milestones by month, for sponsors and steering committees.

Two report lists (core.ranked_table): the phases (outline level 1) and the milestones in date order, each with its
dates. A phase row (B-1009) also carries its owner (Tasks > Owner), its progress (duration-weighted, as MS Project
rolls it up) with the count beside it ("2 of 3 done"), a RAG from its finish slip against the planned finish (the
Status thresholds), its next milestone and a link to its row on Tasks. To their right is a month timeline drawn purely with conditional formatting, as on
the Gantt: the month header is the only formula in it, so the timeline adds nothing to recalculation. A phase bar
spans its forecast months (green for the done share, amber when it finishes late against plan) inside an outline of
the planned months; a milestone fills the month it falls in (red when its deadline is missed).
"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import column_index_from_string, get_column_letter

from . import rag
from .core import HEADER_FILLS, NOTE_FONT, Col, Show, Tab, fill, font, name, ranked_table, title
from .tasks import H

SHEET = "Roadmap"
MONTHS = 18             # timeline months (scroll with Settings > RoadmapStart)
PHASE_ROWS = 10
MILESTONE_ROWS = 15
FIRST_COL = 2           # the lists' hidden _Task column
MONTH_ROW = 5           # month header, shared by both lists' timelines
COLOURS = {"bar": "5B9BD5", "done": "70AD47", "late": "F4B183", "planned": "1F3864", "milestone": "7030A0",
           "missed": "C00000", "status": "FF0000"}
SHORT = "dd-mmm-yy"

PHASES = [Show("WBS", "tsk_WBS", width=6), Show("Phase", "tsk_Name", width=32), Show("Owner", "tsk_Owner", width=14),
          Show("Planned Start", "tsk_PlanStart", "date", width=10, fmt=SHORT),
          Show("Planned Finish", "tsk_PlanFinish", "date", width=10, fmt=SHORT),
          Show("Start", "tsk_FcStart", "date", width=10, fmt=SHORT),
          Show("Finish", "tsk_FcFinish", "date", width=10, fmt=SHORT),
          Show("Progress", "tsk_Progress", "num", width=9, fmt="0%"),
          Show("Finish Var", "tsk_Variance", "num", width=8, fmt="+0;-0;0")]
MILESTONES = [Show("WBS", "tsk_WBS", width=6), Show("Milestone", "tsk_Name", width=32),
              Show("Planned", "tsk_PlanFinish", "date", width=10, fmt=SHORT),
              Show("Deadline", "tsk_Deadline", "date", width=10, fmt=SHORT),
              Show("Forecast", "tsk_FcFinish", "date", width=10, fmt=SHORT),
              Show("Deadline Check", "tsk_DeadlineCheck", width=10),
              Show("Status", "tsk_Status", width=9),
              Show("Finish Var", "tsk_Variance", "num", width=8, fmt="+0;-0;0")]


def _phase_extras():
    """Count of detail tasks done, RAG, next milestone and a link, for each phase row (tsk_Phase = the phase's WBS)."""
    p = Tab("RoadmapPhases")
    blank = f'{p.row("_Task")}=""'
    wbs = p.row("WBS")
    tasks = f'COUNTIFS(tsk_Phase,{wbs},tsk_IsLeaf,1,tsk_Duration,">0"'          # milestones aren't tasks to do
    done = f'IF({blank},"",{tasks},tsk_Progress,">=1")&" of "&{tasks})&" done")'
    nxt = p.row("_Milestone")
    head = H("Name")
    link = (f'IF({blank},"",HYPERLINK("#"&ADDRESS(ROW({head})+{p.row("_Task")},COLUMN({head}),1,1,"Tasks"),'
            f'"Go to phase"))')
    return [
        Col("Tasks Done", width=10, formula=done, center=True,
            note="Detail tasks complete out of the phase's tasks (milestones left out). Progress weighs them by duration."),
        Col("RAG", width=7, formula=f'IF({blank},"",{rag.rag_text(rag.slip(p.row("Finish Var")))})', center=True,
            note="From the phase's finish slip against its planned finish (Settings > RAG slip thresholds)."),
        Col("Next Milestone", width=26, formula=(
            f'IF(OR({blank},{nxt}=""),"",INDEX(tsk_Name,{nxt})&" ("&TEXT(INDEX(tsk_FcFinish,{nxt}),"dd-mmm")&")")'),
            note="The phase's earliest milestone not yet done, with its forecast date."),
        Col("Link", width=10, formula=link, note="Jumps to the phase's row on Tasks."),
        Col("_Milestone", "help", 4, formula=(
            f'IF({blank},"",IFERROR(MATCH(_xlfn.MINIFS(tsk_MilestoneKey,tsk_Phase,{wbs}),tsk_MilestoneKey,0),""))'),
            note="Row on Tasks of the phase's next milestone (smallest open milestone key in the phase)."),
    ]


def build_roadmap(wb):
    ws = wb.create_sheet(SHEET)
    title(ws, None)
    ws["A1"] = '=ProjectName&"  -  Roadmap"'
    ws["A2"] = ('="Status date "&TEXT(StatusDate,"dd-mmm-yyyy")&"   Forecast finish "&TEXT(ProjectFinish,"dd-mmm-yyyy")'
                '&" ("&TEXT(FinishVarianceDays,"+0;-0;0")&" working days vs plan)"')
    ws["A2"].font = font(bold=True, color="404040")
    ws.column_dimensions["A"].width = 2
    phases_top = MONTH_ROW + 2
    phases = ranked_table(ws, "RoadmapPhases", phases_top, FIRST_COL, "tsk_PhaseKey", PHASE_ROWS, PHASES,
                          label='"Phases"', extra=_phase_extras())
    ws.column_dimensions[phases["_Milestone"]].hidden = True
    rag.colour(ws, f"{phases['RAG']}{phases_top + 1}:{phases['RAG']}{phases_top + PHASE_ROWS}",
               f"{phases['RAG']}{phases_top + 1}")
    milestones_top = phases_top + PHASE_ROWS + 3
    milestones = ranked_table(ws, "RoadmapMilestones", milestones_top, FIRST_COL, "tsk_MilestoneOrder",
                              MILESTONE_ROWS, MILESTONES, label='"Milestones (in planned order)"')
    widths = {}                                               # the lists share columns: the wider of each pair
    for table, cols in ((phases, PHASES + _phase_extras()), (milestones, MILESTONES)):
        for c in cols:
            letter = table[c.header if isinstance(c, Show) else c.name]
            widths[letter] = max(widths.get(letter, 0), c.width)
    for letter, width in widths.items():
        if not ws.column_dimensions[letter].hidden:
            ws.column_dimensions[letter].width = width
    spacer = max(column_index_from_string(x) for t in (phases, milestones) for x in t.values()) + 1
    first, last = spacer + 1, spacer + MONTHS
    ws.column_dimensions[get_column_letter(spacer)].width = 1
    name(wb, "rm_Origin", f"{SHEET}!${get_column_letter(spacer)}${MONTH_ROW}")
    month = "EDATE(RoadmapStart,COLUMN()-COLUMN(rm_Origin)-1)"
    for col in range(first, last + 1):
        ws.column_dimensions[get_column_letter(col)].width = 4.6
        head = ws.cell(MONTH_ROW, col, f"={month}")
        head.number_format, head.fill = "mmm", HEADER_FILLS["calc"]
        head.font, head.alignment = font(size=8, bold=True, color="FFFFFF"), Alignment(horizontal="center")
        year = ws.cell(MONTH_ROW - 1, col, f'=IF(OR(COLUMN()-COLUMN(rm_Origin)=1,MONTH({month})=1),'
                                           f'TEXT({month},"yyyy"),"")')
        year.font = font(size=8, bold=True, color="404040")
    _phase_bars(ws, phases, first, last, phases_top + 1, phases_top + PHASE_ROWS)
    _milestone_marks(ws, milestones, first, last, milestones_top + 1, milestones_top + MILESTONE_ROWS)
    _legend(ws, milestones_top + MILESTONE_ROWS + 2)
    ws.freeze_panes = ws.cell(MONTH_ROW + 1, FIRST_COL + 3)
    ws.page_setup.orientation = "landscape"                 # one page when printed or saved as PDF
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.gridLines = False
    return ws


def _month_cf(first):
    """A1 pieces for conditional formatting (Excel requires A1 there): this column's month and the next month."""
    m = f"{get_column_letter(first)}${MONTH_ROW}"
    return m, f"EDATE({m},1)"


def _phase_bars(ws, L, first, last, top, bottom):
    grid = f"{get_column_letter(first)}{top}:{get_column_letter(last)}{bottom}"
    m, nxt = _month_cf(first)
    at = lambda col: f"${L[col]}{top}"
    covers = lambda a, b: f"AND(ISNUMBER({at(a)}),ISNUMBER({at(b)}),{m}<={at(b)},{nxt}>{at(a)})"
    done_to = f"{at('Start')}+({at('Finish')}-{at('Start')}+1)*N({at('Progress')})"
    planned = Side(style="medium", color=COLOURS["planned"])
    for rule in (
        FormulaRule(formula=[f"AND({covers('Start', 'Finish')},N({at('Progress')})>0,{m}<{done_to})"],
                    fill=fill(COLOURS["done"])),
        FormulaRule(formula=[f"AND({covers('Start', 'Finish')},N({at('Finish Var')})>0)"], fill=fill(COLOURS["late"])),
        FormulaRule(formula=[covers("Start", "Finish")], fill=fill(COLOURS["bar"])),
        FormulaRule(formula=[covers("Planned Start", "Planned Finish")], border=Border(top=planned, bottom=planned)),
    ):
        ws.conditional_formatting.add(grid, rule)
    _status_line(ws, grid, m, nxt)


def _milestone_marks(ws, L, first, last, top, bottom):
    grid = f"{get_column_letter(first)}{top}:{get_column_letter(last)}{bottom}"
    m, nxt = _month_cf(first)
    at = lambda col: f"${L[col]}{top}"
    in_month = lambda col: f"AND(ISNUMBER({at(col)}),{at(col)}>={m},{at(col)}<{nxt})"
    for rule in (
        FormulaRule(formula=[f'AND({in_month("Forecast")},LEFT({at("Deadline Check")},4)="Late")'],
                    fill=fill(COLOURS["missed"])),
        FormulaRule(formula=[in_month("Forecast")], fill=fill(COLOURS["milestone"])),
        FormulaRule(formula=[in_month("Planned")],
                    border=Border(top=Side(style="medium", color=COLOURS["planned"]),
                                  bottom=Side(style="medium", color=COLOURS["planned"]))),
    ):
        ws.conditional_formatting.add(grid, rule)
    deadline = Side(style="medium", color=COLOURS["missed"])
    ws.conditional_formatting.add(grid, FormulaRule(formula=[in_month("Deadline")], border=Border(right=deadline)))
    _status_line(ws, grid, m, nxt)
    ws.conditional_formatting.add(f"{L['Deadline Check']}{top}:{L['Deadline Check']}{bottom}", FormulaRule(
        formula=[f'LEFT({at("Deadline Check")},4)="Late"'], font=Font(color="C00000", bold=True)))


def _status_line(ws, grid, m, nxt):
    side = Side(style="thin", color=COLOURS["status"])
    ws.conditional_formatting.add(grid, FormulaRule(formula=[f"AND(StatusDate>={m},StatusDate<{nxt})"],
                                                    border=Border(left=side, right=side)))


def _legend(ws, row):
    planned = Side(style="medium", color=COLOURS["planned"])
    items = [("Forecast", COLOURS["bar"], None), ("Done", COLOURS["done"], None), ("Late vs plan", COLOURS["late"], None),
             ("Planned (outline)", None, Border(top=planned, bottom=planned)), ("Milestone", COLOURS["milestone"], None),
             ("Deadline missed", COLOURS["missed"], None),
             ("Status date month", None, Border(left=Side(style="thin", color=COLOURS["status"])))]
    for i, (text, colour, border) in enumerate(items):          # two rows of swatch + label pairs, under the lists
        r, c = row + i // 4, FIRST_COL + 1 + 2 * (i % 4)
        swatch = ws.cell(r, c)
        if colour:
            swatch.fill = fill(colour)
        if border:
            swatch.border = border
        ws.cell(r, c + 1, text).font = font(size=8)
    ws.cell(row + 2, FIRST_COL + 1, "Months: Settings > RoadmapStart. Lists: drag a table's corner down for more "
                                    "rows. The timeline is conditional formatting only.").font = NOTE_FONT
