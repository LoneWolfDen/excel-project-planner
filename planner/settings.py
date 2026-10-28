"""Settings (one named cell per setting), dropdown lists and the holiday calendar."""
import datetime as dt

from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .core import (BOX, CALC_FILL, CALC_FONT, DATE, INPUT_FONT, NOTE_FONT, Col, add_table, fill, font, has,
                   header_cell, missing, name, release, title)
from .tasks import CONSTRAINTS

# Working-week presets -> Mon..Sun mask for WORKDAY.INTL (1 = not working); "Custom" uses the seven day rows
WEEKS = {"Mon-Fri": "0000011", "Sun-Thu": "0000110", "Mon-Sat": "0000001"}
DAYS = ("WorkMon", "WorkTue", "WorkWed", "WorkThu", "WorkFri", "WorkSat", "WorkSun")


def week_preset(settings):
    """The Working week preset matching the Yes / No day settings in a dict (missing days = Mon-Fri), or Custom."""
    mask = "".join("0" if settings.get(d, "Yes" if i < 5 else "No") == "Yes" else "1" for i, d in enumerate(DAYS))
    return next((w for w, m in WEEKS.items() if m == mask), "Custom")


# the settings a plan needs: typed on the Start sheet (the rest are on Settings, with defaults)
ESSENTIAL = ("ProjectName", "ProjectManager", "ProjectStartDate", "WorkingWeek", "HoursPerDay", "ProjectType")
START = "Start"
START_ROW = 4                      # first essential setting's row on Start
STATUS = "Status"                  # the status date is typed on the Status sheet (the weekly page), at this row
STATUS_ROW = 5
PROMPTED = ESSENTIAL + ("StatusDateInput",)

LISTS = {  # table name, column name, named range, values
    "L_DepType": ("ListDepType", "Link Type", ["FS", "SS", "FF", "SF"]),
    "L_Mode": ("ListMode", "Task Mode", ["Auto", "Manual"]),
    "L_WorkingWeek": ("ListWorkingWeek", "Working Week", list(WEEKS) + ["Custom"]),
    "L_Constraint": ("ListConstraint", "Constraint", CONSTRAINTS),
    "L_YesNo": ("ListYesNo", "Yes / No", ["Yes", "No"]),
    "L_ResType": ("ListResType", "Resource Type", ["Work", "Material", "Cost"]),
    "L_Period": ("ListPeriod", "Gantt Days per Column", [1, 7, 30, 91]),
    "L_RAG": ("ListRAG", "RAG", ["Green", "Amber", "Red"]),
    "L_LogType": ("ListLogType", "Log Type", ["Risk", "Assumption", "Issue", "Dependency", "Decision", "Change"]),
    "L_RAIDStatus": ("ListRAIDStatus", "RAID Status", ["Open", "In Progress", "Closed"]),
    "L_DecisionStatus": ("ListDecisionStatus", "Decision Status", ["Proposed", "Approved", "Rejected", "Superseded"]),
    "L_ChangeStatus": ("ListChangeStatus", "Change Status", ["Submitted", "Under Review", "Approved", "Rejected",
                                                             "Implemented"]),
}


def settings_rows(blank, max_sprints):
    """(name, label, value, number format, is input, note). Values that are formulas reference other names only."""
    custom = "&".join(f'IF({d}="Yes","0","1")' for d in
                      ("WorkMon", "WorkTue", "WorkWed", "WorkThu", "WorkFri", "WorkSat", "WorkSun"))
    work_week = "".join(f'IF(WorkingWeek="{w}","{mask}",' for w, mask in WEEKS.items()) + custom + ")" * len(WEEKS)
    return [
        ("ProjectName", "Project name", "New Project" if blank else "Customer Portal Implementation", "@", True,
         "Shown on every sheet title."),
        ("ProjectManager", "Project manager", None if blank else "Priya Sharma", "@", True, ""),
        ("ProjectStartDate", "Project start date", None if blank else dt.date(2026, 9, 7), DATE, True,
         "Required. Tasks without predecessors start here (moved to the next working day if needed)."),
        ("ProjectStart", "Project start (in use)", "=MAX(1,N(ProjectStartDate))", DATE, False,
         "Copy of the date above (dates show as 1900 until it is entered)."),
        ("StatusDateInput", "Status date", None if blank else dt.date(2026, 10, 25), DATE, True,
         "The date progress is reported as of. Type it, or press Ctrl+; for today. It is a typed date rather than "
         "TODAY() on purpose: TODAY() would make Excel recalculate the whole schedule after every edit."),
        ("StatusDate", "Status date (in use)", "=IF(N(StatusDateInput)>0,StatusDateInput,ProjectStart)", DATE, False,
         "The project start until a status date is entered. Drives Status, the red status-date line on the Gantt "
         "and rescheduling of late work."),
        ("StatusWorkday", "First working day on/after status date", "=WORKDAY.INTL(StatusDate-1,1,WorkWeek,Holidays)",
         DATE, False, ""),
        ("RescheduleAfterStatus", "Reschedule incomplete work after status date?", "Yes", "@", True,
         "Yes = like MS Project 'move incomplete parts after status date': unstarted tasks can't start before the status "
         "date and in-progress tasks finish their remaining duration after it, pushing all successors."),
        ("HoursPerDay", "Working hours per day", 8, "0.0", True, "Work (hrs) = Duration x Hours/day x Units."),
        ("ProjectType", "Project type", DEFAULT_TYPE, "@", True,
         "Sets the defaults for the stages (Settings > Stages), the sprint length and the RAG slip thresholds: "
         + ", ".join(PRESETS) + ". Pick it first; each default stays editable (type over it)."),
        ("DefaultUnits", "Default assignment units", 1, "0%", True,
         "Used when a task has a resource but Units is blank. 100% = full time."),
        ("WorkingWeek", "Working week", "Mon-Fri", "@", True,
         "Mon-Fri, Sun-Thu or Mon-Sat. Custom = the seven day rows on Settings."),
        ("WorkMon", "Monday is a working day", "Yes", "@", True,
         "Used when Working week (Start sheet) is Custom: Yes / No for each day."),
        ("WorkTue", "Tuesday is a working day", "Yes", "@", True, ""),
        ("WorkWed", "Wednesday is a working day", "Yes", "@", True, ""),
        ("WorkThu", "Thursday is a working day", "Yes", "@", True, ""),
        ("WorkFri", "Friday is a working day", "Yes", "@", True, ""),
        ("WorkSat", "Saturday is a working day", "No", "@", True, ""),
        ("WorkSun", "Sunday is a working day", "No", "@", True, ""),
        ("WorkWeek", "Weekend mask (calculated)", f"={work_week}", "@", False,
         "7-character Mon..Sun string used by WORKDAY.INTL (1 = non-working)."),
        ("ViewFrom", "Show the Gantt and Resource Report from", "=MAX(ProjectStart,StatusDate-14)", DATE, True,
         "Defaults to two weeks before the status date (the project start until then). Overwrite with any date to "
         "scroll both views at once; GanttStart and ResourceWeekStart below follow it unless you overwrite them."),
        ("GanttStart", "Gantt timeline start", "=ViewFrom-WEEKDAY(ViewFrom,3)", DATE, True,
         "Defaults to the Monday of the View from week. Overwrite with any date to scroll the timeline alone."),
        ("GanttPeriod", "Gantt days per column", 1, "0", True,
         "1 = daily timeline, 7 = weekly, 30 = monthly (calendar months), 91 = quarterly (calendar quarters)."),
        ("OutlineShowLevel", "Show tasks down to outline level", 9, "0", True,
         "Tasks > Show (DETAIL section) is Yes for rows at this level or above: filter Show = Yes to see only phases (1), "
         "phases and tasks (2) ... Reapply the filter (Data > Reapply) after changing it."),
        ("LiveGrids", "Live resource grids", "Yes", "@", True,
         "No = the Resource Report's weekly and sprint grids per person (and Tasks > Over-allocated) stop "
         "calculating, for fast bulk entry on a big plan; Status says so while they are off. Set Yes afterwards."),
        ("RoadmapStart", "Roadmap first month", "=DATE(YEAR(ProjectStart),MONTH(ProjectStart),1)", "mmm-yy", True,
         "The Roadmap sheet's first month. Overwrite with the first of any month to scroll it."),
        ("Sprint1Start", "Sprint 1 start date", "=ProjectStart", DATE, True,
         "Defaults to the project start. Overwrite to align with your sprint calendar."),
        ("SprintLength", "Sprint length (calendar days)", f"={preset('Sprint Length', 14)}", "0", True,
         "14 = two-week sprints. The default comes from the Project type (Start); type a number to set your own."),
        ("SprintCount", "Number of sprints", min(12, max_sprints), "0", True,
         f"Up to the number of rows in the sprint table (Agile sheet) ({max_sprints}; add rows to the table for more). "
         "Tasks outside the sprint range are labelled 'Backlog'."),
        ("SprintPrefix", "Sprint name prefix", "Sprint ", "@", True, "Sprint names = prefix & number."),
        ("FocusFactor", "Sprint focus (share of time on delivery)", 0.8, "0%", True,
         "Of each person's available days in a sprint, the share that is delivery (1 person-day = 1 story point); "
         "the rest is ceremonies and overhead. A Focus % on Resources overrides it per person (0% = advisory)."),
        ("SprintLastEnd", "Last sprint end (calculated)", "=INDEX(spr_End,SprintCount)", DATE, False, ""),
        ("CurrentSprint", "Sprint containing the status date (calculated)",
         '=IF(OR(StatusDate<Sprint1Start,StatusDate>SprintLastEnd),"",INDEX(spr_Name,MATCH(StatusDate,spr_Start,1)))',
         "@", False, "Blank outside the sprint calendar. The Agile sheet shows its backlog and daily burndown."),
        ("CurrentSprintStart", "Current sprint start (calculated)",
         '=IF(CurrentSprint="","",INDEX(spr_Start,MATCH(CurrentSprint,spr_Name,0)))', DATE, False, ""),
        ("CurrentSprintEnd", "Current sprint end (calculated)",
         '=IF(CurrentSprint="","",INDEX(spr_End,MATCH(CurrentSprint,spr_Name,0)))', DATE, False, ""),
        ("TotalPoints", "Total story points (calculated)", "=SUM(tsk_LeafPoints)", "0.#", False,
         "Story points on all detail tasks: the release scope for burndown and burn-up."),
        ("MaxPeople", "Most people on one task (calculated)", "=MAX(tsk_People)", "0", False,
         "Per-resource totals skip the person slots no task uses, so a plan with one person per task stays fast."),
        ("CriticalSlackDays", "Critical when total slack is at most (days)", 0, "0", True,
         "0 = only tasks with no slack (or negative slack from a missed deadline) are critical, like MS Project."),
        ("ProjectFinish", "Forecast project finish (calculated)", "=MAX(tsk_LeafFcstFinish)", DATE, False,
         "The critical path's backward pass starts here (the detail tasks' early finish)."),
        ("PlannedFinish", "Planned project finish (calculated)", "=MAX(tsk_PlanFinish)", DATE, False,
         "Baseline finish once a baseline is set."),
        ("BaselineSet", "Baseline set (calculated)", "=COUNT(tsk_BaseStart)>0", "@", False,
         "TRUE once any task has a Baseline Start: the Gantt then outlines only the tasks that have one."),
        ("FinishVarianceDays", "Project finish variance (calculated, working days)",
         "=IF(ProjectFinish>=PlannedFinish,NETWORKDAYS.INTL(PlannedFinish,ProjectFinish,WorkWeek,Holidays)-1,"
         "1-NETWORKDAYS.INTL(ProjectFinish,PlannedFinish,WorkWeek,Holidays))", "+0;-0;0", False,
         "Forecast finish minus planned finish. Positive = late."),
        ("LookaheadDays", "Look-ahead window (working days)", 10, "0", True,
         "Status lists tasks starting or due within this many working days of the status date."),
        ("LookaheadEnd", "Look-ahead window ends (calculated)", "=WORKDAY.INTL(StatusDate-1,LookaheadDays,WorkWeek,Holidays)",
         DATE, False, ""),
        ("RAGAmberSlipDays", "RAG: amber when the finish slips by (working days)", f"={preset('Amber Slip', 3)}", "0", True,
         "Status sheet thresholds. A finish slip of at least this many working days is amber."),
        ("RAGRedSlipDays", "RAG: red when the finish slips by (working days)", f"={preset('Red Slip', 10)}", "0", True,
         "Both slip thresholds default from the Project type (Start); type a number to set your own."),
        ("RAGAmberIndex", "RAG: amber when SPI or CPI is below", 0.95, "0.00", True,
         "SPI = earned value / planned value; CPI = earned value / actual cost. 1.00 = on plan."),
        ("RAGRedIndex", "RAG: red when SPI or CPI is below", 0.85, "0.00", True, ""),
        ("RAIDHighScore", "RAID: high rating from score", 15, "0", True,
         "Score = probability x impact (1-25). At or above this a risk or issue is rated High."),
        ("RAIDMediumScore", "RAID: medium rating from score", 8, "0", True, ""),
        ("ResourceWeekStart", "Resource Report first week", "=MAX(2,ViewFrom-WEEKDAY(ViewFrom,3))", DATE, True,
         "Defaults to the Monday of the View from week. Overwrite with any Monday to scroll the weekly grids alone."),
        ("NextFreeUID", "Next free task UID (calculated)", "=MAX(tsk_UID)+1", "0", False,
         "Give a newly inserted task row this UID (see Start > Inserting tasks)."),
    ] + (RELEASE_ROWS if release() else [])


# Release 1.0 settings
RELEASE_ROWS = [
    ("TakeBaseline", "Re-baseline every task (BL Source)", "No", "@", True,
     "No = the BL Source columns keep each task's pasted baseline and give the current dates only to tasks without "
     "one: copy BL Source and Paste Values onto Baseline to baseline new tasks. Yes = BL Source is the current "
     "schedule for every task: the same copy and paste re-baselines the whole plan. Set it back to No afterwards."),
    ("DuplicateUIDs", "Tasks sharing a UID (calculated)", '=COUNTIF(tsk_Issues,"*is used twice*")', "0", False,
     "Should be 0. Each such task says so in Tasks > Issues; give one of them Settings > NextFreeUID."),
    ("MissingUIDs", "Named tasks without a UID (calculated)", '=COUNTIF(tsk_Issues,"No UID*")', "0", False,
     "Should be 0: links can't point at a task without a UID."),
]


def build_settings(wb, blank, max_sprints, values=None):
    """values: {setting name: value} replacing the defaults (an import sets the project's own)."""
    st = wb.create_sheet("Settings")
    title(st, "Project Settings",
          "Blue = edit me. Grey = calculated. Column D shows the defined name every formula uses for that setting. "
          "The settings every plan needs are on the Start sheet.")
    start = wb.create_sheet(START)                  # the essentials go there; build_start adds the guide
    status = wb.create_sheet(STATUS)                # the status date goes there; build_status adds the rest
    start.column_dimensions["A"].width = 3
    for col, (text, width) in enumerate([("Setting", 34), ("Value", 24), ("Named range", 24), ("Notes", 80)], start=2):
        header_cell(st.cell(4, col), text)
        st.column_dimensions[chr(64 + col)].width = width
    st.column_dimensions["A"].width = 3
    rows, cells = {}, {}
    # B-913: the sprint calendar's settings only with the Agile module; the rest whenever they can calculate
    kept = [row for row in settings_rows(blank, max_sprints)
            if not missing(row[2]) and (has("agile") or row[0] not in SPRINT_SETTINGS)]
    settings = [row for row in kept if row[0] in ESSENTIAL]
    settings.sort(key=lambda row: ESSENTIAL.index(row[0]))
    settings += [row for row in kept if row[0] not in ESSENTIAL]
    on_settings = 0
    for n, label, value, fmt, is_input, note in settings:
        if values and n in values:
            value = values[n]
        if n in ESSENTIAL:
            ws, r = start, START_ROW + ESSENTIAL.index(n)
        elif n == "StatusDateInput":
            ws, r = status, STATUS_ROW
        else:
            ws, r = st, 5 + on_settings
            on_settings += 1
            ws.cell(r, 4, n).font = font(color="7F7F7F")
            ws.cell(r, 5, note).font = NOTE_FONT
        rows[n], cells[n] = r, (ws, f"C{r}")
        ws.cell(r, 2, label).font = font(bold=n in PROMPTED)
        c = ws.cell(r, 3, value)
        c.number_format = fmt
        c.font = INPUT_FONT if is_input else CALC_FONT
        if not is_input:
            c.fill = CALC_FILL
        c.border = BOX
        name(wb, n, f"{ws.title}!$C${r}")
        if n in PROMPTED:                           # the note shows when the cell is selected
            prompt = DataValidation(type="list", formula1="=L_WorkingWeek") if n == "WorkingWeek" else \
                DataValidation(type="list", formula1="=L_ProjectType") if n == "ProjectType" else \
                DataValidation(type="decimal", operator="between", formula1="1", formula2="24") if n == "HoursPerDay" \
                else DataValidation(allow_blank=True)
            prompt.showInputMessage, prompt.promptTitle, prompt.prompt = True, label[:32], note[:255]
            ws.add_data_validation(prompt)
            prompt.add(f"C{r}")

    yes_no = DataValidation(type="list", formula1="=L_YesNo", allow_blank=False)
    st.add_data_validation(yes_no)
    for n in ("RescheduleAfterStatus", "LiveGrids", "WorkMon", "WorkTue", "WorkWed", "WorkThu", "WorkFri", "WorkSat",
              "WorkSun") + (("TakeBaseline",) if "TakeBaseline" in rows else ()):
        yes_no.add(f"C{rows[n]}")
    period = DataValidation(type="list", formula1="=L_Period")
    st.add_data_validation(period)
    period.add(f"C{rows['GanttPeriod']}")
    levels = DataValidation(type="whole", operator="between", formula1="1", formula2="9")
    st.add_data_validation(levels)
    levels.add(f"C{rows['OutlineShowLevel']}")
    if "SprintCount" in rows:
        sprints = DataValidation(type="whole", operator="between", formula1="1", formula2=str(max_sprints))
        st.add_data_validation(sprints)
        sprints.add(f"C{rows['SprintCount']}")
    for n in ("ProjectStartDate", "StatusDateInput"):
        ws, cell = cells[n]
        ws.conditional_formatting.add(cell, FormulaRule(formula=[f"N({cell})=0"], fill=fill("FFFF00")))
    return st


SPRINT_SETTINGS = {"Sprint1Start", "SprintLength", "SprintCount", "SprintPrefix", "TotalPoints"}
LISTS_COL = 7                      # Settings column G: the dropdown lists, side by side, in a collapsed group


def build_lists(wb):
    """The dropdown lists, on Settings to the right of the settings, in a collapsed column group (+ above column F)."""
    ws = wb["Settings"]
    ws.sheet_properties.outlinePr.summaryRight = False
    header = ws.cell(4, LISTS_COL - 1, "Dropdown lists (+)")
    header.font = font(bold=True, color="1F4E78")
    ws.column_dimensions[get_column_letter(LISTS_COL - 1)].width = 18
    for i, (n, (table, column, values)) in enumerate(LISTS.items()):
        col = LISTS_COL + i * 2
        add_table(ws, table, 4, col, [Col(column, "in", 16, values=values)], len(values), autofilter=False)
        name(wb, n, f"{table}[{column}]")
    last = _presets(wb, ws, LISTS_COL + 2 * len(LISTS))
    for c in range(LISTS_COL, last + 1):
        dim = ws.column_dimensions[get_column_letter(c)]
        dim.outlineLevel, dim.hidden = 1, True
    ws.cell(len(max((v[2] for v in LISTS.values()), key=len)) + 6, LISTS_COL,
            "Dropdown sources. Add a value by typing it in the row under a list: the table and its named range grow."
            ).font = NOTE_FONT
    return ws


# Progress by stage: each stage's % Complete (weighted milestones: the same stage is the same progress everywhere)
STAGES = [("Not Started", 0), ("Design", 0.15), ("In Development", 0.4), ("Test Ready", 0.7), ("In Testing", 0.85),
          ("Complete", 1)]
# B-1018: what a Project type sets by default (each stays editable): the stages, sprint length, RAG slip thresholds
PRESETS = {
    "Software / Agile": {"stages": STAGES, "Sprint Length": 14, "Amber Slip": 3, "Red Slip": 10},
    "Construction": {"stages": [("Not Started", 0), ("Mobilised", 0.1), ("Under Way", 0.5), ("Inspection", 0.85),
                                ("Snagging", 0.95), ("Complete", 1)], "Sprint Length": 28, "Amber Slip": 5, "Red Slip": 15},
    "IT infrastructure": {"stages": [("Not Started", 0), ("Design", 0.15), ("Build", 0.5), ("Testing", 0.8),
                                     ("Cutover", 0.95), ("Complete", 1)], "Sprint Length": 14, "Amber Slip": 2,
                          "Red Slip": 5},
}
DEFAULT_TYPE = "Software / Agile"
N_STAGES = max(len(p["stages"]) for p in PRESETS.values())


def preset(column, fallback):
    """The Project type's value in the ProjectPresets table (the fallback if the type isn't one of them)."""
    return f"IFERROR(INDEX(ProjectPresets[{column}],MATCH(ProjectType,ProjectPresets[Project Type],0)),{fallback})"


def stages_of(project_type):
    """The (stage, %) rows a Project type starts with: for converting a typed % to a Stage when a plan comes in."""
    return PRESETS.get(project_type or DEFAULT_TYPE, PRESETS[DEFAULT_TYPE])["stages"]


def _presets(wb, ws, col):
    """The ProjectPresets table, beside the dropdown lists: one row per Project type, editable like any list (add a
    row for your own kind of project). Returns its last column."""
    names = list(PRESETS)
    rows = [PRESETS[n] for n in names]
    stage = lambda k: [p["stages"][k][0] if k < len(p["stages"]) else None for p in rows]
    pct = lambda k: [p["stages"][k][1] if k < len(p["stages"]) else None for p in rows]
    cols = [Col("Project Type", "in", 18, values=names),
            *[Col(h, "in", 10, values=[p[h] for p in rows]) for h in ("Sprint Length", "Amber Slip", "Red Slip")],
            *[Col(f"Stage {k + 1}", "in", 14, values=stage(k)) for k in range(N_STAGES)],
            *[Col(f"% {k + 1}", "in", 7, "0%", values=pct(k)) for k in range(N_STAGES)]]
    add_table(ws, "ProjectPresets", 4, col, cols, len(names) + 3, autofilter=False)
    name(wb, "L_ProjectType", "ProjectPresets[Project Type]")
    return col + len(cols) - 1


def build_stages(wb, stages=None):
    """The Stages table on Settings, below the settings: the Stage dropdown on Tasks and each stage's % Complete.
    Editable for other kinds of work (rename a stage, change its %, add rows); keep the first at 0% and the last at
    100%, since milestones use those two. stages: (name, %) rows replacing the default."""
    ws = wb["Settings"]
    top = ws.max_row + 3
    ws.cell(top - 1, 2, "STAGES: pick one per task on Tasks > Stage; its % Complete counts as the task's progress "
                        "(a typed % Complete overrides it).").font = font(bold=True, color="1F4E78")
    if stages:                                         # an import or upgrade: the plan's own stages, typed
        names, pcts = [n for n, _ in stages], [p for _, p in stages]
    else:                                              # B-1018: the Project type's stages, until someone types over one
        k = "ROW()-ROW(StageTable[[#Headers],[Stage]])"
        row = "MATCH(ProjectType,ProjectPresets[Project Type],0)"
        names = [f'=IFERROR(INDEX(ProjectPresets[[Stage 1]:[Stage {N_STAGES}]],{row},{k})&"","")'] * N_STAGES
        pcts = [f'=N(IFERROR(INDEX(ProjectPresets[[% 1]:[% {N_STAGES}]],{row},{k}),0))'] * N_STAGES
    add_table(ws, "StageTable", top, 2, [
        Col("Stage", "in", 34, values=names),               # the settings' column widths
        Col("% Complete", "in", 24, "0%", values=pcts),
    ], len(names), autofilter=False)
    name(wb, "stg_Name", "StageTable[Stage]")
    name(wb, "stg_Pct", "StageTable[% Complete]")
    name(wb, "L_Stage", "StageTable[Stage]")
    return ws


def nearest_stage(pct, stages=None):
    """The stage whose % is closest to pct (the lower one on a tie), for a % from an older workbook or MS Project."""
    return min(stages or STAGES, key=lambda s: (abs(s[1] - pct), s[1]))


HOLIDAYS = [(dt.date(2026, 12, 24), "Christmas Eve (example - replace with your calendar)"),
            (dt.date(2026, 12, 25), "Christmas Day (example)"),
            (dt.date(2026, 12, 31), "New Year's Eve (example)"),
            (dt.date(2027, 1, 1), "New Year's Day (example)")]


def build_holidays(wb, blank, n_rows=100, holidays=None):
    """The holiday table on Settings, below the settings. holidays: (date, description) rows replacing the sample."""
    ws = wb["Settings"]
    top = ws.max_row + 3
    ws.cell(top - 1, 2, "HOLIDAYS: every date here is skipped by all scheduling formulas (named range: Holidays). "
                        "Add rows to the table as needed.").font = font(bold=True, color="1F4E78")
    data = holidays if holidays is not None else [] if blank else HOLIDAYS
    n_rows = max(n_rows, len(data) + 20)
    add_table(ws, "HolidayTable", top, 2, [
        Col("Date", "in", 34, DATE, values=[d for d, _ in data]),          # the settings' column widths
        Col("Description", "in", 24, values=[t for _, t in data]),
    ], n_rows)
    name(wb, "Holidays", "HolidayTable[Date]")
    return ws
