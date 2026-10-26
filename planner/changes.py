"""Changes sheet: the change request log (Excel Table "ChangeLog").

Each request records its schedule and cost impact as estimated when raised; the Status Report totals the approved
and pending ones. Approving a change does not edit the plan: the PM updates the tasks (and re-baselines if the
sponsor agrees), so the plan and its history stay explicit.
"""
import datetime as dt

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from .core import DATE, Col, Tab, add_table, col_range, linked_task, name, title

D = dt.date
STATUSES = ["Submitted", "Under Review", "Approved", "Rejected", "Implemented"]
# Ref, Raised, Requested by, Title, Reason, Days, Cost, Status, Decided, Decided by, Task UID
SAMPLE = [
    (1, D(2026, 9, 28), "Marketing", "Add a newsletter sign-up to the portal", "Supports the launch campaign.", 2, 1360,
     "Approved", D(2026, 10, 2), "Sponsor", 17),
    (2, D(2026, 10, 22), "Priya Sharma", "Second developer on the core API for two weeks",
     "Recovers most of the API slip (RAID 3).", -4, 6800, "Under Review", None, None, 13),
    (3, D(2026, 10, 12), "Sales", "Live chat on every page", "Nice to have; not needed for go-live.", 8, 9600,
     "Rejected", D(2026, 10, 16), "Steering committee", None),
]
NAMES = {"Status": "chg_Status", "Schedule Impact": "chg_Days", "Cost Impact": "chg_Cost"}


def build_changes(wb, blank, n_rows=100):
    ws = wb.create_sheet("Changes")
    title(ws, "Change Request Log",
          "Schedule impact in working days (negative = saves time) and cost impact as estimated when raised. "
          "Approving a request doesn't change the plan: update the tasks, and re-baseline if agreed.")
    ws.column_dimensions["A"].width = 3
    data = [] if blank else SAMPLE
    t = Tab("ChangeLog")
    r = t.row
    values = lambda i: [row[i] for row in data]
    letters = add_table(ws, "ChangeLog", 4, 2, [
        Col("Ref", "in", 6, center=True, values=list(range(1, n_rows + 1))),
        Col("Raised", "in", 12, DATE, values=values(1)),
        Col("Requested By", "in", 16, values=values(2)),
        Col("Title", "in", 44, values=values(3)),
        Col("Reason", "in", 40, values=values(4)),
        Col("Schedule Impact", "in", 9, "+0;-0;0", center=True, values=values(5), note="Working days (negative = saves time)."),
        Col("Cost Impact", "in", 11, "#,##0", values=values(6)),
        Col("Status", "in", 12, center=True, values=values(7)),
        Col("Decided", "in", 12, DATE, values=values(8)),
        Col("Decided By", "in", 16, values=values(9)),
        Col("Waiting (days)", width=9, center=True, note="Days an undecided request has waited, to the status date.",
            formula=(f'IF(OR({r("Title")}="",N({r("Raised")})=0,AND({r("Status")}<>"Submitted",'
                     f'{r("Status")}<>"Under Review")),"",MAX(0,StatusDate-{r("Raised")}))')),
        Col("Task UID", "in", 8, center=True, values=values(10)),
        Col("Task", width=28, formula=linked_task(r("Task UID"))),
    ], n_rows)
    for column, defined in NAMES.items():
        name(wb, defined, f"ChangeLog[{column}]")
    first, last = 5, 4 + n_rows
    rule = DataValidation(type="list", formula1="=L_ChangeStatus")
    ws.add_data_validation(rule)
    rule.add(col_range(letters["Status"], first, last))
    status = letters["Status"]
    for word, colour in (("Approved", "548235"), ("Implemented", "548235"), ("Rejected", "A6A6A6"),
                         ("Submitted", "C55A11"), ("Under Review", "C55A11")):
        ws.conditional_formatting.add(col_range(status, first, last), FormulaRule(
            formula=[f'{status}{first}="{word}"'], font=Font(color=colour, bold=True)))
    ws.freeze_panes = f"E{first}"
    return ws
