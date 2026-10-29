"""Resource sheet (Excel Table "ResourceTable"): who is on the project, how much of them (Allocation %), when (Start /
End), how much of their sprint time is delivery (Focus %), and their leave by month. Day rates are not here: they are
on the Cost sheet, which can be deleted before the plan is shared."""
from openpyxl.worksheet.datavalidation import DataValidation

from .core import DATE, Col, Tab, add_table, col_range, has, name, release, title, v11, v12
from .leave import leave_columns, month_labels, name_leave
from .tasks import people

# name, initials, group, allocation %, day rate (for the Cost sheet's rate card), email
RESOURCES = [("Priya Sharma", "PS", "Project Manager", 0.5, 760, "priya@example.com"),
             ("Morgan Patel", "MP", "Business Analyst", 1, 560, "morgan@example.com"),
             ("Alex Chen", "AC", "Solution Architect", 0.5, 880, "alex@example.com"),
             ("Casey Kim", "CK", "UX Designer", 1, 600, "casey@example.com"),
             ("Jordan Lee", "JL", "Backend Developer", 1, 680, "jordan@example.com"),
             ("Sam Rivera", "SR", "Frontend Developer", 1, 640, "sam@example.com"),
             ("Taylor Morgan", "TM", "QA Engineer", 1, 520, "taylor@example.com")]
# the sample's advisory roles: allocated, but no story-point capacity (Focus % = 0)
FOCUS = {"Priya Sharma": 0, "Alex Chen": 0}

NAMES = {"Resource Name": "res_Name", "Initials": "res_Initials", "Group": "res_Group", "Type": "res_Type",
         "Allocation %": "res_Allocation", "Start": "res_Start", "End": "res_End", "Focus %": "res_FocusInput",
         "_Focus": "res_Focus", "Email": "res_Email"}


def by_resource(field, resource):
    """A per-resource total of one per-person value (tasks.PERSON_FIELDS): one SUMIFS per person on a task."""
    parts = [f"SUMIFS({value},{who},{resource})" for value, who in zip(people(field), people("Name"))]
    # in brackets, so a caller can divide or multiply the whole total
    return "(" + "+".join([parts[0]] + [f"IF(MaxPeople>={k},{p},0)" for k, p in enumerate(parts[1:], start=2)]) + ")"


def count_by_resource(resource, criteria=""):
    """Tasks this resource is on (any of a task's people), with extra COUNTIFS criteria such as ',tsk_IsLeaf,1'."""
    parts = [f"COUNTIFS({who},{resource}{criteria})" for who in people("Name")]
    return "(" + "+".join([parts[0]] + [f"IF(MaxPeople>={k},{p},0)" for k, p in enumerate(parts[1:], start=2)]) + ")"


def named_in(column, person):
    """Tasks whose RACI cell names this person, alone or in a comma-separated list (exact names, not parts)."""
    return "+".join(f'COUNTIF({column},{pattern})' for pattern in
                    (person, f'{person}&",*"', f'"*, "&{person}', f'"*, "&{person}&",*"'))


def _raci_columns(nm):
    """Template 1.1 (F-1101): how many tasks each person is accountable for, consulted on and informed of
    (Responsible is the Tasks column and Resources' Tasks count)."""
    return [Col(header, "calc", 11, "0", f'IF({nm}="","",{named_in(column, nm)})', note=note)
            for header, column, note in (
                ("Accountable For", "tsk_Accountable", "Tasks with this person as Accountable (RACI)."),
                ("Consulted On", "tsk_Consulted", "Tasks that name this person in Consulted (RACI)."),
                ("Informed Of", "tsk_Informed", "Tasks that name this person in Informed (RACI)."))]


def build_resources(wb, blank, n_rows=50, resources=None, leave=None, months=None, n_sprints=26, n_grid=25):
    """resources: (name, initials, group, allocation, day rate, email[, type, start, end, focus]) rows replacing the
    sample (an import); the day rate goes to the Cost sheet. leave: {person: {month header: days text}}; months: the
    month headers (one column each, after the details)."""
    ws = wb.create_sheet("Resources")
    title(ws, "Resource Sheet", "One row per person: who they are, their availability and rate, then their leave by "
                                'month, e.g. "16, 17(0.5)" under Nov-26 = all of 16 November and half of the 17th. '
                                "Leave lowers capacity (Resource Report, Agile) but doesn't move tasks. Add rows for "
                                'more people; add a month by inserting a column among the months (right-click a month > Insert > Table Columns '
                                'to the Left) and typing its name, e.g. "Jan-29".')
    ws.column_dimensions["A"].width = 3
    data = resources if resources is not None else [] if blank else \
        [d + ("Work", None, None, FOCUS.get(d[0])) for d in RESOURCES]
    extra = lambda i: [d[i] if len(d) > i else None for d in data]
    t = Tab("ResourceTable")
    nm = t.row("Resource Name")
    letters = add_table(ws, "ResourceTable", 5, 2, [
        Col("Resource Name", "in", 22, values=[d[0] for d in data]),
        Col("Initials", "in", 9, values=[d[1] for d in data]),
        Col("Group", "in", 20, values=[d[2] for d in data]),
        Col("Type", "in", 10, values=[d[6] if len(d) > 6 else "Work" for d in data], note="Work, Material or Cost (MS Project resource types)."),
        Col("Allocation %", "in", 11, "0%", values=[d[3] for d in data],
            note="How much of this person is on this project: 100% = full time, 70% = 7 of every 10 days."),
        Col("Start", "in", 12, DATE, values=extra(7), note="Optional: first day on the project (blank = from the start)."),
        Col("End", "in", 12, DATE, values=extra(8), note="Optional: last day on the project (blank = to the end)."),
        Col("Focus %", "in", 9, "0%", values=extra(9),
            note="Share of their sprint time that is delivery (story points); the rest is ceremonies and overhead. "
                 "Blank = Settings > FocusFactor. 0% = an advisory role: allocated, but no points."),
        Col("Email", "in", 26, values=[d[5] for d in data]),
        Col("Tasks", "calc", 9, "0", f'IF({nm}="","",{count_by_resource(nm, ",tsk_IsLeaf,1")})',
            note="Detail tasks this resource is on (Tasks > Resources)."),
        Col("Work", "calc", 13, "#,##0.0", f'IF({nm}="","",{by_resource("Work", nm)})',
            note="Assigned work (hours)."),
        Col("Notes", "in", 30),
        *(_raci_columns(nm) if v11() else []),
        Col("_Focus", "help", 6, "0%", f'IF({t.row("Focus %")}="",FocusFactor,N({t.row("Focus %")}))',
            note="Focus % in use (the setting when blank)."),
    ] + leave_columns(data, blank, leave, months), n_rows)
    ws.column_dimensions[letters["_Focus"]].hidden = True
    name_leave(wb, months or month_labels())
    for column, defined in NAMES.items():
        name(wb, defined, f"ResourceTable[{column}]")
    dv = DataValidation(type="list", formula1="=L_ResType")
    ws.add_data_validation(dv)
    dv.add(col_range(letters["Type"], 6, 5 + n_rows))
    below = 5 + n_rows + 5
    if release():                                      # Release 1.0: the leave day grid, under the people
        from .leave import build_leave_grid
        # template 1.2 types leave only on the grid: an import's or upgrade's leave goes there
        below = build_leave_grid(wb, ws, below, blank, entries=leave if v12() else None) + 5
    if has("agile"):                                   # sprint capacity needs the sprint calendar (B-913)
        capacity_grids(wb, ws, below, n_sprints, n_grid)
    ws.freeze_panes = "C6"
    return ws


WD = "WorkWeek,Holidays"


def capacity_grids(wb, ws, top, n_sprints, n_people):
    """Each person's sprint capacity (B-1006), below the table: available days = working days of the sprint within
    their Start / End, less holidays and their leave, x Allocation %; delivery capacity in story points (1 person-day
    = 1 point) = available days x Focus %. Rows follow ResourceTable, columns the sprint calendar (Agile), each cell
    found by its position from a named anchor. They read people, the calendar and leave only, so editing a task
    never recalculates them."""
    days_top = top
    points_top = top + n_people + 5
    for kind, at, heading in (("Days", days_top, "AVAILABLE DAYS by sprint: working days on the project, less "
                                                 "leave, x Allocation %"),
                              ("Points", points_top, "DELIVERY CAPACITY (story points) by sprint: available days x "
                                                     "Focus % (1 person-day = 1 point)")):
        _capacity_grid(wb, ws, kind, at, n_sprints, n_people, heading)
    _sprint_leave_grid(wb, ws, points_top + n_people + 5, n_sprints, n_people)


def _capacity_grid(wb, ws, kind, top, n_sprints, n_people, heading):
    from openpyxl.utils import get_column_letter as L
    from .core import BOX, CALC_FONT, font, header_cell
    ws.cell(top - 2, 2, heading).font = font(bold=True, color="1F4E78")
    origin = f"cap_{kind}Origin"
    name(wb, origin, f"Resources!$B${top}")
    header_cell(ws.cell(top, 2), "Resource")
    i = f"(ROW()-ROW({origin})-1)"                 # person number (row 1 below the sprint-start row)
    j = f"(COLUMN()-COLUMN({origin}))"            # sprint number
    sprint = f"INDEX(spr_Name,{j})"
    first_col, last_col = 3, 2 + n_sprints
    for c in range(first_col, last_col + 1):
        header_cell(ws.cell(top, c), None)
        ws.cell(top, c).value = f'=IFERROR({sprint}&"","")'
        d = ws.cell(top + 1, c, f'=IFERROR(IF({sprint}="","",INDEX(spr_Start,{j})),"")')
        d.number_format, d.font = "dd-mmm", font(italic=True, size=9)
    ws.cell(top + 1, 2, "Sprint start").font = font(italic=True, size=9)
    person = f"INDEX(res_Name,{i})"
    if kind == "Days":
        start, end = f"INDEX(spr_Start,{j})", f"INDEX(spr_End,{j})"
        lo = f"MAX({start},N(INDEX(res_Start,{i})))"
        hi = f"MIN({end},IF(N(INDEX(res_End,{i}))=0,{end},INDEX(res_End,{i})))"
        worked = f"IF({hi}<{lo},0,NETWORKDAYS.INTL({lo},{hi},{WD}))"
        value = (f'IF(INDEX(res_Type,{i})<>"Work",0,MAX(0,{worked}-INDEX(spr_LeaveGrid,{i},{j}))'
                 f'*INDEX(res_Allocation,{i}))')
    else:
        value = f"N(INDEX(cap_Days,{i},{j}))*INDEX(res_Focus,{i})"
    first_row, last_row = top + 2, top + 1 + n_people
    for r in range(first_row, last_row + 1):
        ws.cell(r, 2, f'=IFERROR({person}&"","")').font = font(bold=True)
        for c in range(first_col, last_col + 1):
            cell = ws.cell(r, c, f'=IFERROR(IF(OR({person}="",{sprint}=""),"",{value}),"")')
            cell.number_format, cell.border, cell.font = "0.0;;-", BOX, CALC_FONT
    name(wb, f"cap_{kind}", f"Resources!${L(first_col)}${first_row}:${L(last_col)}${last_row}")


def _sprint_leave_grid(wb, ws, top, n_sprints, n_resources):
    """Hidden grid: working days of leave per person per sprint (the leave months), read by the capacity grids above
    and the Resource Report's sprint utilisation. Each person's Resources row is looked up once, in column C;
    each sprint's month columns come from SprintTable."""
    from openpyxl.utils import get_column_letter
    from .agile import LEAVE_MONTHS
    from .core import CALC_FONT, font
    from .leave import leave_days, leave_row, month_start
    origin = "spr_LeaveOrigin"
    name(wb, origin, f"'{ws.title}'!$B${top}")
    ws.cell(top, 2, "Leave days by sprint (helper for the capacity grids)").font = font(italic=True)
    i = f"(ROW()-ROW({origin}))"
    j = f"(COLUMN()-COLUMN({origin})-1)"
    resource = f"INDEX(res_Name,{i})"
    start, end = f"INDEX(spr_Start,{j})", f"INDEX(spr_End,{j})"
    months = [(month_start(start, k), f"INDEX(spr_LeaveCol{k},{j})") for k in range(LEAVE_MONTHS)]
    days = leave_days(start, end, f"INDEX(spr_LeaveRow,{i})", months)
    L = get_column_letter
    last_col = L(3 + n_sprints)
    for r in range(top + 1, top + 1 + n_resources):
        ws.cell(r, 3, f'=IFERROR(IF({resource}="",0,{leave_row(resource)}),0)').font = CALC_FONT
        for c in range(4, 4 + n_sprints):
            ws.cell(r, c, f'=IFERROR(IF(OR({resource}="",INDEX(spr_Name,{j})=""),0,{days}),0)').font = CALC_FONT
    name(wb, "spr_LeaveRow", f"'{ws.title}'!$C${top + 1}:$C${top + n_resources}")
    name(wb, "spr_LeaveGrid", f"'{ws.title}'!$D${top + 1}:${last_col}${top + n_resources}")
    for r in range(top, top + 1 + n_resources):
        ws.row_dimensions[r].hidden = True
