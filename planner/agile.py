"""Agile sheet: the sprint calendar (Excel Table "SprintTable") with each sprint's work, capacity and story points,
then the release figures, the current sprint's backlog and burndown, and the charts.

Per-sprint numbers (planned work, capacity after leave, committed, done, velocity, release burndown) are SprintTable
columns, one formula each; the sections below add the project-level forecast and the current sprint. Story points
are optional: with none entered the figures show zeros and a count of unestimated tasks. Each person's load and
utilisation by sprint are on the Resource Report, next to the weekly grids.

The sections share the sheet's columns, so their widths are set once (WIDTHS); column A holds the backlog list's
hidden helper column, so the title starts in column B.
"""
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .core import CALC_FONT, DATE, NOTE_FONT, TITLE_FONT, Col, Show, Tab, add_table, fill, font, name, ranked_table
from .leave import leave_days, leave_row, month_column, month_start, team_leave_hours
from .workload import work_between

SHEET = "Agile"
HEADER_ROW = 5          # SprintTable's header row (the Status sheet's sprint chart reads the table by position)
LEAVE_MONTHS = 3        # calendar months a sprint can touch (any sprint up to about two months long)
WD = "WorkWeek,Holidays"
GOALS = ["Charter & kickoff", "Requirements", "Design & architecture", "Foundations: DB + components",
         "Core APIs + screens", "Payments + screens", "Feature complete", "System testing", "UAT",
         "Release & hypercare", "Hypercare", "Close-out"]
# B..H: SprintTable's first columns; below it the backlog, the release figures and the burndown use the same columns
WIDTHS = {"A": 3, "B": 6, "C": 11, "D": 44, "E": 16, "F": 12, "G": 11, "H": 16}

# defined name, label, formula, format
KPIS = [
    ("agl_Total", "Total story points (scope)", "TotalPoints", "0.#"),
    ("agl_Done", "Done by the status date", 'SUMIFS(tsk_LeafPoints,tsk_DoneDate,"<="&StatusDate)', "0.#"),
    ("agl_Remaining", "Remaining", "agl_Total-agl_Done", "0.#"),
    ("agl_Velocity", "Velocity (average of the last 3 closed sprints)",
     'IFERROR(LOOKUP(2,1/(spr_Velocity<>""),spr_Velocity),"")', "0.0"),
    ("agl_SprintsLeft", "Sprints needed at that velocity",
     'IF(N(agl_Velocity)>0,ROUNDUP(agl_Remaining/agl_Velocity,0),"")', "0"),
    ("agl_LastClosed", "Sprints closed", 'COUNTIF(spr_Status,"Closed")', "0"),
    ("agl_FinishSprint", "Forecast to finish in",
     'IF(agl_SprintsLeft="","",IF(agl_LastClosed+agl_SprintsLeft<=SprintCount,'
     'INDEX(spr_Name,agl_LastClosed+agl_SprintsLeft),"after "&INDEX(spr_Name,SprintCount)))', "@"),
    ("agl_FinishDate", "Forecast finish (velocity)",
     'IF(agl_SprintsLeft="","",IF(agl_LastClosed+agl_SprintsLeft<=SprintCount,'
     'INDEX(spr_End,agl_LastClosed+agl_SprintsLeft),""))', DATE),
    ("agl_ScheduleFinish", "Forecast finish (schedule)", "ProjectFinish", DATE),
    ("agl_Current", "Current sprint", 'IF(CurrentSprint="","(outside sprint range)",CurrentSprint)', "@"),
    ("agl_CurCommitted", "Current sprint: committed points", "SUMIFS(tsk_LeafPoints,tsk_Sprint,CurrentSprint)", "0.#"),
    ("agl_CurDone", "Current sprint: done so far",
     'SUMIFS(tsk_LeafPoints,tsk_Sprint,CurrentSprint,tsk_DoneDate,"<="&StatusDate)', "0.#"),
    ("agl_Unestimated", "Tasks without story points (not milestones)",
     'COUNTIFS(tsk_IsLeaf,1,tsk_StoryPoints,"",tsk_Duration,">0")', "0"),
]
BACKLOG_ROWS = 20
BURNDOWN_DAYS = 31


def build_agile(wb, blank, n_sprints, n_resources, n_chart_sprints=12):
    """Returns (sheet, {SprintTable column: letter}) for the Status sheet's sprint chart."""
    ws = wb.create_sheet(SHEET)
    ws["B1"] = '=ProjectName&"  -  Agile: sprints, story points and burndown"'
    ws["B1"].font = TITLE_FONT
    ws["B2"] = ("Sprint calendar from Settings (Sprint1Start, SprintLength, SprintCount); add rows to the table for more "
                "sprints. Each task's work is spread evenly over its forecast working days, then summed per sprint. "
                "Enter Story Points (and optionally a Sprint) on Tasks > AGILE. Per-person load: Resource Report.")
    ws["B2"].font = NOTE_FONT
    L = _sprint_table(wb, ws, blank, n_sprints)
    top = HEADER_ROW + n_sprints + 4
    kpi_top = top
    top = _release(wb, ws, top)
    top = _backlog(ws, top)
    burndown_top = _burndown(ws, top)
    _leave_grid(wb, ws, burndown_top + BURNDOWN_DAYS + 4, n_sprints, n_resources)
    for col, width in WIDTHS.items():
        ws.column_dimensions[col].width = width
    _charts(ws, L, min(n_chart_sprints, n_sprints), kpi_top, burndown_top)
    ws.freeze_panes = "A4"
    return ws, L


def _sprint_table(wb, ws, blank, n_sprints):
    ws.cell(HEADER_ROW - 1, 2, "SPRINTS").font = font(bold=True, color="1F4E78")
    s = Tab("SprintTable")
    r = s.row
    has = f'{r("Sprint")}=""'
    cols = [
        Col("No", width=WIDTHS["B"], formula=f"ROW()-ROW({s.head('No')})", center=True),
        Col("Sprint", width=WIDTHS["C"], formula=f'IF({r("No")}<=SprintCount,SprintPrefix&{r("No")},"")'),
        Col("Goal", "in", WIDTHS["D"], values=[] if blank else GOALS),
        Col("Start", width=WIDTHS["E"], fmt=DATE, formula=f"Sprint1Start+({r('No')}-1)*SprintLength"),
        Col("End", width=WIDTHS["F"], fmt=DATE, formula=f"{r('Start')}+SprintLength-1"),
        Col("Working Days", width=WIDTHS["G"], formula=f"NETWORKDAYS.INTL({r('Start')},{r('End')},{WD})", center=True),
        Col("Tasks Starting", width=WIDTHS["H"], formula=f'IF({has},"",COUNTIF(tsk_Sprint,{r("Sprint")}))', center=True),
        Col("Planned Work", width=11, fmt="#,##0.0", formula=f'IF({has},"",{work_between(r("_FirstWD"), r("_LastWD"))})',
            note="Hours of task work falling in the sprint."),
        Col("Capacity", width=11, fmt="#,##0.0",
            note="Sum of Work resources' Max Units x HoursPerDay x working days, less the leave of the people in the "
                 "Resource Report's sprint grids (Resources sheet, leave months).",
            formula=(f'IF({has},"",SUMIFS(res_MaxUnits,res_Type,"Work")*HoursPerDay*{r("Working Days")}'
                     f'-{team_leave_hours("spr_LeaveGrid", r("No"))})')),
        Col("Utilisation", width=10, fmt="0%",
            formula=f'IF(OR({has},N({r("Capacity")})=0),"",{r("Planned Work")}/{r("Capacity")})'),
        Col("Status", width=10, center=True,
            formula=f'IF({has},"",IF(StatusDate>{r("End")},"Closed",IF(StatusDate>={r("Start")},"Active","Future")))'),
        # story points. Done = detail tasks finished in the sprint; remaining = scope minus everything done by the
        # sprint end (or the status date), so the columns are a release burndown
        Col("Points Capacity", "in", 9, "0.#", center=True,
            note="Optional: story points the team can take on in the sprint. Committed Points turn red above it."),
        Col("Committed Points", width=10, fmt="0.#", center=True,
            formula=f'IF({has},"",SUMIFS(tsk_LeafPoints,tsk_Sprint,{r("Sprint")}))',
            note="Story points of the tasks in this sprint (Tasks > Sprint)."),
        Col("Done Points", width=9, fmt="0.#", center=True,
            formula=(f'IF(OR({has},{r("Status")}="Future"),"",SUMIFS(tsk_LeafPoints,tsk_DoneDate,">="&{r("Start")},'
                     f'tsk_DoneDate,"<="&{r("End")}))'),
            note="Story points of tasks finished during the sprint (by Actual Finish)."),
        Col("Velocity", width=9, fmt="0.0", center=True,
            formula=(f'IF(OR({has},{r("Status")}<>"Closed"),"",AVERAGE(INDEX({s.col("Done Points")},MAX(1,{r("No")}-2))'
                     f':INDEX({s.col("Done Points")},{r("No")})))'),
            note="Average Done Points of this and the two previous sprints (closed sprints only)."),
        Col("Remaining Points", width=10, fmt="0.#", center=True,
            formula=(f'IF(OR({has},{r("Status")}="Future"),"",TotalPoints-SUMIFS(tsk_LeafPoints,tsk_DoneDate,'
                     f'"<="&MIN({r("End")},StatusDate)))'),
            note="Release burndown: story points not yet done at the sprint end (or the status date)."),
        Col("Forecast Remaining", width=10, fmt="0.#", center=True,
            formula=f'IF({has},"",TotalPoints-SUMIFS(tsk_LeafPoints,tsk_FcFinish,"<="&{r("End")}))',
            note="Points still open at the sprint end if tasks finish on their forecast dates."),
        Col("Ideal Remaining", width=10, fmt="0.#", center=True,
            formula=f'IF({has},"",TotalPoints*MAX(0,1-{r("No")}/SprintCount))',
            note="Straight line from the total scope to zero at the last sprint."),
        # charts plot "" as zero; #N/A leaves a gap, so the burndown chart reads this copy (hidden)
        Col("_ChartRemaining", "help", 8, formula=f'IF({r("Remaining Points")}="",NA(),{r("Remaining Points")})'),
        Col("_FirstWD", "help", 6, formula=f"NETWORKDAYS.INTL(ProjectStart,{r('Start')},{WD})",
            note="The sprint's first day as a working-day number from the project start (for the work totals)."),
        Col("_LastWD", "help", 6, formula=f"NETWORKDAYS.INTL(ProjectStart,{r('End')},{WD})"),
    ] + [Col(f"_LeaveCol{k}", "help", 6, formula=month_column(month_start(r("Start"), k)),
             note="Resources leave column of a month the sprint touches (for the leave look-ups).")
         for k in range(LEAVE_MONTHS)]
    L = add_table(ws, "SprintTable", HEADER_ROW, 2, cols, n_sprints, autofilter=False)
    for column, defined in {"Sprint": "spr_Name", "Start": "spr_Start", "End": "spr_End", "Working Days": "spr_WorkDays",
                            "Planned Work": "spr_Work", "Capacity": "spr_Capacity", "Status": "spr_Status",
                            "Committed Points": "spr_Committed", "Done Points": "spr_Done", "Velocity": "spr_Velocity",
                            "Remaining Points": "spr_Remaining", "_FirstWD": "spr_FirstWD", "_LastWD": "spr_LastWD"}.items():
        name(wb, defined, f"SprintTable[{column}]")
    for k in range(LEAVE_MONTHS):
        name(wb, f"spr_LeaveCol{k}", f"SprintTable[_LeaveCol{k}]")
    first, last = HEADER_ROW + 1, HEADER_ROW + n_sprints
    cf = ws.conditional_formatting
    start, end = L["No"], L["Ideal Remaining"]
    ws.column_dimensions.group(L["_ChartRemaining"], L[f"_LeaveCol{LEAVE_MONTHS - 1}"], hidden=True, outline_level=1)
    cf.add(f"{start}{first}:{end}{last}", FormulaRule(formula=[f'${L["Sprint"]}{first}=""'], font=Font(color="D9D9D9")))
    cf.add(f"{start}{first}:{end}{last}", FormulaRule(formula=[f'${L["Status"]}{first}="Active"'], fill=fill("FFF2CC"),
                                                      font=Font(bold=True)))
    cf.add(f"{L['Utilisation']}{first}:{L['Utilisation']}{last}",
           FormulaRule(formula=[f"AND(ISNUMBER(${L['Utilisation']}{first}),${L['Utilisation']}{first}>1)"],
                       font=Font(color="C00000", bold=True), fill=fill("FFC7CE")))
    pc, cp = L["Points Capacity"], L["Committed Points"]
    cf.add(f"{cp}{first}:{cp}{last}", FormulaRule(
        formula=[f"AND(N(${pc}{first})>0,N(${cp}{first})>${pc}{first})"], font=Font(color="C00000", bold=True),
        fill=fill("FFC7CE")))
    return L


def _release(wb, ws, top):
    """Release figures in D:E (named cells agl_*, read by Status and History)."""
    ws.cell(top - 1, 4, "RELEASE").font = font(bold=True, color="1F4E78")
    add_table(ws, "AgileKPIs", top, 4, [
        Col("Measure", width=WIDTHS["D"], values=[k[1] for k in KPIS]),
        Col("Value", width=WIDTHS["E"], values=[f"={k[2]}" for k in KPIS]),
    ], len(KPIS), autofilter=False)
    for i, (defined, _, _, fmt) in enumerate(KPIS):
        r = top + 1 + i
        ws.cell(r, 4).font = font(bold=True)
        ws.cell(r, 5).number_format = fmt
        ws.cell(r, 5).alignment = Alignment(horizontal="right")
        name(wb, defined, f"{SHEET}!$E${r}")
    return top + len(KPIS) + 3


def _backlog(ws, top):
    """The current sprint's tasks (the list's hidden _Task column is column A)."""
    ranked_table(ws, "SprintBacklog", top, 1, "tsk_CurrentSprintKey", BACKLOG_ROWS, [
        Show("UID", "tsk_UID", "num", width=WIDTHS["B"]),
        Show("Story Points", "tsk_StoryPoints", "num", width=WIDTHS["C"], fmt="0.#;;-"),
        Show("Task", "tsk_Name", width=WIDTHS["D"]), Show("Finish", "tsk_FcFinish", "date", WIDTHS["E"], DATE),
        Show("Progress", "tsk_Progress", "num", WIDTHS["F"], "0%"), Show("Status", "tsk_Status", width=WIDTHS["G"]),
        Show("Resources", "tsk_ResourceNames", width=WIDTHS["H"]),
    ], label='"CURRENT SPRINT BACKLOG ("&agl_Current&")"')
    return top + BACKLOG_ROWS + 3


def _burndown(ws, top):
    """Daily burndown of the current sprint, one row per calendar day, in D:I."""
    b = Tab("SprintBurndown")
    day = b.row("Day")
    k = f"ROW()-ROW({b.head('Day')})"
    ws.cell(top - 1, 4, "CURRENT SPRINT BURNDOWN (story points of the tasks in the sprint)").font = \
        font(bold=True, color="1F4E78")
    L = add_table(ws, "SprintBurndown", top, 4, [
        Col("Day", width=WIDTHS["D"], fmt=DATE,
            formula=f'IF(OR(CurrentSprint="",{k}>SprintLength),"",CurrentSprintStart+{k}-1)'),
        Col("Remaining", width=WIDTHS["E"], fmt="0.#", formula=(
            f'IF(OR({day}="",{day}>StatusDate),"",agl_CurCommitted-SUMIFS(tsk_LeafPoints,tsk_Sprint,CurrentSprint,'
            f'tsk_DoneDate,"<="&{day}))'),
            note="Committed points not yet done at the end of the day (blank after the status date)."),
        Col("Ideal", width=WIDTHS["F"], fmt="0.#", formula=(
            f'IF({day}="","",agl_CurCommitted*MAX(0,1-NETWORKDAYS.INTL(CurrentSprintStart,{day},WorkWeek,Holidays)/'
            f'MAX(1,NETWORKDAYS.INTL(CurrentSprintStart,CurrentSprintEnd,WorkWeek,Holidays))))'),
            note="Straight line over the sprint's working days."),
        # chart copies: text labels (blank rows stay blank) and #N/A for days not reached, which charts leave out
        Col("_ChartDay", "help", WIDTHS["G"], formula=f'IF({day}="","",TEXT({day},"dd-mmm"))'),
        Col("_ChartRemaining", "help", WIDTHS["H"], formula=f'IF({b.row("Remaining")}="",NA(),{b.row("Remaining")})'),
        Col("_ChartIdeal", "help", 11, formula=f'IF({b.row("Ideal")}="",NA(),{b.row("Ideal")})'),
    ], BURNDOWN_DAYS, autofilter=False)
    first = f'{L["_ChartDay"]}{top + 1}'
    ws.conditional_formatting.add(f"{first}:{L['_ChartIdeal']}{top + BURNDOWN_DAYS}",
                                  FormulaRule(formula=[f"ISNA({first})"], font=Font(color="FFFFFF")))
    return top


def _charts(ws, L, n, kpi_top, burndown_top):
    col = lambda key: ws[f"{L[key]}{HEADER_ROW}"].column
    first, last = HEADER_ROW, HEADER_ROW + n            # SprintTable header row and the first n sprints (charts need A1)
    cats = Reference(ws, min_col=col("Sprint"), min_row=first + 1, max_row=last)

    release = LineChart()
    release.title, release.y_axis.title = "Release burndown (story points)", "Points"
    for key, label in (("_ChartRemaining", "Remaining"), ("Forecast Remaining", "Forecast"), ("Ideal Remaining", "Ideal")):
        release.series.append(Series(Reference(ws, min_col=col(key), min_row=first + 1, max_row=last), title=label))
    release.set_categories(cats)
    release.height, release.width = 8, 16
    ws.add_chart(release, f"J{kpi_top - 1}")

    velocity = BarChart()
    velocity.title = "Committed vs done points per sprint"
    for key in ("Committed Points", "Done Points"):
        velocity.add_data(Reference(ws, min_col=col(key), min_row=first, max_row=last), titles_from_data=True)
    velocity.set_categories(cats)
    velocity.height, velocity.width = 8, 16
    ws.add_chart(velocity, f"J{kpi_top + 16}")

    sprint = LineChart()
    sprint.title = "Current sprint burndown"
    for c, label in ((8, "Remaining"), (9, "Ideal")):       # _ChartRemaining (H), _ChartIdeal (I)
        sprint.series.append(Series(Reference(ws, min_col=c, min_row=burndown_top + 1,
                                              max_row=burndown_top + BURNDOWN_DAYS), title=label))
    sprint.set_categories(Reference(ws, min_col=7, min_row=burndown_top + 1, max_row=burndown_top + BURNDOWN_DAYS))
    sprint.height, sprint.width = 8, 16
    for chart in (release, sprint):
        chart.visible_cells_only = False                    # the chart copies sit in hidden / white columns
        for s in chart.series:
            s.smooth = False
    ws.add_chart(sprint, f"J{burndown_top}")


def _leave_grid(wb, ws, top, n_sprints, n_resources):
    """Hidden grid: working days of leave per person per sprint (Resources sheet, leave months), read by the sprint
    table's Capacity and the Resource Report's sprint utilisation. Each person's Resources row is looked up once, in
    column C; each sprint's month columns come from SprintTable."""
    origin = "spr_LeaveOrigin"
    name(wb, origin, f"'{SHEET}'!$B${top}")
    ws.cell(top, 2, "Leave days by sprint (helper for Capacity and the sprint utilisation grid)").font = font(italic=True)
    i = f"(ROW()-ROW({origin}))"
    j = f"(COLUMN()-COLUMN({origin})-1)"
    resource = f"INDEX(res_Name,{i})"
    start, end = f"INDEX(spr_Start,{j})", f"INDEX(spr_End,{j})"
    months = [(month_start(start, k), f"INDEX(spr_LeaveCol{k},{j})") for k in range(LEAVE_MONTHS)]
    days = leave_days(start, end, f"INDEX(spr_LeaveRow,{i})", months)
    L = get_column_letter
    last_col = L(3 + n_sprints)
    for r in range(top + 1, top + 1 + n_resources):
        ws.cell(r, 3, f'=IFERROR(IF({resource}="",0,{leave_row(resource)}),0)').font = CALC_FONT
        for c in range(4, 4 + n_sprints):
            ws.cell(r, c, f'=IFERROR(IF(OR({resource}="",INDEX(spr_Name,{j})=""),0,{days}),0)').font = CALC_FONT
    name(wb, "spr_LeaveRow", f"'{SHEET}'!$C${top + 1}:$C${top + n_resources}")
    name(wb, "spr_LeaveGrid", f"'{SHEET}'!$D${top + 1}:${last_col}${top + n_resources}")
    for r in range(top, top + 1 + n_resources):
        ws.row_dimensions[r].hidden = True
