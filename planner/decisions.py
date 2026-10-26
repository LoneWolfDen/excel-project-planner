"""Decisions sheet: the decision log (Excel Table "DecisionLog").

A record of what was decided, by whom and why, so the reasoning survives staff changes and audits. All inputs
except the linked task's name and how long a proposal has waited.
"""
import datetime as dt

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from .core import DATE, Col, Tab, add_table, col_range, linked_task, name, title

D = dt.date
STATUSES = ["Proposed", "Approved", "Rejected", "Superseded"]
# Ref, Date, Decision, Options and reasons, Decided by, Status, Task UID, RAID ref
SAMPLE = [
    (1, D(2026, 9, 14), "Use the existing identity provider for customer log-in",
     "Options: build our own, buy a new service, reuse. Reuse is cheapest and already security-approved.",
     "Steering committee", "Approved", 12, None),
    (2, D(2026, 10, 9), "Move the reporting screens to phase 2",
     "Keeps the Build deadline (06-Nov). Reporting is not needed for go-live.", "Sponsor", "Approved", 17, 6),
    (3, D(2026, 10, 22), "Add a second developer to the core API for two weeks",
     "Recovers about 4 of the 6 days' slip. Cost in change request CR 2.", "Sponsor", "Proposed", 13, 3),
]


def build_decisions(wb, blank, n_rows=100):
    ws = wb.create_sheet("Decisions")
    title(ws, "Decision Log", "One row per decision. Link the task (UID) and RAID item it relates to.")
    ws.column_dimensions["A"].width = 3
    data = [] if blank else SAMPLE
    t = Tab("DecisionLog")
    r = t.row
    values = lambda i: [row[i] for row in data]
    letters = add_table(ws, "DecisionLog", 4, 2, [
        Col("Ref", "in", 6, center=True, values=list(range(1, n_rows + 1))),
        Col("Date", "in", 12, DATE, values=values(1)),
        Col("Decision", "in", 44, values=values(2)),
        Col("Options and Reasons", "in", 56, values=values(3)),
        Col("Decided By", "in", 18, values=values(4)),
        Col("Status", "in", 11, center=True, values=values(5)),
        Col("Waiting (days)", width=9, center=True, note="Days a proposed decision has waited, to the status date.",
            formula=(f'IF(OR({r("Decision")}="",{r("Status")}<>"Proposed",N({r("Date")})=0),"",'
                     f'MAX(0,StatusDate-{r("Date")}))')),
        Col("Task UID", "in", 8, center=True, values=values(6)),
        Col("Task", width=28, formula=linked_task(r("Task UID"))),
        Col("RAID Ref", "in", 8, center=True, values=values(7)),
    ], n_rows)
    name(wb, "dec_Status", "DecisionLog[Status]")
    first, last = 5, 4 + n_rows
    rule = DataValidation(type="list", formula1="=L_DecisionStatus")
    ws.add_data_validation(rule)
    rule.add(col_range(letters["Status"], first, last))
    ws.conditional_formatting.add(col_range(letters["Status"], first, last), FormulaRule(
        formula=[f'{letters["Status"]}{first}="Proposed"'], font=Font(color="C55A11", bold=True)))
    ws.freeze_panes = f"E{first}"
    return ws
