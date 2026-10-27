"""Instructions sheet."""
from openpyxl.styles import Alignment

from .core import CALC_FILL, HEADER_FILLS, HEADER_FONT, INPUT_FONT, font


def text(sizes):
    return [
        ("REQUIREMENTS", None),
        ("Excel version", "Excel 2019 or Microsoft 365 (Windows or Mac), or Excel for the web. No macros needed."),
        ("HOW IT WORKS", None),
        ("1. Settings", "Set project start and status date (both typed dates; Ctrl+; enters today), working week "
                        "(Mon-Sun Yes/No), hours/day, Gantt scale and sprint calendar. Add holidays on 'Holidays'. "
                        "Update the status date whenever you report progress."),
        ("2. Resources", "List people/roles with Max Units and hourly rate. They appear in the Resource dropdown on Tasks and "
                         "in sprint capacity. More than one person on a task: put the first on Tasks (Resource, Units) "
                         "and add a row on the Assignments sheet (Task UID, Resource, Units) for each of the others."),
        ("3. Tasks", "One row per task. Enter Level (1 = phase/summary, 2 = task, 3 = sub-task...), Task Name, Duration in "
                     "working days (0 = milestone), and up to four predecessors by UID with link type (FS, SS, FF, SF) and lag in "
                     "days (negative = lead). A row becomes a SUMMARY automatically when the rows beneath it have a higher "
                     "level; its dates, work, cost and progress roll up. Optional Deadline: Deadline Check shows 'Late by Nd'. "
                     "Hover over a column header for what it means."),
        ("4. Set the baseline", "When the plan is agreed: on Tasks, select the 5 'BL Source' columns (all task rows), Copy, then "
                                "select the first 'Baseline Start' cell and Paste Special > Values. Planned Start/Finish now stay "
                                "frozen (like MS Project's Baseline) and the variance columns compare against it. To re-baseline, "
                                "repeat; to clear it, delete the 5 Baseline columns' contents."),
        ("5. Track progress", "Enter Actual Start, Actual Finish and % Complete. Forecast Start/Finish then use actuals and push "
                              "every dependent task. Start / Finish / Duration / Work / Cost Var and Status flag slippage. "
                              "Optionally enter Actual Cost (cost to date) on detail tasks; blank = Cost x Progress."),
        ("Critical path", "Late Start / Late Finish come from a backward pass from the forecast project finish (and any "
                          "deadline). Total Slack = how far a task can slip before it delays the project or a deadline; Free "
                          "Slack = before it delays a successor. Critical = slack <= Settings > CriticalSlackDays on an "
                          "unfinished task. A missed deadline gives negative slack to the task and the chain feeding it."),
        ("6. Gantt", "Read-only view: dark outline = planned dates, green = done, blue = remaining on track, red = remaining "
                     "on the critical path, orange = remaining but late vs plan, dark grey = summary, purple = milestone, red cell = deadline, red border = status date. "
                     "Change GanttStart / GanttPeriod (1 = days, 7 = weeks) on Settings to scroll or zoom."),
        ("7. Sprints", "Sprint calendar from Settings. Each task is tagged with the sprint its forecast start falls in. Work "
                       "hours are prorated across sprints per resource, compared against capacity (Max Units x Hours/day x "
                       "working days); over-allocations turn red. Optional agile columns: Story Points and Planned "
                       "Sprint on Tasks (Planned Sprint wins over the automatic sprint), Points Capacity on Sprints; "
                       "the table then shows committed and done points, velocity and the release burndown."),
        ("8. Dashboard", "KPIs (including overall RAG, SPI, CPI, missed deadlines, link problems), status counts, upcoming "
                         "milestones and sprint load vs capacity."),
        ("9. Status Report", "For steering meetings: overall RAG and RAG by area (schedule, SPI, CPI, deadlines, critical "
                             "path, plan quality), calculated from the thresholds on Settings (RAG...). Type a PM RAG to "
                             "override one and say why in Comment. Earned value (BAC, PV, EV, AC, SPI, CPI, EAC, VAC, TCPI) "
                             "and a commentary box for the period."),
        ("10. Schedule Report", "Phase summary with RAG, all milestones, late tasks, tasks starting and due in the "
                                "look-ahead window (Settings > LookaheadDays) and open critical tasks. Lists update with "
                                "every edit; drag a list's bottom corner down to show more rows."),
        ("11. Resource Report", "Per resource: work, done, remaining, cost, actual cost, late and critical tasks, peak "
                                "week and weeks over capacity. Weekly hours and utilisation grids by week from "
                                "Settings > ResourceWeekStart (red = over capacity), and a team hours vs capacity chart."),
        ("12. Agile Report", "Story points: scope, done, remaining, velocity (last 3 closed sprints) and the sprint it "
                             "forecasts finishing in; the current sprint's backlog and daily burndown; release burndown "
                             "and committed vs done charts."),
        ("Leave", "Leave sheet: one row per person, one column per month (Jan-26 ... Dec-28; add a month by typing e.g. "
                  "Jan-29 in the next header cell). Type the days off: '16, 17(0.5)' under Nov-26 = all of 16 November "
                  "and half of the 17th. Weekends and holidays are ignored. Leave lowers capacity on the Resource "
                  "Report and Sprints and can make a task Over-allocated; it does not move task dates."),
        ("13. RAID", "Risks, assumptions, issues and dependencies: owner, due date, probability and impact (1-5). Score "
                     "= probability x impact (issues count as probability 5); Rating from Settings > RAIDHighScore / "
                     "RAIDMediumScore; Overdue when not closed past its due date. Link a task by UID. The Status Report "
                     "lists the top open items and rates RAID amber when any open item is High."),
        ("14. Decisions and Changes", "Decision log (what, options and reasons, who, status) and change request log "
                                      "(schedule and cost impact, status, who decided). The Status Report totals "
                                      "approved and pending changes. Approving a change doesn't edit the plan: update "
                                      "the tasks, and re-baseline if the sponsor agrees."),
        ("15. History", "On each reporting date, copy the SnapshotNow row on History and Paste Special > Values onto the "
                        "first empty row of StatusHistory. Trend charts: forecast vs planned finish, earned value "
                        "S-curve, SPI / CPI and a story-point burn-up that shows scope changes."),
        ("Earned value", "Set a baseline first: without one, Budget = current Cost and PV follows the current plan, so "
                         "SPI and CPI only show progress against the latest plan. PV = budget of work planned before the "
                         "status date; EV = Budget x Progress; AC = Actual Cost. SPI = EV/PV, CPI = EV/AC, "
                         "EAC = AC + (BAC - EV)/CPI."),
        ("CELL COLOURS", None),
        ("Blue text", "Input - type here."),
        ("Black text on grey", "Formula - do not overwrite. Summary rows grey-out inputs that are ignored."),
        ("Dark blue header", "Input column.  Dark grey header = calculated.  Light grey header = helper column (hidden, grouped)."),
        ("RULES (same as MS Project)", None),
        ("ID and UID", "ID = position of the task (renumbers when rows move, like MS Project). UID = permanent number that "
                       "predecessors refer to, so links survive inserting, deleting and moving rows. A duplicate or missing UID "
                       "turns red."),
        ("Inserting tasks", "Tasks is an Excel Table: right-click a row > Insert > Table Rows Above (or type in the row under "
                            "the table). Every formula fills in automatically. Set the new task's UID to Settings > NextFreeUID "
                            "(also shown at the top of Tasks)."),
        ("Outline levels", "Each row's Level may be at most 1 higher than the row above. A '?' in WBS (red) means a level jump."),
        ("Predecessors", "Enter the predecessor's UID. A predecessor must be a detail task listed ABOVE its successor "
                         "(order tasks roughly by date). Links that break this, or point to an empty row or a missing UID, are "
                         "ignored and explained in the Link Check column (red). Summary rows ignore Duration / Predecessors / "
                         "Actuals - they roll up. The rule keeps the schedule and the critical path consistent and rules out "
                         "circular references."),
        ("Mode", "Auto = scheduled from predecessors, project start and the constraint. Manual = pinned to the "
                 "Constraint Date."),
        ("Constraints", "Pick an MS Project constraint and enter its Constraint Date (a date alone = Start No Earlier Than). "
                        "Start / Finish No Earlier Than push the task later; Start / Finish No Later Than and Must Start / "
                        "Finish On win over predecessors (as in MS Project), so a conflict shows as negative slack. As Late "
                        "As Possible moves the task later by its free slack: it finishes just in time for its successors "
                        "and moves nothing else (its planned dates move by the same days)."),
        ("Rescheduling", "Settings > RescheduleAfterStatus = Yes moves incomplete work to after the status date (like MS "
                         "Project 'Update Project'), so late tasks push successors."),
        ("Capacity", f"Tables grow when you add rows: Tasks starts with {sizes.tasks} rows, Resources {sizes.resources}, "
                     f"Assignments {sizes.assignments}, "
                     f"Sprints {sizes.sprints}, Holidays 100. The Gantt and MSP_Tasks sheets show the first {sizes.tasks} "
                     f"tasks (the Gantt says so when there are more); the Gantt timeline has {sizes.gantt_columns} columns; "
                     f"the per-resource grids on Sprints and Resource Report show the first {sizes.grid_resources} "
                     f"resources and the Resource Report {sizes.weeks} weeks. To change these sizes, rebuild with "
                     "scripts/build_planner.py --tasks / --resources / --assignments / --grid-resources / --sprints / "
                     "--gantt-columns / --weeks (see README)."),
        ("Speed", "An edit recalculates in well under a second for typical plans (about 0.05 s for 200 tasks). Each change "
                  "to the build is timed in CI; see docs/ARCHITECTURE.md in the repository. Keep the status date typed "
                  "(not =TODAY()), which would recalculate everything after every edit."),
        ("MOVING TO MICROSOFT PROJECT", None),
        ("Best: MS Project XML", "Save this workbook, then convert it with the browser page docs/export.html in the "
                                 "repository (nothing to install; works on a phone; the file stays on your device) or "
                                 "python scripts/export_mspdi.py. In MS Project: File > Open > the .xml > 'Open as a new "
                                 "project'. It carries the calendar and holidays, links by UID with lags, constraints, "
                                 "deadlines, actuals, the baseline, resources, assignments, Sprint (Text1) and Story "
                                 "Points (Number1), so MS Project schedules the plan the same way."),
        ("From MS Project", "MS Project > File > Save As > XML, then choose the .xml on docs/export.html (or run "
                            "python scripts/import_mspdi.py): you get a new planner workbook with the project's tasks, "
                            "links, calendar, holidays, resources, assignments, baseline and leave. Anything the planner "
                            "can't hold is listed and noted on the task."),
        ("Or: Import Wizard", "The MSP_Tasks / MSP_Resources sheets below. Carries less (no calendar or holidays)."),
        ("One-time setting", "MS Project > File > Options > Trust Center > Trust Center Settings > Legacy Formats > 'Allow "
                             "loading files with legacy or non-default file formats'."),
        ("Import", "In MS Project: File > Open > choose this .xlsx > Import Wizard > New map > 'As a new project' > tick Tasks "
                   "(and Resources) > Tasks source worksheet = MSP_Tasks, Resources = MSP_Resources. Column headers already "
                   "use MS Project field names, so the wizard maps them automatically. Map 'Predecessors' (by ID) for a new "
                   "project; 'Unique ID' and 'Unique ID Predecessors' are for merging into an existing one."),
        ("Tip: Start/Finish", "For auto-scheduled tasks map Duration + Predecessors (+ Constraint Type/Date) and set Start/Finish "
                              "to '(not mapped)' - importing Start/Finish makes MS Project add fixed 'Start No Earlier Than' "
                              "constraints. Map Baseline Start/Finish to keep the plan, and Actual Start/Finish/% Complete for "
                              "progress."),
        ("Rows", "Blank MSP_Tasks rows import as blank tasks - delete them in MS Project, or filter the table to non-blank "
                 "names and copy those rows to a new sheet before importing."),
        ("From MS Project to Excel", "MS Project > File > Save As > Excel Workbook, then paste Name / Outline Level / Duration "
                                     "/ Predecessors etc. into the Tasks input columns (split '3FS+2d' into UID, Type, Lag)."),
        ("NAMES AND TABLES (Formulas > Name Manager)", None),
        ("Tables", "Tasks, ResourceTable, AssignmentTable, SprintTable, HolidayTable, GanttView, MSPTasks, MSPResources, KPIs, StatusSummary, "
                   "NextMilestones, OverallStatus, StatusRAG, EarnedValue, StatusNotes, PhaseSummary, MilestoneReport, "
                   "LateTasks, StartingSoon, DueSoon, CriticalTasks, ResourceSummary, AgileKPIs, SprintBacklog, SprintBurndown, "
                   "RAIDLog, DecisionLog, ChangeLog, TopRAID, ChangeControl, SnapshotNow, StatusHistory and the List* "
                   "dropdown tables. Inside a table, formulas read [@Column] (this row)."),
        ("Settings", "ProjectStart, StatusDate, StatusWorkday, WorkWeek, Holidays, HoursPerDay, DefaultUnits, "
                     "RescheduleAfterStatus, GanttStart, GanttPeriod, Sprint1Start, SprintLength, SprintCount, SprintPrefix, "
                     "SprintLastEnd, CriticalSlackDays, ProjectFinish, PlannedFinish, FinishVarianceDays, LookaheadDays, "
                     "RAGAmberSlipDays, RAGRedSlipDays, RAGAmberIndex, RAGRedIndex, ResourceWeekStart, CurrentSprint, TotalPoints, "
                     "RAIDHighScore, RAIDMediumScore, NextFreeUID. Report values: ProjectRAG, "
                     "evm_BAC, evm_PV, evm_EV, evm_AC, evm_SPI, evm_CPI, evm_EAC ..."),
        ("Task columns", "tsk_ID, tsk_UID, tsk_WBS, tsk_Name, tsk_Duration, tsk_PlanStart, tsk_FcFinish, tsk_Status, tsk_Resource, "
                         "tsk_Work, tsk_Cost ... each is Tasks[Column], so it grows with the table. Other sheets read Tasks "
                         "only through these names."),
        ("Resources / Sprints", "res_Name, res_MaxUnits, res_Rate, ...  spr_Name, spr_Start, spr_End, spr_WorkDays, spr_Capacity, "
                                "spr_Work"),
        ("Key formulas", "Calculated start = WORKDAY.INTL(MIN(MAX(ProjectStart, [@_ConstraintLow], predecessor drivers), "
                         "[@_ConstraintHigh]) - 1, 1, WorkWeek, Holidays).  Predecessor driver (FS) = WORKDAY.INTL(OFFSET(Tasks[[#Headers],[_LeafCalcFinish]], "
                         "[@_Pred1Pos], 0), 1 + Lag, WorkWeek, Holidays).  Planned Start = Baseline Start if set, "
                         "else calculated.  Summary start = MIN(INDEX(Tasks[_LeafCalcStart], [@_Row]+1):INDEX(Tasks[_LeafCalcStart], "
                         "[@_Row]+[@_Descendants])), where _Descendants = COUNTIF(Tasks[WBS], [@WBS] & \".*\")."),
    ]


def build_instructions(wb, sizes):
    ws = wb.create_sheet("Instructions")
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 110
    ws["B1"] = "Project Planner - MS Project-compatible Excel template"
    ws["B1"].font = font(size=18, bold=True, color="1F4E78")
    r = 3
    for label, body in text(sizes):
        if body is None:
            r += 1
            ws.cell(r, 2, label).font = HEADER_FONT
            ws.cell(r, 2).fill = HEADER_FILLS["in"]
            ws.cell(r, 3).fill = HEADER_FILLS["in"]
        else:
            ws.cell(r, 2, label).font = font(bold=True)
            ws.cell(r, 2).alignment = Alignment(vertical="top")
            c = ws.cell(r, 3, body)
            c.font = INPUT_FONT if label == "Blue text" else font()
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if label == "Black text on grey":
                c.fill = CALC_FILL
        r += 1
    return ws
