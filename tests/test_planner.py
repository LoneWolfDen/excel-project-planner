"""Tests for the Excel project planner.

Builds the workbook with scripts/build_planner.py, recalculates it with LibreOffice
(headless) and checks the scheduling results against hand-worked dates.

Run:  python -m unittest discover -s tests -v
Needs openpyxl and LibreOffice (`soffice` on PATH); recalculation tests are skipped without it.
"""
import datetime as dt
import io
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / "scripts" / "build_planner.py"
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
STATUS_DATE = dt.date(2026, 10, 25)
ERRORS = ("#NAME?", "#VALUE!", "#REF!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:")
FIRST = 6  # first task row on Tasks

# Tasks headers used by the tests (columns are found by header, not letter)
H = {
    "ID": "ID", "UID": "UID", "WBS": "WBS", "Level": "Level", "Name": "Task Name", "Dur": "Duration",
    "P1": "Predecessors", "Res": "Resources", "AS": "Actual Start", "AF": "Actual Finish", "Pct": "% Complete",
    "PS": "Planned Start", "PF": "Planned Finish", "FS": "Start", "FF": "Finish",
    "SVar": "Start Var", "Var": "Finish Var", "DLChk": "Deadline Check", "Link": "Link Check", "Prog": "Progress",
    "Status": "Status", "Work": "Work", "Over-allocated": "Over-allocated", "Sprint": "Sprint", "Pred": "Predecessors by ID",
    "DVar": "Duration Var", "WVar": "Work Var",
}
SOURCE = ["BL Source Start", "BL Source Finish", "BL Source Duration", "BL Source Work"]
BASELINE = ["Baseline Start", "Baseline Finish", "Baseline Duration", "Baseline Work"]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(ROOT))
import perf_check  # noqa: E402
from planner import log, mspdi, settings  # noqa: E402


def build(path, *flags):
    subprocess.run([sys.executable, str(BUILD), str(path), *flags], check=True, capture_output=True)


def columns(ws):
    return {ws.cell(5, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(5, c).value}


def named_cell(wb, name):
    sheet, ref = wb.defined_names[name].attr_text.split("!")
    return wb[sheet.strip("'")][ref.replace("$", "")]


def recalc(src, outdir, edit=None, status_date=True):
    """Optionally edit inputs (edit(wb) with formulas loaded), then let LibreOffice recalculate a copy."""
    wb = openpyxl.load_workbook(src)
    if status_date:
        named_cell(wb, "StatusDateInput").value = STATUS_DATE   # deterministic results
    if edit:
        edit(wb)
    stage = Path(tempfile.mkdtemp(dir=outdir))
    staged = stage / src.name
    wb.save(staged)
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(stage / "out"), str(staged)],
                   check=True, capture_output=True, timeout=900)
    return openpyxl.load_workbook(stage / "out" / src.name, data_only=True)


def d(value):
    return value.date() if isinstance(value, dt.datetime) else value


class Tasks:
    """Read a recalculated Tasks sheet by task ID and header key."""

    def __init__(self, wb):
        self.ws = wb["Tasks"]
        self.col = columns(self.ws)

    def row_of(self, task_id):
        c = self.col[H["ID"]]
        for r in range(FIRST, self.ws.max_row + 1):
            if self.ws.cell(r, c).value == task_id:
                return r
        raise KeyError(task_id)

    def get(self, task_id, key):
        return d(self.ws.cell(self.row_of(task_id), self.col[H.get(key, key)]).value)


def set_input(wb, task_id, key, value):
    """Set an input cell on the unrecalculated workbook (sample task n is on row n + 5)."""
    ws = wb["Tasks"]
    ws.cell(task_id + FIRST - 1, columns(ws)[H.get(key, key)]).value = value


def errors_in(wb):
    """Error values anywhere, except #N/A in _Chart... table columns (it is how a chart leaves a point out)."""
    chart_cells = set()
    for ws in wb.worksheets:
        for table in ws.tables.values():
            cells = list(ws[table.ref])
            for i, head in enumerate(cells[0]):
                if str(head.value).startswith("_Chart"):
                    chart_cells.update((ws.title, row[i].coordinate) for row in cells[1:])
    return [f"{ws.title}!{c.coordinate}={c.value}" for ws in wb.worksheets for r in ws.iter_rows() for c in r
            if isinstance(c.value, str) and c.value.startswith(ERRORS)
            and not (c.value == "#N/A" and (ws.title, c.coordinate) in chart_cells)]


class Tmp(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.dir = Path(cls.tmp.name)
        cls.path = cls.dir / "Project_Planner.xlsx"
        build(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()


class StructureTest(Tmp):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.wb = openpyxl.load_workbook(cls.path)

    def test_sheets(self):
        self.assertEqual(self.wb.sheetnames, ["Start", "Status", "Roadmap", "Schedule Report", "Resource Report",
                                              "Agile", "Log", "History", "Settings", "Tasks", "Gantt", "Resources", "Cost"])

    def test_named_ranges(self):
        names = set(self.wb.defined_names)
        for n in ("ProjectStart", "ProjectStartDate", "StatusDate", "WorkWeek", "Holidays", "HoursPerDay", "GanttStart",
                  "GanttPeriod", "Sprint1Start", "SprintLength", "NextFreeUID", "tsk_ID", "tsk_UID", "tsk_PlanStart",
                  "tsk_CalcStart", "tsk_FcFinish", "tsk_Resource", "tsk_Deadline", "tsk_DeadlineCheck", "tsk_LinkCheck",
                  "tsk_BaseStart", "tsk_BaseWork", "res_Name", "res_Allocation", "res_Focus", "spr_Name", "spr_Start",
                  "cost_Rate", "cost_RateName", "evm_CPI", "cap_Days", "cap_Points", "ShownSprint", "sch_SPI",
                  "gantt_Origin", "load_Grid"):
            self.assertIn(n, names)
        # column names are structured references, so they grow with the table
        self.assertEqual(self.wb.defined_names["tsk_Duration"].attr_text, "Tasks[Duration]")
        self.assertEqual(self.wb.defined_names["Holidays"].attr_text, "HolidayTable[Date]")

    def test_tables_have_calculated_columns(self):
        tasks = self.wb["Tasks"].tables["Tasks"]
        calculated = {c.name for c in tasks.tableColumns if c.calculatedColumnFormula is not None}
        self.assertIn("Finish", calculated)
        self.assertNotIn("Duration", calculated)                 # inputs are not calculated columns
        for sheet, table in (("Resources", "ResourceTable"), ("Agile", "SprintTable"), ("Gantt", "GanttView"),
                             ("Settings", "HolidayTable"), ("Settings", "ListRAG")):
            self.assertIn(table, self.wb[sheet].tables)

    def test_no_a1_references_in_formulas(self):
        report = perf_check.static_report(self.path)
        self.assertEqual(report["a1_references"], [])

    def test_formulas_within_excel_limits(self):
        report = perf_check.static_report(self.path)
        self.assertEqual(report["excel_limit_problems"], [])
        self.assertEqual(perf_check.excel_limit_problems('=IF(A,"(",B'), ["unbalanced parentheses, quotes or brackets"])
        self.assertEqual(perf_check.excel_limit_problems('=IF(Tasks[[#This Row],[x)]]=1,"a)",2)'), [])

    def test_a1_lint_catches_cell_references(self):
        names = {"ProjectStart"}
        self.assertEqual(perf_check.a1_references("=SUM(A1:B2)+Settings!$C$5+ProjectStart", names), ["A1:B2", "Settings!$C$5"])
        self.assertEqual(perf_check.a1_references("=Tasks[[#This Row],[Duration]]*2", names), [])

    def test_volatile_functions_only_where_allowed(self):
        # OFFSET only in Tasks, for the few columns that read another row's calculated value; no TODAY, INDIRECT ...
        volatile = perf_check.static_report(self.path)["volatile"]
        self.assertEqual(volatile, {"Tasks": {"OFFSET": 2200}})
        ws = self.wb["Tasks"]
        offset_cols = {ws.cell(5, c).value for c in range(1, ws.max_column + 1)
                       if "OFFSET(" in str(ws.cell(FIRST, c).value)}
        drivers = {f"_Pred{k}{x}Driver" for k in (1, 2, 3, 4) for x in ("Plan", "Fcst")}
        self.assertEqual(offset_cols, {"WBS", *drivers, "_SuccessorFinishLimit", "Free Slack"})

    def test_titles_not_in_hidden_columns(self):
        # report lists hide their _Task helper column; it must never be the column a sheet's title is in
        for ws in self.wb.worksheets:
            for r in (1, 2):
                for c in ws[r]:
                    if c.value is not None and ws.column_dimensions[c.column_letter].hidden:
                        self.fail(f"{ws.title}!{c.coordinate} is in a hidden column")

    def test_gantt_timeline_has_no_formulas(self):
        ws = self.wb["Gantt"]
        origin = named_cell(self.wb, "gantt_Origin")
        for row in ws.iter_rows(min_row=origin.row + 1, max_row=origin.row + 30, min_col=origin.column + 1,
                                max_col=origin.column + 150):
            for c in row:
                self.assertIsNone(c.value, c.coordinate)

    def test_scheduling_formulas_use_names(self):
        ws = self.wb["Tasks"]
        drivers = [ws.cell(7, c).value for c in range(1, ws.max_column + 1)
                   if isinstance(ws.cell(7, c).value, str) and "CHOOSE(" in ws.cell(7, c).value]
        self.assertTrue(drivers)
        for token in ("Tasks[[#Headers],[_LeafCalcFinish]]", "Tasks[[#This Row],[_Pred1Pos]]", "WorkWeek", "Holidays"):
            self.assertIn(token, drivers[0])
        # every row of a calculated column carries the same formula text: nothing points at a particular row
        for name in ("Start", "_Pred1PlanDriver", "WBS"):
            col = columns(ws)[name]
            self.assertEqual(len({ws.cell(r, col).value for r in range(FIRST, FIRST + 200)}), 1, name)

    def test_uids_prefilled(self):
        ws = self.wb["Tasks"]
        c = columns(ws)["UID"]
        self.assertEqual([ws.cell(r, c).value for r in range(FIRST, FIRST + 200)], list(range(1, 201)))

    def test_views_read_tasks_by_name_not_row(self):
        # Gantt and the reports use INDEX(tsk_..., n), so rows inserted on Tasks still show up
        for sheet in ("Gantt", "Status", "Schedule Report", "Resource Report", "Agile", "Log", "History"):
            for row in self.wb[sheet].iter_rows(max_row=120, max_col=25):
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        self.assertNotIn("Tasks!", c.value, f"{sheet}!{c.coordinate}")

    def test_functions_need_at_most_excel_2019(self):
        # Excel 2019 / Microsoft 365 is the stated requirement. Only MINIFS/MAXIFS from Excel 2019 are used (stored
        # with the _xlfn. prefix); Microsoft 365-only functions are banned, and LibreOffice couldn't check them anyway.
        allowed = {"_xlfn.MINIFS(", "_xlfn.MAXIFS("}
        banned = ("XLOOKUP", "XMATCH", "FILTER(", "SEQUENCE(", "UNIQUE(", "SORT(", "LET(", "LAMBDA(")
        for ws in self.wb.worksheets:
            for r in ws.iter_rows(max_row=12):
                for c in r:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        for b in banned:
                            self.assertNotIn(b, c.value, f"{ws.title}!{c.coordinate}")
                        for fn in __import__("re").findall(r"_xlfn\.[A-Z.]+\(", c.value):
                            self.assertIn(fn, allowed, f"{ws.title}!{c.coordinate}")


class CommittedWorkbooksTest(unittest.TestCase):
    """The .xlsx files in the repo must match what the build script produces (run the build commands in README)."""

    def compare(self, committed, *flags):
        with tempfile.TemporaryDirectory() as tmp:
            fresh = Path(tmp) / committed.name
            build(fresh, *flags)
            a, b = openpyxl.load_workbook(fresh), openpyxl.load_workbook(committed)
            self.assertEqual(a.sheetnames, b.sheetnames)
            diffs = []
            for ws in a.worksheets:
                other = b[ws.title]
                for row in ws.iter_rows():
                    for c in row:
                        x, y = d(c.value), d(other[c.coordinate].value)
                        if x != y and not (x in (None, "") and y in (None, "")):
                            diffs.append(f"{ws.title}!{c.coordinate}: build={x!r} committed={y!r}")
            self.assertEqual(diffs[:10], [], f"{committed.name} is out of date ({len(diffs)} cells differ)")
            # LibreOffice's save drops calculated-column formulas; the build restores them so Excel auto-fills rows
            with zipfile.ZipFile(committed) as z:
                tables = [z.read(n).decode() for n in z.namelist() if n.startswith("xl/tables/")]
            tasks = next(t for t in tables if 'displayName="Tasks"' in t)
            self.assertIn("<calculatedColumnFormula>", tasks)
            # B-1001: LibreOffice's save drops these too; Excel must recalculate automatically and on open
            with zipfile.ZipFile(committed) as z:
                calc = re.search(r"<calcPr[^>]*>", z.read("xl/workbook.xml").decode()).group(0)
            self.assertIn('calcMode="auto"', calc)
            self.assertIn('fullCalcOnLoad="1"', calc)

    def test_sample_workbook_up_to_date(self):
        self.compare(ROOT / "Project_Planner.xlsx", "--baselined")

    def test_blank_workbook_up_to_date(self):
        self.compare(ROOT / "Project_Planner_Blank.xlsx", "--blank")


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class ScheduleTest(Tmp):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.wb = recalc(cls.path, cls.dir)
        cls.t = Tasks(cls.wb)

    def val(self, task_id, key):
        return self.t.get(task_id, key)

    def test_no_formula_errors(self):
        self.assertEqual(errors_in(self.wb), [])

    def test_wbs_and_outline(self):
        self.assertEqual(self.val(1, "WBS"), "1")
        self.assertEqual(self.val(13, "WBS"), "3.1.2")
        self.assertEqual(self.val(30, "WBS"), "5.3")

    def test_finish_to_start(self):
        self.assertEqual(self.val(2, "PS"), dt.date(2026, 9, 7))
        self.assertEqual(self.val(2, "PF"), dt.date(2026, 9, 9))
        self.assertEqual(self.val(3, "PS"), dt.date(2026, 9, 10))

    def test_start_to_start_with_lag(self):
        self.assertEqual(self.val(7, "PS"), dt.date(2026, 9, 17))   # task 6 starts Mon 14-Sep, SS+3

    def test_finish_to_finish_with_lag(self):
        self.assertEqual(self.val(18, "PF"), dt.date(2026, 10, 29))  # task 17 finishes 28-Oct, FF+1
        self.assertEqual(self.val(18, "PS"), dt.date(2026, 10, 27))

    def test_two_predecessors_take_latest(self):
        self.assertEqual(self.val(9, "PS"), dt.date(2026, 10, 1))   # max(task 7, task 8 finish)

    def test_weekend_skipped(self):
        self.assertEqual(self.val(6, "PF"), dt.date(2026, 9, 23))   # 8 days from Mon 14-Sep

    def test_holidays_skipped(self):
        # Hypercare: 10 days from Wed 16-Dec skipping 24, 25, 31 Dec and 1 Jan
        self.assertEqual(self.val(26, "PF"), dt.date(2027, 1, 4))

    def test_start_no_earlier_than(self):
        self.assertEqual(self.val(22, "PS"), dt.date(2026, 12, 7))

    def test_manual_task(self):
        self.assertEqual(self.val(28, "PS"), dt.date(2026, 11, 16))

    def test_summary_rollup(self):
        self.assertEqual(self.val(10, "PS"), dt.date(2026, 10, 2))
        self.assertEqual(self.val(10, "PF"), dt.date(2026, 10, 30))
        self.assertEqual(self.val(11, "Work"), 168)   # 40 + 80 + 48 hrs
        self.assertAlmostEqual(self.val(1, "Prog"), 1.0)

    def test_actuals_drive_forecast(self):
        self.assertEqual(self.val(2, "FF"), dt.date(2026, 9, 10))   # actual finish used
        self.assertEqual(self.val(13, "FF"), dt.date(2026, 10, 30))  # in progress, 50% of 10 days left after status
        self.assertEqual(self.val(14, "FS"), dt.date(2026, 11, 2))   # successor pushed
        self.assertEqual(self.val(13, "Status"), "Behind")
        self.assertEqual(self.val(14, "Status"), "Late Start")
        self.assertEqual(self.val(10, "Var"), 6)
        self.assertEqual(self.val(14, "SVar"), 6)

    def test_predecessor_text(self):
        self.assertEqual(self.val(7, "Pred"), "6SS+3d")
        self.assertEqual(self.val(17, "Pred"), "16,13SS+2d")
        self.assertEqual(self.val(18, "Pred"), "17FF+1d")

    def test_deadlines(self):
        self.assertEqual(self.val(19, "DLChk"), "Late by 1d")   # forecast Mon 09-Nov vs deadline Fri 06-Nov
        self.assertEqual(self.val(25, "DLChk"), "OK")
        self.assertIsNone(self.val(2, "DLChk"))
        rag = {r["Area"]: r for r in table_rows(self.wb, "Status", "StatusRAG")}
        self.assertEqual(rag["Deadlines"]["Value"], 1)
        gantt = self.wb["Gantt"]
        col = {c.value: c.column for c in gantt[5] if c.value}
        self.assertEqual(d(gantt.cell(self.t.row_of(19), col["_Deadline"]).value), dt.date(2026, 11, 6))

    def test_critical_path(self):
        # hand-worked: UAT's Start No Earlier Than (07-Dec) drives the finish; everything after it is critical
        for task in (22, 24, 25, 26, 29, 30):
            self.assertEqual(self.val(task, "Total Slack"), 0, task)
            self.assertEqual(self.val(task, "Critical"), "Yes", task)
        self.assertEqual(self.val(21, "Late Finish"), dt.date(2026, 12, 4))   # day before UAT's latest start
        self.assertEqual(self.val(21, "Total Slack"), 11)
        self.assertEqual(self.val(23, "Total Slack"), 18)                      # training, feeding the readiness review
        self.assertEqual(self.val(28, "Total Slack"), 33)                      # no successors: slack to project finish
        self.assertIsNone(self.val(28, "Critical"))
        # Build complete's 06-Nov deadline is missed by a day: -1 slack back up the chain feeding it
        for task in (13, 14, 19):
            self.assertEqual(self.val(task, "Total Slack"), -1, task)
        self.assertEqual(self.val(17, "Late Finish"), dt.date(2026, 11, 5))   # FF+1 into accessibility review
        self.assertEqual(self.val(18, "Total Slack"), 0)
        # free slack: solution architecture can slip 7 days before the sign-off milestone needs it
        self.assertEqual(self.val(7, "Total Slack"), 9)
        self.assertEqual(self.val(7, "Free Slack"), 7)
        self.assertEqual(self.val(10, "Critical"), "Yes")                       # summary: a child is critical
        self.assertEqual(self.dashboard("Open critical tasks"), 11)

    def test_issues_in_sample(self):
        # no bad links; the issues are the sample's over-allocations (3) and its missed deadline
        self.assertEqual(self.dashboard("Tasks with issues (see Tasks > Issues)"), 4)

    def test_baseline_not_set_by_default(self):
        self.assertEqual(self.dashboard("Baseline"), "Not set (planned = calculated)")
        self.assertIsNone(self.val(2, "DVar"))

    def test_sprint_and_cost(self):
        self.assertEqual(self.val(2, "Sprint"), "Sprint 1")
        cost = {r["ID"]: r["Cost"] for r in table_rows(self.wb, "Cost", "TaskCost") if r["ID"] not in (None, "")}
        self.assertEqual(cost[2], 3 * 760)                                  # 3 days at Priya's day rate
        self.assertEqual(self.dashboard("Current sprint"), "Sprint 4")

    def dashboard(self, label):
        return key_figures(self.wb)[label]


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class BaselineTest(Tmp):
    """Frozen baseline: copy the source block to the baseline block (values), then change the plan."""

    def test_baseline_freezes_plan_and_measures_variance(self):
        first = recalc(self.path, self.dir)
        ws1 = first["Tasks"]
        col = columns(ws1)
        snapshot = {(r, i): ws1.cell(r, col[s]).value for r in range(FIRST, FIRST + 30) for i, s in enumerate(SOURCE)}

        def set_baseline_then_slip(wb):
            ws = wb["Tasks"]
            for (r, i), v in snapshot.items():
                ws.cell(r, col[BASELINE[i]]).value = v      # what Paste Special > Values does
            set_input(wb, 2, "Dur", 5)                      # charter now takes 5 days instead of 3

        wb = recalc(self.path, self.dir, set_baseline_then_slip)
        self.assertEqual(errors_in(wb), [])
        t = Tasks(wb)
        # planned dates stay on the baseline (baseline taken from the forecast, which uses actuals)
        self.assertEqual(t.get(3, "PS"), dt.date(2026, 9, 11))
        self.assertEqual(t.get(2, "PF"), dt.date(2026, 9, 10))
        # variances against the baseline
        self.assertEqual(t.get(2, "DVar"), 2)
        self.assertEqual(t.get(2, "WVar"), 16)
        cost = {r["ID"]: r for r in table_rows(wb, "Cost", "TaskCost") if r["ID"] not in (None, "")}
        self.assertEqual(cost[2]["Cost Var"], 16 * 95)       # 2 more days of Priya than the baseline
        self.assertEqual(t.get(1, "DVar"), 0)       # summary duration from dates; actual finish unchanged
        self.assertEqual(t.get(3, "Var"), 0)        # completed tasks keep their actual dates
        self.assertEqual(key_figures(wb)["Baseline"], "Set on 30 tasks")

    def test_unstarted_task_slip_shows_against_baseline(self):
        first = recalc(self.path, self.dir)
        ws1 = first["Tasks"]
        col = columns(ws1)
        snapshot = {(r, i): ws1.cell(r, col[s]).value for r in range(FIRST, FIRST + 30) for i, s in enumerate(SOURCE)}

        def edit(wb):
            ws = wb["Tasks"]
            for (r, i), v in snapshot.items():
                ws.cell(r, col[BASELINE[i]]).value = v
            set_input(wb, 21, "Dur", 10)   # system testing grows from 8 to 10 days

        t = Tasks(recalc(self.path, self.dir, edit))
        self.assertEqual(t.get(21, "PF"), dt.date(2026, 11, 19))   # frozen
        self.assertEqual(t.get(21, "FF"), dt.date(2026, 11, 23))
        self.assertEqual(t.get(21, "Var"), 2)
        self.assertEqual(t.get(22, "Var"), 0)                       # UAT's Start No Earlier Than absorbs the slip


def table_rows(wb, sheet, table):
    """Rows of an Excel Table as dicts keyed by header (blank list rows included)."""
    ws = wb[sheet]
    cells = list(ws[ws.tables[table].ref])
    head = [c.value for c in cells[0]]
    return [{h: d(c.value) for h, c in zip(head, row)} for row in cells[1:]]


def key_figures(wb):
    """The Status sheet's KEY FIGURES, {measure: value}."""
    return {r["Measure"]: r["Value"] for r in table_rows(wb, "Status", "KPIs")}


def working_days(a, b, holidays):
    """NETWORKDAYS: Mon-Fri working days from a to b inclusive, skipping holidays."""
    n, day = 0, a
    while day <= b:
        n += day.weekday() < 5 and day not in holidays
        day += dt.timedelta(days=1)
    return n


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class ReportTest(Tmp):
    """Status sheet and Schedule Report, checked against values worked out in Python from the Tasks sheet."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.wb = recalc(cls.path, cls.dir)

    def leaves(self, wb):
        """Detail tasks as dicts by header, from the recalculated Tasks sheet."""
        return [r for r in table_rows(wb, "Tasks", "Tasks") if r["_IsLeaf"] == 1]

    def expected_evm(self, wb):
        """Earned value recomputed from the Tasks sheet and the rate card: cost = work x day rate per person,
        budget = that cost at the baseline's work once a baseline is set, AC = work done x day rate."""
        holidays = {d(r["Date"]) for r in table_rows(wb, "Settings", "HolidayTable") if r["Date"]}
        rates = {r["Resource"]: r["Day Rate"] for r in table_rows(wb, "Cost", "RateCard") if r["Resource"]}
        hours = named_cell(wb, "HoursPerDay").value
        bac = pv = ev = ac = 0
        for t in self.leaves(wb):
            cost = sum(w / hours * rates[n] for n, w, _ in person_slots(t))
            work = sum(w for _, w, _ in person_slots(t))
            base = t["Baseline Work"]
            budget = (base * cost / work if work else cost) if base else cost
            ps, pf = t["Planned Start"], t["Planned Finish"]
            if pf < STATUS_DATE:
                planned = 1
            elif ps >= STATUS_DATE:
                planned = 0
            else:
                planned = working_days(ps, STATUS_DATE - dt.timedelta(days=1), holidays) / working_days(ps, pf, holidays)
            bac += budget
            pv += budget * planned
            ev += budget * t["Progress"]
            ac += sum(dn / hours * rates[n] for n, _, dn in person_slots(t))
        return bac, pv, ev, ac

    def evm(self, wb, n):
        return named_cell(wb, f"evm_{n}").value

    def test_no_formula_errors(self):
        self.assertEqual(errors_in(self.wb), [])

    def test_earned_value(self):
        bac, pv, ev, ac = self.expected_evm(self.wb)
        self.assertAlmostEqual(bac, 58040 + 10980)          # = total cost: no baseline yet (10,980 = the second people on 4 tasks)
        self.assertAlmostEqual(ac, ev)                      # no baseline: budget = cost, so work done x rate = EV
        for n, want in (("BAC", bac), ("PV", pv), ("EV", ev), ("AC", ac)):
            self.assertAlmostEqual(self.evm(self.wb, n), want, places=6, msg=n)
        self.assertAlmostEqual(self.evm(self.wb, "SPI"), ev / pv)
        self.assertAlmostEqual(self.evm(self.wb, "CPI"), ev / ac)
        self.assertAlmostEqual(self.evm(self.wb, "EAC"), ac + (bac - ev) / (ev / ac))
        self.assertAlmostEqual(self.evm(self.wb, "TCPI"), (bac - ev) / (bac - ac))
        # the phase table: schedule SPI = progress / planned progress, both duration-weighted (no money there)
        phases = [r for r in table_rows(self.wb, "Schedule Report", "PhaseSummary") if r["Phase"]]
        self.assertEqual([r["Phase"] for r in phases], ["Initiation", "Requirements & Design", "Build", "Test & Release",
                                                        "Project Close"])
        for r in phases:
            if r["Planned"]:
                self.assertAlmostEqual(r["SPI"], r["Progress"] / r["Planned"], msg=r["Phase"])
        card = [r for r in table_rows(self.wb, "Cost", "RateCard") if r["Resource"]]
        self.assertAlmostEqual(sum(r["Actual Cost"] for r in card), ac, places=6)

    def test_earned_value_uses_the_baseline(self):
        first = self.wb["Tasks"]
        col = columns(first)
        snapshot = {(r, i): first.cell(r, col[s]).value for r in range(FIRST, FIRST + 30) for i, s in enumerate(SOURCE)}

        def baseline_then_grow(wb):
            for (r, i), v in snapshot.items():
                wb["Tasks"].cell(r, col[BASELINE[i]]).value = v
            set_input(wb, 21, "Dur", 10)                    # cost grows, the budget doesn't
        wb = recalc(self.path, self.dir, baseline_then_grow)
        bac, pv, ev, ac = self.expected_evm(wb)
        self.assertAlmostEqual(self.evm(wb, "BAC"), bac, places=6)   # the baseline work, not the grown duration
        self.assertAlmostEqual(self.evm(wb, "PV"), pv, places=6)
        self.assertAlmostEqual(self.evm(wb, "EAC"), ac + (bac - ev) / (ev / ac), places=6)

    def test_rag(self):
        rag = {r["Area"]: r for r in table_rows(self.wb, "Status", "StatusRAG")}
        self.assertEqual({a: r["RAG"] for a, r in rag.items()}, {
            "Schedule": "Green", "Schedule performance": "Red", "Deadlines": "Red",
            "Critical path": "Amber", "RAID": "Amber", "Plan quality": "Green"})
        self.assertEqual(rag["Plan quality"]["Value"], 0)    # "?" in COUNTIF is a wildcard; the formula escapes it
        self.assertEqual(rag["Critical path"]["Value"], 5)
        self.assertEqual(table_rows(self.wb, "Status", "OverallStatus")[0]["RAG"], "Red")

    def test_pm_override_and_thresholds(self):
        def edit(wb):
            ws = wb["Status"]
            for row in ws.iter_rows(min_row=1, max_row=40):
                if row[1].value in ("Schedule performance", "Deadlines"):
                    row[5].value = "Amber"                  # PM RAG
            named_cell(wb, "RAGRedIndex").value = 0.5
        wb = recalc(self.path, self.dir, edit)
        rag = {r["Area"]: r for r in table_rows(wb, "Status", "StatusRAG")}
        self.assertEqual(rag["Deadlines"]["Calculated"], "Red")
        self.assertEqual(rag["Deadlines"]["RAG"], "Amber")
        self.assertEqual(rag["Schedule performance"]["Calculated"], "Amber")   # SPI 0.78 with red below 0.5
        overall = table_rows(wb, "Status", "OverallStatus")[0]
        self.assertEqual((overall["RAG"], overall["Value"]), ("Amber", "0 red, 4 amber"))

    def listed(self, table):
        return [r["Task"] if "Task" in r else r["Milestone"] for r in table_rows(self.wb, "Schedule Report", table)
                if r["_Task"] not in (None, "")]

    def test_schedule_lists(self):
        tasks = self.leaves(self.wb)
        end = named_cell(self.wb, "LookaheadEnd").value
        end = d(end) if not isinstance(end, (int, float)) else dt.date(1899, 12, 30) + dt.timedelta(days=end)
        self.assertEqual(end, dt.date(2026, 11, 6))
        order = lambda rows, key: [t["Task Name"] for t in sorted(rows, key=key)]
        self.assertEqual(self.listed("MilestoneReport"), order(
            [t for t in tasks if t["Duration"] == 0], lambda t: (t["Planned Finish"], t["ID"])))
        late = [t for t in tasks if t["Progress"] < 1 and (t["Finish Var"] > 0 or t["Status"] == "Late Start")]
        self.assertEqual(len(late), 7)
        self.assertEqual(self.listed("LateTasks"), order(late, lambda t: (-max(0, t["Finish Var"]), t["ID"])))
        starting = [t for t in tasks if t["Actual Start"] is None and STATUS_DATE <= t["Start"] <= end]
        self.assertEqual(self.listed("StartingSoon"), order(starting, lambda t: (t["Start"], t["ID"])))
        self.assertEqual(self.listed("StartingSoon"), ["Payment gateway integration", "Accessibility review"])
        due = [t for t in tasks if t["Progress"] < 1 and STATUS_DATE <= t["Finish"] <= end]
        self.assertEqual(self.listed("DueSoon"), order(due, lambda t: (t["Finish"], t["ID"])))
        critical = [t for t in tasks if t["Critical"] == "Yes"]
        self.assertEqual(len(critical), 11)
        self.assertEqual(self.listed("CriticalTasks"), order(critical, lambda t: (t["Start"], t["ID"])))
        late_rows = table_rows(self.wb, "Schedule Report", "LateTasks")
        self.assertEqual(late_rows[0]["Resource"], "Jordan Lee")
        self.assertIn(late_rows[4]["Resource"], (None, ""))          # a milestone has no resource: blank, not 0
        self.assertIn(late_rows[-1]["Task"], (None, ""))


def serial(value):
    """A date read back from LibreOffice, whether stored as a date or as a serial number."""
    return d(value) if not isinstance(value, (int, float)) else dt.date(1899, 12, 30) + dt.timedelta(days=value)


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class AgileTest(Tmp):
    """Story points, velocity and burndown, against values worked out in Python from the Tasks sheet."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.wb = recalc(cls.path, cls.dir)

    def sprints(self, wb):
        return [r for r in table_rows(wb, "Agile", "SprintTable") if r["Sprint"]]

    def kpi(self, wb, n):
        return named_cell(wb, f"agl_{n}").value

    def test_no_formula_errors(self):
        self.assertEqual(errors_in(self.wb), [])

    def test_velocity_and_release_burndown(self):
        tasks = [r for r in table_rows(self.wb, "Tasks", "Tasks") if r["_IsLeaf"] == 1]
        total = sum(t["Story Points"] or 0 for t in tasks)
        self.assertEqual(total, 82)
        sprints = self.sprints(self.wb)
        done = []
        for s in sprints:
            if s["Status"] == "Future":
                break
            done.append(sum(t["Story Points"] or 0 for t in tasks
                            if t["Actual Finish"] and s["Start"] <= t["Actual Finish"] <= s["End"]))
        self.assertEqual(done, [5, 13, 15, 0])
        self.assertEqual([s["Done Points"] for s in sprints[:4]], done)
        self.assertEqual([s["Velocity"] for s in sprints[:3]], [5, 9, 11])     # rolling average of up to 3 sprints
        self.assertIn(sprints[3]["Velocity"], (None, ""))                       # the active sprint isn't closed yet
        self.assertEqual([s["Remaining Points"] for s in sprints[:4]], [77, 64, 49, 49])
        committed = {s["Sprint"]: s["Committed Points"] for s in sprints}
        for name, points in committed.items():
            self.assertEqual(points, sum(t["Story Points"] or 0 for t in tasks if t["Sprint"] == name), name)
        self.assertEqual((self.kpi(self.wb, "Remaining"), self.kpi(self.wb, "Velocity")), (49, 11))
        self.assertEqual((self.kpi(self.wb, "SprintsLeft"), self.kpi(self.wb, "FinishSprint")), (5, "Sprint 8"))
        self.assertEqual(serial(self.kpi(self.wb, "FinishDate")), dt.date(2026, 12, 27))

    def test_current_sprint_backlog_and_burndown(self):
        backlog = [r["Task"] for r in table_rows(self.wb, "Agile", "SprintBacklog") if r["_Task"] not in (None, "")]
        self.assertEqual(backlog, ["Feature screens"])
        days = [r for r in table_rows(self.wb, "Agile", "SprintBurndown") if r["Day"]]
        self.assertEqual((serial(days[0]["Day"]), serial(days[-1]["Day"]), len(days)),
                         (dt.date(2026, 10, 19), dt.date(2026, 11, 1), 14))
        self.assertEqual(days[6]["Remaining"], 8)                               # 25-Oct, the status date
        self.assertIn(days[7]["Remaining"], (None, ""))                         # nothing after the status date
        self.assertAlmostEqual(days[0]["Ideal"], 8 * 9 / 10)                    # 1 of 10 working days gone
        self.assertEqual(days[11]["Ideal"], 0)                                  # last working day (Fri 30-Oct)

    def test_planned_sprint_override_and_unestimated_tasks(self):
        def edit(wb):
            set_input(wb, 14, "Planned Sprint", "Sprint 4")    # payment gateway pulled into the current sprint
            set_input(wb, 21, "Story Points", None)            # system testing not estimated yet
        wb = recalc(self.path, self.dir, edit)
        self.assertEqual(errors_in(wb), [])
        t = Tasks(wb)
        self.assertEqual(t.get(14, "Sprint"), "Sprint 4")
        committed = {s["Sprint"]: s["Committed Points"] for s in self.sprints(wb)}
        self.assertEqual(committed["Sprint 4"], 13)
        self.assertEqual((self.kpi(wb, "Total"), self.kpi(wb, "Unestimated")), (74, 1))
        backlog = [r["Task"] for r in table_rows(wb, "Agile", "SprintBacklog") if r["_Task"] not in (None, "")]
        self.assertEqual(backlog, ["Payment gateway integration", "Feature screens"])


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class ResourceReportTest(Tmp):
    """Weekly hours: each task's work spread evenly over its forecast working days, recomputed in Python."""

    START = dt.date(2026, 9, 7)        # first grid week = the project start, so 26 weeks cover the whole plan

    def grid(self, wb, kind):
        ws = wb["Resource Report"]
        origin = named_cell(wb, f"rw_{kind}Origin")
        names = [ws.cell(origin.row + 2 + i, origin.column).value for i in range(25)]
        return {n: [ws.cell(origin.row + 2 + i, origin.column + 2 + w).value for w in range(26)]
                for i, n in enumerate(names) if n}

    def expected(self, wb):
        holidays = {d(r["Date"]) for r in table_rows(wb, "Settings", "HolidayTable") if r["Date"]}
        hours = {}
        tasks = {t["UID"]: t for t in table_rows(wb, "Tasks", "Tasks") if t["_IsLeaf"] == 1}
        people = [(t, t["_Resource1"], t["_LeafWork"]) for t in tasks.values() if t["_Resource1"]]
        people += [(t, t[f"_Res{k}"], t[f"_Res{k}Work"]) for t in tasks.values() for k in (2, 3, 4)
                   if t[f"_Res{k}"]]                                         # everyone else in the Resources cell
        for t, person, work in people:
            if not work:
                continue
            days, day = [], t["Start"]
            while day <= max(t["Finish"], t["Start"]):
                if day.weekday() < 5 and day not in holidays:
                    days.append(day)
                day += dt.timedelta(days=1)
            row = hours.setdefault(person, [0] * 26)
            for day in days:
                week = (day - self.START).days // 7
                if 0 <= week < 26:
                    row[week] += work / len(days)
        return hours

    def check(self, wb):
        self.assertEqual(errors_in(wb), [])
        load, want = self.grid(wb, "Load"), self.expected(wb)
        resources = {r["Resource Name"]: r for r in table_rows(wb, "Resources", "ResourceTable") if r["Resource Name"]}
        for name, row in load.items():
            got = [v or 0 for v in row]
            for w, (a, b) in enumerate(zip(got, want.get(name, [0] * 26))):
                self.assertAlmostEqual(a, b, places=6, msg=f"{name} week {w + 1}")
            self.assertAlmostEqual(sum(got), resources[name]["Work"] or 0, places=6, msg=name)   # whole plan in view
        return load

    def test_weekly_hours_and_summary(self):
        wb = recalc(self.path, self.dir, lambda wb: setattr(named_cell(wb, "ResourceWeekStart"), "value", self.START))
        self.check(wb)
        summary = {r["Resource"]: r for r in table_rows(wb, "Resource Report", "ResourceSummary") if r["Resource"]}
        jordan = summary["Jordan Lee"]
        # 20 h of it on Solution architecture (the second person there, done)
        self.assertEqual((jordan["Work"], jordan["Work Done"], jordan["Remaining"]), (228, 100, 128))
        # over capacity in 2 weeks only because of leave: 30-Oct, and 21-23 Dec around the holidays
        self.assertEqual((jordan["Late Tasks"], jordan["Open Critical"], jordan["Weeks Over"]), (2, 3, 2))
        card = [r for r in table_rows(wb, "Cost", "RateCard") if r["Resource"]]
        self.assertAlmostEqual(sum(r["Actual Cost"] or 0 for r in card), named_cell(wb, "evm_AC").value, places=6)

    def test_over_allocation(self):
        def edit(wb):
            named_cell(wb, "ResourceWeekStart").value = self.START
            set_input(wb, 14, "Res", "Jordan Lee[200%]")   # payment gateway at 200%: Jordan is double-booked
        wb = recalc(self.path, self.dir, edit)
        load = self.check(wb)
        util = self.grid(wb, "Util")
        ws = wb["Resource Report"]
        origin = named_cell(wb, "rw_LoadOrigin")
        week_days = [ws.cell(origin.row + 1, origin.column + 2 + w).value for w in range(26)]
        leave = leave_by_day(wb)
        weeks = [self.START + dt.timedelta(days=7 * w) for w in range(26)]
        off = [leave_between(leave, "Jordan Lee", w, w + dt.timedelta(days=6)) for w in weeks]
        over = sum(1 for h, n, o in zip(load["Jordan Lee"], week_days, off) if (h or 0) > 8 * (n - o))
        self.assertGreater(over, 0)
        summary = {r["Resource"]: r for r in table_rows(wb, "Resource Report", "ResourceSummary") if r["Resource"]}
        self.assertEqual(summary["Jordan Lee"]["Weeks Over"], over)
        self.assertAlmostEqual(summary["Jordan Lee"]["Peak Use"], max(u or 0 for u in util["Jordan Lee"]))


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class PeopleTest(Tmp):
    """Several people in a task's Resources cell: names or initials, units, work, cost, over-allocation."""

    def test_resources_cell(self):
        def edit(wb):
            # payment gateway (6 days): Jordan full time, Priya, Morgan by initials at 50%, Taylor at 200%, and a fifth
            set_input(wb, 14, "Res", "Jordan Lee, Priya Sharma, mp[50%], Taylor Morgan[200%], Sam Rivera")
            set_input(wb, 6, "Res", "Nobody")

        wb = recalc(self.path, self.dir, edit)
        self.assertEqual(errors_in(wb), [])
        t = Tasks(wb)
        self.assertEqual(t.get(14, "Resource Names"),
                         "Jordan Lee, Priya Sharma, Morgan Patel[50%], Taylor Morgan[200%] +more")
        self.assertEqual(t.get(14, "Work"), 48 + 48 + 24 + 96)
        self.assertIn("Only the first 4 people in Resources are used.", t.get(14, "Issues"))
        # Taylor is booked 200% in the week of 02-Nov through this task only
        self.assertEqual(t.get(14, "Over-allocated"), "Yes")
        self.assertEqual((t.get(6, "Issues"), t.get(6, "Work")), ("Nobody is not on Resources.", 0))
        resources = {r["Resource Name"]: r for r in table_rows(wb, "Resources", "ResourceTable") if r["Resource Name"]}
        self.assertEqual((resources["Taylor Morgan"]["Tasks"], resources["Taylor Morgan"]["Work"]), (3, 112 + 96))
        leaves = [r for r in table_rows(wb, "Tasks", "Tasks") if r["_IsLeaf"] == 1]
        self.assertAlmostEqual(sum(r["Work"] for r in resources.values() if r["Work"]), sum(r["Work"] for r in leaves))
        card = [r for r in table_rows(wb, "Cost", "RateCard") if r["Resource"]]
        costs = [r for r in table_rows(wb, "Cost", "TaskCost") if r["_LeafCost"]]
        self.assertAlmostEqual(sum(r["Task Cost"] or 0 for r in card), sum(r["_LeafCost"] for r in costs), places=6)
        self.assertEqual(named_cell(wb, "MaxPeople").value, 4)


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class ViewsTest(Tmp):
    """Outline levels, the Gantt's month / quarter scales and the Roadmap."""

    def test_outline_grouping_and_show_filter(self):
        raw = openpyxl.load_workbook(self.path)
        for sheet in ("Tasks", "Gantt"):
            ws = raw[sheet]
            self.assertFalse(ws.sheet_properties.outlinePr.summaryBelow, sheet)       # summary rows above
            self.assertEqual([ws.row_dimensions[FIRST + i].outline_level for i in range(13)],
                             [0, 1, 1, 1, 0, 1, 1, 1, 1, 0, 1, 2, 2], sheet)          # sample levels 1, 2, 2, 2, 1 ...
        wb = recalc(self.path, self.dir, lambda wb: setattr(named_cell(wb, "OutlineShowLevel"), "value", 1))
        self.assertEqual(errors_in(wb), [])
        shown = [r["Task Name"] for r in table_rows(wb, "Tasks", "Tasks") if r["Show"] == "Yes"]
        self.assertEqual(shown, ["Initiation", "Requirements & Design", "Build", "Test & Release", "Project Close"])
        gantt = [r["Show"] for r in table_rows(wb, "Gantt", "GanttView")[:4]]
        self.assertEqual(gantt, ["Yes", "No", "No", "No"])

    def gantt_header(self, wb, n):
        ws, origin = wb["Gantt"], named_cell(wb, "gantt_Origin")
        return [[ws.cell(r, origin.column + 1 + k).value for k in range(n)] for r in (3, 4, origin.row)]

    def test_month_and_quarter_scales(self):
        for period, months, letters, years in (
                # View from = two weeks before the 25-Oct status date: the Gantt starts in the week of 05-Oct
                (30, [(2026, 10), (2026, 11), (2026, 12), (2027, 1), (2027, 2)], ["O", "N", "D", "J", "F"],
                 ["2026", "", "", "2027", ""]),
                (91, [(2026, 10), (2027, 1), (2027, 4), (2027, 7), (2027, 10)], ["Q4", "Q1", "Q2", "Q3", "Q4"],
                 ["2026", "2027", "", "", ""])):
            wb = recalc(self.path, self.dir, lambda wb: setattr(named_cell(wb, "GanttPeriod"), "value", period))
            self.assertEqual(errors_in(wb), [])
            row3, row4, dates = self.gantt_header(wb, 5)
            self.assertEqual([(d(x).year, d(x).month, d(x).day) for x in dates], [(y, m, 1) for y, m in months])
            self.assertEqual((row3, [v or "" for v in row4]), (letters, years))

    def test_roadmap(self):
        self.wb = recalc(self.path, self.dir)
        t = Tasks(self.wb)
        phases = [r for r in table_rows(self.wb, "Roadmap", "RoadmapPhases") if r["Phase"]]
        self.assertEqual([r["Phase"] for r in phases], ["Initiation", "Requirements & Design", "Build", "Test & Release",
                                                        "Project Close"])
        build = phases[2]
        self.assertEqual((d(build["Start"]), d(build["Finish"]), build["Finish Var"]),
                         (t.get(10, "FS"), t.get(10, "FF"), t.get(10, "Var")))
        milestones = [r["Milestone"] for r in table_rows(self.wb, "Roadmap", "RoadmapMilestones") if r["Milestone"]]
        self.assertEqual(milestones, ["Kickoff complete", "Design sign-off", "Build complete", "Go-live", "Project closed"])
        ws, origin = self.wb["Roadmap"], named_cell(self.wb, "rm_Origin")
        months = [d(ws.cell(origin.row, origin.column + k).value) for k in range(1, 19)]
        self.assertEqual(months[0], dt.date(2026, 9, 1))                    # the project start's month
        self.assertEqual(months[-1], dt.date(2028, 2, 1))
        raw = openpyxl.load_workbook(self.path)["Roadmap"]
        self.assertEqual((raw.page_setup.fitToWidth, raw.page_setup.fitToHeight), (1, 1))     # one page
        timeline = [raw.cell(r, c).value for r in range(origin.row + 1, origin.row + 40)
                    for c in range(origin.column + 1, origin.column + 19)]
        self.assertTrue(all(v is None for v in timeline))                    # conditional formatting only


class EssentialsTest(Tmp):
    """The Start sheet, the Tasks sheet's essential view and its input prompts (structure only: no recalculation)."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.raw = openpyxl.load_workbook(cls.path)

    def test_start_sheet_holds_the_essential_settings(self):
        self.assertEqual(self.raw.sheetnames[0], "Start")
        for n in ("ProjectName", "ProjectManager", "ProjectStartDate", "WorkingWeek", "HoursPerDay"):
            self.assertTrue(self.raw.defined_names[n].attr_text.startswith("Start!"), n)
        self.assertTrue(self.raw.defined_names["GanttPeriod"].attr_text.startswith("Settings!"))
        links = [c.hyperlink.location for row in self.raw["Start"].iter_rows(max_row=30) for c in row if c.hyperlink]
        self.assertIn("'Tasks'!A1", links)
        self.assertIn("'Status'!A1", links)                              # the weekly update starts there

    def test_status_sheet_is_the_weekly_page(self):
        self.assertEqual(self.raw.defined_names["StatusDateInput"].attr_text, "Status!$C$5")
        ws = self.raw["Status"]
        links = [c.hyperlink.location for row in ws.iter_rows(max_row=12) for c in row if c.hyperlink]
        self.assertEqual(links, ["'Tasks'!A1", "'Log'!A1", "'History'!A1"])
        self.assertEqual(ws.print_area, "'Status'!$B$1:$H$" + str(ws.tables["StatusNotes"].ref.split(":")[1][1:]))
        snapshot = openpyxl.utils.cell.range_boundaries(ws.tables["SnapshotNow"].ref)
        notes = openpyxl.utils.cell.range_boundaries(ws.tables["StatusNotes"].ref)
        self.assertGreater(snapshot[1], notes[3])                       # below the printed page

    def test_gantt_panel_shows_the_essentials(self):
        ws = self.raw["Gantt"]
        head = {c.value: c.column_letter for c in ws[5] if c.value}
        hidden = lambda h: bool(ws.column_dimensions[head[h]].hidden)
        self.assertEqual([h for h in head if not h.startswith("_") and not hidden(h)][:7],
                         ["ID", "Task", "Start", "Finish", "%", "Resources", "Show"])
        self.assertTrue(all(hidden(h) for h in ("WBS", "Dur", "Planned Start", "Planned Finish", "Status")))
        self.assertEqual(ws.column_dimensions[head["Planned Start"]].width, 10)   # widths kept inside the group

    def test_optional_columns_are_grouped_and_collapsed(self):
        ws = self.raw["Tasks"]
        col = columns(ws)
        state = lambda header: (ws.column_dimensions[openpyxl.utils.get_column_letter(col[header])].outlineLevel,
                                bool(ws.column_dimensions[openpyxl.utils.get_column_letter(col[header])].hidden))
        for header in ("Task Name", "Duration", "Predecessors", "Resources", "Stage", "Actual Start", "Finish",
                       "Status", "Critical", "Issues", "Deadline", "Finish Var", "Work", "Sprint", "Duration Var"):
            self.assertEqual(state(header), (0, False), header)          # essentials and each section's key column
        for header in ("Mode", "Constraint", "Late Start", "Total Slack", "% Complete", "Resource Names", "Story Points",
                       "BL Source Start", "Baseline Start", "Baseline Work"):
            self.assertEqual(state(header), (1, True), header)           # collapsed sections
        self.assertEqual(state("_Row"), (2, True))                       # helpers: one level deeper
        self.assertFalse(ws.sheet_properties.outlinePr.summaryRight)     # +/- on the key column, left of its group

    def test_input_prompts(self):
        ws = self.raw["Tasks"]
        col = columns(ws)
        prompts = {}
        for dv in ws.data_validations.dataValidation:
            for rng in dv.sqref.ranges:
                prompts[rng.min_col] = (dv.showInputMessage, dv.promptTitle)
        for header, title in (("Task Name", "Task name"), ("Duration", "Duration (working days)"), ("Level", "Outline level"),
                              ("Predecessors", "Predecessors (by UID)"), ("Resources", "Resources"),
                              ("Actual Finish", "Actual Finish")):
            self.assertEqual(prompts[col[header]], (True, title), header)


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class IssuesTest(Tmp):
    def test_issues_column_says_what_is_wrong(self):
        def edit(wb):
            set_input(wb, 5, "UID", 3)                 # the same UID as task 3
            set_input(wb, 6, "Res", "Nobody")          # not on Resources
            set_input(wb, 7, "P1", 7)                  # a link to itself
            set_input(wb, 9, "Level", 4)               # two levels below the row above

        wb = recalc(self.path, self.dir, edit)
        self.assertEqual(errors_in(wb), [])
        t = Tasks(wb)
        self.assertIn("UID 3 is used twice.", t.get(3, "Issues"))
        self.assertIn("UID 3 is used twice.", t.get(5, "Issues"))
        self.assertEqual(t.get(6, "Issues"), "Nobody is not on Resources.")
        self.assertEqual(t.get(7, "Issues"), "Pred 1: links to itself.")
        self.assertIn("Level is more than 1 below the row above.", t.get(9, "Issues"))
        self.assertEqual(t.get(19, "Issues"), "Misses its deadline (late by 1d).")
        self.assertEqual(t.get(13, "Issues"), "Over-allocated: see the Resource Report.")
        self.assertIn(t.get(2, "Issues"), (None, ""))

    def test_working_week_preset(self):
        wb = recalc(self.path, self.dir, lambda wb: setattr(named_cell(wb, "WorkingWeek"), "value", "Sun-Thu"))
        self.assertEqual(errors_in(wb), [])
        self.assertEqual(named_cell(wb, "WorkWeek").value, "0000110")
        days = {d(r["Start"]).weekday() for r in table_rows(wb, "Tasks", "Tasks")
                if r["_IsLeaf"] == 1 and not r["Actual Start"]}
        self.assertFalse(days & {4, 5})                                    # nothing new starts on a Friday or Saturday
        out = io.BytesIO()
        wb.save(out)
        ns = "{http://schemas.microsoft.com/project}"
        week = {int(w.find(ns + "DayType").text): w.find(ns + "DayWorking").text
                for w in ET.fromstring(mspdi.convert(out.getvalue())).iter(ns + "WeekDay")
                if w.find(ns + "DayType").text != "0"}
        self.assertEqual(week, {1: "1", 2: "1", 3: "1", 4: "1", 5: "1", 6: "0", 7: "0"})   # Sun-Thu in MS Project too


class StageConversionTest(unittest.TestCase):
    def test_a_percent_becomes_the_nearest_stage(self):
        from planner.tasks import take_stage
        self.assertEqual(take_stage({"Pct": 1}), {"Stage": "Complete"})              # exact: no override kept
        self.assertEqual(take_stage({"Pct": 0.7}), {"Stage": "Test Ready"})
        self.assertEqual(take_stage({"Pct": 0.5}), {"Stage": "In Development", "Pct": 0.5})   # 40% is nearest
        self.assertEqual(take_stage({"Pct": 0.25}), {"Stage": "Design", "Pct": 0.25})
        self.assertEqual(take_stage({"Pct": None, "AS": 1}), {"Pct": None, "AS": 1})    # nothing to convert
        self.assertEqual(take_stage({"Pct": 0.9, "Stage": "Design"}), {"Pct": 0.9, "Stage": "Design"})
        self.assertEqual(take_stage({"Pct": 0.5}, [("Open", 0), ("Half", 0.5), ("Done", 1)]), {"Stage": "Half"})


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class StageTest(Tmp):
    """Progress by Stage (Settings > Stages), with % Complete as an optional override."""

    def test_stage_sets_progress_and_issues_ask_for_dates(self):
        def edit(wb):
            set_input(wb, 13, "Pct", None)                        # Core API: In Development = 40%
            set_input(wb, 14, "Stage", "Test Ready")              # not started yet: 70% without an Actual Start
            set_input(wb, 18, "Stage", "Complete")                # 100% without an Actual Finish
            set_input(wb, 19, "Stage", "Design")                  # a milestone half way
            set_input(wb, 12, "Stage", "In Testing")              # finished (Actual Finish) but not Complete
            set_input(wb, 16, "Stage", "Typo")
            ws = wb["Settings"]
            table = ws.tables["StageTable"]
            c1, r1, _, _ = openpyxl.utils.cell.range_boundaries(table.ref)
            ws.cell(r1 + 4, c1 + 1).value = 0.75                  # Test Ready is now 75%

        wb = recalc(self.path, self.dir, edit)
        self.assertEqual(errors_in(wb), [])
        t = Tasks(wb)
        self.assertEqual((t.get(13, "Prog"), t.get(14, "Prog"), t.get(18, "Prog")), (0.4, 0.75, 1))
        self.assertEqual(t.get(12, "Prog"), 1)                       # an Actual Finish counts as 100%
        self.assertIn("Started: enter the Actual Start.", t.get(14, "Issues"))
        self.assertIn("Complete: enter the Actual Finish.", t.get(18, "Issues"))
        self.assertIn("A milestone is Not Started or Complete.", t.get(19, "Issues"))
        self.assertIn("Actual Finish is set: set the Stage to Complete.", t.get(12, "Issues"))
        self.assertIn("Stage Typo is not on Settings > Stages.", t.get(16, "Issues"))
        self.assertEqual(t.get(13, "FF"), dt.date(2026, 11, 2))      # 6 of 10 days left: 26-Oct to 2-Nov

    def test_sample_uses_stages(self):
        wb = openpyxl.load_workbook(self.path)
        stages = {r["Stage"]: r["% Complete"] for r in table_rows(wb, "Settings", "StageTable")}
        self.assertEqual(stages, dict(settings.STAGES))
        t = {r["UID"]: r for r in table_rows(wb, "Tasks", "Tasks")}
        self.assertEqual((t[2]["Stage"], t[2]["% Complete"]), ("Complete", None))
        self.assertEqual((t[13]["Stage"], t[13]["% Complete"]), ("In Development", 0.5))     # the override kept


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class SprintCalendarTest(Tmp):
    """B-1003: sprint dates follow the Settings cadence until a sprint is moved; the sprints below follow it."""

    def set(self, wb, sprint, column, value):
        ws = wb["Agile"]
        c1, r1 = openpyxl.utils.cell.range_boundaries(ws.tables["SprintTable"].ref)[:2]
        head = {ws.cell(r1, c).value: c for c in range(c1, c1 + 40)}
        ws.cell(r1 + sprint, head[column]).value = value

    def test_moving_a_sprint_moves_the_ones_after_it(self):
        def edit(wb):
            self.set(wb, 5, "Set End", dt.date(2026, 11, 22))          # sprint 5: three weeks instead of two
            self.set(wb, 9, "Set Start", dt.date(2027, 1, 4))          # sprint 9 after a holiday break
        wb = recalc(self.path, self.dir, edit)
        self.assertEqual(errors_in(wb), [])
        s = {r["No"]: r for r in table_rows(wb, "Agile", "SprintTable")}
        self.assertEqual((serial(s[4]["End"]), serial(s[5]["Start"]), serial(s[5]["End"])),
                         (dt.date(2026, 11, 1), dt.date(2026, 11, 2), dt.date(2026, 11, 22)))
        self.assertEqual((serial(s[6]["Start"]), serial(s[8]["End"])), (dt.date(2026, 11, 23), dt.date(2027, 1, 3)))
        self.assertEqual((serial(s[9]["Start"]), serial(s[9]["End"])), (dt.date(2027, 1, 4), dt.date(2027, 1, 17)))
        self.assertEqual(s[5]["Working Days"], 15)
        tasks = {r["UID"]: r for r in table_rows(wb, "Tasks", "Tasks")}
        self.assertEqual(tasks[14]["Sprint"], "Sprint 5")               # starts 02-Nov: in the longer sprint

    def test_cadence_without_overrides(self):
        wb = openpyxl.load_workbook(self.path)
        ws = wb["Agile"]
        head = [c.value for c in ws[openpyxl.utils.cell.range_boundaries(ws.tables["SprintTable"].ref)[1]]]
        self.assertIn("Set Start", head)
        self.assertIn("Set End", head)


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class PlannedVsCurrentTest(unittest.TestCase):
    """B-1002: the baselined sample (the committed workbook) draws planned dates as an outline and slip in red."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.dir = Path(cls.tmp.name)
        cls.path = cls.dir / "Project_Planner.xlsx"
        build(cls.path, "--baselined")

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_baselined_sample_shows_slip(self):
        def add_task(wb):                                             # scope added after the baseline
            set_input(wb, 31, "Level", 2)
            set_input(wb, 31, "Name", "Data migration rehearsal")
            set_input(wb, 31, "Dur", 3)
            set_input(wb, 31, "P1", 26)
        wb = recalc(self.path, self.dir, add_task)
        self.assertEqual(errors_in(wb), [])
        self.assertTrue(named_cell(wb, "BaselineSet").value)
        t = {r["UID"]: r for r in table_rows(wb, "Tasks", "Tasks")}
        self.assertEqual((serial(t[13]["Planned Start"]), serial(t[13]["Planned Finish"])),
                         (dt.date(2026, 10, 9), dt.date(2026, 10, 20)))            # first estimate: 8 days
        self.assertGreater(t[30]["Finish Var"], 0)                                 # the project has slipped
        g = {r["ID"]: r for r in table_rows(wb, "Gantt", "GanttView") if r["ID"]}
        self.assertTrue(g[13]["_ShowPlan"])
        self.assertFalse(g[31]["_ShowPlan"])                          # added after the baseline: no outline
        self.assertEqual(key_figures(wb)["Baseline"], "Set on 30 tasks")

    def test_outline_rules(self):
        ws = openpyxl.load_workbook(self.path)["Gantt"]
        borders = [r.dxf.border for cf in ws.conditional_formatting for r in cf.rules if r.dxf and r.dxf.border]
        edges = {tuple(bool(getattr(b, s) and getattr(b, s).style) for s in ("left", "right", "top", "bottom"))
                 for b in borders}
        # closed box on a one-day plan, start, end, middle; plus the status-date line
        self.assertTrue({(True, True, True, True), (True, False, True, True), (False, True, True, True),
                         (False, False, True, True)} <= edges)


def person_slots(t):
    """[(resource, work hours, work done hours)] of a Tasks row: the first person and people 2-4."""
    slots = [(t["_Resource1"], t["_LeafWork"], t["_LeafWorkDone"])]
    slots += [(t[f"_Res{k}"], t[f"_Res{k}Work"], t[f"_Res{k}WorkDone"]) for k in (2, 3, 4)]
    return [(n, w or 0, d or 0) for n, w, d in slots if n]


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class CostSheetTest(Tmp):
    """B-1004: all money on the Cost sheet, from day rates, and nothing else reads it."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.wb = recalc(cls.path, cls.dir)

    def test_task_and_resource_cost_from_day_rates(self):
        wb = self.wb
        self.assertEqual(errors_in(wb), [])
        rates = {r["Resource"]: r["Day Rate"] for r in table_rows(wb, "Cost", "RateCard") if r["Resource"]}
        self.assertEqual(rates["Priya Sharma"], 760)
        hours = named_cell(wb, "HoursPerDay").value
        tasks = [t for t in table_rows(wb, "Tasks", "Tasks") if t["Task Name"]]
        cost = {r["ID"]: r for r in table_rows(wb, "Cost", "TaskCost") if r["ID"] not in (None, "")}
        ac = 0
        for t in tasks:
            if t["_IsLeaf"] != 1:
                continue
            want = sum(w / hours * rates[n] for n, w, _ in person_slots(t))
            self.assertAlmostEqual(cost[t["ID"]]["Cost"], want, places=6, msg=t["Task Name"])
            ac += sum(d / hours * rates[n] for n, _, d in person_slots(t))
        self.assertEqual(cost[2]["Cost"], 24 / 8 * 760)                      # 3 days of Priya
        phase = [t for t in tasks if t["Level"] == 1][1]                       # Requirements & Design
        children = [t for t in tasks if t["WBS"].startswith(phase["WBS"] + ".") and t["_IsLeaf"] == 1]
        self.assertAlmostEqual(cost[phase["ID"]]["Cost"], sum(cost[c["ID"]]["Cost"] for c in children))
        self.assertAlmostEqual(named_cell(wb, "evm_AC").value, ac, places=6)    # work done x day rates
        card = {r["Resource"]: r for r in table_rows(wb, "Cost", "RateCard") if r["Resource"]}
        self.assertAlmostEqual(sum(r["Actual Cost"] for r in card.values()), ac, places=6)
        self.assertAlmostEqual(sum(r["Task Cost"] for r in card.values()), named_cell(wb, "cost_Task").value)
        # resource cost = allocated days x rate: Priya is on the whole project at 50%
        start, finish = named_cell(wb, "ProjectStart").value, named_cell(wb, "ProjectFinish").value
        holidays = {d(r["Date"]) for r in table_rows(wb, "Settings", "HolidayTable") if r["Date"]}
        days = working_days(d(start), d(finish), holidays) * 0.5
        self.assertAlmostEqual(card["Priya Sharma"]["Allocated Days"], days)
        self.assertAlmostEqual(card["Priya Sharma"]["Resource Cost"], days * 760)
        self.assertAlmostEqual(named_cell(wb, "cost_Gap").value,
                               named_cell(wb, "cost_Resource").value - named_cell(wb, "cost_Task").value)

    def test_deleting_the_cost_sheet_leaves_the_plan_working(self):
        def delete_cost(wb):
            tables = set(wb["Cost"].tables)
            wb.remove(wb["Cost"])
            for n in list(wb.defined_names):
                text = wb.defined_names[n].attr_text
                if "Cost!" in text or any(text.startswith(t + "[") for t in tables):
                    del wb.defined_names[n]
            self.assertEqual(perf_check.missing_references(wb), [])
        wb = recalc(self.path, self.dir, delete_cost)
        self.assertNotIn("Cost", wb.sheetnames)
        self.assertEqual(errors_in(wb), [])
        self.assertEqual(key_figures(wb)["Detail tasks"], 23)

    def test_no_money_outside_the_cost_sheet(self):
        raw = openpyxl.load_workbook(self.path)
        for ws in raw.worksheets:
            if ws.title == "Cost":
                continue
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        for token in ("evm_", "cost_", "RateCard", "TaskCost"):
                            self.assertNotIn(token, c.value, f"{ws.title}!{c.coordinate}")


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class CapacityTest(Tmp):
    """B-1006 / B-1007: capacity per person per sprint from dates, leave, allocation and focus; the sprint picker."""

    def grid(self, wb, name):
        sheet, ref = wb.defined_names[name].attr_text.split("!")
        cells = wb[sheet.strip("'")][ref.replace("$", "")]
        people = [r["Resource Name"] for r in table_rows(wb, "Resources", "ResourceTable") if r["Resource Name"]]
        return {p: [c.value for c in row] for p, row in zip(people, cells)}

    def test_capacity_by_person_and_sprint(self):
        def edit(wb):
            ws = wb["Resources"]
            head = {c.value: c.column for c in ws[5]}
            row = next(r for r in range(6, 20) if ws.cell(r, head["Resource Name"]).value == "Casey Kim")
            ws.cell(row, head["End"]).value = dt.date(2026, 10, 23)          # leaves mid-sprint 4
            named_cell(wb, "ShowSprintInput").value = "Sprint 3"
        wb = recalc(self.path, self.dir, edit)
        self.assertEqual(errors_in(wb), [])
        days, points = self.grid(wb, "cap_Days"), self.grid(wb, "cap_Points")
        leave = self.grid(wb, "spr_LeaveGrid")
        sprints = [s for s in table_rows(wb, "Agile", "SprintTable") if s["Sprint"]]
        people = {r["Resource Name"]: r for r in table_rows(wb, "Resources", "ResourceTable") if r["Resource Name"]}
        holidays = {d(r["Date"]) for r in table_rows(wb, "Settings", "HolidayTable") if r["Date"]}
        for j, s in enumerate(sprints[:8]):
            for p, r in people.items():
                end = min(d(s["End"]), d(r["End"]) if r["End"] else d(s["End"]))
                worked = working_days(d(s["Start"]), end, holidays) if end >= d(s["Start"]) else 0
                want = max(0, worked - (leave[p][j] or 0)) * r["Allocation %"]
                self.assertAlmostEqual(days[p][j] or 0, want, places=6, msg=(p, s["Sprint"]))
                focus = 0.8 if r["Focus %"] is None else r["Focus %"]
                self.assertAlmostEqual(points[p][j] or 0, want * focus, places=6, msg=(p, s["Sprint"]))
            self.assertAlmostEqual(s["Available Days"], sum(days[p][j] or 0 for p in people), places=6)
            self.assertAlmostEqual(s["Capacity"], sum(points[p][j] or 0 for p in people), places=6)
            self.assertAlmostEqual(s["Advisory / Overhead"], s["Available Days"] - s["Capacity"], places=6)
        self.assertEqual(points["Priya Sharma"][0], 0)                           # advisory: Focus % = 0
        self.assertEqual(days["Casey Kim"][4], 0)                                # gone by sprint 5
        backlog = [r["Task"] for r in table_rows(wb, "Agile", "SprintBacklog") if r["_Task"] not in (None, "")]
        tasks = {t["Task Name"]: t for t in table_rows(wb, "Tasks", "Tasks") if t["Task Name"]}
        self.assertEqual(sorted(backlog), sorted(n for n, t in tasks.items()
                                                 if t["_IsLeaf"] == 1 and t["Sprint"] == "Sprint 3"))
        burndown = [r for r in table_rows(wb, "Agile", "SprintBurndown") if r["Day"]]
        self.assertEqual(serial(burndown[0]["Day"]), d(sprints[2]["Start"]))


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class LogsTest(Tmp):
    """The Log (RAID items, decisions and changes), the Status sheet's summaries of it, and status snapshots."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.wb = recalc(cls.path, cls.dir)

    def raid(self, wb):
        return {r["Ref"]: r for r in table_rows(wb, "Log", "Log") if r["Title"] and r["Type"] in log.RAID}

    def test_no_formula_errors(self):
        self.assertEqual(errors_in(self.wb), [])

    def test_raid_scoring_and_links(self):
        raid = self.raid(self.wb)
        # score = probability x impact; an issue counts as probability 5
        self.assertEqual({k: (r["Score"], r["Rating"]) for k, r in raid.items()},
                         {1: (16, "High"), 2: (9, "Medium"), 3: (20, "High"), 4: (8, "Medium"), 5: (15, "High"),
                          6: (6, "Low")})
        self.assertEqual([k for k, r in raid.items() if r["Overdue"] == "Overdue"], [3])   # due 23-Oct, not closed
        self.assertEqual((raid[1]["Task"], serial(raid[1]["Task Finish"])),
                         ("Payment gateway integration", dt.date(2026, 11, 9)))
        top = [r["Item"] for r in table_rows(self.wb, "Status", "TopRAID") if r["_Task"] not in (None, "")]
        self.assertEqual([t.split(":")[0] for t in top], ["Issue 3", "Risk 1", "Dependency 5", "Risk 2", "Assumption 4"])

    def test_raid_edits(self):
        def edit(wb):
            ws = wb["Log"]
            head = {c.value: c.column for c in ws[4]}
            ws.cell(6, head["Task UID"]).value = 999          # Ref 2 now points at a missing task
            ws.cell(7, head["Status"]).value = "Closed"       # the issue is resolved
            named_cell(wb, "RAIDHighScore").value = 16
        wb = recalc(self.path, self.dir, edit)
        self.assertEqual(errors_in(wb), [])
        raid = self.raid(wb)
        self.assertEqual(raid[2]["Task"], "UID 999 not found")
        self.assertIn(raid[3]["Overdue"], (None, ""))
        self.assertEqual(raid[5]["Rating"], "Medium")          # 15 is below the new threshold
        rag = {r["Area"]: r for r in table_rows(wb, "Status", "StatusRAG")}
        self.assertEqual(rag["RAID"]["Value"], 1)              # only Risk 1 (16) is still high and open
        top = [r["Item"] for r in table_rows(wb, "Status", "TopRAID") if r["_Task"] not in (None, "")]
        self.assertNotIn("Issue 3", " ".join(top))

    def test_change_control_summary(self):
        cc = {r["Measure"]: r["Value"] for r in table_rows(self.wb, "Status", "ChangeControl")}
        self.assertEqual(cc, {
            "Change requests awaiting decision": 1, "Change requests approved or implemented": 1,
            "Approved schedule impact (working days)": 2, "Approved cost impact": 1360,
            "Cost impact awaiting decision": 6800, "Decisions awaiting approval": 1, "Overdue log items": 1})
        waiting = {r["Ref"]: r["Waiting (days)"] for r in table_rows(self.wb, "Log", "Log") if r["Title"]}
        self.assertEqual((waiting[11], waiting[9]), (3, 3))   # raised 22-Oct, status date 25-Oct, still open
        self.assertIn(waiting[10], (None, ""))                 # approved
        self.assertIn(waiting[3], (None, ""))                  # an issue: waiting applies to decisions and changes

    def test_snapshot_matches_the_reports(self):
        snap = table_rows(self.wb, "Status", "SnapshotNow")[0]
        self.assertEqual(serial(snap["Status Date"]), STATUS_DATE)
        for field, n in (("% Done", "sch_Progress"), ("% Planned", "sch_Planned"), ("SPI", "sch_SPI"),
                         ("Points Done", "agl_Done")):
            self.assertAlmostEqual(snap[field], named_cell(self.wb, n).value, msg=field)
        self.assertEqual((snap["Open Risks"], snap["Open Issues"], snap["Pending Changes"], snap["RAG"]),
                         (2, 1, 1, "Red"))
        history = [r for r in table_rows(self.wb, "History", "StatusHistory") if r["Status Date"]]
        self.assertEqual(len(history), 4)
        self.assertEqual(list(history[0]), list(snap))          # same columns, so a pasted row lines up


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class ImportWizardTest(unittest.TestCase):
    """--import-wizard-sheets: MSP_Tasks / MSP_Resources, for MS Project's Import Wizard (not built by default)."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        path = Path(cls.tmp.name) / "Project_Planner.xlsx"
        build(path, "--import-wizard-sheets")
        cls.wb = recalc(path, Path(cls.tmp.name))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_sheets_built_on_request(self):
        self.assertEqual(self.wb.sheetnames[-2:], ["MSP_Tasks", "MSP_Resources"])
        self.assertEqual(errors_in(self.wb), [])

    def test_msp_headers_match_project_fields(self):
        headers = [c.value for c in self.wb["MSP_Tasks"][1]]
        for h in ("ID", "Unique ID", "Name", "Outline Level", "Duration", "Predecessors", "Unique ID Predecessors",
                  "Resource Names", "% Complete", "Actual Start", "Actual Finish", "Baseline Start", "Baseline Finish",
                  "Baseline Duration", "Baseline Work", "Deadline"):
            self.assertIn(h, headers)

    def test_msp_export(self):
        mx = self.wb["MSP_Tasks"]
        col = {c.value: c.column for c in mx[1]}
        row = lambda task_id: task_id + 1
        self.assertEqual(mx.cell(row(6), col["Duration"]).value, "8d")
        self.assertEqual(mx.cell(row(18), col["Resource Names"]).value, "Casey Kim[50%]")
        self.assertEqual(mx.cell(row(22), col["Constraint Type"]).value, "Start No Earlier Than")
        self.assertEqual(mx.cell(row(19), col["Unique ID Predecessors"]).value, "14,18")
        self.assertEqual(d(mx.cell(row(19), col["Deadline"]).value), dt.date(2026, 11, 6))
        self.assertEqual((mx.cell(row(17), col["Text1"]).value, mx.cell(row(17), col["Number1"]).value), ("Sprint 4", 8))
        self.assertIn(mx.cell(row(1), col["Text1"]).value, (None, ""))          # summaries carry no sprint


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class ConstraintTest(Tmp):
    """The eight MS Project constraint types, hand-checked on the sample (holidays 24/25/31-Dec and 1-Jan)."""

    def msp_task(self, wb, uid):
        """The task's MS Project XML element, exported from the recalculated workbook (its cached values)."""
        out = io.BytesIO()
        wb.save(out)
        ns = "{http://schemas.microsoft.com/project}"
        root = ET.fromstring(mspdi.convert(out.getvalue()))
        task = next(t for t in root.iter(ns + "Task") if t.find(ns + "UID").text == str(uid))
        return lambda tag: task.find(ns + tag).text if task.find(ns + tag) is not None else None

    def test_late_and_lower_bounds(self):
        def edit(wb):
            set_input(wb, 23, "Constraint", "As Late As Possible")
            set_input(wb, 24, "Constraint", "Must Finish On")              # no date: ignored, and explained
            set_input(wb, 29, "Constraint", "Finish No Earlier Than")
            set_input(wb, 29, "Constraint Date", dt.date(2027, 1, 15))
            set_input(wb, 30, "Constraint", "Must Start On")
            set_input(wb, 30, "Constraint Date", dt.date(2027, 1, 20))

        wb = recalc(self.path, self.dir, edit)
        self.assertEqual(errors_in(wb), [])
        t = Tasks(wb)
        # training's free slack is 18 days (early finish 18-Nov, readiness review starts 15-Dec): shown 08-14 Dec and
        # nothing after it moves. The plan (04-10 Nov) moves by the same 18 days, so the finish variance stays 6.
        self.assertEqual([t.get(23, k) for k in ("FS", "FF", "PS", "PF", "Var", "Free Slack")],
                         [dt.date(2026, 12, 8), dt.date(2026, 12, 14), dt.date(2026, 11, 30), dt.date(2026, 12, 4), 6, 18])
        self.assertEqual(t.get(24, "FS"), dt.date(2026, 12, 15))
        self.assertEqual((t.get(24, "FS"), t.get(24, "Link")),
                         (dt.date(2026, 12, 15), "Constraint ignored: enter a Constraint Date."))
        self.assertEqual((t.get(29, "FS"), t.get(29, "FF"), t.get(29, "Total Slack")),
                         (dt.date(2027, 1, 14), dt.date(2027, 1, 15), 3))               # can slip to 20-Jan
        self.assertEqual(t.get(30, "FS"), dt.date(2027, 1, 20))
        self.assertEqual(d(named_cell(wb, "ProjectFinish").value), dt.date(2027, 1, 20))
        self.assertEqual(self.msp_task(wb, 23)("ConstraintType"), "1")                  # As Late As Possible
        self.assertEqual((self.msp_task(wb, 30)("ConstraintType"), self.msp_task(wb, 30)("ConstraintDate")[:10]),
                         ("2", "2027-01-20"))                                       # Must Start On

    def test_upper_bounds_win_over_links(self):
        def edit(wb):
            set_input(wb, 21, "Constraint", "Finish No Later Than")
            set_input(wb, 21, "Constraint Date", dt.date(2026, 11, 18))
            set_input(wb, 26, "Constraint", "Start No Later Than")
            set_input(wb, 26, "Constraint Date", dt.date(2026, 12, 14))

        wb = recalc(self.path, self.dir, edit)
        self.assertEqual(errors_in(wb), [])
        t = Tasks(wb)
        # testing (8 days) must finish by 18-Nov, so it starts 09-Nov although the build finishes that day
        self.assertEqual((t.get(21, "FS"), t.get(21, "FF"), t.get(21, "Late Finish")),
                         (dt.date(2026, 11, 9), dt.date(2026, 11, 18), dt.date(2026, 11, 18)))
        # hypercare starts by 14-Dec, before go-live (15-Dec): go-live and the tasks driving it show negative slack
        self.assertEqual(t.get(26, "FS"), dt.date(2026, 12, 14))
        self.assertEqual((t.get(25, "Late Finish"), t.get(25, "Total Slack"), t.get(22, "Total Slack")),
                         (dt.date(2026, 12, 11), -2, -2))


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class LinkCheckTest(Tmp):
    def test_bad_links_are_explained_not_circular(self):
        def edit(wb):
            set_input(wb, 7, "P1", 7)      # itself
            set_input(wb, 13, "P1", "12, 11")    # its own summary (Backend)
            set_input(wb, 14, "P1", 999)   # missing
            set_input(wb, 6, "P1", "4, 150")     # empty row (below, too)
            set_input(wb, 8, "P1", "6, 21")      # a task further down: predecessors must be above
            set_input(wb, 21, "P1", "19, 15")    # another phase's summary (Frontend)

        wb = recalc(self.path, self.dir, edit)
        self.assertEqual(errors_in(wb), [])
        t = Tasks(wb)
        self.assertEqual(t.get(7, "Link"), "Pred 1: links to itself.")
        self.assertEqual(t.get(13, "Link"), "Pred 2: UID 11 is a summary: link to one of its tasks.")
        self.assertEqual(t.get(14, "Link"), "Pred 1: UID 999 not found.")
        self.assertEqual(t.get(6, "Link"), "Pred 2: UID 150 is an empty row.")
        self.assertEqual(t.get(8, "Link"), "Pred 2: UID 21 is below this task: move it above.")
        self.assertEqual(t.get(21, "Link"), "Pred 2: UID 15 is a summary: link to one of its tasks.")
        self.assertEqual(t.get(8, "PS"), dt.date(2026, 9, 24))     # the valid link still applies
        # task 14 lost its link (UID 999), so the build milestone moves to 29-Oct (task 18) and testing follows it
        self.assertEqual(t.get(21, "PS"), dt.date(2026, 10, 30))
        self.assertEqual(t.get(7, "PS"), dt.date(2026, 9, 7))      # bad link ignored -> project start
        self.assertEqual(t.get(13, "PS"), dt.date(2026, 10, 9))    # good link still applies
        self.assertIn(t.get(7, "Pred"), (None, ""))
        issues = [r for r in table_rows(wb, "Tasks", "Tasks") if r["Issues"]]
        self.assertEqual(key_figures(wb)["Tasks with issues (see Tasks > Issues)"], len(issues))
        self.assertGreaterEqual(len(issues), 6)                       # the six bad links, at least


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class InsertRowTest(Tmp):
    """Inserting a task mid-list renumbers IDs but links follow UIDs.

    Every task row carries identical relative formulas, so shifting the input cells down one row is exactly what
    Excel's 'Insert Copied Cells' of a blank task row produces.
    """

    def test_insert_task_keeps_links(self):
        def insert(wb):
            ws = wb["Tasks"]
            inputs = [c for c in range(1, ws.max_column + 1) if (ws.cell(5, c).fill.start_color.rgb or "").endswith("1F4E78")]
            at = 14 + FIRST - 1                          # insert above "Payment gateway integration"
            last = FIRST + 199
            for r in range(last - 1, at - 1, -1):
                for c in inputs:
                    ws.cell(r + 1, c).value = ws.cell(r, c).value
            for c in inputs:
                ws.cell(at, c).value = None
            new = {"UID": 201, "Level": 3, "Name": "Security review", "Dur": 2, "P1": 13}
            col = columns(ws)
            for k, v in new.items():
                ws.cell(at, col[H[k]]).value = v

        wb = recalc(self.path, self.dir, insert)
        self.assertEqual(errors_in(wb), [])
        t = Tasks(wb)
        self.assertEqual(t.get(14, "Name"), "Security review")
        self.assertEqual(t.get(14, "WBS"), "3.1.3")
        self.assertEqual(t.get(14, "PS"), dt.date(2026, 10, 23))
        self.assertEqual(t.get(15, "Name"), "Payment gateway integration")
        self.assertEqual(t.get(15, "WBS"), "3.1.4")
        self.assertEqual(t.get(15, "PS"), dt.date(2026, 10, 23))     # still follows UID 13
        self.assertEqual(t.get(20, "Name"), "Build complete")
        self.assertEqual(t.get(20, "Pred"), "15,19")                  # by ID: shifted
        self.assertEqual(t.get(18, "Pred"), "17,13SS+2d")             # Feature screens now ID 18
        self.assertEqual(t.get(11, "Work"), 168 + 0)                  # no resource on new task
        gantt = wb["Gantt"]
        task = next(c.column for c in gantt[5] if c.value == "Task")
        self.assertEqual(gantt.cell(14 + FIRST - 1, task).value.strip(), "Security review")


import lo_session  # noqa: E402


@unittest.skipUnless(SOFFICE and lo_session.AVAILABLE, "LibreOffice / python3-uno not installed")
class RealRowEditTest(Tmp):
    """Insert and delete real rows, as a person would in Excel, and check every existing formula still works.

    This is what catches formulas that only look robust: a range that ends at the row above, for example, does
    not grow when a row is inserted directly above, so the row below silently loses sight of the new task.
    """

    def edit(self, fn):
        prepared = self.dir / f"prepared_{fn.__name__}.xlsx"
        wb = openpyxl.load_workbook(self.path)
        named_cell(wb, "StatusDateInput").value = STATUS_DATE
        wb.save(prepared)
        out = self.dir / f"edited_{fn.__name__}.xlsx"
        cols = columns(wb["Tasks"])
        with lo_session.Calc(prepared) as calc:
            fn(calc, cols)
            calc.save(out)
        return openpyxl.load_workbook(out, data_only=True)

    def test_insert_task_directly_above_and_link_to_it(self):
        def insert(calc, cols):
            row = 14 + FIRST - 1                                    # "Payment gateway integration"
            calc.insert_row_like_excel("Tasks", row, max(cols.values()))
            for c in cols.values():                                  # clear the copied inputs, keep the formulas
                if not calc.sheet("Tasks").getCellByPosition(c - 1, row - 1).getFormula().startswith("="):
                    calc.set("Tasks", c - 1, row, None)
            for key, value in {"UID": 201, "Level": 3, "Name": "Security review", "Dur": 2, "P1": 13}.items():
                calc.set("Tasks", cols[H[key]] - 1, row, value)
            calc.set("Tasks", cols[H["P1"]] - 1, row + 1, 201)      # the task below now follows the new one

        wb = self.edit(insert)
        self.assertEqual(errors_in(wb), [])
        t = Tasks(wb)
        self.assertEqual(t.get(14, "Name"), "Security review")
        self.assertEqual(t.get(14, "WBS"), "3.1.3")
        self.assertEqual(t.get(14, "PF"), dt.date(2026, 10, 26))
        self.assertEqual(t.get(15, "Name"), "Payment gateway integration")
        self.assertIsNone(t.get(15, "Link"))
        self.assertEqual(t.get(15, "PS"), dt.date(2026, 10, 27))   # follows the new task
        self.assertEqual(t.get(20, "Pred"), "15,19")                 # IDs renumbered in the MSP text
        gantt = wb["Gantt"]
        task = next(c.column for c in gantt[5] if c.value == "Task")
        self.assertEqual(gantt.cell(14 + FIRST - 1, task).value.strip(), "Security review")

    def test_delete_a_predecessor(self):
        def delete(calc, cols):
            calc.delete_row("Tasks", 14 + FIRST - 1)                # remove "Payment gateway integration" (UID 14)

        wb = self.edit(delete)
        self.assertEqual(errors_in(wb), [])
        t = Tasks(wb)
        self.assertEqual(t.get(18, "Name"), "Build complete")
        self.assertEqual(t.get(18, "Link"), "Pred 1: UID 14 not found.")
        self.assertEqual(t.get(18, "PS"), dt.date(2026, 10, 29))   # still follows its other predecessor


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class RippleTest(Tmp):
    """A late actual finish must push every successor's forecast dates."""

    def test_slip_ripples_to_successors(self):
        t = Tasks(recalc(self.path, self.dir, lambda wb: set_input(wb, 13, "AF", dt.date(2026, 11, 4))))
        self.assertEqual(t.get(14, "FS"), dt.date(2026, 11, 5))    # next working day
        self.assertEqual(t.get(14, "FF"), dt.date(2026, 11, 12))
        self.assertEqual(t.get(19, "FS"), dt.date(2026, 11, 12))   # milestone
        self.assertEqual(t.get(21, "FS"), dt.date(2026, 11, 13))   # testing
        self.assertEqual(t.get(14, "PF"), dt.date(2026, 10, 30))   # planned untouched

    def test_status_date_moves_unstarted_work(self):
        def edit(wb):
            set_input(wb, 13, "AS", None)
            set_input(wb, 13, "Pct", None)
            set_input(wb, 13, "Stage", None)

        t = Tasks(recalc(self.path, self.dir, edit))
        self.assertEqual(t.get(13, "FS"), dt.date(2026, 10, 26))   # first working day after 25-Oct
        self.assertEqual(t.get(13, "FF"), dt.date(2026, 11, 6))
        self.assertEqual(t.get(14, "FS"), dt.date(2026, 11, 9))


def parse_leave(text):
    """{day of month: fraction} from a Leave cell such as "16, 17(0.5)" (semicolons and spaces allowed)."""
    out = {}
    for token in str(text or "").replace(";", ",").replace(" ", "").split(","):
        if not token:
            continue
        day, _, rest = token.partition("(")
        out[int(day)] = float(rest.rstrip(")")) if rest else 1.0
    return out


MONTH_NAMES = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def leave_by_day(wb):
    """{(person, date): working days of leave}, worked out in Python from the Leave sheet, working days only."""
    holidays = {d(r["Date"]) for r in table_rows(wb, "Settings", "HolidayTable") if r["Date"]}
    out = {}
    for row in table_rows(wb, "Resources", "ResourceTable"):
        if not row["Resource Name"]:
            continue
        for header, text in row.items():
            if not re.match(r"^[A-Z][a-z]{2}-\d\d$", str(header)) or not text:      # the leave month columns
                continue
            month, year = MONTH_NAMES.index(header[:3]) + 1, 2000 + int(header[-2:])
            for day, fraction in parse_leave(text).items():
                when = dt.date(year, month, day)
                if when.weekday() < 5 and when not in holidays:
                    out[(row["Resource Name"], when)] = fraction
    return out


def leave_between(leave, person, start, end):
    return sum(f for (p, when), f in leave.items() if p == person and start <= when <= end)


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class LeaveTest(Tmp):
    """Leave lowers capacity (not task dates): weekly and sprint grids, team capacity and the Over-allocated flag."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.wb = recalc(cls.path, cls.dir)

    def grid(self, wb, name, origin_rows, people):
        """{person: [values]} of a named grid whose rows follow ResourceTable."""
        sheet, ref = wb.defined_names[name].attr_text.split("!")
        top_left, bottom_right = ref.replace("$", "").split(":")
        rows = [[c.value for c in row] for row in wb[sheet.strip("'")][top_left:bottom_right]]
        return dict(zip(people, rows))

    def people(self, wb):
        return [r["Resource Name"] for r in table_rows(wb, "Resources", "ResourceTable") if r["Resource Name"]]

    def check(self, wb):
        self.assertEqual(errors_in(wb), [])
        leave, people = leave_by_day(wb), self.people(wb)
        first = serial(named_cell(wb, "ResourceWeekStart").value)
        weekly = self.grid(wb, "rw_Leave", 0, people)
        for person in people:
            for w, got in enumerate(weekly[person][:26]):
                start = first + dt.timedelta(days=7 * w)
                self.assertAlmostEqual(got or 0, leave_between(leave, person, start, start + dt.timedelta(days=6)),
                                       msg=f"{person} week of {start}")
        sprints = [s for s in table_rows(wb, "Agile", "SprintTable") if s["Sprint"]]
        by_sprint = self.grid(wb, "spr_LeaveGrid", 0, people)
        units = {r["Resource Name"]: r["Allocation %"] for r in table_rows(wb, "Resources", "ResourceTable") if r["Resource Name"]}
        for j, s in enumerate(sprints):
            days = {p: leave_between(leave, p, s["Start"], s["End"]) for p in people}
            for p in people:
                self.assertAlmostEqual(by_sprint[p][j] or 0, days[p], msg=f"{p} {s['Sprint']}")
            want = sum(units[p] * (s["Working Days"] - days[p]) for p in people)   # nobody starts or ends mid-plan
            self.assertAlmostEqual(s["Available Days"], want, msg=s["Sprint"])
        return leave, weekly

    def test_sample_leave(self):
        leave, weekly = self.check(self.wb)
        self.assertEqual(weekly["Casey Kim"][6:8], [4.5, 5])           # weeks of 16 and 23 November
        self.assertEqual(leave_between(leave, "Jordan Lee", dt.date(2026, 12, 21), dt.date(2026, 12, 27)), 2.5)
        sprints = {s["Sprint"]: s for s in table_rows(self.wb, "Agile", "SprintTable") if s["Sprint"]}
        self.assertEqual(sprints["Sprint 6"]["Available Days"], 60 - 9.5)  # Casey off 9.5 days
        summary = {r["Resource"]: r for r in table_rows(self.wb, "Resource Report", "ResourceSummary") if r["Resource"]}
        self.assertEqual((summary["Casey Kim"]["Leave Days"], summary["Casey Kim"]["Weeks Over"]), (9.5, 1))

    def test_over_allocated_flag(self):
        people = self.people(self.wb)
        util = self.grid(self.wb, "rw_Util", 0, people)
        first = serial(named_cell(self.wb, "ResourceWeekStart").value)
        want = []
        for t in table_rows(self.wb, "Tasks", "Tasks"):
            names = [n for n in (t["_Resource1"], t["_Res2"], t["_Res3"], t["_Res4"]) if n]   # everyone on the task
            if t["_IsLeaf"] != 1 or not names or t["Progress"] >= 1:
                continue
            a = max(0, (t["Start"] - first).days // 7)
            b = min(25, (t["Finish"] - first).days // 7)
            if any((u or 0) > 1 for n in names for u in util[n][a:b + 1] if isinstance(u, (int, float))):
                want.append(t["Task Name"])
        flagged = [t["Task Name"] for t in table_rows(self.wb, "Tasks", "Tasks") if t["Over-allocated"] == "Yes"]
        self.assertEqual(flagged, want)
        self.assertEqual(flagged, ["Core API endpoints", "Training & user guides", "Hypercare support"])

    def test_leave_entries(self):
        def edit(wb):
            ws = wb["Resources"]
            head = {c.value: c.column for c in ws[5]}
            row = {ws.cell(r, head["Resource Name"]).value: r for r in range(6, 20)}
            ws.cell(row["Casey Kim"], head["Nov-26"]).value = "16;17 (0.25) , 21, 22"     # 21-22 Nov is a weekend
            ws.cell(row["Jordan Lee"], head["Dec-26"]).value = "24, 25, 28(0.5)"         # 24-25 Dec are holidays
            ws.cell(row["Taylor Morgan"], head["Nov-26"]).value = "9,10,11,12,13"
        wb = recalc(self.path, self.dir, edit)
        leave, weekly = self.check(wb)
        self.assertEqual(weekly["Casey Kim"][6], 1.25)
        self.assertEqual(leave_between(leave, "Jordan Lee", dt.date(2026, 12, 21), dt.date(2026, 12, 31)), 0.5)
        self.assertEqual(weekly["Taylor Morgan"][5], 5)
        t = Tasks(wb)
        self.assertEqual(t.get(21, "Over-allocated"), "Yes")            # system testing, Taylor away all week
        self.assertEqual(t.get(21, "FF"), dt.date(2026, 11, 19))        # capacity only: dates don't move


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class BuildOptionsTest(unittest.TestCase):
    """--resources, --grid-resources, --sprints, --gantt-columns and --weeks size the tables and grids."""

    def check(self, flags, resources, grid, sprints, gantt, weeks):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Project_Planner.xlsx"
            build(path, *flags, "--import-wizard-sheets")
            raw = openpyxl.load_workbook(path)
            report = perf_check.static_report(path)
            self.assertEqual((report["a1_references"], report["excel_limit_problems"]), ([], []))
            rows = lambda sheet, table: openpyxl.utils.cell.range_boundaries(raw[sheet].tables[table].ref)[3] - \
                openpyxl.utils.cell.range_boundaries(raw[sheet].tables[table].ref)[1]
            self.assertEqual((rows("Resources", "ResourceTable"), rows("MSP_Resources", "MSPResources"),
                              rows("Agile", "SprintTable")), (resources, resources, sprints))
            origin = named_cell(raw, "gantt_Origin")
            gantt_ws = raw["Gantt"]
            self.assertIsNotNone(gantt_ws.cell(origin.row, origin.column + gantt).value)    # last timeline date
            self.assertIsNone(gantt_ws.cell(origin.row, origin.column + gantt + 1).value)
            for grid_name in ("load_Grid", "rw_Load", "rw_Util"):
                ref = raw.defined_names[grid_name].attr_text.split("!")[1].replace("$", "")
                c1, r1, c2, r2 = openpyxl.utils.cell.range_boundaries(ref)
                self.assertEqual(r2 - r1 + 1, grid, grid_name)
                if grid_name != "load_Grid":
                    self.assertEqual(c2 - c1 + 1, weeks, grid_name)
            wb = recalc(path, tmp)
            self.assertEqual(errors_in(wb), [])
            self.assertEqual(named_cell(wb, "SprintCount").value, min(12, sprints))

    def test_smaller(self):
        self.check(["--resources", "10", "--grid-resources", "8", "--sprints", "6", "--gantt-columns", "40",
                    "--weeks", "8"], 10, 8, 6, 40, 8)

    def test_larger(self):
        self.check(["--resources", "80", "--grid-resources", "40", "--sprints", "40", "--gantt-columns", "220",
                    "--weeks", "52"], 80, 40, 40, 220, 52)

    def test_invalid_sizes_are_refused(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = subprocess.run([sys.executable, str(BUILD), str(Path(tmp) / "x.xlsx"), "--grid-resources", "60",
                                     "--resources", "50"], capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("--grid-resources can't be more than --resources", result.stderr)


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class BlankTemplateTest(unittest.TestCase):
    def test_blank_template(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Project_Planner_Blank.xlsx"
            build(path, "--blank")
            wb = recalc(path, tmp, status_date=False)
            self.assertEqual(errors_in(wb), [])
            t = Tasks(wb)
            self.assertTrue(all(t.ws.cell(r, t.col[H["Name"]]).value is None for r in range(FIRST, FIRST + 200)))
            self.assertIsNone(wb["Resources"]["B6"].value)
            self.assertEqual([r["Date"] for r in table_rows(wb, "Settings", "HolidayTable") if r["Date"]], [])
            self.assertIsNone(named_cell(wb, "ProjectStartDate").value)      # required input, highlighted yellow
            self.assertIsNone(named_cell(wb, "StatusDateInput").value)
            self.assertEqual(key_figures(wb)["Detail tasks"], 0)


if __name__ == "__main__":
    unittest.main()
