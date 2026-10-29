"""Template 1.3 (the --release 1.3 build, release-1.3/): template 1.2 plus a set-up checklist on Start that ticks
itself. The other editions don't change with any of this.

Run:  python -m unittest tests.test_release_13 -v      (the recalculation tests need LibreOffice)
"""
import datetime as dt
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(ROOT))
from test_planner import SOFFICE, CommittedWorkbooksTest, errors_in, named_cell, recalc  # noqa: E402
from test_release import head, small  # noqa: E402

RELEASE = ROOT / "release-1.3"
SAMPLE, BLANK, SIMPLE = (RELEASE / n for n in ("Project_Planner.xlsx", "Project_Planner_Blank.xlsx",
                                               "Project_Planner_Simple.xlsx"))
STEPS = ["1. Your project", "2. People", "3. Tasks", "4. Check", "5. Baseline", "6. First status"]


def make(**k):
    from planner.build import make_workbook
    return make_workbook(release_edition="1.3", **k)


def classic_again():
    from planner.build import make_workbook
    make_workbook(sizes=small())


def checklist(wb):
    """{step label: Done / To do} and the Progress line, read through the start_Steps / start_Done names."""
    sheet, ref = wb.defined_names["start_Steps"].attr_text.split("!")
    ws = wb[sheet]
    first, last = (int(x.split("$")[-1]) for x in ref.split(":"))
    steps = {ws.cell(r, 2).value: ws.cell(r, 4).value for r in range(first, last + 1)}
    return steps, ws.cell(first - 1, 3).value


class SetupStructureTest(unittest.TestCase):
    """The checklist: six linked steps, each with a Done / To do formula on defined names."""

    @classmethod
    def setUpClass(cls):
        cls.wb = make(baselined=True)
        cls.simple = make(baselined=True, modules=())

    @classmethod
    def tearDownClass(cls):
        classic_again()

    def test_versions(self):
        self.assertEqual(self.wb.defined_names["TemplateVersion"].attr_text, '"1.3"')
        self.assertEqual(self.wb.defined_names["PlannerVersion"].attr_text, '"2.3"')

    def test_six_linked_steps(self):
        steps, progress = checklist(self.wb)
        self.assertEqual(list(steps), STEPS)
        sheet, ref = self.wb.defined_names["start_Steps"].attr_text.split("!")
        first = int(ref.split(":")[0].split("$")[-1])
        ws = self.wb[sheet]
        self.assertEqual([ws.cell(first + k, 3).hyperlink.location for k in range(6)],
                         ["'Settings'!A1", "'Resources'!A1", "'Tasks'!A1", "'Check'!A1", "'TakeSnapshot'!A1",
                          "'Status'!A1"])
        for k in range(6):
            self.assertTrue(str(ws.cell(first + k, 4).value).startswith("=IF("), STEPS[k])
        self.assertIn("start_Done", progress)
        self.assertNotIn("QUICK START", [c.value for c in ws["B"]])          # the checklist replaces it

    def test_simple_falls_back_to_tasks(self):
        """Simple has no Check or TakeSnapshot sheet: those steps link to Tasks and read the Issues column."""
        sheet, ref = self.simple.defined_names["start_Steps"].attr_text.split("!")
        first = int(ref.split(":")[0].split("$")[-1])
        ws = self.simple[sheet]
        self.assertEqual(ws.cell(first + 3, 3).hyperlink.location, "'Tasks'!A1")
        self.assertIn("tsk_Issues", ws.cell(first + 3, 4).value)
        self.assertEqual(ws.cell(first + 4, 3).hyperlink.location, "'Tasks'!A1")

    def test_next_step_highlighted(self):
        sheet = self.wb.defined_names["start_Done"].attr_text.split("!")[0]
        rules = [r.formula[0] for cf in self.wb[sheet].conditional_formatting for r in cf.rules]
        self.assertTrue(any("COUNTIF(" in f and '"To do"' in f for f in rules), rules)


class CommittedTest(unittest.TestCase):
    compare = CommittedWorkbooksTest.compare

    def test_sample(self):
        self.compare(SAMPLE, "--release", "1.3", "--baselined")

    def test_blank(self):
        self.compare(BLANK, "--release", "1.3", "--blank")

    def test_simple(self):
        self.compare(SIMPLE, "--release", "1.3", "--baselined", "--simple")


class CommittedSampleTest(unittest.TestCase):

    def test_no_errors(self):
        for path in (SAMPLE, BLANK, SIMPLE):
            self.assertEqual(errors_in(openpyxl.load_workbook(path, data_only=True)), [], path.name)

    def test_sample_and_blank(self):
        """The sample is set up but for Check (its plan has issues to show); the blank one has everything to do."""
        steps, progress = checklist(openpyxl.load_workbook(SAMPLE, data_only=True))
        self.assertEqual([s for s, v in steps.items() if v != "Done"], ["4. Check"])
        self.assertEqual(progress, "5 of 6 steps done. Next: 4. Check")
        steps, progress = checklist(openpyxl.load_workbook(BLANK, data_only=True))
        self.assertEqual(set(steps.values()), {"To do"})
        self.assertEqual(progress, "0 of 6 steps done. Next: 1. Your project")


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class SetupTicksTest(unittest.TestCase):
    """Filling in a blank workbook ticks the steps, without touching the checklist."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        src = Path(cls.tmp) / "Blank.xlsx"
        make(blank=True, sizes=small()).save(src)

        def fill(wb):
            named_cell(wb, "ProjectName").value = "Office move"
            named_cell(wb, "ProjectStartDate").value = dt.date(2026, 9, 7)
            ws = wb["Resources"]
            ws.cell(6, head(ws)["Resource Name"]).value = "Sam"
            ws = wb["Tasks"]
            ws.cell(6, head(ws)["Task Name"]).value = "Pack the boxes"
            ws.cell(6, head(ws)["Duration"]).value = 3
        cls.before = recalc(src, cls.tmp, status_date=False)
        cls.after = recalc(src, cls.tmp, edit=fill)

    @classmethod
    def tearDownClass(cls):
        classic_again()

    def test_ticks(self):
        steps, _ = checklist(self.before)
        self.assertEqual(set(steps.values()), {"To do"})
        steps, progress = checklist(self.after)
        for step in ("1. Your project", "2. People", "3. Tasks", "6. First status"):
            self.assertEqual(steps[step], "Done", step)
        self.assertEqual(steps["5. Baseline"], "To do")
        nxt = next(s for s in STEPS if steps[s] == "To do")
        self.assertEqual(progress, f"{sum(v == 'Done' for v in steps.values())} of 6 steps done. Next: {nxt}")


if __name__ == "__main__":
    unittest.main()
