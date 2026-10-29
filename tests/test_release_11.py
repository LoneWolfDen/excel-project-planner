"""Template 1.1 (the --release 1.1 build, its workbooks in release-1.1/): RACI (F-1101), external dependencies on
Status (F-1102), the Sandbox what-if (F-1103), the automatic summary (F-1104) and three-point estimates with P50 /
P80 (F-1105), on top of everything Release 1.0 has (test_release.py). Release 1.0 and the classic workbooks don't
change with any of this.

Run:  python -m unittest tests.test_release_11 -v      (the recalculation tests need LibreOffice)
"""
import io
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(ROOT))
import perf_check  # noqa: E402
from test_planner import SOFFICE, CommittedWorkbooksTest, d, errors_in, named_cell, recalc, table_rows  # noqa: E402
from test_release import head, put, small, task_row, tasks_by_uid  # noqa: E402

RELEASE = ROOT / "release-1.1"
SAMPLE, BLANK, SIMPLE = (RELEASE / n for n in ("Project_Planner.xlsx", "Project_Planner_Blank.xlsx",
                                               "Project_Planner_Simple.xlsx"))
V10 = ROOT / "release-1.0" / "Project_Planner.xlsx"
SCHEDULE = ("Start", "Finish", "Planned Start", "Planned Finish", "Late Start", "Late Finish", "Total Slack",
            "Critical", "Work", "Resource Names", "Predecessors by ID")


def make(**k):
    from planner.build import make_workbook
    return make_workbook(release_edition="1.1", **k)


def classic_again():
    from planner.build import make_workbook
    make_workbook(sizes=small())


class StructureTest(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.path = Path(cls.tmp.name) / "v11.xlsx"
        make(baselined=True).save(cls.path)
        cls.wb = openpyxl.load_workbook(cls.path)
        cls.report = perf_check.static_report(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()
        classic_again()

    def test_versions(self):
        names = self.wb.defined_names
        self.assertEqual((names["TemplateVersion"].attr_text, names["PlannerVersion"].attr_text), ('"1.1"', '"2.1"'))
        self.assertIn("template 1.1", self.wb["Status"].oddFooter.left.text)

    def test_formula_rules(self):
        r = self.report
        self.assertEqual((r["a1_references"], r["missing_references"], r["excel_limit_problems"]), ([], [], []))
        self.assertEqual(r["volatile"], {"Tasks": {"OFFSET": 2200}})
        self.assertLessEqual(sum(v["rules"] for v in r["cf"].values()), 120)
        self.assertLess(sum(r["formulas"].values()), 50000)

    def test_new_columns(self):
        ws = self.wb["Tasks"]
        col = head(ws)
        for h in ("Accountable", "Consulted", "Informed", "Expected", "Optimistic", "Pessimistic",
                  "Sandbox Duration", "_PertVariance", "_PertShift"):
            self.assertIn(h, col, h)
        hidden = lambda h: ws.column_dimensions[ws.cell(5, col[h]).column_letter].hidden
        self.assertFalse(hidden("Accountable") or hidden("Expected"))           # the +/- columns stay visible
        self.assertTrue(hidden("Consulted") and hidden("Optimistic") and hidden("Sandbox Duration"))
        r = task_row(self.wb, 13)
        self.assertEqual([ws.cell(r, col[h]).value for h in ("Accountable", "Optimistic", "Pessimistic",
                                                             "Sandbox Duration")], ["Jordan Lee", 8, 16, 7])
        res = head(self.wb["Resources"])
        self.assertTrue({"Accountable For", "Consulted On", "Informed Of"} <= set(res))
        log = head(self.wb["Log"], 4)
        self.assertTrue({"Expected", "Task Start", "Gap (days)", "_DepKey"} <= set(log))
        for n in ("Sandbox", "PertP50", "PertP80", "sts_Narrative", "log_DepKey", "tsk_Accountable"):
            self.assertIn(n, self.wb.defined_names, n)
        self.assertIn("TopDependencies", self.wb["Status"].tables)

    def test_every_new_input_has_a_prompt(self):
        ws = self.wb["Tasks"]
        col = head(ws)
        prompted = {c for dv in ws.data_validations.dataValidation if dv.showInputMessage and dv.prompt
                    for rng in str(dv.sqref).split() for c in [openpyxl.utils.cell.range_boundaries(rng)[0]]}
        for h in ("Accountable", "Consulted", "Informed", "Optimistic", "Pessimistic", "Sandbox Duration"):
            self.assertIn(col[h], prompted, h)


class ModulesTest(unittest.TestCase):

    @classmethod
    def tearDownClass(cls):
        classic_again()

    def test_combinations(self):
        from planner.core import MODULES, RELEASE_MODULES
        names = list(MODULES) + list(RELEASE_MODULES)
        for combo in [(), ("log",), ("check",), ("check", "log"), tuple(n for n in names if n != "log"),
                      tuple(names)]:
            wb = make(sizes=small(), modules=combo)
            self.assertEqual(perf_check.missing_references(wb), [], combo)

    def test_editions_dont_leak(self):
        """Building 1.1 then 1.0 then classic in one process: each is its own edition."""
        from planner.build import make_workbook
        make(sizes=small())
        v10 = make_workbook(sizes=small(), release_edition=True)
        self.assertNotIn("Accountable", head(v10["Tasks"]))
        self.assertNotIn("Sandbox", v10.defined_names)
        self.assertEqual(v10.defined_names["TemplateVersion"].attr_text, '"1.0"')
        classic = make_workbook(sizes=small())
        self.assertNotIn("Pred1", head(classic["Tasks"]))


class CommittedTest(unittest.TestCase):
    compare = CommittedWorkbooksTest.compare

    def test_sample(self):
        self.compare(SAMPLE, "--release", "1.1", "--baselined")

    def test_blank(self):
        self.compare(BLANK, "--release", "1.1", "--blank")

    def test_simple(self):
        self.compare(SIMPLE, "--release", "1.1", "--baselined", "--simple")


class CommittedSampleTest(unittest.TestCase):
    """The committed 1.1 sample (Sandbox = No): no errors, the same plan as Release 1.0, and the new figures."""

    @classmethod
    def setUpClass(cls):
        cls.wb = openpyxl.load_workbook(SAMPLE, data_only=True)
        cls.t = tasks_by_uid(cls.wb)

    def test_no_errors(self):
        for path in (SAMPLE, BLANK, SIMPLE):
            self.assertEqual(errors_in(openpyxl.load_workbook(path, data_only=True)), [], path.name)

    def test_same_plan_as_release_10(self):
        old = tasks_by_uid(openpyxl.load_workbook(V10, data_only=True))
        self.assertEqual([(u, k) for u in old for k in SCHEDULE if old[u][k] != self.t[u][k]], [])

    def test_pert(self):
        t = self.t
        self.assertEqual(t[13]["Expected"], round((8 + 4 * 10 + 16) / 6, 1))
        crit = [r for r in t.values() if r["Critical"] == "Yes" and r["Expected"] not in (None, "")]
        shift = sum(r["Expected"] - r["_Days"] for r in crit)
        sd = sum(((r["Pessimistic"] - r["Optimistic"]) / 6) ** 2 for r in crit) ** 0.5
        hol = {d(r["Date"]) for r in table_rows(self.wb, "Settings", "HolidayTable") if r["Date"]}
        finish = d(named_cell(self.wb, "ProjectFinish").value)

        def add(day, n):
            import datetime as dt
            while n:
                day += dt.timedelta(days=1 if n > 0 else -1)
                if day.weekday() < 5 and day not in hol:
                    n += -1 if n > 0 else 1
            return day
        self.assertEqual(d(named_cell(self.wb, "PertP50").value), add(finish, round(shift)))
        self.assertEqual(d(named_cell(self.wb, "PertP80").value), add(finish, round(shift + 0.84 * sd)))
        kpis = {r["Measure"]: r["Value"] for r in table_rows(self.wb, "Status", "KPIs")}
        self.assertEqual(d(kpis["Finish P80 (three-point estimates)"]), d(named_cell(self.wb, "PertP80").value))

    def test_dependencies(self):
        deps = [r for r in table_rows(self.wb, "Status", "TopDependencies") if r["Dependency"]]
        self.assertEqual([r["Dependency"].split(":")[0] for r in deps], ["Dependency 13", "Dependency 5"])
        self.assertLess(deps[0]["Gap (days)"], 0)                         # arrives after the task should start
        log = {r["Ref"]: r for r in table_rows(self.wb, "Log", "Log")}
        self.assertEqual(d(log[13]["Task Start"]), d(self.t[14]["Start"]))

    def test_raci_counts(self):
        res = {r["Resource Name"]: r for r in table_rows(self.wb, "Resources", "ResourceTable") if r["Resource Name"]}
        self.assertEqual((res["Priya Sharma"]["Accountable For"], res["Priya Sharma"]["Informed Of"]), (2, 2))
        self.assertEqual(res["Jordan Lee"]["Consulted On"], 1)               # "Jordan Lee, Sam Rivera"
        self.assertEqual(res["Sam Rivera"]["Consulted On"], 1)

    def test_narrative(self):
        text = named_cell(self.wb, "sts_Narrative").value
        self.assertTrue(text.startswith("Overall "), text)
        for part in ("Forecast finish 06-Jan-2027", "Next milestone: ", "Top risk: ",
                     "1 external dependency arrives after"):
            self.assertIn(part, text)
        self.assertNotIn("SANDBOX", text)


class P2Test(unittest.TestCase):
    """The v1.1 P2 items: risk exposure in History (F-1106), the capacity chart (F-1107), fewer conditional-formatting
    rules with the same meaning (F-1109)."""

    @classmethod
    def tearDownClass(cls):
        classic_again()

    def test_risk_exposure(self):
        wb = openpyxl.load_workbook(SAMPLE, data_only=True)
        log = [r for r in table_rows(wb, "Log", "Log") if r["Title"]]
        exposure = sum(r["Score"] for r in log if r["_Open"] == 1 and isinstance(r["Score"], (int, float)))
        snap = table_rows(wb, "TakeSnapshot", "SnapshotRow")[0]
        self.assertEqual(snap["Risk Exposure"], exposure)
        self.assertEqual(table_rows(wb, "Status", "SnapshotNow")[0]["Risk Exposure"], exposure)
        history = [r["Risk Exposure"] for r in table_rows(wb, "History", "StatusHistory") if r["Status Date"]]
        self.assertEqual(history, [44, 50, 47, 52])
        titles = [c.title.tx.rich.p[0].r[0].t for c in openpyxl.load_workbook(SAMPLE)["History"]._charts]
        self.assertIn("Risk exposure (open RAID score)", titles)

    def test_capacity_chart(self):
        from planner.build import make_workbook
        charts = lambda wb: [c.title.tx.rich.p[0].r[0].t for c in wb["Agile"]._charts]
        self.assertIn("Available vs planned days per sprint", charts(make(sizes=small())))
        self.assertNotIn("Available vs planned days per sprint",
                         charts(make_workbook(sizes=small(), release_edition=True)))

    def test_merged_rules_mean_the_same(self):
        """Every cell gets the same rules, in the same order, with and without the merge."""
        from unittest import mock
        from openpyxl.utils.cell import range_boundaries
        from planner import release

        def per_cell(wb):
            out = {}
            for ws in wb.worksheets:
                for cf in ws.conditional_formatting:
                    for rule in cf.rules:
                        for rng in str(cf.sqref).split():
                            c1, r1, c2, r2 = range_boundaries(rng)
                            for r in range(r1, min(r2, r1 + 3) + 1):        # the first rows: enough to compare
                                for c in range(c1, c2 + 1):
                                    out.setdefault((ws.title, r, c), []).append(
                                        (rule.priority, rule.type, tuple(rule.formula or ()), repr(rule.dxf)))
            return {k: [x[1:] for x in sorted(v)] for k, v in out.items()}
        merged = make(sizes=small())
        with mock.patch.object(release, "merge_rules", lambda ws: None):
            plain = make(sizes=small())
        count = lambda wb: sum(len(cf.rules) for ws in wb.worksheets for cf in ws.conditional_formatting)
        self.assertLess(count(merged), count(plain))
        self.assertEqual(per_cell(merged), per_cell(plain))

    def test_committed_rule_count(self):
        report = perf_check.static_report(SAMPLE)
        # 106 after F-1109, plus R1-18's highlight of leave typed twice on the day grid
        self.assertLessEqual(sum(v["rules"] for v in report["cf"].values()), 107)


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class SandboxTest(unittest.TestCase):
    """Settings > Sandbox = Yes: the Sandbox Durations drive the plan, and Status, Check and the summary say so."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.dir = Path(cls.tmp.name)
        path = cls.dir / "v11.xlsx"
        make(baselined=True).save(path)

        def edit(wb):
            named_cell(wb, "Sandbox").value = "Yes"
            put(wb, 17, {"Optimistic": 14, "Sandbox Duration": "2days"})
        cls.wb = recalc(path, cls.dir, edit)
        cls.t = tasks_by_uid(cls.wb)
        cls.base = tasks_by_uid(openpyxl.load_workbook(SAMPLE, data_only=True))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()
        classic_again()

    def test_no_errors(self):
        self.assertEqual(errors_in(self.wb), [])

    def test_sandbox_drives_the_plan(self):
        self.assertEqual(self.t[13]["_Days"], 7)
        self.assertLess(d(self.t[13]["Finish"]), d(self.base[13]["Finish"]))
        self.assertLess(d(self.t[14]["Start"]), d(self.base[14]["Start"]))      # its successor moves too
        self.assertEqual(self.t[17]["_Days"], self.base[17]["_Days"])        # an unreadable sandbox value: ignored
        self.assertIn("Sandbox Duration 2days can't be read", self.t[17]["Issues"])
        self.assertIn("Optimistic should be at most the Duration", self.t[17]["Issues"])

    def test_it_is_flagged(self):
        self.assertIn("SANDBOX ON", self.wb["Status"]["B3"].value)
        self.assertIn("SANDBOX", named_cell(self.wb, "sts_Narrative").value)
        counts = {r[0].value: r[1].value for r in self.wb["Check"].iter_rows(min_row=5, max_row=16, min_col=2,
                                                                             max_col=3)}
        self.assertEqual(counts["Sandbox is on (what-if dates)"], 1)


class ToolsTest(unittest.TestCase):

    @classmethod
    def tearDownClass(cls):
        classic_again()

    def test_upgrade_keeps_the_11_inputs(self):
        from planner import upgrade
        wb = openpyxl.load_workbook(SAMPLE)
        put(wb, 21, {"Informed": "Sponsor"})
        named_cell(wb, "Sandbox").value = "Yes"
        buf = io.BytesIO()
        wb.save(buf)
        new = openpyxl.load_workbook(io.BytesIO(upgrade.convert(buf.getvalue())[0]))
        self.assertEqual(new.defined_names["PlannerVersion"].attr_text, '"2.1"')
        col, r = head(new["Tasks"]), task_row(new, 21)
        self.assertEqual([new["Tasks"].cell(r, col[h]).value for h in ("Informed", "Optimistic", "Pred1")],
                         ["Sponsor", 6, 19])
        self.assertEqual(named_cell(new, "Sandbox").value, "Yes")
        log = {x["Ref"]: x for x in table_rows(new, "Log", "Log")}
        import datetime as dt
        self.assertEqual(d(log[13]["Expected"]), dt.date(2026, 11, 4))

    def test_release_10_moves_up(self):
        from planner import upgrade
        new = openpyxl.load_workbook(io.BytesIO(upgrade.convert(V10.read_bytes(), release="1.1")[0]))
        self.assertEqual(new.defined_names["TemplateVersion"].attr_text, '"1.1"')
        col = head(new["Tasks"])
        self.assertIn("Accountable", col)
        self.assertEqual(new["Tasks"].cell(task_row(new, 17), col["Pred2"]).value, 13)


if __name__ == "__main__":
    unittest.main()
