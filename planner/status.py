"""Status sheet: the weekly page. The status date and this week's checklist at the top, then everything a steering
meeting reads: overall RAG and RAG by area, key figures, tasks by status, next milestones, the top open RAID items,
the schedule watch lists (late, due and starting in the look-ahead, open critical: B-1008, which replaced the
Schedule Report), change control and the period commentary, printable on one page. The snapshot row to paste onto History sits at
the bottom. Money is on the Cost sheet, which nothing here reads, so it can be deleted before sharing.

Each RAG is calculated from the plan against the Settings thresholds. The project manager can override it (PM RAG)
and explain why (Comment); the RAG column shows the override when there is one, which is how most PMOs report.
The schedule measures have defined names (sch_Progress, sch_Planned, sch_SPI: duration-weighted, so they need no
cost) that History reuses. The status date cell itself is written by planner.settings (StatusDateInput), like every other setting.
"""
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

from . import rag
from .core import (DATE, HEADER_FILLS, HEADER_FONT, NOTE_FONT, TITLE_FONT, Col, Show, Tab, add_table, fill, font, has,
                   missing, name, ranked_table)
from .history import fields
from .settings import STATUS as SHEET, STATUS_ROW
from .tasks import STATUS_COLOURS

WIDTHS = {"A": 3, "B": 36, "C": 56, "D": 16, "E": 12, "F": 10, "G": 10, "H": 50}
MONEY = "#,##0"

# this week's steps, under the status date: (label, text, sheet to link to)
WEEKLY = [
    ("2. Progress", "On Tasks: the Stage, Actual Start and Actual Finish for the tasks that moved.", "Tasks"),
    ("3. Log", "Update risks, issues, decisions and change requests.", "Log"),
    ("4. RAG", "Review the RAG below; where you disagree, pick a PM RAG and say why in Comment.", None),
    ("5. Commentary", "Write this period's notes in the COMMENTARY table below.", None),
    ("6. Snapshot", "Copy the SnapshotNow row at the bottom of this sheet and Paste Special > Values under the last "
                    "row of History's table.", "History"),
]

# area, measure, value, severity of the value (0 green, 1 amber, 2 red), number format
AREAS = [
    ("Schedule", "Forecast finish vs plan (working days)", "FinishVarianceDays", rag.slip("{v}"), "+0;-0;0"),
    ("Schedule performance", "SPI = progress / progress planned by now (duration-weighted)", "sch_SPI",
     rag.index("{v}"), "0.00"),
    ("Deadlines", "Tasks forecast to miss their deadline", 'COUNTIF(tsk_DeadlineCheck,"Late*")', rag.count("{v}", 2), "0"),
    ("Critical path", "Critical tasks forecast to finish late",
     'COUNTIFS(tsk_Critical,"Yes",tsk_IsLeaf,1,tsk_Variance,">0")', rag.count("{v}", 1), "0"),
    ("RAID", "Open high-rated risks, issues, dependencies, assumptions",
     'COUNTIFS(log_Rating,"High",log_Open,1)', rag.count("{v}", 1), "0"),
    ("Plan quality", "Ignored links and outline errors (see Tasks)", 'COUNTIF(tsk_LinkCheck,"?*")+COUNTIF(tsk_WBS,"~?")',
     rag.count("{v}", 1), "0"),
]

# label, formula, format, where it comes from
KPIS = [
    ("Planned start", "MIN(tsk_PlanStart)", DATE, "Tasks"),
    ("Planned finish", "MAX(tsk_PlanFinish)", DATE, "Tasks (the baseline once it is set)"),
    ("Forecast finish", "MAX(tsk_FcFinish)", DATE, "Tasks"),
    ("Overall progress (duration-weighted)", "IFERROR(SUM(tsk_WtDone)/SUM(tsk_WtDur),0)", "0%", "Tasks"),
    ("Planned progress by status date", "IFERROR(SUM(tsk_WtPlanned)/SUM(tsk_WtDur),0)", "0%", "Tasks"),
    ("Schedule performance (SPI)", 'IF(sch_Planned>0,sch_Progress/sch_Planned,"")', "0.00",
     "Progress / planned progress"),
    ("Detail tasks", "SUM(tsk_IsLeaf)", "0", "Tasks"),
    ("Milestones", 'SUMPRODUCT(tsk_IsLeaf*(tsk_Duration=0)*(tsk_Name<>""))', "0", "Tasks"),
    ("Total work (hrs)", "SUMPRODUCT(tsk_Work,tsk_IsLeaf)", "#,##0", "Tasks"),
    ("Baseline", 'IF(COUNT(tsk_BaseStart)=0,"Not set (planned = calculated)","Set on "&COUNT(tsk_BaseStart)&" tasks")',
     "@", "Tasks > BASELINE"),
    ("Open critical tasks", 'COUNTIFS(tsk_Critical,"Yes",tsk_IsLeaf,1)', "0", "Tasks"),
    ("Tasks with issues (see Tasks > Issues)", 'COUNTIF(tsk_Issues,"?*")', "0", "Tasks > Issues"),
    ("Tasks with an over-allocated resource", 'COUNTIF(tsk_OverAllocated,"Yes")', "0", "Resource Report"),
    ("Current sprint", 'IF(CurrentSprint="","(outside sprint range)",CurrentSprint)', "@", "Agile"),
    ("Story points done / total", 'IF(TotalPoints=0,"-",agl_Done&" / "&TotalPoints)', "@", "Agile"),
    ("Open risks / issues",
     'COUNTIFS(log_Type,"Risk",log_Open,1)&" / "&COUNTIFS(log_Type,"Issue",log_Open,1)', "@", "Log"),
]
ALERT_IF_POSITIVE = {"Tasks with issues (see Tasks > Issues)", "Tasks with an over-allocated resource"}



def changes(status):
    """Change requests in one of the given statuses (the Log's Change rows)."""
    return "+".join(f'COUNTIFS(log_Type,"Change",log_Status,"{s}")' for s in status)


def impact(column, status):
    return "+".join(f'SUMIFS({column},log_Type,"Change",log_Status,"{s}")' for s in status)


PENDING, AGREED = ("Submitted", "Under Review"), ("Approved", "Implemented")
# label, formula, format
CHANGE_CONTROL = [
    ("Change requests awaiting decision", changes(PENDING), "0"),
    ("Change requests approved or implemented", changes(AGREED), "0"),
    ("Approved schedule impact (working days)", impact("log_Days", AGREED), "+0;-0;0"),
    ("Approved cost impact", impact("log_Cost", AGREED), MONEY),
    ("Cost impact awaiting decision", impact("log_Cost", PENDING), MONEY),
    ("Decisions awaiting approval", 'COUNTIFS(log_Type,"Decision",log_Status,"Proposed")', "0"),
    ("Overdue log items", 'COUNTIF(log_Overdue,"Overdue")', "0"),
]
TOP_RAID_ROWS = 8

NOTES = ["Summary", "Achievements this period", "Planned for next period", "Issues and escalations",
         "Decisions needed", "Risks to watch"]


def _heading(ws, row, text):
    ws.cell(row, 2, text).font = font(bold=True, color="1F4E78")


def _rag_columns(table, measures, values, severities):
    t = Tab(table)
    return [
        Col("Area", width=WIDTHS["B"], values=measures[0]),
        Col("Measure", width=WIDTHS["C"], values=measures[1]),
        Col("Value", width=WIDTHS["D"], values=values),
        Col("Calculated", width=WIDTHS["E"], values=severities, center=True,
            note="RAG from the thresholds on Settings (RAGAmberSlipDays, RAGRedSlipDays, RAGAmberIndex, RAGRedIndex)."),
        Col("PM RAG", "in", WIDTHS["F"], center=True, note="Optional override by the project manager: Green, Amber or Red."),
        Col("RAG", width=WIDTHS["G"], center=True,
            formula=f'IF({t.row("PM RAG")}<>"",{t.row("PM RAG")},{t.row("Calculated")})',
            note="Reported RAG: the PM override if there is one, otherwise the calculated RAG."),
        Col("Comment", "in", WIDTHS["H"], note="Why: especially when the PM RAG differs from the calculated one."),
    ]


def build_status(wb, agile_ws=None, sprint_letters=None, n_chart_sprints=12):
    """agile_ws / sprint_letters: the Agile sheet's SprintTable, for the sprint load chart (left out without it)."""
    ws = wb[SHEET]                                  # created by build_settings, which wrote the status date
    # the title sits in column B: column A holds the top-RAID list's hidden helper column
    ws["B1"] = '=ProjectName&"  -  Status"'
    ws["B1"].font = TITLE_FONT
    ws["B2"] = ('="Project manager: "&ProjectManager&"   Forecast finish "&TEXT(ProjectFinish,"ddd dd-mmm-yyyy")&'
                '" ("&TEXT(FinishVarianceDays,"+0;-0;0")&" working days vs plan)"')
    ws["B2"].font = NOTE_FONT
    ws["B3"] = ('=IF(LiveGrids="No","Resource grids are OFF (Settings > Live resource grids): the Resource Report '
                'and Tasks > Over-allocated are not up to date.","")')
    ws["B3"].font = font(bold=True, color="C00000")
    for col, width in WIDTHS.items():
        ws.column_dimensions[col].width = width
    top = _this_week(ws)
    top = _rag(wb, ws, top)
    top = _key_figures(wb, ws, top)
    if has("log"):
        top = _top_raid(ws, top)
    top = _schedule_watch(ws, top)
    chart_at = top
    if has("log"):
        top = _change_control(ws, top)
    end = _commentary(ws, top)
    if agile_ws is not None:
        _sprint_chart(ws, agile_ws, sprint_letters, n_chart_sprints, chart_at)
    if has("history"):
        _snapshot(ws, end + 3)
    ws.print_area = f"B1:H{end}"                    # one page: everything above the snapshot row
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.freeze_panes = "A4"
    return ws


def _this_week(ws):
    """The status date (row STATUS_ROW, written by settings) as step 1, then the rest of the weekly update."""
    r = STATUS_ROW - 1
    for c in (2, 3):
        ws.cell(r, c).fill = HEADER_FILLS["in"]
    ws.cell(r, 2, "THIS WEEK").font = HEADER_FONT
    ws.cell(r, 3, "Every reporting date, top to bottom").font = HEADER_FONT
    ws.cell(STATUS_ROW, 2).value = "1. Status date"
    ws.cell(STATUS_ROW, 4, "Type the reporting date (Ctrl+; = today)").font = NOTE_FONT
    for i, (label, body, sheet) in enumerate(s for s in WEEKLY if not missing(s[2])):
        r = STATUS_ROW + 1 + i
        ws.cell(r, 2, label).font = font(bold=True)
        c = ws.cell(r, 3, body + (f"  Go to {sheet}." if sheet else ""))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if sheet:                                   # a link inside the workbook (a hyperlink, not a formula)
            c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{sheet}'!A1")
            c.font = font(color="1F4E78", underline="single")
        ws.row_dimensions[r].height = 30
    return STATUS_ROW + len(WEEKLY) + 3


def _live(rows, at):
    """The rows whose formula (column `at`) this build can calculate (B-913: optional modules)."""
    return [r for r in rows if not missing(r[at])]


def _rag(wb, ws, top):
    shown = _live(AREAS, 2)
    areas = Tab("StatusRAG")
    worst = f'MAX(IF(COUNTIF({areas.col("RAG")},"Red")>0,2,0),IF(COUNTIF({areas.col("RAG")},"Amber")>0,1,0))'
    overall_value = f'=COUNTIF({areas.col("RAG")},"Red")&" red, "&COUNTIF({areas.col("RAG")},"Amber")&" amber"'
    overall = top
    add_table(ws, "OverallStatus", overall, 2, _rag_columns(
        "OverallStatus", (["Overall"], ["Worst of the areas below"]), [overall_value], [f"={rag.rag_text(worst)}"]), 1,
        autofilter=False)
    name(wb, "ProjectRAG", "OverallStatus[RAG]")
    top = overall + 3
    add_table(ws, "StatusRAG", top, 2, _rag_columns(
        "StatusRAG", ([a[0] for a in shown], [a[1] for a in shown]), [f"={a[2]}" for a in shown],
        [f"={rag.rag_text(a[3].format(v=areas.row('Value')))}" for a in shown]), len(shown), autofilter=False)
    for i, a in enumerate(shown):
        ws.cell(top + 1 + i, 4).number_format = a[4]
        ws.cell(top + 1 + i, 4).alignment = Alignment(horizontal="right")
    for r in range(overall + 1, top + 1 + len(shown)):
        ws.cell(r, 2).font = font(bold=True)
    rag.colour(ws, f"E{overall + 1}:G{top + len(shown)}", f"E{overall + 1}")
    pick = DataValidation(type="list", formula1="=L_RAG")
    ws.add_data_validation(pick)
    pick.add(f"F{overall + 1}")
    pick.add(f"F{top + 1}:F{top + len(shown)}")
    return top + len(shown) + 3


# key figures other sheets read by name
KPI_NAMES = {"Overall progress (duration-weighted)": "sch_Progress", "Planned progress by status date": "sch_Planned",
             "Schedule performance (SPI)": "sch_SPI"}


def _key_figures(wb, ws, top):
    kpis = [k for k in _live(KPIS, 1) if not missing(k[3])]      # its source sheet left out: nothing to count
    """Key figures (B:D) with detail tasks by status beside them (E:F) and its chart below that, then the next
    milestones."""
    _heading(ws, top - 1, "KEY FIGURES")
    add_table(ws, "KPIs", top, 2, [
        Col("Measure", width=WIDTHS["B"], values=[k[0] for k in kpis]),
        Col("Source", width=WIDTHS["C"], values=[k[3] for k in kpis]),
        Col("Value", width=WIDTHS["D"], values=[f"={k[1]}" for k in kpis]),
    ], len(kpis), autofilter=False)
    for i, (label, _, fmt, _) in enumerate(kpis):
        r = top + 1 + i
        ws.cell(r, 2).font = font(bold=True)
        value = ws.cell(r, 4)
        value.number_format = fmt
        value.font = font(bold=True)
        value.alignment = Alignment(horizontal="right")
        if label in KPI_NAMES:
            name(wb, KPI_NAMES[label], f"{SHEET}!$D${r}")
        if label in ALERT_IF_POSITIVE:
            ws.conditional_formatting.add(f"D{r}", FormulaRule(formula=[f"D{r}>0"], font=Font(color="C00000", bold=True),
                                                               fill=fill("FFC7CE")))
            ws.conditional_formatting.add(f"D{r}", FormulaRule(formula=[f"D{r}<=0"], font=Font(color="548235", bold=True)))

    status = Tab("StatusSummary")
    add_table(ws, "StatusSummary", top, 5, [
        Col("Status", width=WIDTHS["E"], values=list(STATUS_COLOURS)),
        Col("Count", width=WIDTHS["F"], formula=f"COUNTIFS(tsk_Status,{status.row('Status')},tsk_IsLeaf,1)"),
    ], len(STATUS_COLOURS), autofilter=False)
    for i, colour in enumerate(STATUS_COLOURS.values()):
        ws.cell(top + 1 + i, 5).font = font(bold=True, color=colour)
        ws.cell(top + 1 + i, 6).font = font(bold=True)
    chart = BarChart()
    chart.type, chart.style, chart.title, chart.legend = "bar", 10, "Detail tasks by status", None
    chart.add_data(Reference(ws, min_col=6, min_row=top, max_row=top + len(STATUS_COLOURS)), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=5, min_row=top + 1, max_row=top + len(STATUS_COLOURS)))
    chart.height, chart.width = 6, 12
    ws.add_chart(chart, f"E{top + len(STATUS_COLOURS) + 2}")

    top += len(kpis) + 3
    nxt = Tab("NextMilestones")
    k = f"(ROW()-ROW({nxt.head('Next milestones (forecast)')}))"
    nth = f"SMALL(tsk_MilestoneKey,{k})"
    add_table(ws, "NextMilestones", top, 2, [
        Col("Next milestones (forecast)", width=WIDTHS["B"],
            formula=f'IFERROR(INDEX(tsk_Name,MATCH({nth},tsk_MilestoneKey,0)),"")'),
        Col("Forecast date", width=WIDTHS["C"], fmt=DATE,
            formula=f'IFERROR(INDEX(tsk_FcFinish,MATCH({nth},tsk_MilestoneKey,0)),"")'),
    ], 5, autofilter=False)
    for r in range(top + 1, top + 6):
        ws.cell(r, 3).alignment = Alignment(horizontal="left")
    return top + 5 + 3


def _top_raid(ws, top):
    """Top open RAID items by score (the list's hidden _Task column is the margin, column A)."""
    ranked_table(ws, "TopRAID", top, 1, "log_TopKey", TOP_RAID_ROWS, [
        Show("Item", "log_Label", width=WIDTHS["B"]), Show("Response", "log_Response", width=WIDTHS["C"]),
        Show("Owner", "log_Owner", width=WIDTHS["D"]), Show("Due", "log_Due", "date", WIDTHS["E"], "dd-mmm-yy"),
        Show("Score", "log_Score", "num", WIDTHS["F"]), Show("Rating", "log_Rating", width=WIDTHS["G"]),
        Show("Task", "log_Task", width=WIDTHS["H"]),
    ], largest=True, label='"TOP OPEN RAID ITEMS (by score)"')
    ws.column_dimensions["A"].width = WIDTHS["A"]
    rag.colour(ws, f"G{top + 1}:G{top + TOP_RAID_ROWS}", f"G{top + 1}")
    return top + TOP_RAID_ROWS + 3


VAR = "+0;-0;0"
TASK, RESOURCE = Show("Task", "tsk_Name", width=WIDTHS["B"]), Show("Resources", "tsk_ResourceNames", width=WIDTHS["C"])
FINISH, START = (Show(h, n, "date", WIDTHS["D"], "ddd dd-mmm") for h, n in (("Finish", "tsk_FcFinish"),
                                                                            ("Start", "tsk_FcStart")))
CRITICAL = Show("Critical", "tsk_Critical", width=WIDTHS["G"])
WATCH_ROWS = 5
# table, key, largest first, label, columns B-G (the hidden _Task column is the margin, column A)
WATCH = [
    ("LateTasks", "tsk_LateKey", True, '"LATE TASKS (most days late first)"', [
        TASK, RESOURCE, FINISH, Show("Finish Var", "tsk_Variance", "num", WIDTHS["E"], VAR),
        Show("Status", "tsk_Status", width=WIDTHS["F"]), CRITICAL]),
    ("DueSoon", "tsk_DueSoonKey", False, '"DUE BY "&UPPER(TEXT(LookaheadEnd,"ddd dd-mmm"))', [
        TASK, RESOURCE, FINISH, Show("Progress", "tsk_Progress", "num", WIDTHS["E"], "0%"),
        Show("Deadline Check", "tsk_DeadlineCheck", width=WIDTHS["F"]), CRITICAL]),
    ("StartingSoon", "tsk_StartSoonKey", False,
     '"STARTING BY "&UPPER(TEXT(LookaheadEnd,"ddd dd-mmm"))&" (look-ahead "&LookaheadDays&" working days)"', [
        TASK, RESOURCE, START, Show("Duration", "tsk_Duration", "num", WIDTHS["E"]),
        Show("Predecessors", "tsk_Predecessors", width=WIDTHS["F"]), CRITICAL]),
    ("CriticalTasks", "tsk_CriticalKey", False, '"OPEN CRITICAL TASKS (by start)"', [
        TASK, RESOURCE, START, Show("Total Slack", "tsk_TotalSlack", "num", WIDTHS["E"]),
        Show("Deadline Check", "tsk_DeadlineCheck", width=WIDTHS["F"]), FINISH]),
]


def _schedule_watch(ws, top):
    """The first few tasks of each schedule list. Tasks works out which tasks belong on each list and in what order
    (its _LateKey, _DueSoonKey, ... columns); these only look them up. Filter Tasks for the whole list."""
    for table, key, largest, label, shows in WATCH:
        letters = ranked_table(ws, table, top, 1, key, WATCH_ROWS, shows, largest, label)
        first, last = top + 1, top + WATCH_ROWS
        if "Finish Var" in letters:
            c = letters["Finish Var"]
            ws.conditional_formatting.add(f"{c}{first}:{c}{last}", FormulaRule(
                formula=[f"AND(ISNUMBER({c}{first}),{c}{first}>0)"], font=Font(color="C00000", bold=True)))
        if "Critical" in letters:
            c = letters["Critical"]
            ws.conditional_formatting.add(f"{c}{first}:{c}{last}", FormulaRule(
                formula=[f'{c}{first}="Yes"'], font=Font(color="FFFFFF", bold=True), fill=fill("C00000")))
        if "Deadline Check" in letters:
            c = letters["Deadline Check"]
            ws.conditional_formatting.add(f"{c}{first}:{c}{last}", FormulaRule(
                formula=[f'LEFT({c}{first},4)="Late"'], font=Font(color="C00000", bold=True)))
        top = last + 3
    return top


def _change_control(ws, top):
    _heading(ws, top - 1, "CHANGE CONTROL AND DECISIONS (from the Log)")
    add_table(ws, "ChangeControl", top, 2, [
        Col("Measure", width=WIDTHS["B"], values=[c[0] for c in CHANGE_CONTROL]),
        Col("Source", width=WIDTHS["C"], values=["Log: Change rows"] * 5 + ["Log: Decision rows", "Log: every type"]),
        Col("Value", width=WIDTHS["D"], values=[f"={c[1]}" for c in CHANGE_CONTROL]),
    ], len(CHANGE_CONTROL), autofilter=False)
    for i, (_, _, fmt) in enumerate(CHANGE_CONTROL):
        ws.cell(top + 1 + i, 4).number_format = fmt
        ws.cell(top + 1 + i, 4).font = font(bold=True)
    return top + len(CHANGE_CONTROL) + 3


def _commentary(ws, top):
    _heading(ws, top - 1, "COMMENTARY for this period (type in the Notes column)")
    add_table(ws, "StatusNotes", top, 2, [
        Col("Section", width=WIDTHS["B"], values=NOTES),
        Col("Notes", "in", WIDTHS["C"]),
    ], len(NOTES), autofilter=False)
    for i in range(len(NOTES)):
        ws.cell(top + 1 + i, 3).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(top + 1 + i, 2).alignment = Alignment(vertical="top")
        ws.row_dimensions[top + 1 + i].height = 36
    return top + len(NOTES)


def _sprint_chart(ws, agile_ws, L, n, at):
    """Planned days vs available days for the first n sprints (charts need A1 ranges into SprintTable)."""
    from .agile import HEADER_ROW
    col = lambda key: agile_ws[f"{L[key]}{HEADER_ROW}"].column
    chart = BarChart()
    chart.title = "Sprint planned vs available person-days"
    for key in ("Planned Days", "Available Days"):
        chart.add_data(Reference(agile_ws, min_col=col(key), min_row=HEADER_ROW, max_row=HEADER_ROW + n),
                       titles_from_data=True)
    chart.set_categories(Reference(agile_ws, min_col=col("Sprint"), min_row=HEADER_ROW + 1, max_row=HEADER_ROW + n))
    chart.height, chart.width = 7, 16
    ws.add_chart(chart, f"E{at}")


def _snapshot(ws, top):
    """This status date's numbers in one row (formulas), to paste as values onto History's StatusHistory."""
    _heading(ws, top - 1, "SNAPSHOT: copy this row and Paste Special > Values under the last row of History's table")
    add_table(ws, "SnapshotNow", top, 2, [Col(h, width=11, fmt=fmt, formula=f, center=True) for _, (h, f, fmt) in fields()],
              1, autofilter=False)
    for col, width in WIDTHS.items():                # the snapshot shares the page's columns
        ws.column_dimensions[col].width = width
