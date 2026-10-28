"""Backlog sheet (template 1.2, B-405; optional module "backlog", not in Simple): the product backlog, the work not
yet planned, with a priority and story points, before it becomes tasks.

A row is an item: priority (MoSCoW), points, owner, a status (New, Ready, Planned, Done) and notes. "On Tasks" says
Yes once a task of the same name is on the Tasks sheet, so promoting an item is a copy of its name (and points) to a
new task row; the totals above the table show the points by priority and how much is ready. Nothing on any other
sheet reads it, so it can be deleted.
"""
from openpyxl.worksheet.datavalidation import DataValidation

from .core import BOX, CALC_FILL, CALC_FONT, LABEL_FONT, NOTE_FONT, TITLE_FONT, Col, Tab, add_table, col_range, \
    header_cell, name

SHEET = "Backlog"
ROWS = 100
PRIORITIES = ["Must", "Should", "Could", "Won't"]
STATES = ["New", "Ready", "Planned", "Done"]
SAMPLE = [("Saved payment cards", "Must", 8, "Jordan Lee", "Ready", "Needs the payment provider's vault."),
          ("Order history export (CSV)", "Should", 3, "Sam Rivera", "Ready", ""),
          ("Reporting screens (phase 2)", "Should", 13, "Morgan Patel", "New", "Moved out of release 1 (Decision 8)."),
          ("Dark mode", "Could", 5, "Casey Kim", "New", ""),
          ("Live chat on every page", "Won't", 8, "Sales", "New", "Rejected for go-live (Change 12)."),
          ("Accessibility review", "Must", 3, "Casey Kim", "Planned", "On Tasks as UID 18.")]


def build_backlog(wb, blank):
    ws = wb.create_sheet(SHEET)
    ws["B1"] = '=ProjectName&"  -  Product backlog"'
    ws["B1"].font = TITLE_FONT
    ws["B2"] = ("Work not yet planned. To plan an item, add a task with the same name on Tasks (and copy its points): "
                "On Tasks then says Yes. Set its Status to Planned.")
    ws["B2"].font = NOTE_FONT
    for col, width in {"A": 3, "B": 34, "C": 10}.items():
        ws.column_dimensions[col].width = width
    data = [] if blank else SAMPLE
    top = 5 + len(PRIORITIES) + 3
    header_cell(ws.cell(4, 2), "Priority")
    header_cell(ws.cell(4, 3), "Points")
    for i, p in enumerate(PRIORITIES + ["Ready to plan"], start=5):
        ws.cell(i, 2, p).font = LABEL_FONT
        formula = ('=SUMIFS(bkl_Points,bkl_Status,"Ready")' if p == "Ready to plan" else
                   f'=SUMIFS(bkl_Points,bkl_Priority,"{p}")')
        c = ws.cell(i, 3, formula)
        c.font, c.fill, c.border, c.number_format = CALC_FONT, CALC_FILL, BOX, "0.#"
    t = Tab("ProductBacklog")
    item = t.row("Item")
    letters = add_table(ws, "ProductBacklog", top, 2, [
        Col("Item", "in", 40, values=[d[0] for d in data], note="What the team would build."),
        Col("Priority", "in", 10, values=[d[1] for d in data], center=True, note="Must, Should, Could or Won't."),
        Col("Story Points", "in", 9, "0.#", values=[d[2] for d in data], center=True),
        Col("Owner", "in", 18, values=[d[3] for d in data]),
        Col("Status", "in", 10, values=[d[4] for d in data], center=True, note="New, Ready, Planned or Done."),
        Col("Notes", "in", 44, values=[d[5] for d in data]),
        Col("On Tasks", width=9, center=True, formula=f'IF({item}="","",IF(COUNTIF(tsk_Name,{item})>0,"Yes",""))',
            note="Yes once a task of the same name is on Tasks."),
    ], ROWS)
    for column, defined in (("Priority", "bkl_Priority"), ("Story Points", "bkl_Points"), ("Status", "bkl_Status")):
        name(wb, defined, f"ProductBacklog[{column}]")
    first, last = top + 1, top + ROWS
    for key, values in (("Priority", PRIORITIES), ("Status", STATES)):
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"')
        ws.add_data_validation(dv)
        dv.add(col_range(letters[key], first, last))
    dv = DataValidation(type="list", formula1="=res_Name", showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(col_range(letters["Owner"], first, last))
    ws.freeze_panes = ws.cell(first, 3)
    return ws
