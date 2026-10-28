"""Cost sheet: the rate card, each task's cost, resource cost against task cost, earned value and the cost history.

Rates are often sensitive and not every audience needs cost, so all money lives here, and the dependency runs one
way: this sheet reads Tasks, Resources and Settings, and **nothing else reads this sheet**. Delete it before
sharing the plan and every other sheet still calculates (tests delete it and recalculate).

- Rate card: a Day Rate per person (typed against their name, so moving rows on Resources can't misalign them).
- Task cost = each person's work on the task, in days, x their day rate. Actual cost = work done (work x
  progress) x the rate, as MS Project calculates it.
- Resource cost = the person's allocated days on the project (working days between their Start / End, or the
  project's, x Allocation %) x the rate: what the team costs to run, advisory roles included. The gap to task cost
  is time not assigned to tasks (overhead, ceremonies, advisory work).
- Earned value: budget = cost at the baseline's work once a baseline is set (Baseline Work x the task's rate per
  hour), otherwise the current cost; planned value = budget x the share planned by the status date; earned value
  = budget x progress.

The task rows follow Tasks by position, like the Gantt: every cell is a formula, so inserting a task on Tasks
just recalculates them.
"""
import datetime as dt

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from . import rag
from .core import DATE, NOTE_FONT, TITLE_FONT, Col, Tab, add_table, fill, font, name, v12
from .resources import by_resource
from .tasks import PEOPLE, people

SHEET = "Cost"
MONEY = "#,##0"
WD = "WorkWeek,Holidays"
HISTORY_ROWS = 60
FIXED_ROWS = 30                    # template 1.2 (B-302): rows of the FixedCosts table

# defined name, label, formula, meaning, format
EVM = [
    ("evm_BAC", "Budget at completion (BAC)", "SUM(TaskCost[_LeafBudget])",
     "Cost at the baseline's work once a baseline is set, otherwise the current cost.", MONEY),
    ("evm_PV", "Planned value (PV)", "SUM(TaskCost[_LeafPV])", "Budget of the work planned before the status date.",
     MONEY),
    ("evm_EV", "Earned value (EV)", "SUM(TaskCost[_LeafEV])", "Budget of the work done (budget x progress).", MONEY),
    ("evm_AC", "Actual cost (AC)", "SUM(TaskCost[_LeafActual])", "Work done x day rates.", MONEY),
    ("evm_SV", "Schedule variance (SV)", "evm_EV-evm_PV", "EV - PV. Negative = behind plan.", MONEY),
    ("evm_CV", "Cost variance (CV)", "evm_EV-evm_AC", "EV - AC. Negative = over cost.", MONEY),
    ("evm_SPI", "Schedule performance index (SPI, cost-weighted)", 'IF(evm_PV>0,evm_EV/evm_PV,"")',
     "EV / PV. Below 1 = behind (the Status sheet's SPI is duration-weighted).", "0.00"),
    ("evm_CPI", "Cost performance index (CPI)", 'IF(evm_AC>0,evm_EV/evm_AC,"")', "EV / AC. Below 1 = over cost.",
     "0.00"),
    ("evm_EAC", "Estimate at completion (EAC)", "IF(N(evm_CPI)>0,evm_AC+(evm_BAC-evm_EV)/evm_CPI,evm_AC+evm_BAC-evm_EV)",
     "AC + (BAC - EV) / CPI: forecast cost at the current CPI.", MONEY),
    ("evm_ETC", "Estimate to complete (ETC)", "evm_EAC-evm_AC", "EAC - AC: forecast cost of the remaining work.",
     MONEY),
    ("evm_VAC", "Variance at completion (VAC)", "evm_BAC-evm_EAC", "BAC - EAC. Negative = forecast overrun.", MONEY),
    ("evm_TCPI", "To-complete performance index (TCPI)", 'IF(evm_BAC>evm_AC,(evm_BAC-evm_EV)/(evm_BAC-evm_AC),"")',
     "CPI the remaining work needs to finish on budget.", "0.00"),
    ("cost_Task", "Task cost (assigned work)", "SUM(RateCard[Task Cost])",
     "Every task's work x day rates: what the plan's work costs.", MONEY),
    ("cost_Resource", "Resource cost (allocated days)", "SUM(RateCard[Resource Cost])",
     "Everyone's allocated days x day rates, advisory roles included: what the team costs to run.", MONEY),
    ("cost_Gap", "Not assigned to tasks", "cost_Resource-cost_Task",
     "Resource cost - task cost: overhead, ceremonies, advisory time (negative = more work than allocation).", MONEY),
]
# header, formula (names on this sheet only), number format
HISTORY = [("Status Date", "StatusDate", DATE), ("BAC", "evm_BAC", MONEY), ("PV", "evm_PV", MONEY),
           ("EV", "evm_EV", MONEY), ("AC", "evm_AC", MONEY), ("CPI", "evm_CPI", "0.00"), ("EAC", "evm_EAC", MONEY)]


def rate(person):
    """A person's day rate from the rate card (0 when they have none)."""
    return f"IFERROR(INDEX(cost_Rate,MATCH({person},cost_RateName,0)),0)"


def per_use(person):
    """Template 1.2 (B-302): a person's cost per use, charged once for each task they are on (MS Project's)."""
    return f"IFERROR(N(INDEX(cost_PerUse,MATCH({person},cost_RateName,0))),0)"


def build_cost(wb, blank, n_tasks, resources, n_rates=50):
    """resources: the (name, initials, group, allocation, day rate, ...) rows the workbook was built with; their
    names and day rates fill the rate card."""
    ws = wb.create_sheet(SHEET)
    ws["B1"] = '=ProjectName&"  -  Cost"'
    ws["B1"].font = TITLE_FONT
    ws["B2"] = ("Day rates, task and resource cost, earned value. Nothing else in the workbook reads this sheet: "
                "delete it before sharing the plan with anyone who shouldn't see rates.")
    ws["B2"].font = NOTE_FONT
    ws.column_dimensions["A"].width = 3
    top = _rate_card(wb, ws, 5, resources, n_rates)
    if v12():
        top = _fixed_costs(wb, ws, top)
    top = _summary(wb, ws, top)
    top = _history(wb, ws, top, blank)
    _task_costs(wb, ws, top, n_tasks)
    ws.freeze_panes = "C4"
    return ws


def _rate_card(wb, ws, top, resources, n_rows):
    ws.cell(top - 1, 2, "RATE CARD: one row per person on Resources (type the name exactly as there)").font = \
        font(bold=True, color="1F4E78")
    t = Tab("RateCard")
    who = t.row("Resource")
    blank = f'{who}=""'
    row = f"MATCH({who},res_Name,0)"
    start = f"MAX(ProjectStart,N(INDEX(res_Start,{row})))"
    end = f"IF(N(INDEX(res_End,{row}))=0,ProjectFinish,MIN(ProjectFinish,INDEX(res_End,{row})))"
    allocated = f"IF({end}<{start},0,NETWORKDAYS.INTL({start},{end},{WD}))*INDEX(res_Allocation,{row})"
    day = t.row("Day Rate")
    add_table(ws, "RateCard", top, 2, [
        Col("Resource", "in", 30, values=[r[0] for r in resources]),
        Col("Day Rate", "in", 11, MONEY, values=[r[4] for r in resources], note="Cost of one working day."),
    ] + ([Col("Per Use", "in", 9, MONEY, note="Optional: charged once on every task the person is on (a call-out "
                                              "fee, a licence). MS Project's Cost/Use.")] if v12() else []) + [
        Col("Check", width=16, formula=f'IF({blank},"",IF(ISNA({row}),"Not on Resources",""))'),
        Col("Allocated Days", width=10, fmt="0.0", center=True, formula=f'IF({blank},"",IFERROR({allocated},0))',
            note="Working days between their Start and End (or the project's) x Allocation %."),
        Col("Task Days", width=10, fmt="0.0", center=True,
            formula=f'IF({blank},"",{by_resource("Work", who)}/HoursPerDay)', note="Days of task work assigned."),
        Col("Task Cost", width=11, fmt=MONEY, formula=f'IF({blank},"",{t.row("Task Days")}*N({day}))'),
        Col("Actual Cost", width=11, fmt=MONEY, formula=f'IF({blank},"",{by_resource("Done", who)}/HoursPerDay*N({day}))',
            note="Work done (work x progress) x day rate."),
        Col("Resource Cost", width=11, fmt=MONEY, formula=f'IF({blank},"",{t.row("Allocated Days")}*N({day}))',
            note="Allocated days x day rate: what the person costs on the project, tasks or not."),
        Col("Gap", width=11, fmt=MONEY, formula=f'IF({blank},"",{t.row("Resource Cost")}-{t.row("Task Cost")})',
            note="Resource cost - task cost: time not assigned to tasks."),
    ], max(n_rows, len(resources) + 10))
    name(wb, "cost_RateName", "RateCard[Resource]")
    name(wb, "cost_Rate", "RateCard[Day Rate]")
    if v12():
        name(wb, "cost_PerUse", "RateCard[Per Use]")
    last = top + max(n_rows, len(resources) + 10)
    check = get_column_letter(5 if v12() else 4)
    ws.conditional_formatting.add(f"{check}{top + 1}:{check}{last}", FormulaRule(formula=[f'{check}{top + 1}<>""'],
                                                                      font=Font(color="C00000", bold=True)))
    return last + 4


def _fixed_costs(wb, ws, top):
    """Template 1.2 (B-302): costs that aren't work (a venue, a licence, materials), typed against a task's UID so
    moving rows on Tasks can't misalign them. Several rows may share a UID; they add up on that task, accrue with
    its progress (MS Project's prorated accrual) and are part of its budget."""
    ws.cell(top - 1, 2, "FIXED COSTS: type the task's UID (Tasks > UID); several rows may share one").font = \
        font(bold=True, color="1F4E78")
    t = Tab("FixedCosts")
    uid = t.row("UID")
    add_table(ws, "FixedCosts", top, 2, [
        Col("UID", "in", 30, "0", note="The task's UID (Tasks > UID): a detail task."),
        Col("Amount", "in", 11, MONEY),
        Col("Description", "in", 30),
        Col("Task", width=30, formula=f'IF({uid}="","",IFERROR(INDEX(tsk_Name,MATCH({uid},tsk_UID,0)),"Not on Tasks"))'),
    ], FIXED_ROWS)
    name(wb, "cost_FixedUID", "FixedCosts[UID]")
    name(wb, "cost_FixedAmount", "FixedCosts[Amount]")
    last = top + FIXED_ROWS
    ws.conditional_formatting.add(f"E{top + 1}:E{last}", FormulaRule(formula=[f'E{top + 1}="Not on Tasks"'],
                                                                      font=Font(color="C00000", bold=True)))
    return last + 4


def _summary(wb, ws, top):
    ws.cell(top - 1, 2, "COST AND EARNED VALUE at the status date").font = font(bold=True, color="1F4E78")
    add_table(ws, "CostSummary", top, 2, [
        Col("Measure", width=30, values=[e[1] for e in EVM]),
        Col("Value", width=11, values=[f"={e[2]}" for e in EVM]),
        Col("Meaning", width=16, values=[e[3] for e in EVM]),
    ], len(EVM), autofilter=False)
    for i, (defined, _, _, _, fmt) in enumerate(EVM):
        r = top + 1 + i
        ws.cell(r, 3).number_format = fmt
        ws.cell(r, 3).font = font(bold=True)
        name(wb, defined, f"{SHEET}!$C${r}")
    r = top + len(EVM) + 2
    ws.cell(r, 2, "Cost RAG (CPI against Settings > RAGAmberIndex / RAGRedIndex)").font = font(bold=True)
    c = ws.cell(r, 3, f"={rag.rag_text(rag.index('evm_CPI'))}")
    c.alignment = Alignment(horizontal="center")
    name(wb, "cost_RAG", f"{SHEET}!$C${r}")
    rag.colour(ws, f"C{r}", f"C{r}")
    chart = BarChart()
    chart.type, chart.style, chart.title, chart.legend = "bar", 10, "Earned value (BAC, PV, EV, AC)", None
    chart.add_data(Reference(ws, min_col=3, min_row=top + 1, max_row=top + 4))
    chart.set_categories(Reference(ws, min_col=2, min_row=top + 1, max_row=top + 4))
    chart.height, chart.width = 7, 16
    ws.add_chart(chart, f"F{top}")
    return r + 4


def sample_history():
    """Four weekly cost snapshots before the sample's status date (illustrative)."""
    out = []
    bac = 53600
    for when, earned, planned, cost_ratio in ((dt.date(2026, 9, 27), 0.17, 0.19, 1.02),
                                              (dt.date(2026, 10, 4), 0.27, 0.31, 1.04),
                                              (dt.date(2026, 10, 11), 0.36, 0.43, 1.05),
                                              (dt.date(2026, 10, 18), 0.44, 0.54, 1.06)):
        ev, pv = round(bac * earned), round(bac * planned)
        ac = round(ev * cost_ratio)
        cpi = ev / ac
        out.append([when, bac, pv, ev, ac, round(cpi, 2), round(ac + (bac - ev) / cpi)])
    return out


def _history(wb, ws, top, blank):
    """This status date's cost numbers in one row (formulas), pasted as values under the last row of CostHistory on
    each reporting date, as the Status sheet's snapshot is pasted onto History."""
    ws.cell(top - 1, 2, "COST SNAPSHOT: copy this row and Paste Special > Values under the last row of CostHistory "
                        "below (with the Status sheet's snapshot, each reporting date)").font = font(bold=True,
                                                                                                    color="1F4E78")
    add_table(ws, "CostSnapshot", top, 2, [Col(h, width=30 if fmt == DATE else 11, fmt=fmt, formula=f, center=True)
                                            for h, f, fmt in HISTORY], 1, autofilter=False)
    top += 4
    data = [] if blank else sample_history()
    add_table(ws, "CostHistory", top, 2, [Col(h, "in", 30 if fmt == DATE else 11, fmt, center=True,
                                              values=[row[i] for row in data])
                                          for i, (h, _, fmt) in enumerate(HISTORY)], HISTORY_ROWS)
    col = {h: 2 + i for i, (h, _, _) in enumerate(HISTORY)}
    dates = Reference(ws, min_col=col["Status Date"], min_row=top + 1, max_row=top + HISTORY_ROWS)
    for title_text, headers, anchor, y_format in (("Earned value S-curve", ["PV", "EV", "AC"], f"K{top}", None),
                                                  ("CPI", ["CPI"], f"K{top + 15}", "0.00")):
        c = LineChart()
        c.title = title_text
        for h in headers:
            c.add_data(Reference(ws, min_col=col[h], min_row=top, max_row=top + HISTORY_ROWS), titles_from_data=True)
        for s in c.series:
            s.smooth = False
        c.set_categories(dates)
        c.x_axis.number_format = "dd-mmm"
        if y_format:
            c.y_axis.number_format = y_format
        c.height, c.width = 7, 14
        ws.add_chart(c, anchor)
    return top + HISTORY_ROWS + 4


def _task_costs(wb, ws, top, n_rows):
    """One row per Tasks row, by position. Detail tasks: each person's work in days x their day rate; summaries add
    up the detail rows below them (Tasks' _Descendants)."""
    ws.cell(top - 1, 2, "TASK COSTS (follow the Tasks sheet row by row)").font = font(bold=True, color="1F4E78")
    t = Tab("TaskCost")
    n = t.row("_Row")
    blank = f'IFERROR(INDEX(tsk_Name,{n})="",TRUE())'
    task = lambda col: f"INDEX({col},{n})"
    leaf = f"{task('tsk_IsLeaf')}=1"
    slot = lambda field, k: task(people(field)[k - 1])

    def per_person(field):
        parts = [f'IF({slot("Name", k)}="",0,{slot(field, k)}/HoursPerDay*{rate(slot("Name", k))})'
                 for k in range(1, PEOPLE + 1)]
        return "+".join([parts[0]] + [f"IF(MaxPeople>={k},{p},0)" for k, p in enumerate(parts[1:], start=2)])

    def rolled(helper):
        below = f"INDEX({t.col(helper)},{n}+1):INDEX({t.col(helper)},{n}+{task('tsk_Descendants')})"
        return f'IF({blank},"",IF({task("tsk_IsSummary")},SUM({below}),{t.row(helper)}))'

    work, base = task("tsk_LeafTotalWork"), task("tsk_BaseWork")
    labour = f"({t.row('_LeafCost')}-{t.row('_LeafExtra')})" if v12() else t.row("_LeafCost")
    budget = f"IF(N({base})>0,IF({work}>0,{base}*{labour}/{work},{labour}),{labour})"
    extra, extra_done = [], []
    if v12():
        # B-302: each person's cost per use (accrued at the task's start, as MS Project does) and the task's fixed
        # costs (accrued with progress); the budget keeps them whole, only work cost scales with the baseline
        uses = [f'IF({slot("Name", k)}="",0,{per_use(slot("Name", k))})' for k in range(1, PEOPLE + 1)]
        uses = "+".join([uses[0]] + [f"IF(MaxPeople>={k},{p},0)" for k, p in enumerate(uses[1:], start=2)])
        fixed = f"SUMIFS(cost_FixedAmount,cost_FixedUID,{task('tsk_UID')})"
        progress = task("tsk_LeafProgress")
        extra = [Col("_LeafExtra", "help", 9, formula=f"IF({blank},0,IF({leaf},{uses}+{fixed},0))"),
                 Col("_ExtraDone", "help", 9, formula=f"IF({blank},0,IF({leaf},IF({progress}>0,{uses},0)"
                                                      f"+{fixed}*{progress},0))")]
        budget = f"{budget}+{t.row('_LeafExtra')}"
    add_table(ws, "TaskCost", top, 2, [
        Col("Task", width=30, formula=f'IF({blank},"",REPT("    ",{task("tsk_Level")}-1)&{task("tsk_Name")})'),
        Col("ID", width=11, center=True, formula=f'IF({blank},"",{task("tsk_ID")})'),
        Col("Cost", width=11, fmt=MONEY, formula=rolled("_LeafCost"),
            note="Work x day rates (summaries add up)" + (", plus cost per use and fixed costs." if v12() else ".")),
        Col("Actual Cost", width=11, fmt=MONEY, formula=rolled("_LeafActual"),
            note="Work done x day rates" + (", plus per-use and fixed costs accrued." if v12() else ".")),
        Col("Budget", width=11, fmt=MONEY, formula=rolled("_LeafBudget"),
            note="Cost at the baseline's work once a baseline is set, otherwise the cost."),
        Col("PV", width=11, fmt=MONEY, formula=rolled("_LeafPV"), note="Budget x share planned by the status date."),
        Col("EV", width=11, fmt=MONEY, formula=rolled("_LeafEV"), note="Budget x progress."),
        Col("Cost Var", width=11, fmt=MONEY, formula=f'IF({blank},"",{t.row("Cost")}-{t.row("Budget")})',
            note="Cost - budget. Positive = more than budgeted."),
        Col("_Row", "help", 5, formula=f"ROW()-ROW({t.head('Task')})"),
        Col("_LeafCost", "help", 9, formula=f"IF({blank},0,IF({leaf},{per_person('Work')},0))"
                                            + (f"+{t.row('_LeafExtra')}" if v12() else "")),
        Col("_LeafActual", "help", 9, formula=f"IF({blank},0,IF({leaf},{per_person('Done')},0))"
                                              + (f"+{t.row('_ExtraDone')}" if v12() else "")),
        Col("_LeafBudget", "help", 9, formula=f"IF({blank},0,IF({leaf},{budget},0))"),
        Col("_LeafPV", "help", 9, formula=f"{t.row('_LeafBudget')}*IF({blank},0,{task('tsk_PlannedPct')})"),
        Col("_LeafEV", "help", 9, formula=f"{t.row('_LeafBudget')}*IF({blank},0,{task('tsk_LeafProgress')})"),
    ] + extra, n_rows, autofilter=False)
    first = get_column_letter(10)
    ws.column_dimensions.group(first, get_column_letter(15 + len(extra)), hidden=True, outline_level=1)
    ws.conditional_formatting.add(f"B{top + 1}:I{top + n_rows}", FormulaRule(
        formula=[f'AND($B{top + 1}<>"",INDEX(tsk_IsSummary,ROW()-ROW($B${top})))'], font=Font(bold=True),
        fill=fill("DDEBF7")))
