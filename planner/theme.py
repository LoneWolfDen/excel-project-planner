"""Template 1.2's look (Fluent blue, docs/BACKLOG.md v1.2): one pass over the finished workbook that remaps every
colour and font the builders use to a Microsoft Fluent palette, then the sheet-level touches (tab colours by role,
no gridlines on the report pages, zoom, the Start page as a landing page).

Doing it as a final pass keeps every builder unchanged, so Release 1.0 and 1.1 look exactly as before, and the
mapping is in one table here. It runs after protection: unlocking reads the classic input blue (core.INPUT_FONT).
Colours keep their meaning (blue = type here, grey = calculated, green / amber / red = RAG); Fluent's tints are
lighter and its text colours meet WCAG AA contrast on white.
"""
from copy import copy

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT_NAME = "Segoe UI"                 # Windows' UI font; Excel substitutes a similar one on Mac, web and iPhone
BRAND, BRAND_DARK, BRAND_TINT = "0F6CBD", "0F548C", "EBF3FC"
# classic colour -> Fluent colour, for fonts, fills, borders and conditional formats alike
COLOURS = {
    # blues: inputs, headers, titles, links, the summary rows and the Gantt's remaining work
    "0000FF": BRAND, "1F4E78": BRAND_DARK, "2F75B5": BRAND, "0563C1": BRAND, "DDEBF7": BRAND_TINT,
    "5B9BD5": "479EF5", "9DC3E6": "B4D6FA", "C9D5E8": "D6E4F5",
    # neutrals: calculated cells and their headers, helper headers, notes, borders
    "F2F2F2": "F5F5F5", "595959": "424242", "404040": "424242", "A6A6A6": "9E9E9E", "7F7F7F": "707070",
    "D9D9D9": "E0E0E0", "EDEDED": "F0F0F0",
    # status colours: red, amber, green (text and tint), orange, purple
    "C00000": "C50F1F", "9C0006": "C50F1F", "FFC7CE": "FDE7E9", "E06666": "E37D80",
    "9C5700": "8A5D00", "FFEB9C": "FFF4CE", "FFF2CC": "FFF4CE", "FFFF00": "FFF4CE",
    "548235": "107C10", "006100": "0E700E", "C6EFCE": "DFF6DD", "70AD47": "54B054",
    "C55A11": "CA5010", "F4B183": "F7B183", "7030A0": "8764B8", "E4DFEC": "EFE6F7",
}
# sheet tab colours by role: start, the plan, reports, logs and history, settings, money
TABS = {"Start": BRAND_DARK, "Tasks": BRAND, "Resources": BRAND, "Status": "038387", "Check": "C50F1F",
        "Roadmap": "038387", "Resource Report": "038387", "Agile": "038387", "Board": "038387",
        "Backlog": BRAND, "Log": "CA5010", "History": "CA5010", "TakeSnapshot": "CA5010", "Settings": "616161",
        "Cost": "8A5D00"}
PAGES = ("Start", "Status", "Check", "Roadmap", "Resource Report", "Agile", "TakeSnapshot", "History", "Cost",
         "Board")                     # report-like sheets: no gridlines


def _rgb(colour):
    """The six-digit colour of an openpyxl Color, or None (theme / indexed colours are left alone)."""
    rgb = getattr(colour, "rgb", None)
    return rgb[-6:].upper() if isinstance(rgb, str) and len(rgb) >= 6 else None


def _map(colour):
    rgb = _rgb(colour)
    return COLOURS.get(rgb) if rgb else None


def _font(f):
    new = copy(f)
    if f.name in ("Arial", None):
        new.name = FONT_NAME
    to = _map(f.color) if f.color is not None else None
    if to:
        new.color = to
    return new


def _fill(f):
    if f is None or f.fill_type != "solid":
        return copy(f)
    to = _map(f.start_color)
    return PatternFill("solid", start_color=to, end_color=to) if to else copy(f)


def _side(s):
    if s is None or s.style is None:
        return s
    to = _map(s.color)
    return Side(style=s.style, color=to) if to else s


def _border(b):
    return Border(left=_side(b.left), right=_side(b.right), top=_side(b.top), bottom=_side(b.bottom))


def _cells(ws):
    """Restyle every styled cell; each distinct style is converted once (a sheet has a few dozen)."""
    cache = {}
    for row in ws.iter_rows():
        for c in row:
            if not c.has_style:
                continue
            key = (c._style.fontId, c._style.fillId, c._style.borderId)     # openpyxl's style ids
            if key not in cache:
                cache[key] = (_font(c.font), _fill(c.fill), _border(c.border))
            c.font, c.fill, c.border = cache[key]


def _formats(ws):
    """Conditional formats: the same colour table, on fonts, fills and borders."""
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            d = rule.dxf
            if d is None:
                continue
            if d.font is not None and d.font.color is not None and _map(d.font.color):
                d.font = copy(d.font)
                d.font.color = _map(d.font.color)
            if d.fill is not None:
                for attr in ("fgColor", "bgColor"):
                    to = _map(getattr(d.fill, attr, None))
                    if to:
                        d.fill = copy(d.fill)
                        setattr(d.fill, attr, to)
            if d.border is not None:
                d.border = _border(d.border)


def _start_page(ws):
    """Start as a landing page: a large title, a subtitle, roomy step rows with a tinted label column."""
    ws["B1"].font = Font(name=FONT_NAME, size=24, bold=True, color=BRAND_DARK)
    ws.row_dimensions[1].height = 40
    ws["B2"].font = Font(name=FONT_NAME, size=11, color="616161")
    ws.row_dimensions[2].height = 22
    tint = PatternFill("solid", start_color=BRAND_TINT, end_color=BRAND_TINT)
    for r in range(3, ws.max_row + 1):
        label = ws.cell(r, 2)
        if isinstance(label.value, str) and label.value[:2] in ("1.", "2.", "3.", "4.", "5.") and \
                ws.cell(r, 3).hyperlink is not None:
            label.fill = tint
            label.font = Font(name=FONT_NAME, size=11, bold=True, color=BRAND_DARK)
            label.alignment = Alignment(vertical="center", indent=1)
            ws.row_dimensions[r].height = 36
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0      # prints one page wide
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def apply(wb):
    for ws in wb.worksheets:
        _cells(ws)
        _formats(ws)
        if ws.title in TABS:
            ws.sheet_properties.tabColor = TABS[ws.title]
        if ws.title in PAGES:
            ws.sheet_view.showGridLines = False
            ws.sheet_view.zoomScale = 100
        for table in ws.tables.values():                 # light banding on rows the builders don't fill
            if table.tableStyleInfo is not None:
                table.tableStyleInfo.showRowStripes = True
    if "Start" in wb.sheetnames:
        _start_page(wb["Start"])
