"""Sprints sheet: sprint calendar (Excel Table "SprintTable") and resource-by-sprint load and utilisation grids.

Grid cells find their resource and sprint by position, ROW()/COLUMN() minus a named anchor cell, so every cell in
a grid holds the same formula and nothing refers to a fixed address.
"""
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .core import (BOX, CALC_FONT, DATE, LABEL_FONT, NOTE_FONT, Col, Tab, add_table, fill, font, header_cell, name)

WD = "WorkWeek,Holidays"
GOALS = ["Charter & kickoff", "Requirements", "Design & architecture", "Foundations: DB + components",
         "Core APIs + screens", "Payments + screens", "Feature complete", "System testing", "UAT",
         "Release & hypercare", "Hypercare", "Close-out"]


def work_between(start, end, resource=None):
    """Hours of task work scheduled from date start to date end (inclusive), counting working days only.

    Uses the cumulative form W(D) = SUMIFS(work, finish < D) + D * SUMIFS(work per day, running on D)
    - SUMIFS(work before start, running on D), with D in working days since the project start. Six SUMIFS per cell
    are far cheaper than a SUMPRODUCT over every task, and SUMIFS needs no array formulas."""
    res = f',tsk_Resource,{resource}' if resource else ""

    def upto(d):
        running = f'tsk_StartWD,"<="&{d},tsk_FinishWD,">="&{d}{res}'
        return (f'(SUMIFS(tsk_LeafWork,tsk_FinishWD,"<"&{d},tsk_FinishWD,">0"{res})'
                f'+{d}*SUMIFS(tsk_WorkPerWD,{running})-SUMIFS(tsk_WorkBeforeStart,{running}))')
    first = f"NETWORKDAYS.INTL(ProjectStart,{start},{WD})"
    last = f"NETWORKDAYS.INTL(ProjectStart,{end},{WD})"
    return f"({upto(last)}-{upto(f'({first}-1)')})"


def build_sprints(wb, blank, n_sprints, n_resources):
    ws = wb.create_sheet("Sprints")
    ws["A1"] = '=ProjectName&"  -  Sprint & Capacity Planning"'
    ws["A1"].font = font(size=16, bold=True, color="1F4E78")
    ws["A2"] = ("Sprint calendar from Settings (Sprint1Start, SprintLength, SprintCount); add rows to the table for more "
                "sprints. Each task's work is spread evenly over its forecast working days, then summed per sprint.")
    ws["A2"].font = NOTE_FONT
    s = Tab("SprintTable")
    r = s.row
    has = f'{r("Sprint")}=""'
    cols = [
        Col("No", width=5, formula=f"ROW()-ROW({s.head('No')})", center=True),
        Col("Sprint", width=12, formula=f'IF({r("No")}<=SprintCount,SprintPrefix&{r("No")},"")'),
        Col("Start", width=14, fmt=DATE, formula=f"Sprint1Start+({r('No')}-1)*SprintLength"),
        Col("End", width=14, fmt=DATE, formula=f"{r('Start')}+SprintLength-1"),
        Col("Working Days", width=9, formula=f"NETWORKDAYS.INTL({r('Start')},{r('End')},{WD})", center=True),
        Col("Goal", "in", 36, values=[] if blank else GOALS),
        Col("Tasks Starting", width=9, formula=f'IF({has},"",COUNTIF(tsk_Sprint,{r("Sprint")}))', center=True),
        Col("Planned Work", width=11, fmt="#,##0.0", formula=f'IF({has},"",{work_between(r("Start"), r("End"))})',
            note="Hours of task work falling in the sprint."),
        Col("Capacity", width=11, fmt="#,##0.0", note="Sum of Work resources' Max Units x HoursPerDay x working days.",
            formula=f'IF({has},"",SUMIFS(res_MaxUnits,res_Type,"Work")*HoursPerDay*{r("Working Days")})'),
        Col("Utilisation", width=10, fmt="0%",
            formula=f'IF(OR({has},N({r("Capacity")})=0),"",{r("Planned Work")}/{r("Capacity")})'),
        Col("Status", width=12, center=True,
            formula=f'IF({has},"",IF(StatusDate>{r("End")},"Closed",IF(StatusDate>={r("Start")},"Active","Future")))'),
    ]
    L = add_table(ws, "SprintTable", 5, 1, cols, n_sprints, autofilter=False)
    for column, defined in {"Sprint": "spr_Name", "Start": "spr_Start", "End": "spr_End", "Working Days": "spr_WorkDays",
                            "Planned Work": "spr_Work", "Capacity": "spr_Capacity"}.items():
        name(wb, defined, f"SprintTable[{column}]")
    first, last = 6, 5 + n_sprints
    cf = ws.conditional_formatting
    cf.add(f"A{first}:K{last}", FormulaRule(formula=[f'${L["Sprint"]}{first}=""'], font=Font(color="D9D9D9")))
    cf.add(f"A{first}:K{last}", FormulaRule(formula=[f'${L["Status"]}{first}="Active"'], fill=fill("FFF2CC"), font=Font(bold=True)))
    cf.add(f"{L['Utilisation']}{first}:{L['Utilisation']}{last}",
           FormulaRule(formula=[f"AND(ISNUMBER(${L['Utilisation']}{first}),${L['Utilisation']}{first}>1)"],
                       font=Font(color="C00000", bold=True), fill=fill("FFC7CE")))

    load_top = last + 3
    util_top = load_top + n_resources + 5
    _grid(wb, ws, "load", load_top, n_sprints, n_resources,
          "Resource load by sprint (hours)  -  red = over capacity (Max Units x HoursPerDay x working days)")
    _grid(wb, ws, "util", util_top, n_sprints, n_resources, "Resource utilisation by sprint (%)")
    ws.freeze_panes = "D6"
    ws.column_dimensions["B"].width = 16
    return ws, L


def _grid(wb, ws, kind, top, n_sprints, n_resources, heading):
    """Rows = resources, columns = sprints. The origin cell (resource header) anchors every position."""
    ws.cell(top - 1, 1, heading).font = font(bold=True, size=12, color="1F4E78")
    origin = f"{kind}_Origin"
    name(wb, origin, f"Sprints!$B${top}")
    header_cell(ws.cell(top, 2), "Resource")
    header_cell(ws.cell(top, 3), "Max Units")
    i = f"(ROW()-ROW({origin})-1)"                 # resource number (row 1 below the dates row)
    j = f"(COLUMN()-COLUMN({origin})-1)"            # sprint number
    sprint_name = f"INDEX(spr_Name,{j})"
    for c in range(4, 4 + n_sprints):
        header_cell(ws.cell(top, c), None)
        ws.cell(top, c).value = f'=IFERROR({sprint_name}&"","")'
        d = ws.cell(top + 1, c, f'=IFERROR(IF({sprint_name}="","",INDEX(spr_Start,{j})),"")')
        d.number_format = "dd-mmm"
        d.font = font(italic=True, size=9)
        ws.column_dimensions[get_column_letter(c)].width = 10
    ws.cell(top + 1, 2, "Sprint start").font = font(italic=True, size=9)
    resource = f"INDEX(res_Name,{i})"
    units = f"INDEX(res_MaxUnits,{i})"
    first_row, last_row = top + 2, top + 1 + n_resources
    last_col = get_column_letter(3 + n_sprints)
    for rr in range(first_row, last_row + 1):
        ws.cell(rr, 2, f'=IFERROR({resource}&"","")').font = font(bold=True)
        u = ws.cell(rr, 3, f'=IFERROR(IF({resource}="","",{units}),"")')
        u.number_format = "0%"
        for c in range(4, 4 + n_sprints):
            if kind == "load":
                hours = work_between(f"INDEX(spr_Start,{j})", f"INDEX(spr_End,{j})", resource)
                v = f'=IFERROR(IF(OR({resource}="",{sprint_name}=""),"",{hours}),"")'
                fmt = "#,##0;;-"
            else:
                load = f"INDEX(load_Grid,{i},{j})"
                v = (f'=IFERROR(IF(OR({resource}="",{sprint_name}="",N({units})=0),"",'
                     f'{load}/({units}*HoursPerDay*INDEX(spr_WorkDays,{j}))),"")')
                fmt = "0%;;-"
            cell = ws.cell(rr, c, v)
            cell.number_format = fmt
            cell.border = BOX
            cell.font = CALC_FONT
    grid = f"D{first_row}:{last_col}{last_row}"
    if kind == "load":
        name(wb, "load_Grid", f"Sprints!$D${first_row}:${last_col}${last_row}")
        ws.conditional_formatting.add(grid, FormulaRule(
            formula=[f"AND(ISNUMBER(D{first_row}),D{first_row}>$C{first_row}*HoursPerDay*"
                     f"INDEX(spr_WorkDays,COLUMN()-COLUMN({origin})-1))"],
            fill=fill("FFC7CE"), font=Font(color="9C0006", bold=True)))
    else:
        ws.conditional_formatting.add(grid, ColorScaleRule(
            start_type="num", start_value=0, start_color="FFFFFF", mid_type="num", mid_value=0.8, mid_color="C6EFCE",
            end_type="num", end_value=1.2, end_color="F8696B"))
