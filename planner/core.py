"""Shared building blocks: styles, defined names and Excel Tables.

Every formula the planner writes refers to cells through Excel Tables (structured references such as
``Tasks[[#This Row],[Duration]]``, shown by Excel as ``[@Duration]``) or defined names (``ProjectStart``,
``tsk_Duration``). No formula contains an A1 address, so inserting or moving rows and columns never breaks it.
Conditional formatting, data validation and charts are the only places that use A1 ranges, because Excel does
not accept structured references there; Excel shifts those ranges itself when rows or columns are inserted.
"""
from dataclasses import dataclass, field

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo

FONT = "Arial"
DATE = "ddd dd-mmm-yy"
SHORT_DATE = "dd-mmm-yy"


def font(**k):
    k.setdefault("name", FONT)
    k.setdefault("size", 10)
    return Font(**k)


def fill(colour):
    return PatternFill("solid", start_color=colour, end_color=colour)


INPUT_FONT = font(color="0000FF")
CALC_FONT = font()
HELPER_FONT = font(color="404040")
CALC_FILL = fill("F2F2F2")
HEADER_FILLS = {"in": fill("1F4E78"), "calc": fill("595959"), "help": fill("A6A6A6")}
HEADER_FONT = font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D9D9D9")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center")
TITLE_FONT = font(size=16, bold=True, color="1F4E78")
NOTE_FONT = font(italic=True, color="595959")
LABEL_FONT = font(bold=True, color="1F4E78")


def name(wb, defined_name, ref):
    """Workbook-level defined name, e.g. name(wb, "ProjectStart", "Settings!$C$7")."""
    wb.defined_names[defined_name] = DefinedName(defined_name, attr_text=ref)


def title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    if note:
        ws["A2"] = note
        ws["A2"].font = NOTE_FONT


def header_cell(cell, text, kind="calc"):
    cell.value = text
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILLS[kind]
    cell.alignment = WRAP_CENTER
    cell.border = BOX


@dataclass
class Col:
    """One column of an Excel Table.

    name     header text, which is also the structured-reference name ([@name])
    kind     "in" (user input, blue), "calc" (formula, grey) or "help" (hidden helper formula)
    formula  formula for every data row (same text in every row: a calculated column), or None for inputs
    """
    name: str
    kind: str = "calc"
    width: float = 10
    fmt: str = None
    formula: str = None
    center: bool = False
    note: str = None
    key: str = None
    values: list = field(default_factory=list)   # optional constants for input columns, row by row


class Tab:
    """Structured-reference helpers for one Excel Table."""

    def __init__(self, table):
        self.table = table

    def row(self, col):       # the value in this row: shown by Excel as [@col]
        return f"{self.table}[[#This Row],[{col}]]"

    def col(self, col):       # the whole data column
        return f"{self.table}[{col}]"

    def head(self, col):      # the header cell, used as an anchor for OFFSET and ROW
        return f"{self.table}[[#Headers],[{col}]]"

    def upto(self, col):      # header cell down to this row (an expanding range that never includes later rows)
        return f"{self.head(col)}:{self.row(col)}"




def add_table(ws, table_name, header_row, first_col, cols, n_rows, style="TableStyleLight1", autofilter=True):
    """Write a styled Excel Table with calculated columns. Returns {column name: column letter}."""
    letters = {}
    last_row = header_row + n_rows
    for i, c in enumerate(cols):
        letters[c.name] = get_column_letter(first_col + i)
        if c.key:
            letters[c.key] = letters[c.name]

    for i, c in enumerate(cols):
        col = first_col + i
        letter = letters[c.name]
        header_cell(ws.cell(header_row, col), c.name, c.kind)
        if c.note:
            ws.cell(header_row, col).comment = Comment(c.note, "Planner")
        ws.column_dimensions[letter].width = c.width
        for r in range(header_row + 1, last_row + 1):
            cell = ws.cell(r, col)
            cell.border = BOX
            if c.kind == "in":
                cell.font = INPUT_FONT
                k = r - header_row - 1
                if k < len(c.values) and c.values[k] is not None:
                    cell.value = c.values[k]
            else:
                cell.font = HELPER_FONT if c.kind == "help" else CALC_FONT
                cell.fill = CALC_FILL
                k = r - header_row - 1
                if c.formula:
                    cell.value = f"={c.formula}"
                elif k < len(c.values):        # per-row content (labels or formulas), not a calculated column
                    cell.value = c.values[k]
            if c.fmt:
                cell.number_format = c.fmt
            if c.center:
                cell.alignment = CENTER
    ref = f"{get_column_letter(first_col)}{header_row}:{get_column_letter(first_col + len(cols) - 1)}{last_row}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=False)
    if autofilter:
        table.autoFilter = AutoFilter(ref=ref)
    table.tableColumns = [
        TableColumn(id=i + 1, name=c.name,
                    calculatedColumnFormula=TableFormula(attr_text=c.formula) if c.formula else None)
        for i, c in enumerate(cols)
    ]
    ws.add_table(table)
    return letters


def protect_view(ws):
    """B-1019: protect a view sheet without a password, so a stray keystroke can't overwrite a formula. Input cells
    (the blue INPUT_FONT: PM RAG, comments, notes, the status date) stay unlocked; selecting, copying, filtering,
    sorting and resizing or unhiding rows and columns still work. Review > Unprotect Sheet lifts it (no password).
    Tasks is never protected: Excel can't add table rows on a protected sheet."""
    blue = INPUT_FONT.color.rgb
    for row in ws.iter_rows():
        for cell in row:
            if cell.font is not None and cell.font.color is not None and cell.font.color.rgb == blue:
                cell.protection = Protection(locked=False)
    p = ws.protection
    p.sheet = True
    p.formatColumns = p.formatRows = p.autoFilter = p.sort = False      # False = allowed
    p.selectLockedCells = p.selectUnlockedCells = False


# B-913: optional modules. A build includes any of them; each builder leaves out, or blanks, whatever reads a module
# that isn't there (set_modules is called once per build, by build.make_workbook).
MODULES = {"agile": "Agile", "log": "Log", "history": "History", "resources": "Resource Report", "cost": "Cost"}
_MODULE_NAMES = {"agile": r"\b(spr_|agl_|cap_|ShownSprint|ShowSprintInput|CurrentSprint|SprintLastEnd|TotalPoints)",
                 "log": r"\blog_", "resources": r"\b(rw_|load_Grid|spr_LeaveGrid)", "cost": r"\b(evm_|cost_)"}
ACTIVE = set(MODULES)
# Release 1.0 (the --release build, its own workbooks in release-1.0/): two more optional sheets, never in Simple
RELEASE_MODULES = {"check": "Check", "snapshot": "TakeSnapshot"}
RELEASE = {"on": False, "version": None}
# template versions of the release line; each later one adds to the one before. "m365" is template 1.1 rebuilt with
# Microsoft 365 functions (LAMBDA, LET, FILTER): its own folder, so a problem there never touches 1.1 (F-1108)
EDITIONS = ("1.0", "1.1", "m365", "1.2")
# each edition's feature level: 0 = Release 1.0, 1 = the v1.1 features, 2 = template 1.2 (design, UX, engine items)
LEVELS = {"1.0": 0, "1.1": 1, "m365": 1, "1.2": 2}
_MODULE_NAMES.update({"check": r"\bchk_", "snapshot": r"\bsnap_"})


def edition_of(value):
    """The template version a build option means: False / None = classic, True = "1.0", or "1.0" / "1.1"."""
    if value in (None, False, ""):
        return None
    version = "1.0" if value is True else str(value)
    if version not in EDITIONS:
        raise ValueError(f"unknown release {version} (choose from {', '.join(EDITIONS)})")
    return version


def set_release(value):
    """Build a release-line edition (1.0: pickers, the leave day grid, Check, TakeSnapshot ...; 1.1 adds RACI,
    dependencies on Status, Sandbox, the narrative and PERT) or the classic workbook. Called once per build, before
    set_modules, by build.make_workbook."""
    RELEASE["version"] = edition_of(value)
    RELEASE["on"] = RELEASE["version"] is not None


def release():
    return RELEASE["on"]


def v11():
    """True for template 1.1 and later: the v1.1 features (docs/BACKLOG.md F-1101 - F-1107)."""
    return release() and LEVELS[RELEASE["version"]] >= 1


def v12():
    """True for template 1.2: the Fluent design, the UX items and the engine items (docs/BACKLOG.md v1.2)."""
    return release() and LEVELS[RELEASE["version"]] >= 2


def m365():
    """True for the Microsoft 365 edition (F-1108): named LAMBDA functions and FILTER lists. Excel 2019 can't open
    it correctly, and LibreOffice 24.2 can't calculate it, so it is tested by structure and checked in Excel."""
    return release() and RELEASE["version"] == "m365"


def all_modules():
    """The optional sheets this edition can have: {module: sheet name}."""
    return {**MODULES, **RELEASE_MODULES} if release() else dict(MODULES)


def set_modules(modules):
    """The modules this build includes (None = all)."""
    known = all_modules()
    unknown = set(modules or ()) - set(known)
    if unknown:
        raise ValueError(f"unknown modules: {', '.join(sorted(unknown))} (choose from {', '.join(known)})")
    ACTIVE.clear()
    ACTIVE.update(known if modules is None else modules)


def has(module):
    return module in ACTIVE


def missing(text):
    """True when a formula (or a sheet name) needs a module this build leaves out."""
    import re
    text = str(text or "")
    known = all_modules()
    gone = [m for m in known if m not in ACTIVE]
    return any(text == known[m] or (m in _MODULE_NAMES and re.search(_MODULE_NAMES[m], text)) for m in gone)


def col_range(letter, first_row, last_row):
    return f"{letter}{first_row}:{letter}{last_row}"


@dataclass
class Show:
    """One column of a ranked list: header text, the tsk_ name to read, and how to show it."""
    header: str
    source: str
    kind: str = "text"        # "text", "date" (blank when empty) or "num"
    width: float = 14
    fmt: str = None


def ranked_table(ws, table_name, header_row, first_col, key, n_rows, shows, largest=False, label=None, extra=()):
    """A report list: row k shows the task with the k-th smallest (or largest) value in the helper column `key`.

    Keys are unique and blank for tasks not on the list, so SMALL / LARGE then MATCH finds the task. The match is
    done once per row, in the _Task column (first, so every list on a sheet hides the same column); the visible
    columns only INDEX by it. `extra` adds calculated Cols after the shown ones. `label` is an Excel text expression
    shown above the table with the count of tasks on the list. To show more rows, drag the table's corner down: the
    formulas fill in. Returns {column name: letter}."""
    t = Tab(table_name)
    k = f"ROW()-ROW({t.head('_Task')})"
    nth = f"{'LARGE' if largest else 'SMALL'}({key},{k})"
    pos = t.row("_Task")

    def show(s):
        value = f"INDEX({s.source},{pos})"
        if s.kind == "date":
            value = f'IF(N({value})=0,"",{value})'
        elif s.kind == "text":
            value = f'{value}&""'
        return f'IF({pos}="","",{value})'

    cols = [Col("_Task", "help", 4, formula=f'IFERROR(MATCH({nth},{key},0),"")',
                note="Row of the task in the Tasks table (helper, hidden).")]
    cols += [Col(s.header, width=s.width, fmt=s.fmt, formula=show(s), center=s.kind != "text") for s in shows]
    cols += list(extra)
    letters = add_table(ws, table_name, header_row, first_col, cols, n_rows, autofilter=False)
    ws.column_dimensions[letters["_Task"]].hidden = True
    if label:
        rows = f'ROWS({t.col("_Task")})'
        cell = ws.cell(header_row - 1, first_col + 1)
        cell.value = f'={label}&": "&COUNT({key})&IF(COUNT({key})>{rows}," (showing the first "&{rows}&")","")'
        cell.font = LABEL_FONT
    return letters


def linked_task(uid, column="tsk_Name"):
    """A task's value looked up by UID, for logs that link to a task: blank when no UID, a message when it's
    missing. UID is an input column, so the MATCH is not volatile; only the value read may change."""
    return f'IF(N({uid})=0,"",IFERROR(INDEX({column},MATCH({uid},tsk_UID,0)),"UID "&{uid}&" not found"))'
