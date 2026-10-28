"""The Microsoft 365 edition (F-1108, `--release m365`, release-1.1-m365/): template 1.1 with named LAMBDA functions
and FILTER lists. LibreOffice 24.2, which recalculates every other test, can't calculate LAMBDA, LET or FILTER, so
these tests check the structure: the functions are defined as Excel stores them, the formulas call them, the inputs
are the 1.1 sample's, nothing points at a missing name, and the other editions don't change. Its values are
checked in Excel with docs/TESTING_IN_EXCEL.md section 8.

Run:  python -m unittest tests.test_m365 -v
"""
import io
import re
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(ROOT))
import perf_check  # noqa: E402
from test_planner import CommittedWorkbooksTest, d, named_cell  # noqa: E402
from test_release import head, small, task_row  # noqa: E402

RELEASE = ROOT / "release-1.1-m365"
SAMPLE, BLANK, SIMPLE = (RELEASE / n for n in ("Project_Planner.xlsx", "Project_Planner_Blank.xlsx",
                                               "Project_Planner_Simple.xlsx"))


def make(**k):
    from planner.build import make_workbook
    return make_workbook(release_edition="m365", **k)


def classic_again():
    from planner.build import make_workbook
    make_workbook(sizes=small())


class StructureTest(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.path = Path(cls.tmp.name) / "m365.xlsx"
        make(baselined=True).save(cls.path)
        cls.wb = openpyxl.load_workbook(cls.path)
        cls.report = perf_check.static_report(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()
        classic_again()

    def test_versions(self):
        names = self.wb.defined_names
        self.assertEqual(names["TemplateVersion"].attr_text, '"1.1 M365"')
        self.assertEqual(names["PlannerVersion"].attr_text, '"2.1-m365"')

    def test_named_lambdas(self):
        from planner.m365 import LAMBDAS
        for n in LAMBDAS:
            text = self.wb.defined_names[n].attr_text
            self.assertTrue(text.startswith("_xlfn.LAMBDA(_xlpm."), n)
            self.assertEqual(perf_check.excel_limit_problems("=" + text), [], n)
            params = re.findall(r"_xlpm\.(\w+)", text)
            self.assertTrue(params, n)
        self.assertIn("_xlfn.LET(", self.wb.defined_names["L_PARSE_DURATION"].attr_text)

    def test_formulas_call_them(self):
        ws = self.wb["Tasks"]
        col = head(ws)
        self.assertIn("L_PARSE_DURATION(", ws.cell(6, col["_Days"]).value)
        self.assertIn("L_IS_WORKING_DAY(", ws.cell(6, col["Issues"]).value)
        log = self.wb["Log"]
        self.assertIn("L_WORKING_DAYS(", log.cell(5, head(log, 4)["Gap (days)"]).value)
        self.assertIn("L_WORKING_DAYS(", named_cell(self.wb, "FinishVarianceDays").value)

    def test_check_uses_filter(self):
        ws = self.wb["Check"]
        arrays = [c.value.text for row in ws.iter_rows() for c in row if getattr(c.value, "text", None)]
        self.assertEqual(len(arrays), 3)
        self.assertTrue(all("_xlfn._xlws.FILTER(" in a for a in arrays))
        for key in ("_IssueKey", "_OverAllocKey"):
            self.assertNotIn(key, head(self.wb["Tasks"]))
        self.assertNotIn("_OverdueKey", head(self.wb["Log"], 4))
        self.assertFalse(ws.tables)                                    # no ranked tables left

    def test_formula_rules(self):
        r = self.report
        self.assertEqual((r["a1_references"], r["missing_references"], r["excel_limit_problems"]), ([], [], []))
        self.assertEqual(r["volatile"], {"Tasks": {"OFFSET": 2200}})
        self.assertLess(sum(r["formulas"].values()), 50000)

    def test_same_inputs_as_release_11(self):
        """Every typed Tasks cell is the 1.1 sample's: the plan is the same; only formulas differ."""
        from planner.build import make_workbook
        v11 = make_workbook(release_edition="1.1", baselined=True)
        a, b = self.wb["Tasks"], v11["Tasks"]
        ca, cb = head(a), head(b)
        from planner.core import HEADER_FILLS
        blue = HEADER_FILLS["in"].start_color.rgb
        typed = [h for h, c in cb.items() if b.cell(5, c).fill.start_color.rgb == blue]
        self.assertIn("Pred1", typed)
        self.assertEqual([h for h in typed if h not in ca], [])
        same = lambda x, y: d(x) == d(y) or (x in (None, "") and y in (None, ""))
        diffs = [(h, r) for h in typed for r in range(6, 40) if not same(a.cell(r, ca[h]).value, b.cell(r, cb[h]).value)]
        self.assertEqual(diffs, [])


class EditionsTest(unittest.TestCase):

    @classmethod
    def tearDownClass(cls):
        classic_again()

    def test_other_editions_unchanged(self):
        from planner.build import make_workbook
        make(sizes=small())
        for edition in ("1.0", "1.1"):
            wb = make_workbook(sizes=small(), release_edition=edition)
            self.assertNotIn("L_PARSE_DURATION", wb.defined_names, edition)
            self.assertIn("_IssueKey", head(wb["Tasks"]), edition)
            self.assertNotIn("L_", wb["Tasks"].cell(6, head(wb["Tasks"])["_Days"]).value, edition)
        self.assertNotIn("L_PARSE_DURATION", make_workbook(sizes=small()).defined_names)

    def test_modules(self):
        for combo in ((), ("check",), ("check", "log"), ("log",)):
            self.assertEqual(perf_check.missing_references(make(sizes=small(), modules=combo)), [], combo)

    def test_recalc_is_skipped(self):
        """LibreOffice can't calculate it: the build saves it without --recalc (Excel calculates on open)."""
        import subprocess
        with tempfile.TemporaryDirectory() as tmp:
            out = subprocess.run([sys.executable, str(ROOT / "scripts" / "build_planner.py"), str(Path(tmp) / "m.xlsx"),
                                  "--release", "m365", "--recalc", "--tasks", "40", "--resources", "12",
                                  "--grid-resources", "8", "--sprints", "6", "--gantt-columns", "10", "--weeks", "6"],
                                 capture_output=True, text=True, check=True).stdout
            self.assertIn("without --recalc", out)
            calc = openpyxl.load_workbook(Path(tmp) / "m.xlsx").calculation
            self.assertTrue(calc.fullCalcOnLoad)


class CommittedTest(unittest.TestCase):
    compare = CommittedWorkbooksTest.compare

    def test_sample(self):
        self.compare(SAMPLE, "--release", "m365", "--baselined")

    def test_blank(self):
        self.compare(BLANK, "--release", "m365", "--blank")

    def test_simple(self):
        self.compare(SIMPLE, "--release", "m365", "--baselined", "--simple")


class ToolsTest(unittest.TestCase):

    @classmethod
    def tearDownClass(cls):
        classic_again()

    def test_upgrade_keeps_the_edition(self):
        from planner import upgrade
        new = openpyxl.load_workbook(io.BytesIO(upgrade.convert(SAMPLE.read_bytes())[0]))
        self.assertEqual(new.defined_names["PlannerVersion"].attr_text, '"2.1-m365"')
        self.assertIn("L_WORKING_DAYS", new.defined_names)

    def test_release_11_moves_across(self):
        from planner import upgrade
        new = openpyxl.load_workbook(io.BytesIO(upgrade.convert(
            (ROOT / "release-1.1" / "Project_Planner.xlsx").read_bytes(), release="m365")[0]))
        self.assertEqual(new.defined_names["TemplateVersion"].attr_text, '"1.1 M365"')
        col = head(new["Tasks"])
        self.assertEqual(new["Tasks"].cell(task_row(new, 13), col["Sandbox Duration"]).value, 7)


if __name__ == "__main__":
    unittest.main()
