"""Log sheet: risks, assumptions, issues, dependencies, decisions and change requests in one table (Excel Table "Log").

One list is easier to keep up to date than three: every item has the same core (Ref, Type, Title, Owner, dates,
Status, linked task, latest update), and the columns that only one kind of item needs are greyed out on the other
rows. RAID items are scored (Probability x Impact; an issue has already happened, so it counts as probability 5)
and rated against the Settings thresholds; decisions and change requests record the options and reasons, who
decided and when, and a change's schedule and cost impact. The Status dropdown offers the statuses of the row's
type. The Status sheet lists the top open RAID items, counts high-rated ones in its RAG and totals the changes.
"""
import datetime as dt

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from .core import DATE, Col, Tab, add_table, col_range, fill, has, linked_task, name, title, v11

SHEET = "Log"
D = dt.date
RAID = ["Risk", "Assumption", "Issue", "Dependency"]
TYPES = RAID + ["Decision", "Change"]
CLOSED = ["Closed", "Approved", "Rejected", "Superseded", "Implemented"]     # every other status is still open
# Ref, Type, Title, Owner, Raised, Due, Status, Task UID, Probability, Impact, Options and Reasons, Decided By,
# Decided, Schedule Impact, Cost Impact, Related Ref, Response, Update
SAMPLE = [
    (1, "Risk", "Payment provider certification takes longer than planned", "Jordan Lee", D(2026, 10, 1),
     D(2026, 10, 30), "Open", 14, 4, 4, None, None, None, None, None, None,
     "Start the certification paperwork now; book the provider's test slot.", "Paperwork submitted 20-Oct."),
    (2, "Risk", "UX designer on leave for two weeks in November", "Casey Kim", D(2026, 10, 5), D(2026, 11, 2), "Open",
     18, 3, 3, None, None, None, None, None, None, "Agree cover with the design agency.", None),
    (3, "Issue", "Core API endpoints running 6 days behind plan", "Priya Sharma", D(2026, 10, 20), D(2026, 10, 23),
     "In Progress", 13, None, 4, None, None, None, None, None, 11,
     "Add a second developer for two weeks (Change 11).", "Escalated to sponsor."),
    (4, "Assumption", "Client UAT team available from 7 December", "Morgan Patel", D(2026, 9, 15), D(2026, 11, 20),
     "Open", 22, 2, 4, None, None, None, None, None, None,
     "Confirm names and dates at the November steering committee.", None),
    (5, "Dependency", "Test environment delivered by IT operations by 9 November", "Alex Chen", D(2026, 9, 30),
     D(2026, 11, 9), "Open", 21, 3, 5, None, None, None, None, None, None, "Weekly check-in with IT operations.",
     "Hardware ordered."),
    (6, "Risk", "Scope creep on the reporting screens", "Priya Sharma", D(2026, 9, 10), D(2026, 10, 9), "Closed", 17,
     2, 3, None, None, None, None, None, 8, "Reporting screens moved to phase 2 (Decision 8).", "Closed 09-Oct."),
    (7, "Decision", "Use the existing identity provider for customer log-in", "Alex Chen", D(2026, 9, 14), None,
     "Approved", 12, None, None,
     "Options: build our own, buy a new service, reuse. Reuse is cheapest and already security-approved.",
     "Steering committee", D(2026, 9, 14), None, None, None, None, None),
    (8, "Decision", "Move the reporting screens to phase 2", "Priya Sharma", D(2026, 10, 9), None, "Approved", 17,
     None, None, "Keeps the Build deadline (06-Nov). Reporting is not needed for go-live.", "Sponsor",
     D(2026, 10, 9), None, None, 6, None, None),
    (9, "Decision", "Add a second developer to the core API for two weeks", "Priya Sharma", D(2026, 10, 22),
     D(2026, 10, 30), "Proposed", 13, None, None,
     "Recovers about 4 of the 6 days' slip. Cost in Change 11.", None, None, None, None, 3, None, None),
    (10, "Change", "Add a newsletter sign-up to the portal", "Marketing", D(2026, 9, 28), None, "Approved", 17, None,
     None, "Supports the launch campaign.", "Sponsor", D(2026, 10, 2), 2, 1360, None, None, None),
    (11, "Change", "Second developer on the core API for two weeks", "Priya Sharma", D(2026, 10, 22), D(2026, 10, 30),
     "Under Review", 13, None, None, "Recovers most of the API slip (Issue 3).", None, None, -4, 6800, 3, None, None),
    (12, "Change", "Live chat on every page", "Sales", D(2026, 10, 12), None, "Rejected", None, None, None,
     "Nice to have; not needed for go-live.", "Steering committee", D(2026, 10, 16), 8, 9600, None, None, None),
]
# template 1.1 (F-1102): an external dependency the sample adds, and when each is expected (Ref -> date)
V11_SAMPLE = [(13, "Dependency", "Payment provider sandbox credentials from the vendor", "Jordan Lee",
               D(2026, 10, 12), D(2026, 10, 27), "Open", 14, 3, 4, None, None, None, None, None, None,
               "Chase the vendor's account manager daily.", "Promised for early November.")]
EXPECTED = {5: D(2026, 11, 9), 13: D(2026, 11, 4)}
NAMES = {"Ref": "log_Ref", "Type": "log_Type", "Title": "log_Title", "Owner": "log_Owner", "Due": "log_Due",
         "Status": "log_Status", "Score": "log_Score", "Rating": "log_Rating", "Overdue": "log_Overdue",
         "Task": "log_Task", "Response": "log_Response", "Schedule Impact": "log_Days", "Cost Impact": "log_Cost",
         "_Open": "log_Open", "_Label": "log_Label", "_TopKey": "log_TopKey"}
# columns that only some types use: greyed out on the other rows
ONLY = {("Probability", "Impact", "Score", "Rating", "Response"): RAID,
        ("Options and Reasons", "Decided By", "Decided", "Waiting (days)"): ["Decision", "Change"],
        ("Schedule Impact", "Cost Impact"): ["Change"]}


def is_type(type_cell, types):
    """TRUE when the row's Type is one of types (a sum of comparisons: no array formula, nothing volatile)."""
    return "(" + "+".join(f'({type_cell}="{t}")' for t in types) + ">0)"


ROWS = 150


def _dependency_columns(r, blank_row, data):
    """Template 1.1 (F-1102): when an external dependency is expected, the linked task's forecast start, and the
    working days between them (negative = it arrives after the task should start)."""
    task_start = f'IFERROR(INDEX(tsk_FcStart,MATCH({r("Task UID")},tsk_UID,0)),"")'
    expected = r("Expected")
    return [
        Col("Expected", "in", 12, DATE, values=[EXPECTED.get(row[0]) for row in data],
            note="Dependency: the date the other party delivers it."),
        Col("Task Start", width=12, fmt=DATE, formula=f'IF(N({r("Task UID")})=0,"",{task_start})',
            note="The linked task's forecast start."),
        Col("Gap (days)", width=9, fmt="+0;-0;0", center=True,
            note="Dependency: working days from Expected to the task's start; negative = it arrives too late.",
            formula=(f'IF(OR({blank_row},N({expected})=0,N({r("Task Start")})=0),"",'
                     f'IF({r("Task Start")}>={expected},NETWORKDAYS.INTL({expected},{r("Task Start")},WorkWeek,Holidays)'
                     f'-1,1-NETWORKDAYS.INTL({r("Task Start")},{expected},WorkWeek,Holidays)))')),
    ]


def build_log(wb, blank, n_rows=ROWS):
    ws = wb.create_sheet(SHEET)
    title(ws, "Log: Risks, Assumptions, Issues, Dependencies, Decisions and Changes",
          "One row per item; pick its Type first. Grey cells don't apply to that type. RAID score = Probability x "
          "Impact (1-5 each; an issue counts as probability 5), rated by Settings > RAIDHighScore / RAIDMediumScore. "
          "Link a task by its UID, and another log item by its Ref.")
    ws.column_dimensions["A"].width = 3
    data = [] if blank else SAMPLE + (V11_SAMPLE if v11() else [])
    t = Tab("Log")
    r = t.row
    blank_row = f'{r("Title")}=""'
    raid = is_type(r("Type"), RAID)
    open_ = f'{r("_Open")}=1'
    values = lambda i: [row[i] for row in data]
    letters = add_table(ws, "Log", 4, 2, [
        Col("Ref", "in", 6, center=True, values=list(range(1, n_rows + 1)),
            note="Permanent number, quoted in reports and other items (e.g. Risk 3, Change 11)."),
        Col("Type", "in", 12, values=values(1),
            note="Risk, Assumption, Issue, Dependency (RAID), Decision or Change (a change request)."),
        Col("Title", "in", 44, values=values(2)),
        Col("Owner", "in", 16, values=values(3), note="Who acts on it; for a change request, who asked for it."),
        Col("Raised", "in", 12, DATE, values=values(4)),
        Col("Due", "in", 12, DATE, values=values(5), note="Date the next action, mitigation or decision is due."),
        Col("Status", "in", 12, center=True, values=values(6),
            note="RAID: Open, In Progress, Closed. Decision: Proposed, Approved, Rejected, Superseded. Change: "
                 "Submitted, Under Review, Approved, Rejected, Implemented."),
        Col("Task UID", "in", 8, center=True, values=values(7), note="UID of the task this item affects (optional)."),
        Col("Task", width=28, formula=linked_task(r("Task UID"))),
        Col("Task Finish", width=12, fmt=DATE, formula=(
            f'IF(N({r("Task UID")})=0,"",IFERROR(INDEX(tsk_FcFinish,MATCH({r("Task UID")},tsk_UID,0)),""))'),
            note="The linked task's forecast finish."),
        Col("Overdue", width=9, center=True, note="Still open and the due date is before the status date.",
            formula=f'IF(OR({blank_row},N({r("Due")})=0),"",IF(AND({open_},{r("Due")}<StatusDate),"Overdue",""))'),
        *(_dependency_columns(r, blank_row, data) if v11() else []),
        # RAID
        Col("Probability", "in", 9, center=True, values=values(8), note="RAID: 1 (rare) to 5 (almost certain). "
                                                                       "Ignored for issues."),
        Col("Impact", "in", 8, center=True, values=values(9), note="RAID: 1 (minor) to 5 (severe)."),
        Col("Score", width=7, center=True, note="Probability x Impact; an issue counts as probability 5.",
            formula=(f'IF(OR({blank_row},NOT({raid})),"",'
                     f'IF({r("Type")}="Issue",5,N({r("Probability")}))*N({r("Impact")}))')),
        Col("Rating", width=8, center=True,
            formula=(f'IF({r("Score")}="","",IF({r("Score")}>=RAIDHighScore,"High",'
                     f'IF({r("Score")}>=RAIDMediumScore,"Medium","Low")))')),
        Col("Response", "in", 40, values=values(16), note="RAID: mitigation, contingency or action."),
        # decisions and changes
        Col("Options and Reasons", "in", 44, values=values(10),
            note="Decision: the options considered and why this one. Change: why it is needed."),
        Col("Decided By", "in", 16, values=values(11)),
        Col("Decided", "in", 12, DATE, values=values(12), note="Date it was approved or rejected."),
        Col("Waiting (days)", width=9, center=True, note="Days an open decision or change has waited, to the status date.",
            formula=(f'IF(OR({blank_row},N({r("Raised")})=0,NOT({is_type(r("Type"), ["Decision", "Change"])}),'
                     f'NOT({open_})),"",MAX(0,StatusDate-{r("Raised")}))')),
        Col("Schedule Impact", "in", 9, "+0;-0;0", center=True, values=values(13),
            note="Change: working days (negative = saves time), as estimated when raised."),
        Col("Cost Impact", "in", 11, "#,##0", values=values(14), note="Change: cost, as estimated when raised."),
        Col("Related Ref", "in", 8, center=True, values=values(15), note="Ref of a related log item (optional)."),
        Col("Update", "in", 30, values=values(17), note="Latest progress note."),
        Col("_Open", "help", 6, note="1 while the item is open (its status is not a closed one).",
            formula=(f'IF({blank_row},0,IF(ISNA(MATCH({r("Status")},'
                     f'{{{",".join(chr(34) + s + chr(34) for s in CLOSED)}}},0)),1,0))')),
        Col("_Label", "help", 8, formula=f'IF({blank_row},"",{r("Type")}&" "&{r("Ref")}&": "&{r("Title")})',
            note="Type, Ref and Title in one cell, for the Status sheet."),
        Col("_TopKey", "help", 8, note="Open RAID items ranked by score (the Status sheet's top list).",
            formula=(f'IF(AND({r("Score")}<>"",{open_}),'
                     f'{r("Score")}+1-(ROW()-ROW({t.head("Ref")}))/100000,"")')),
    ] + ([Col("_OverdueKey", "help", 8, note="Overdue items by due date (the Check sheet's list).",
              formula=(f'IF({r("Overdue")}="Overdue",{r("Due")}+(ROW()-ROW({t.head("Ref")}))/100000,"")'))]
         if has("check") else [])
       + ([Col("_DepKey", "help", 8, note="Open dependencies with an Expected date, latest against the task first "
                                         "(the Status sheet's list).",
               formula=(f'IF(AND({r("Type")}="Dependency",{open_},N({r("Expected")})>0),'
                        f'IF({r("Gap (days)")}="",100000,{r("Gap (days)")})+(ROW()-ROW({t.head("Ref")}))/100000,"")'))]
          if v11() else []), n_rows)
    extra = {"_OverdueKey": "log_OverdueKey"} if has("check") else {}
    if v11():
        extra.update({"Expected": "log_Expected", "Gap (days)": "log_DepGap", "_DepKey": "log_DepKey",
                      "Update": "log_Update"})
    for column, defined in {**NAMES, **extra}.items():
        name(wb, defined, f"Log[{column}]")
    last_helper = "_DepKey" if v11() else "_OverdueKey" if has("check") else "_TopKey"
    ws.column_dimensions.group(letters["_Open"], letters[last_helper], hidden=True, outline_level=1)
    first, last = 5, 4 + n_rows
    rng = lambda k: col_range(letters[k], first, last)
    type_cell = f'${letters["Type"]}{first}'
    status = DataValidation(type="list", formula1=(f'=IF({type_cell}="Decision",L_DecisionStatus,'
                                                   f'IF({type_cell}="Change",L_ChangeStatus,L_RAIDStatus))'))
    for rule, keys in ((DataValidation(type="list", formula1="=L_LogType"), ("Type",)), (status, ("Status",)),
                       (DataValidation(type="whole", operator="between", formula1="1", formula2="5"), ("Probability", "Impact")),
                       (DataValidation(type="date", operator="greaterThan", formula1="1"), ("Raised", "Due", "Decided")),
                       (DataValidation(type="list", formula1="=res_Name", showErrorMessage=False), ("Owner",))):
        ws.add_data_validation(rule)
        for k in keys:
            rule.add(rng(k))
    cf = ws.conditional_formatting
    only = {**ONLY, ("Expected", "Task Start", "Gap (days)"): ["Dependency"]} if v11() else ONLY
    if v11():
        expected = DataValidation(type="date", operator="greaterThan", formula1="1")
        ws.add_data_validation(expected)
        expected.add(rng("Expected"))
        gap = letters["Gap (days)"]
        cf.add(rng("Gap (days)"), FormulaRule(formula=[f'AND(ISNUMBER({gap}{first}),{gap}{first}<0)'],
                                              font=Font(color="C00000", bold=True)))
    for keys, types in only.items():                 # not this type's column: grey, like a calculated cell
        rule = FormulaRule(formula=[f'AND({type_cell}<>"",NOT({is_type(type_cell, types)}))'], fill=fill("EDEDED"),
                           font=Font(color="A6A6A6"))
        for k in keys:
            cf.add(rng(k), rule)
    rating = letters["Rating"]
    for word, (bg, fg) in {"High": ("FFC7CE", "9C0006"), "Medium": ("FFEB9C", "9C5700"), "Low": ("C6EFCE", "006100")}.items():
        cf.add(rng("Rating"), FormulaRule(formula=[f'{rating}{first}="{word}"'], fill=fill(bg), font=Font(color=fg, bold=True)))
    cf.add(rng("Overdue"), FormulaRule(formula=[f'{letters["Overdue"]}{first}="Overdue"'], font=Font(color="C00000", bold=True)))
    st = letters["Status"]
    for word, colour in (("Proposed", "C55A11"), ("Submitted", "C55A11"), ("Under Review", "C55A11"),
                         ("Approved", "548235"), ("Implemented", "548235")):
        cf.add(rng("Status"), FormulaRule(formula=[f'{st}{first}="{word}"'], font=Font(color=colour, bold=True)))
    cf.add(f"B{first}:{letters['Update']}{last}", FormulaRule(
        formula=[f'AND(${letters["Title"]}{first}<>"",${letters["_Open"]}{first}=0)'], font=Font(color="A6A6A6")))
    ws.freeze_panes = f"E{first}"
    return ws
