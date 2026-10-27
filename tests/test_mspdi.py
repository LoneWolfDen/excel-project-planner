"""Tests for the MS Project XML (MSPDI) export: planner/mspdi.py, scripts/export_mspdi.py and docs/export.html.

The export reads the values saved in the workbook, so most tests run on the committed Project_Planner.xlsx and need
no LibreOffice. The browser test runs docs/export.html in headless Chromium when Node's playwright and a local copy of
Pyodide are available (PYODIDE_DIR; CI downloads it), and checks it writes exactly what the command line writes.
"""
import datetime as dt
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from planner import mspdi, mspdi_import, tasks, upgrade  # noqa: E402
import build_export_page  # noqa: E402
from test_planner import (BASELINE, FIRST, SOFFICE, SOURCE, build, columns, errors_in, recalc, set_input,  # noqa: E402
                          table_rows)

NS = "{http://schemas.microsoft.com/project}"
SAMPLE = ROOT / "Project_Planner.xlsx"


def parse(xml):
    return ET.fromstring(xml.encode("utf-8"))


def tasks_by_uid(root):
    return {int(t.find(NS + "UID").text): t for t in root.iter(NS + "Task")}


def text(element, tag):
    found = element.find(NS + tag)
    return None if found is None else found.text


def day(value):
    return dt.date.fromisoformat(value[:10])


class ExportTest(unittest.TestCase):
    """The committed sample workbook, exported and checked against its own Tasks sheet."""

    @classmethod
    def setUpClass(cls):
        cls.xml = mspdi.convert(SAMPLE.read_bytes())
        cls.root = parse(cls.xml)
        cls.tasks = tasks_by_uid(cls.root)
        wb = openpyxl.load_workbook(SAMPLE, data_only=True)
        ws = wb["Tasks"]
        col = columns(ws)
        cls.sheet = {}
        for r in range(FIRST, ws.max_row + 1):
            if ws.cell(r, col["Task Name"]).value:
                cls.sheet[ws.cell(r, col["UID"]).value] = {h: ws.cell(r, c).value for h, c in col.items()}

    def test_project_and_calendar(self):
        p = self.root
        self.assertEqual(p.tag, NS + "Project")
        self.assertEqual((text(p, "Title"), text(p, "Manager")), ("Customer Portal Implementation", "Priya Sharma"))
        self.assertEqual(text(p, "StartDate"), "2026-09-07T08:00:00")
        self.assertEqual((text(p, "MinutesPerDay"), text(p, "MinutesPerWeek")), ("480", "2400"))
        days = p.findall(f"{NS}Calendars/{NS}Calendar/{NS}WeekDays/{NS}WeekDay")
        week = {int(text(d, "DayType")): text(d, "DayWorking") for d in days if text(d, "DayType") != "0"}
        self.assertEqual(week, {1: "0", 2: "1", 3: "1", 4: "1", 5: "1", 6: "1", 7: "0"})     # Sunday = 1
        holidays = [day(d.find(f"{NS}TimePeriod/{NS}FromDate").text) for d in days if text(d, "DayType") == "0"]
        self.assertEqual(holidays, [dt.date(2026, 12, 24), dt.date(2026, 12, 25), dt.date(2026, 12, 31),
                                    dt.date(2027, 1, 1)])
        monday = next(d for d in days if text(d, "DayType") == "2")
        shifts = [(text(w, "FromTime"), text(w, "ToTime")) for w in monday.iter(NS + "WorkingTime")]
        self.assertEqual(shifts, [("08:00:00", "12:00:00"), ("13:00:00", "17:00:00")])

    def test_every_task_matches_the_tasks_sheet(self):
        self.assertEqual(sorted(self.tasks), sorted(self.sheet))
        for uid, t in self.tasks.items():
            row = self.sheet[uid]
            self.assertEqual(text(t, "Name"), row["Task Name"], uid)
            self.assertEqual(int(text(t, "ID")), row["ID"], uid)
            self.assertEqual(text(t, "WBS"), str(row["WBS"]), uid)
            self.assertEqual(int(text(t, "OutlineLevel")), row["Level"], uid)
            self.assertEqual(day(text(t, "Start")), row["Forecast Start"].date(), uid)
            self.assertEqual(day(text(t, "Finish")), row["Forecast Finish"].date(), uid)
            self.assertEqual(text(t, "Summary"), "1" if row["_IsSummary"] else "0", uid)
            self.assertEqual(int(text(t, "PercentComplete")), round(row["Progress"] * 100), uid)

    def test_durations_links_and_constraints(self):
        t = self.tasks
        self.assertEqual(text(t[6], "Duration"), "PT64H0M0S")                    # 8 days x 8 hours
        link = lambda uid: [(int(text(l, "PredecessorUID")), int(text(l, "Type")), int(text(l, "LinkLag")))
                            for l in t[uid].findall(NS + "PredecessorLink")]
        self.assertEqual(link(7), [(6, 3, 3 * 8 * 600)])                        # SS + 3 days, in tenths of a minute
        self.assertEqual(link(18), [(17, 0, 1 * 8 * 600)])                      # FF + 1 day
        self.assertEqual(link(9), [(7, 1, 0), (8, 1, 0)])                       # two FS links
        self.assertEqual(link(24), [(22, 1, 0), (23, 1, 0), (21, 1, 0)])        # three
        self.assertEqual(link(1), [])                                           # summaries carry no links
        self.assertEqual((text(t[22], "ConstraintType"), text(t[22], "ConstraintDate")), ("4", "2026-12-07T08:00:00"))
        self.assertEqual((text(t[28], "Manual"), text(t[28], "ConstraintType")), ("1", "0"))
        self.assertEqual(day(text(t[28], "Start")), dt.date(2026, 11, 16))
        self.assertEqual((text(t[4], "Milestone"), text(t[4], "Duration")), ("1", "PT0H0M0S"))
        self.assertEqual(text(t[19], "Deadline"), "2026-11-06T17:00:00")
        self.assertEqual((text(t[2], "ActualStart"), text(t[2], "ActualFinish"), text(t[2], "ActualCost")),
                         ("2026-09-07T08:00:00", "2026-09-10T17:00:00", "3040"))
        self.assertIsNone(t[2].find(NS + "Baseline"))                             # no baseline in the sample

    def test_custom_fields(self):
        values = lambda uid: {text(a, "FieldID"): text(a, "Value") for a in self.tasks[uid].findall(NS + "ExtendedAttribute")}
        self.assertEqual(values(17), {mspdi.TEXT1: "Sprint 4", mspdi.NUMBER1: "8"})
        self.assertEqual(values(4), {mspdi.TEXT1: "Sprint 1"})                     # a milestone has no points
        aliases = {text(a, "FieldID"): text(a, "Alias") for a in self.root.find(NS + "ExtendedAttributes")}
        self.assertEqual(aliases, {mspdi.TEXT1: "Sprint", mspdi.NUMBER1: "Story Points"})

    def test_resources_and_assignments(self):
        resources = {text(r, "Name"): r for r in self.root.iter(NS + "Resource")}
        self.assertEqual(len(resources), 7)
        self.assertEqual((text(resources["Priya Sharma"], "MaxUnits"), text(resources["Priya Sharma"], "StandardRate")),
                         ("0.5", "95"))
        uid_of = {text(r, "UID"): n for n, r in resources.items()}
        assignments = {}
        for a in self.root.iter(NS + "Assignment"):
            assignments.setdefault(int(text(a, "TaskUID")), []).append(
                (uid_of[text(a, "ResourceUID")], float(text(a, "Units")), text(a, "Work")))
        expected = {u: r for u, r in self.sheet.items() if r["Resource"] and not r["_IsSummary"]}
        self.assertEqual(sorted(assignments), sorted(expected))
        self.assertEqual(assignments[18], [("Casey Kim", 0.5, "PT12H0M0S")])     # 3 days at 50%
        self.assertEqual(assignments[2], [("Priya Sharma", 1.0, "PT24H0M0S")])
        # the Assignments sheet's people follow the task's own: 5 days, Alex full time and Jordan at 50%
        self.assertEqual(assignments[7], [("Alex Chen", 1.0, "PT40H0M0S"), ("Jordan Lee", 0.5, "PT20H0M0S")])
        self.assertEqual(text(self.tasks[7], "Work"), "PT60H0M0S")
        self.assertEqual(sum(len(v) for v in assignments.values()), len(expected) + 4)

    def test_elements_in_ms_project_order(self):
        # MS Project reports a file with elements out of order as corrupt; every tag must be in ORDER
        for element in self.root.iter():
            tag = element.tag.replace(NS, "")
            if tag in mspdi.ORDER:
                children = [c.tag.replace(NS, "") for c in element]
                unknown = [c for c in children if c not in mspdi.ORDER[tag]]
                self.assertEqual(unknown, [], tag)
                positions = [mspdi.ORDER[tag].index(c) for c in children]
                self.assertEqual(positions, sorted(positions), tag)

    def test_command_line(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "project.xml"
            subprocess.run([sys.executable, str(ROOT / "scripts" / "export_mspdi.py"), str(SAMPLE), str(out)],
                           check=True, capture_output=True)
            self.assertEqual(out.read_text(encoding="utf-8"), self.xml)


class ExportErrorTest(unittest.TestCase):
    def test_not_a_workbook(self):
        with self.assertRaisesRegex(mspdi.ExportError, "isn't an .xlsx"):
            mspdi.convert(b"not a zip file")

    def test_blank_template_needs_a_start_date(self):
        with self.assertRaisesRegex(mspdi.ExportError, "Project start date is empty"):
            mspdi.convert((ROOT / "Project_Planner_Blank.xlsx").read_bytes())

    def test_workbook_never_calculated(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "raw.xlsx"
            build(path)                                    # openpyxl writes formulas without values
            with self.assertRaisesRegex(mspdi.ExportError, "no calculated values"):
                mspdi.convert(path.read_bytes())

    def test_command_line_reports_the_problem(self):
        result = subprocess.run([sys.executable, str(ROOT / "scripts" / "export_mspdi.py"),
                                 str(ROOT / "Project_Planner_Blank.xlsx"), os.devnull], capture_output=True, text=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("Project start date is empty", result.stderr)


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class ExportEditsTest(unittest.TestCase):
    """Edits made in the workbook, recalculated with LibreOffice, then exported."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.dir = Path(cls.tmp.name)
        cls.path = cls.dir / "Project_Planner.xlsx"
        build(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def export(self, wb):
        out = self.dir / "edited.xlsx"
        wb.save(out)       # data_only workbook: values only, as the export reads them
        return tasks_by_uid(parse(mspdi.convert(out.read_bytes())))

    def test_baseline_and_ignored_links(self):
        first = recalc(self.path, self.dir)
        ws = first["Tasks"]
        col = columns(ws)
        snapshot = {(r, i): ws.cell(r, col[s]).value for r in range(FIRST, FIRST + 30) for i, s in enumerate(SOURCE)}

        def edit(wb):
            for (r, i), v in snapshot.items():
                wb["Tasks"].cell(r, col[BASELINE[i]]).value = v
            set_input(wb, 21, "Dur", 10)
            set_input(wb, 23, "P2", 30)            # a link to a task below: ignored, so not exported
        tasks = self.export(recalc(self.path, self.dir, edit))
        base = tasks[21].find(NS + "Baseline")
        self.assertEqual((text(base, "Number"), day(text(base, "Start")), day(text(base, "Finish"))),
                         ("0", dt.date(2026, 11, 10), dt.date(2026, 11, 19)))
        self.assertEqual((text(base, "Duration"), text(base, "Work"), text(base, "Cost")),
                         ("PT64H0M0S", "PT96H0M0S", "6720"))            # Taylor, and Sam at 50%
        self.assertEqual(text(tasks[21], "Duration"), "PT80H0M0S")               # the current plan, not the baseline
        self.assertEqual([text(l, "PredecessorUID") for l in tasks[23].findall(NS + "PredecessorLink")], ["21"])


# A file written the way MS Project writes MSPDI, covering what the planner can't hold
MS_PROJECT_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Project xmlns="http://schemas.microsoft.com/project">
  <SaveVersion>14</SaveVersion><Name>Office move.xml</Name><Title>Office move</Title><Manager>Ana Silva</Manager>
  <StartDate>2027-03-01T09:00:00</StartDate><CalendarUID>1</CalendarUID><MinutesPerDay>450</MinutesPerDay>
  <StatusDate>2027-03-12T17:00:00</StatusDate>
  <Calendars>
    <Calendar><UID>1</UID><Name>Standard</Name><IsBaseCalendar>1</IsBaseCalendar>
      <WeekDays>
        <WeekDay><DayType>1</DayType><DayWorking>0</DayWorking></WeekDay>
        <WeekDay><DayType>2</DayType><DayWorking>1</DayWorking></WeekDay>
        <WeekDay><DayType>3</DayType><DayWorking>1</DayWorking></WeekDay>
        <WeekDay><DayType>4</DayType><DayWorking>1</DayWorking></WeekDay>
        <WeekDay><DayType>5</DayType><DayWorking>1</DayWorking></WeekDay>
        <WeekDay><DayType>6</DayType><DayWorking>1</DayWorking></WeekDay>
        <WeekDay><DayType>7</DayType><DayWorking>0</DayWorking></WeekDay>
      </WeekDays>
      <Exceptions>
        <Exception><TimePeriod><FromDate>2027-03-26T00:00:00</FromDate><ToDate>2027-03-29T23:59:00</ToDate></TimePeriod>
          <Name>Easter</Name><DayWorking>0</DayWorking></Exception>
        <Exception><TimePeriod><FromDate>2027-03-31T00:00:00</FromDate><ToDate>2027-03-31T23:59:00</ToDate></TimePeriod>
          <Name>Half day</Name><DayWorking>1</DayWorking></Exception>
      </Exceptions>
    </Calendar>
    <Calendar><UID>2</UID><Name>Ben</Name><IsBaseCalendar>0</IsBaseCalendar><BaseCalendarUID>1</BaseCalendarUID>
      <Exceptions>
        <Exception><TimePeriod><FromDate>2027-03-15T00:00:00</FromDate><ToDate>2027-03-17T23:59:00</ToDate></TimePeriod>
          <DayWorking>0</DayWorking></Exception>
      </Exceptions>
    </Calendar>
  </Calendars>
  <Tasks>
    <Task><UID>0</UID><ID>0</ID><Name>Office move</Name><OutlineLevel>0</OutlineLevel><Summary>1</Summary></Task>
    <Task><UID>10</UID><ID>1</ID><Name>Plan</Name><Manual>0</Manual><OutlineLevel>1</OutlineLevel><Summary>1</Summary>
      <PredecessorLink><PredecessorUID>40</PredecessorUID><Type>1</Type></PredecessorLink></Task>
    <Task><UID>11</UID><ID>2</ID><Name>Survey</Name><OutlineLevel>2</OutlineLevel><Duration>PT22H30M0S</Duration>
      <PercentComplete>100</PercentComplete><ActualStart>2027-03-01T09:00:00</ActualStart>
      <ActualFinish>2027-03-03T17:00:00</ActualFinish><ActualCost>900</ActualCost>
      <Baseline><Number>0</Number><Cost>800</Cost><Duration>PT15H0M0S</Duration><Finish>2027-03-02T17:00:00</Finish>
        <Start>2027-03-01T09:00:00</Start><Work>PT15H0M0S</Work><DurationFormat>7</DurationFormat></Baseline>
      <ExtendedAttribute><FieldID>188743767</FieldID><Value>5</Value></ExtendedAttribute></Task>
    <Task><UID>12</UID><ID>3</ID><Name>Floor plan</Name><OutlineLevel>2</OutlineLevel><Duration>PT15H0M0S</Duration>
      <ConstraintType>1</ConstraintType>
      <PredecessorLink><PredecessorUID>11</PredecessorUID><Type>1</Type><LinkLag>14400</LinkLag><LagFormat>8</LagFormat></PredecessorLink>
      <PredecessorLink><PredecessorUID>11</PredecessorUID><Type>3</Type><LinkLag>50</LinkLag><LagFormat>19</LagFormat></PredecessorLink>
    </Task>
    <Task><UID>13</UID><ID>4</ID><Name>Book movers</Name><OutlineLevel>2</OutlineLevel><Duration>PT7H30M0S</Duration>
      <ConstraintType>2</ConstraintType><ConstraintDate>2027-03-10T09:00:00</ConstraintDate>
      <Deadline>2027-03-19T17:00:00</Deadline>
      <PredecessorLink><PredecessorUID>11</PredecessorUID><Type>1</Type><LinkLag>9000</LinkLag><LagFormat>7</LagFormat></PredecessorLink>
      <PredecessorLink><PredecessorUID>12</PredecessorUID><Type>0</Type></PredecessorLink>
      <PredecessorLink><PredecessorUID>40</PredecessorUID><Type>1</Type></PredecessorLink></Task>
    <Task><UID>14</UID><ID>5</ID><Name></Name><IsNull>1</IsNull></Task>
    <Task><UID>20</UID><ID>6</ID><Name>Move day</Name><OutlineLevel>1</OutlineLevel><Duration>PT0H0M0S</Duration>
      <Milestone>1</Milestone><ConstraintType>7</ConstraintType><ConstraintDate>2027-04-02T17:00:00</ConstraintDate>
      <PredecessorLink><PredecessorUID>13</PredecessorUID><Type>1</Type></PredecessorLink>
      <PredecessorLink><PredecessorUID>11</PredecessorUID><Type>1</Type></PredecessorLink>
      <PredecessorLink><PredecessorUID>12</PredecessorUID><Type>1</Type></PredecessorLink>
      <PredecessorLink><PredecessorUID>30</PredecessorUID><Type>1</Type></PredecessorLink>
      <PredecessorLink><PredecessorUID>40</PredecessorUID><Type>1</Type></PredecessorLink></Task>
    <Task><UID>30</UID><ID>7</ID><Name>Supplier call</Name><Manual>1</Manual><OutlineLevel>1</OutlineLevel>
      <Start>2027-03-08T09:00:00</Start><Duration>PT7H30M0S</Duration><Notes>Agreed by phone.</Notes></Task>
    <Task><UID>40</UID><ID>8</ID><Name>Budget sign-off</Name><OutlineLevel>1</OutlineLevel><Duration>PT7H30M0S</Duration></Task>
  </Tasks>
  <Resources>
    <Resource><UID>0</UID><ID>0</ID><Name>Unassigned</Name><IsNull>1</IsNull></Resource>
    <Resource><UID>1</UID><ID>1</ID><Name>Ana Silva</Name><Type>1</Type><Initials>AS</Initials><Group>PMO</Group>
      <MaxUnits>1.00</MaxUnits><StandardRate>60</StandardRate></Resource>
    <Resource><UID>2</UID><ID>2</ID><Name>Ben Okafor</Name><Type>1</Type><MaxUnits>0.50</MaxUnits>
      <StandardRate>40</StandardRate><CalendarUID>2</CalendarUID></Resource>
    <Resource><UID>3</UID><ID>3</ID><Name>Van hire</Name><Type>2</Type></Resource>
  </Resources>
  <Assignments>
    <Assignment><UID>1</UID><TaskUID>11</TaskUID><ResourceUID>2</ResourceUID><Units>0.5</Units><Work>PT11H15M0S</Work></Assignment>
    <Assignment><UID>2</UID><TaskUID>12</TaskUID><ResourceUID>1</ResourceUID><Units>1</Units><Work>PT15H0M0S</Work></Assignment>
    <Assignment><UID>3</UID><TaskUID>12</TaskUID><ResourceUID>2</ResourceUID><Units>0.5</Units><Work>PT7H30M0S</Work></Assignment>
  </Assignments>
</Project>
"""


class ImportTest(unittest.TestCase):
    """MS Project XML into a new workbook (planner/mspdi_import.py, scripts/import_mspdi.py)."""

    def test_reads_what_ms_project_writes(self):
        project, warnings = mspdi_import.read_project(MS_PROJECT_XML.encode())
        rows = {r["UID"]: r for r in project.tasks}
        self.assertEqual(list(rows), [10, 11, 12, 13, 20, 30, 40])           # no project summary (UID 0), no blank row
        self.assertEqual(project.settings["ProjectName"], "Office move")
        self.assertEqual((project.settings["ProjectStartDate"], project.settings["StatusDateInput"]),
                         (dt.date(2027, 3, 1), dt.date(2027, 3, 12)))
        self.assertEqual(project.settings["HoursPerDay"], 7.5)
        self.assertEqual((project.settings["WorkMon"], project.settings["WorkSat"]), ("Yes", "No"))
        self.assertEqual([d for d, _ in project.holidays], [dt.date(2027, 3, d) for d in (26, 27, 28, 29)])
        self.assertEqual(project.leave, {"Ben Okafor": {"Mar-27": "15, 16, 17"}})        # resource calendar
        self.assertEqual([r[0] for r in project.resources], ["Ana Silva", "Ben Okafor", "Van hire"])
        self.assertEqual((project.resources[1][3], project.resources[1][4], project.resources[2][6]), (0.5, 40, "Cost"))
        survey = rows[11]
        self.assertEqual((survey["Dur"], survey["AS"], survey["AF"], survey["Pct"], survey["ACin"], survey["SP"]),
                         (3, dt.date(2027, 3, 1), dt.date(2027, 3, 3), 1, 900, 5))
        self.assertEqual((survey["BS"], survey["BF"], survey["BD"], survey["BW"], survey["BC"]),
                         (dt.date(2027, 3, 1), dt.date(2027, 3, 2), 2, 15, 800))
        self.assertEqual((survey["Res"], survey["Units"]), ("Ben Okafor", 0.5))
        self.assertEqual((rows[12]["Res"], rows[12]["Units"]), ("Ana Silva", None))     # the one with the most work
        self.assertEqual(project.assignments, [(12, "Ben Okafor", 0.5)])                # the other: Assignments sheet
        self.assertEqual((rows[12]["P1"], rows[12]["T1"], rows[12]["L1"]), (11, "FS", 1))   # 1 elapsed day
        self.assertIsNone(rows[12].get("L2"))                                              # % lag dropped
        self.assertEqual([(rows[13][f"P{k}"], rows[13][f"T{k}"]) for k in (1, 2, 3)], [(11, "FS"), (12, "FF"), (40, "FS")])
        self.assertEqual((rows[13]["L1"], rows[13]["DL"]), (2, dt.date(2027, 3, 19)))
        # every constraint type is kept, with its date (none for As Late As Possible)
        self.assertEqual((rows[12]["CType"], rows[12]["Cons"]), ("As Late As Possible", None))
        self.assertEqual((rows[13]["CType"], rows[13]["Cons"]), ("Must Start On", dt.date(2027, 3, 10)))
        self.assertEqual((rows[20]["CType"], rows[20]["Cons"]), ("Finish No Later Than", dt.date(2027, 4, 2)))
        self.assertEqual([rows[20][f"P{k}"] for k in (1, 2, 3, 4)], [13, 11, 12, 30])      # 4 links; the 5th dropped
        self.assertEqual((rows[20]["Dur"], rows[30]["Mode"], rows[30]["Cons"]), (0, "Manual", dt.date(2027, 3, 8)))
        self.assertNotIn("Dur", rows[10])                                                  # summaries roll up
        joined = "\n".join(warnings)
        for expected in ("part-day", "Task 10 (Plan): links on a summary task", "lag in %", "elapsed lag",
                         "Task 20 (Move day): only 4 predecessors are kept; dropped UID 40",
                         "3 links point to a task listed below its successor"):
            self.assertIn(expected, joined)
        self.assertNotIn("constraint", joined)
        self.assertIn("[From MS Project: only 4 predecessors are kept", rows[20]["Notes"])

    def test_bad_files(self):
        for data, message in ((b"not xml", "isn't an XML file"), (b"<a/>", "isn't a Microsoft Project file"),
                              (b'<Project xmlns="http://schemas.microsoft.com/project"/>', "no tasks")):
            with self.assertRaisesRegex(mspdi_import.MSPImportError, message):
                mspdi_import.read_project(data)

    def test_command_line(self):
        with tempfile.TemporaryDirectory() as tmp:
            xml, out = Path(tmp) / "move.xml", Path(tmp) / "move.xlsx"
            xml.write_text(MS_PROJECT_XML, encoding="utf-8")
            result = subprocess.run([sys.executable, str(ROOT / "scripts" / "import_mspdi.py"), str(xml), str(out)],
                                    capture_output=True, text=True, check=True)
            self.assertIn("note: Task 20 (Move day): only 4 predecessors are kept", result.stdout)
            wb = openpyxl.load_workbook(out)
            self.assertEqual(wb["Tasks"].cell(FIRST, columns(wb["Tasks"])["Task Name"]).value, "Plan")


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class ImportRecalcTest(unittest.TestCase):
    """Imported workbooks, recalculated: the sample's round trip schedules identically."""

    def test_round_trip_schedules_the_same(self):
        with tempfile.TemporaryDirectory() as tmp:
            xlsx, warnings = mspdi_import.convert(mspdi.convert(SAMPLE.read_bytes()))
            self.assertEqual(warnings, [])
            path = Path(tmp) / "Project_Planner.xlsx"
            path.write_bytes(xlsx)
            wb = recalc(path, tmp, status_date=False)
            self.assertEqual(errors_in(wb), [])
            original = openpyxl.load_workbook(SAMPLE, data_only=True)
            want = [r for r in table_rows(original, "Tasks", "Tasks") if r["Task Name"]]
            got = [r for r in table_rows(wb, "Tasks", "Tasks") if r["Task Name"]]
            keys = ["UID", "Level", "Task Name", "Mode", "Duration", *[f"Pred{k} {x}" for k in tasks.LINK_NUMBERS
                                                                        for x in ("UID", "Type", "Lag")],
                    "Constraint", "Constraint Date", "Deadline", "Actual Start", "Actual Finish",
                    "% Complete", "Resource", "Units", "Resource Names", "Actual Cost", "Story Points", "Notes", "Forecast Start",
                    "Forecast Finish", "Late Finish", "Total Slack", "Work", "Cost", "Status", "Critical"]
            blank = (None, "", 0)
            diffs = [(w["UID"], k, w[k], g[k]) for w, g in zip(want, got) for k in keys
                     if w[k] != g[k] and not (w[k] in blank and g[k] in blank)]
            self.assertEqual((len(got), diffs), (30, []))
            for n in ("ProjectName", "ProjectManager", "ProjectStartDate", "StatusDateInput", "HoursPerDay", "WorkSat"):
                sheet, ref = wb.defined_names[n].attr_text.split("!")
                a = original[sheet.strip("'")][ref.replace("$", "")].value
                self.assertEqual(wb[sheet.strip("'")][ref.replace("$", "")].value, a, n)
            holidays = lambda w: [r["Date"] for r in table_rows(w, "Holidays", "HolidayTable") if r["Date"]]
            self.assertEqual(holidays(wb), holidays(original))
            people = lambda w: [(r["Resource Name"], r["Max Units"], r["Std Rate"], r["Type"])
                                for r in table_rows(w, "Resources", "ResourceTable") if r["Resource Name"]]
            self.assertEqual(people(wb), people(original))
            extra = lambda w: [(r["Task UID"], r["Resource"], r["Units"], r["Work"])
                               for r in table_rows(w, "Assignments", "AssignmentTable") if r["Task UID"]]
            self.assertEqual(extra(wb), extra(original))

    def test_ms_project_file_builds_a_working_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            xlsx, _ = mspdi_import.convert(MS_PROJECT_XML.encode())
            path = Path(tmp) / "move.xlsx"
            path.write_bytes(xlsx)
            wb = recalc(path, tmp, status_date=False)
            self.assertEqual(errors_in(wb), [])
            rows = {r["UID"]: r for r in table_rows(wb, "Tasks", "Tasks") if r["Task Name"]}
            self.assertEqual(rows[11]["Forecast Finish"], dt.date(2027, 3, 3))       # actual finish
            self.assertEqual((rows[13]["Constraint"], rows[13]["Constraint Date"]), ("Must Start On", dt.date(2027, 3, 10)))
            # Must Start On pins book movers to 10-Mar (over its FF link); unstarted work then moves to the status date
            self.assertEqual((rows[13]["Forecast Start"], rows[13]["Total Slack"]), (dt.date(2027, 3, 12), -2))
            self.assertEqual(rows[13]["Link Check"], "Pred 3: UID 40 is below this task: move it above.")
            self.assertEqual((rows[20]["Constraint"], rows[20]["Late Finish"]), ("Finish No Later Than", dt.date(2027, 3, 15)))
            leave = {r["Resource"]: r["Mar-27"] for r in table_rows(wb, "Leave", "LeaveTable") if r["Resource"]}
            self.assertEqual(leave, {"Ben Okafor": "15, 16, 17"})


class ExportPageTest(unittest.TestCase):
    def test_page_is_up_to_date(self):
        page = (ROOT / "docs" / "export.html").read_text(encoding="utf-8")
        self.assertEqual(page, build_export_page.page(), "run: python scripts/build_export_page.py")

    @unittest.skipUnless(shutil.which("node") and os.environ.get("PYODIDE_DIR"),
                         "needs node with playwright, and PYODIDE_DIR (a local Pyodide 0.26.4)")
    def test_browser_writes_what_the_command_line_writes(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "page.xml"
            result = subprocess.run(["node", str(ROOT / "tests" / "browser" / "export_page.js"),
                                     str(ROOT / "docs" / "export.html"), str(SAMPLE), os.environ["PYODIDE_DIR"], str(out)],
                                    capture_output=True, text=True, timeout=600)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("STATUS Done: 30 tasks, 7 resources.", result.stdout)
            self.assertEqual(out.read_text(encoding="utf-8"), mspdi.convert(SAMPLE.read_bytes()))
            blank = subprocess.run(["node", str(ROOT / "tests" / "browser" / "export_page.js"),
                                    str(ROOT / "docs" / "export.html"), str(ROOT / "Project_Planner_Blank.xlsx"),
                                    os.environ["PYODIDE_DIR"], str(out)], capture_output=True, text=True, timeout=600)
            self.assertIn("Project start date is empty", blank.stdout)
            # MS Project XML in: the page builds the same workbook as the command line
            xml, page_xlsx = Path(tmp) / "move.xml", Path(tmp) / "page.xlsx"
            xml.write_text(MS_PROJECT_XML, encoding="utf-8")
            result = subprocess.run(["node", str(ROOT / "tests" / "browser" / "export_page.js"),
                                     str(ROOT / "docs" / "export.html"), str(xml), os.environ["PYODIDE_DIR"],
                                     str(page_xlsx)], capture_output=True, text=True, timeout=600)
            self.assertIn("STATUS Done: 7 tasks, 3 resources. Notes:", result.stdout)
            cli = Path(tmp) / "cli.xlsx"
            cli.write_bytes(mspdi_import.convert(MS_PROJECT_XML.encode())[0])
            a, b = openpyxl.load_workbook(page_xlsx), openpyxl.load_workbook(cli)
            self.assertEqual(a.sheetnames, b.sheetnames)
            diffs = [(ws.title, c.coordinate) for ws in a.worksheets for row in ws.iter_rows() for c in row
                     if c.value != b[ws.title][c.coordinate].value]
            self.assertEqual(diffs, [])
            # an older workbook in, "Move it to the latest template": the same workbook as the command line
            old, page_new = ROOT / "tests" / "fixtures" / "planner_v1.4.xlsx", Path(tmp) / "new.xlsx"
            result = subprocess.run(["node", str(ROOT / "tests" / "browser" / "export_page.js"),
                                     str(ROOT / "docs" / "export.html"), str(old), os.environ["PYODIDE_DIR"],
                                     str(page_new), "upgrade"], capture_output=True, text=True, timeout=600)
            self.assertIn("STATUS Done: 30 tasks, 7 resources.", result.stdout)
            cli.write_bytes(upgrade.convert(old.read_bytes())[0])
            a, b = openpyxl.load_workbook(page_new), openpyxl.load_workbook(cli)
            diffs = [(ws.title, c.coordinate) for ws in a.worksheets for row in ws.iter_rows() for c in row
                     if c.value != b[ws.title][c.coordinate].value]
            self.assertEqual(diffs, [])


if __name__ == "__main__":
    unittest.main()
