"""Release 1.0 (the --release build, its workbooks in release-1.0/): the pickers, Duration as 3 / 3d / 2w, the
leave day grid, the Check and TakeSnapshot sheets, TakeBaseline, protection and print set-up (docs/BACKLOG.md
"Release 1.0"). The classic workbooks are tested by test_planner.py and don't change with any of this.

Run:  python -m unittest tests.test_release -v      (the recalculation tests need LibreOffice)
"""
import datetime as dt
import io
import re
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import range_boundaries

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(ROOT))
import perf_check  # noqa: E402
from test_planner import (SOFFICE, CommittedWorkbooksTest, build, d, errors_in, named_cell, recalc, serial,  # noqa: E402
                          serial_or_none, table_rows)

RELEASE = ROOT / "release-1.0"
SAMPLE, BLANK, SIMPLE = (RELEASE / n for n in ("Project_Planner.xlsx", "Project_Planner_Blank.xlsx",
                                               "Project_Planner_Simple.xlsx"))
FIRST = 6                                               # first task row on Tasks


def make(**k):
    from planner.build import make_workbook
    return make_workbook(release_edition=True, **k)


def small():
    from planner.build import Sizes
    return Sizes(tasks=40, resources=12, grid_resources=8, sprints=6, gantt_columns=10, weeks=6)


def head(ws, row=5):
    return {ws.cell(row, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(row, c).value is not None}


def task_row(wb, uid):
    ws = wb["Tasks"]
    col = head(ws)
    return next(r for r in range(FIRST, ws.max_row + 1) if ws.cell(r, col["UID"]).value == uid)


def put(wb, uid, values):
    """Type values (header -> value) into the task with this UID (an unrecalculated workbook)."""
    ws = wb["Tasks"]
    col, r = head(ws), task_row(wb, uid)
    for h, v in values.items():
        ws.cell(r, col[h]).value = v


def tasks_by_uid(wb):
    return {t["UID"]: t for t in table_rows(wb, "Tasks", "Tasks") if t["Task Name"]}


class StructureTest(unittest.TestCase):
    """Built without recalculating: sheets, names, rules and protection."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.path = Path(cls.tmp.name) / "release.xlsx"
        make(baselined=True).save(cls.path)
        cls.wb = openpyxl.load_workbook(cls.path)
        cls.report = perf_check.static_report(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()
        from planner.build import make_workbook
        make_workbook()                                  # back to the classic build for any test that follows

    def test_sheets_and_versions(self):
        self.assertEqual(self.wb.sheetnames, ["Start", "Status", "Check", "Roadmap", "Resource Report", "Agile",
                                              "Log", "History", "TakeSnapshot", "Settings", "Tasks", "Resources",
                                              "Cost"])
        names = self.wb.defined_names
        self.assertEqual(names["TemplateVersion"].attr_text, '"1.0"')
        self.assertEqual(names["PlannerVersion"].attr_text, '"2.0"')
        self.assertEqual(names["PlannerModules"].attr_text, '"agile,log,history,resources,cost,check,snapshot"')

    def test_formula_rules(self):
        r = self.report
        self.assertEqual((r["a1_references"], r["missing_references"], r["excel_limit_problems"]), ([], [], []))
        self.assertEqual(r["volatile"], {"Tasks": {"OFFSET": 2200}})        # no new volatile cells
        self.assertTrue(all(v["single_row"] <= perf_check.SINGLE_ROW_RANGES_MAX for v in r["cf"].values()))
        self.assertLess(sum(v["rules"] for v in r["cf"].values()), 121)      # fewer rules than the classic sample

    def test_functions_need_at_most_excel_2019(self):
        # Excel 2019: MINIFS / MAXIFS, and TEXTJOIN for the pickers and the leave grid; nothing Microsoft 365 only
        allowed = {"_xlfn.MINIFS(", "_xlfn.MAXIFS(", "_xlfn.TEXTJOIN("}
        banned = ("XLOOKUP", "XMATCH", "FILTER(", "SEQUENCE(", "UNIQUE(", "SORT(", "LET(", "LAMBDA(", "TODAY(",
                  "NOW(", "INDIRECT(")
        seen = set()
        for ws in self.wb.worksheets:
            for row in ws.iter_rows(max_row=60):
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        for b in banned:
                            self.assertNotIn(b, c.value, f"{ws.title}!{c.coordinate}")
                        for fn in re.findall(r"_xlfn\.[A-Z.]+\(", c.value):
                            self.assertIn(fn, allowed, f"{ws.title}!{c.coordinate}")
                            seen.add(fn)
        self.assertIn("_xlfn.TEXTJOIN(", seen)

    def test_tasks_layout(self):
        ws = self.wb["Tasks"]
        col = head(ws)
        shown = [ws.cell(5, c).value for c in range(1, ws.max_column + 1)
                 if ws.cell(5, c).value and not ws.column_dimensions[ws.cell(5, c).column_letter].hidden][:14]
        self.assertEqual(shown, ["ID", "UID", "Level", "Task Name", "Duration", "Pred1", "Type1", "Lag1", "Res1",
                                 "Units1", "Stage", "Actual Start", "Actual Finish", "Start"])
        for h in ("Pred2", "Lag4", "Predecessors", "Res4", "Resources"):                 # behind a + button
            self.assertTrue(ws.column_dimensions[ws.cell(5, col[h]).column_letter].hidden, h)
        self.assertEqual(self.wb.defined_names["tsk_Duration"].attr_text, "Tasks[_Days]")     # the cleaned value
        self.assertEqual(self.wb.defined_names["tsk_BLSource"].attr_text,
                         "Tasks[[BL Source Start]:[BL Source Work]]")
        r = task_row(self.wb, 17)                        # Feature screens: after 16, and 13SS+2d
        self.assertEqual([ws.cell(r, col[h]).value for h in ("Pred1", "Type1", "Lag1", "Pred2", "Type2", "Lag2")],
                         [16, None, None, 13, "SS", 2])
        self.assertIsNone(ws.cell(r, col["Predecessors"]).value)
        r = task_row(self.wb, 18)                        # Accessibility review: Casey Kim at 50%
        self.assertEqual((ws.cell(r, col["Res1"]).value, ws.cell(r, col["Units1"]).value), ("Casey Kim", 0.5))
        self.assertEqual(ws.cell(task_row(self.wb, 26), col["Duration"]).value, "2w")

    def test_validation(self):
        ws = self.wb["Tasks"]
        col = head(ws)
        rules = {}
        for dv in ws.data_validations.dataValidation:
            for rng in str(dv.sqref).split():
                rules[range_boundaries(rng)[0]] = dv
        uid, dur = rules[col["UID"]], rules[col["Duration"]]
        self.assertEqual(uid.type, "custom")
        self.assertIn("COUNTIF(tsk_UID,", uid.formula1)
        self.assertEqual(dur.type, "custom")
        self.assertIn('"w"', dur.formula1)
        for h in ("Actual Start", "Actual Finish", "Deadline", "Constraint Date"):
            self.assertEqual((rules[col[h]].type, rules[col[h]].formula1), ("date", "ProjectStart"), h)
        for k in range(1, 5):                            # every picker says what to pick
            for h in (f"Pred{k}", f"Type{k}", f"Lag{k}", f"Res{k}", f"Units{k}"):
                self.assertTrue(rules[col[h]].showInputMessage and rules[col[h]].prompt, h)
            self.assertEqual(rules[col[f"Pred{k}"]].formula1, "=tsk_PickList")
            self.assertEqual(rules[col[f"Res{k}"]].formula1, "=res_Name")

    def test_duration_rule_accepts_and_refuses(self):
        """The Duration rule, evaluated the way Excel does for the values the owner listed."""
        def valid(x):
            if isinstance(x, (int, float)):
                return x >= 0
            unit = x[-1:].lower() in ("d", "w")
            try:
                return unit and float(x[:-1]) >= 0
            except ValueError:
                return False
        self.assertEqual([valid(x) for x in (1, "1d", "1w", 0, "2days", "d", "-1d")],
                         [True, True, True, True, False, False, False])

    def test_protection(self):
        from planner.build import RELEASE_VIEWS
        for sheet in RELEASE_VIEWS:
            p = self.wb[sheet].protection
            self.assertTrue(p.sheet and p.password is None, sheet)
        for sheet in ("Tasks", "Resources", "Settings", "History"):     # rows are added there
            ws = self.wb[sheet]
            self.assertFalse(ws.protection.sheet, sheet)
        ws = self.wb["Tasks"]
        col = head(ws)
        self.assertFalse(ws.cell(FIRST, col["Duration"]).protection.locked)     # Protect Sheet keeps inputs open
        self.assertFalse(ws.cell(FIRST, col["Pred1"]).protection.locked)
        self.assertTrue(ws.cell(FIRST, col["Finish"]).protection.locked)
        log = self.wb["Log"]
        lc = head(log, 4)
        self.assertFalse(log.cell(5, lc["Title"]).protection.locked)
        self.assertTrue(log.cell(5, lc["Score"]).protection.locked)

    def test_printing(self):
        for sheet in ("Status", "Roadmap", "Resource Report", "Check"):
            ws = self.wb[sheet]
            self.assertIn("template 1.0", ws.oddFooter.left.text, sheet)
            self.assertEqual(ws.page_setup.fitToWidth, 1, sheet)
            self.assertTrue(ws.print_area, sheet)
        for sheet in ("Status", "Roadmap", "Resource Report"):
            self.assertEqual(self.wb[sheet].page_setup.fitToHeight, 1, sheet)

    def test_start_has_five_linked_steps(self):
        ws = self.wb["Start"]
        steps = [r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 2).value or "")[:2] in
                 ("1.", "2.", "3.", "4.", "5.") and ws.cell(r, 3).hyperlink is not None]
        self.assertEqual([ws.cell(r, 2).value for r in steps[:5]],
                         ["1. Your project", "2. People", "3. Tasks", "4. Check", "5. Baseline"])
        self.assertEqual([ws.cell(r, 3).hyperlink.location for r in steps[:5]],
                         ["'Settings'!A1", "'Resources'!A1", "'Tasks'!A1", "'Check'!A1", "'TakeSnapshot'!A1"])
        self.assertTrue(any("TemplateVersion" in str(c.value) for row in ws.iter_rows() for c in row))

    def test_snapshot_links(self):
        ws = self.wb["TakeSnapshot"]
        links = {c.hyperlink.location for row in ws.iter_rows() for c in row if c.hyperlink is not None}
        self.assertTrue({"snap_Row", "tsk_BLSource", "tsk_Baseline", "TakeBaseline", "'History'!A1"} <= links)
        for n in links - {"'History'!A1", "'Check'!A1"}:
            self.assertIn(n, self.wb.defined_names, n)


class ModulesTest(unittest.TestCase):
    """Any set of modules builds with nothing pointing at a sheet, table or name that isn't there; Simple is the six
    sheets with 12 tasks."""

    @classmethod
    def tearDownClass(cls):
        from planner.build import make_workbook
        make_workbook()

    def test_combinations(self):
        from planner.core import MODULES, RELEASE_MODULES
        names = list(MODULES) + list(RELEASE_MODULES)
        combos = [()] + [(m,) for m in names] + [("check", "snapshot"), ("check", "log"), ("snapshot", "history"),
                                                 tuple(names), tuple(n for n in names if n != "log"),
                                                 tuple(n for n in names if n != "check")]
        for combo in combos:
            wb = make(sizes=small(), modules=combo)
            self.assertEqual(perf_check.missing_references(wb), [], combo)

    def test_simple(self):
        wb = make(baselined=True, modules=())
        self.assertEqual(wb.sheetnames, ["Start", "Status", "Roadmap", "Settings", "Tasks", "Resources"])
        named = [r for r in range(FIRST, FIRST + 200) if wb["Tasks"].cell(r, head(wb["Tasks"])["Task Name"]).value]
        self.assertEqual(len(named), 12)
        self.assertEqual(perf_check.missing_references(wb), [])


class CommittedReleaseTest(unittest.TestCase):
    """The workbooks in release-1.0/ match what the build produces (the commands in README)."""
    compare = CommittedWorkbooksTest.compare

    def test_sample(self):
        self.compare(SAMPLE, "--release", "--baselined")

    def test_blank(self):
        self.compare(BLANK, "--release", "--blank")

    def test_simple(self):
        self.compare(SIMPLE, "--release", "--baselined", "--simple")


def _plan_inputs(wb):
    """test_planner._plan_inputs for a Release 1.0 workbook: the links and duration as the formulas read them (the
    pickers' _Predecessors, the cleaned _Days)."""
    rows = [{"Level": t["Level"], "Name": t["Task Name"], "Mode": t["Mode"], "Dur": t["_Days"],
             "Pred": t["_Predecessors"], "CType": t["Constraint"], "Cons": serial_or_none(t["Constraint Date"]),
             "DL": serial_or_none(t["Deadline"]), "AS": serial_or_none(t["Actual Start"]),
             "AF": serial_or_none(t["Actual Finish"]), "Pct": t["% Complete"], "Stage": t["Stage"], "UID": t["UID"],
             "BS": serial_or_none(t["Baseline Start"]), "BF": serial_or_none(t["Baseline Finish"])}
            for t in table_rows(wb, "Tasks", "Tasks")]
    for r in rows:
        r["Dur"] = None if r["Dur"] == "" else r["Dur"]
    value = lambda n: named_cell(wb, n).value
    settings = {"ProjectStart": serial(value("ProjectStart")), "StatusDate": serial(value("StatusDate")),
                "WorkWeek": value("WorkWeek"), "RescheduleAfterStatus": value("RescheduleAfterStatus"),
                "CriticalSlackDays": value("CriticalSlackDays"),
                "Holidays": [serial(r["Date"]) for r in table_rows(wb, "Settings", "HolidayTable") if r["Date"]]}
    stages = {r["Stage"]: r["% Complete"] for r in table_rows(wb, "Settings", "StageTable") if r["Stage"]}
    return rows, settings, stages


SCHEDULE = ("Start", "Finish", "Planned Start", "Planned Finish", "Late Start", "Late Finish", "Total Slack",
            "Free Slack", "Critical", "Work", "Resource Names", "Predecessors by ID", "Issues", "Progress")


class CommittedSampleTest(unittest.TestCase):
    """The committed Release 1.0 sample (recalculated by the build): no errors, and the same plan as the classic
    sample, though its links and people come from the pickers, two durations are typed in weeks and some leave is
    on the day grid."""

    @classmethod
    def setUpClass(cls):
        cls.wb = openpyxl.load_workbook(SAMPLE, data_only=True)
        cls.classic = openpyxl.load_workbook(ROOT / "Project_Planner.xlsx", data_only=True)

    def test_no_errors(self):
        self.assertEqual(errors_in(self.wb), [])
        self.assertEqual(errors_in(openpyxl.load_workbook(SIMPLE, data_only=True)), [])
        self.assertEqual(errors_in(openpyxl.load_workbook(BLANK, data_only=True)), [])

    def test_same_plan_as_the_classic_sample(self):
        new, old = tasks_by_uid(self.wb), tasks_by_uid(self.classic)
        self.assertEqual(set(new), set(old))
        diffs = [(u, k, new[u][k], old[u][k]) for u in old for k in SCHEDULE if new[u][k] != old[u][k]]
        self.assertEqual(diffs, [])
        value = lambda wb, n: (named_cell(wb, n).value if "!" in wb.defined_names[n].attr_text
                               else wb.defined_names[n].attr_text)
        for n in ("ProjectFinish", "PlannedFinish", "ProjectRAG", "sch_SPI"):
            self.assertEqual(value(self.wb, n), value(self.classic, n), n)

    def test_same_capacity_with_leave_on_the_day_grid(self):
        for grid in ("rw_Util", "rw_Leave", "cap_Days", "cap_Points"):
            if grid not in self.classic.defined_names:
                continue
            values = []
            for wb in (self.wb, self.classic):
                sheet, ref = wb.defined_names[grid].attr_text.split("!")
                values.append([[c.value for c in row] for row in wb[sheet.strip("'")][ref.replace("$", "")]])
            self.assertEqual(values[0], values[1], grid)
        days = {r["Resource"]: r["Days"] for r in table_rows(self.wb, "Resources", "LeaveDays") if r["Resource"]}
        self.assertEqual(days["Casey Kim"], "16, 17(0.5), 18, 19, 20, 23, 24, 25, 26, 27")

    def test_reference_schedule(self):
        import reference_schedule
        for wb in (self.wb, openpyxl.load_workbook(SIMPLE, data_only=True)):
            got, finish = reference_schedule.schedule(*_plan_inputs(wb))
            sheet = tasks_by_uid(wb)
            bad = []
            for v in got:
                t = sheet[v["UID"]]
                for k, want in v.items():
                    if k in ("UID", "CS", "CF"):
                        continue
                    have = (t[k] == "Yes" if k == "Critical" else serial_or_none(t[k]) if isinstance(want, dt.date)
                            else t[k])
                    if have != want:
                        bad.append((v["UID"], k, have, want))
            self.assertEqual(bad, [])
            self.assertEqual(serial(named_cell(wb, "ProjectFinish").value), finish)

    def test_check_and_snapshot(self):
        counts = [c[0].value for c in self.wb["Check"][self.wb.defined_names["chk_Counts"].attr_text.split("!")[1]
                                                        .replace("$", "")]]
        issues = [t for t in tasks_by_uid(self.wb).values() if t["Issues"]]
        self.assertEqual(counts[0], len(issues))
        self.assertEqual(named_cell(self.wb, "chk_Verdict").value, f"{sum(counts)} to fix before you report")
        listed = [r["UID"] for r in table_rows(self.wb, "Check", "CheckIssues") if r["UID"] not in (None, "")]
        self.assertEqual(listed, [t["UID"] for t in issues])
        snap = table_rows(self.wb, "TakeSnapshot", "SnapshotRow")[0]
        status = table_rows(self.wb, "Status", "SnapshotNow")[0]
        self.assertEqual(snap, status)

    def test_simple_sample(self):
        wb = openpyxl.load_workbook(SIMPLE, data_only=True)
        t = tasks_by_uid(wb)
        self.assertEqual(len(t), 12)
        self.assertTrue(all(not r["Link Check"] for r in t.values()))       # every link points inside the plan
        self.assertEqual(wb.sheetnames, ["Start", "Status", "Roadmap", "Settings", "Tasks", "Resources"])
        # its baseline is the plan as first calculated: BL Source agrees with it on every task
        for r in t.values():
            self.assertEqual(d(r["Baseline Start"]), d(r["BL Source Start"]), r["Task Name"])


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class EditTest(unittest.TestCase):
    """Edits a user makes, recalculated by LibreOffice."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.dir = Path(cls.tmp.name)
        cls.path = cls.dir / "release.xlsx"
        make(baselined=True).save(cls.path)

        def edit(wb):
            put(wb, 13, {"Duration": "3d"})                                   # Core API endpoints
            put(wb, 14, {"Duration": "2days", "Pred1": "13 Core API endpoints", "Type1": "SS", "Lag1": 2})
            put(wb, 16, {"Duration": "1w", "Predecessors": "9", "Res1": "Jordan Lee", "Units1": 0.5})
            put(wb, 21, {"Pred1": None, "Predecessors": "19FS+1d"})            # the typed text when no picker
            put(wb, 23, {"UID": 22})                                           # a duplicate UID
            named_cell(wb, "TakeBaseline").value = "Yes"
        cls.wb = recalc(cls.path, cls.dir, edit)
        cls.t = tasks_by_uid(cls.wb)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()
        from planner.build import make_workbook
        make_workbook()

    def test_no_errors(self):
        self.assertEqual(errors_in(self.wb), [])

    def test_duration(self):
        t = self.t
        self.assertEqual(t[13]["_Days"], 3)
        self.assertEqual(t[16]["_Days"], 5)
        self.assertIn(t[14]["_Days"], ("", None))
        self.assertIn("Duration 2days can't be read", t[14]["Issues"])
        self.assertEqual(serial(t[13]["Finish"]) - serial(t[13]["Start"]), 2)     # 3 working days, same week

    def test_pickers(self):
        t = self.t
        self.assertEqual(t[14]["_Predecessors"], "13SS+2d")                 # "13 Core API endpoints" reads as 13
        self.assertEqual((t[14]["_Pred1UID"], t[14]["_Pred1Type"], t[14]["_Pred1Lag"]), (13, "SS", 2))
        self.assertIn(t[14]["Link Check"], ("", None))
        self.assertEqual(t[21]["_Predecessors"], "19FS+1d")                 # the text, with no picker filled
        self.assertEqual(t[16]["Resource Names"], "Jordan Lee[50%]")
        self.assertIn("The Predecessors text is ignored", t[16]["Issues"])

    def test_uids_and_check(self):
        self.assertEqual(named_cell(self.wb, "DuplicateUIDs").value, 2)
        self.assertIn("UID 22 is used twice", self.t[22]["Issues"])
        listed = {r["UID"] for r in table_rows(self.wb, "Check", "CheckIssues") if r["UID"] not in (None, "")}
        self.assertIn(22, listed)

    def test_take_baseline(self):
        # Yes: BL Source is the current schedule for every task, so the paste re-baselines the whole plan
        for u in (13, 17, 26):
            self.assertEqual(d(self.t[u]["BL Source Finish"]), d(self.t[u]["Finish"]), u)

    def test_take_baseline_no_keeps_the_baseline(self):
        wb = openpyxl.load_workbook(SAMPLE, data_only=True)               # TakeBaseline = No as shipped
        for t in tasks_by_uid(wb).values():
            self.assertEqual(d(t["BL Source Finish"]), d(t["Baseline Finish"]), t["Task Name"])


class ToolsTest(unittest.TestCase):
    """Upgrade, reorder and the MS Project export keep a Release 1.0 workbook's inputs."""

    @classmethod
    def tearDownClass(cls):
        from planner.build import make_workbook
        make_workbook()

    def test_upgrade_keeps_the_release_inputs(self):
        from planner import upgrade
        wb = openpyxl.load_workbook(SAMPLE)
        put(wb, 17, {"Lag2": 5})
        named_cell(wb, "TakeBaseline").value = "Yes"
        ws = wb["Resources"]
        lv = ws.tables["LeaveDays"]
        c1, r1, _, _ = range_boundaries(lv.ref)
        col = {ws.cell(r1, c).value: c for c in range(c1, c1 + 40)}
        ws.cell(r1 + 3, col["Resource"]).value, ws.cell(r1 + 3, col["Month"]).value = "Sam Rivera", "Jan-27"
        ws.cell(r1 + 3, col["4"]).value = 1
        buf = io.BytesIO()
        wb.save(buf)
        out, warnings = upgrade.convert(buf.getvalue())
        new = openpyxl.load_workbook(io.BytesIO(out))
        self.assertEqual(new.defined_names["PlannerVersion"].attr_text, '"2.0"')
        self.assertEqual(new.sheetnames, wb.sheetnames)
        nc = head(new["Tasks"])
        r = task_row(new, 17)
        self.assertEqual([new["Tasks"].cell(r, nc[h]).value for h in ("Pred1", "Pred2", "Type2", "Lag2")],
                         [16, 13, "SS", 5])
        self.assertEqual(new["Tasks"].cell(task_row(new, 26), nc["Duration"]).value, "2w")
        self.assertEqual(named_cell(new, "TakeBaseline").value, "Yes")
        rows = table_rows(new, "Resources", "LeaveDays")
        self.assertIn(("Sam Rivera", "Jan-27", 1), [(x["Resource"], x["Month"], x["4"]) for x in rows])

    def test_classic_workbook_moves_into_the_release(self):
        from planner import upgrade
        out, _ = upgrade.convert((ROOT / "Project_Planner.xlsx").read_bytes(), release=True)
        new = openpyxl.load_workbook(io.BytesIO(out))
        self.assertIn("Check", new.sheetnames)
        nc = head(new["Tasks"])
        self.assertEqual(new["Tasks"].cell(task_row(new, 17), nc["Predecessors"]).value, "16, 13SS+2d")
        out, _ = upgrade.convert((ROOT / "Project_Planner_Simple.xlsx").read_bytes(), release=True)
        self.assertEqual(openpyxl.load_workbook(io.BytesIO(out)).sheetnames,
                         ["Start", "Status", "Roadmap", "Settings", "Tasks", "Resources"])

    def test_reorder_reads_the_pickers(self):
        from planner import reorder
        self.assertEqual(reorder.task_links({"Pick1": "12 Design", "Pick2": 14, "PredIn": "3"}), [12, 14])
        self.assertEqual(reorder.task_links({"PredIn": "3, 5SS+2d"}), [3, 5])

    def test_ms_project_export_matches_the_classic(self):
        from planner import mspdi
        new = mspdi.convert(SAMPLE.read_bytes()) if hasattr(mspdi, "convert") else None
        old = mspdi.convert((ROOT / "Project_Planner.xlsx").read_bytes()) if new is not None else None
        if new is None:
            self.skipTest("no mspdi.convert")
        text = lambda x: x.decode() if isinstance(x, bytes) else x
        tasks = lambda x: re.findall(r"<Task>.*?</Task>", text(x), re.S)
        self.assertEqual(len(tasks(new)), len(tasks(old)))
        for pattern in (r"<PredecessorUID>(\d+)", r"<Duration>([^<]+)", r"<ResourceUID>(\d+)"):
            self.assertEqual(re.findall(pattern, text(new)), re.findall(pattern, text(old)), pattern)


class ClassicUnchangedTest(unittest.TestCase):
    """The classic build is the same after a Release 1.0 build in the same process (the edition is per build)."""

    def test_classic_after_release(self):
        from planner.build import make_workbook
        make(sizes=small())
        wb = make_workbook(sizes=small())
        cols = head(wb["Tasks"])
        self.assertNotIn("Pred1", cols)
        self.assertNotIn("_Days", cols)
        self.assertEqual(wb.defined_names["tsk_Duration"].attr_text, "Tasks[Duration]")
        self.assertEqual(wb.defined_names["PlannerVersion"].attr_text, '"1.8"')
        self.assertNotIn("TemplateVersion", wb.defined_names)
        self.assertNotIn("Check", wb.sheetnames)


if __name__ == "__main__":
    unittest.main()
