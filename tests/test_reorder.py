"""The reorder tool (planner/reorder.py, scripts/reorder_planner.py): every predecessor above its successor (B-1015)."""
import io
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from planner import reorder  # noqa: E402
from test_mspdi import table_rows  # noqa: E402
from test_planner import SOFFICE, build, columns, errors_in, recalc  # noqa: E402

FIRST = 6
T = lambda uid, level, name, pred=None: {"UID": uid, "Level": level, "Name": name, "PredIn": pred}


class OrderTest(unittest.TestCase):
    def names(self, tasks):
        return [t["Name"] for t in reorder.order(tasks)[0]]

    def test_moves_inside_a_phase(self):
        plan = [T(1, 1, "P1"), T(2, 2, "a", "3"), T(3, 2, "b"), T(4, 1, "P2"), T(5, 2, "c", "2"), T(6, 2, "d")]
        self.assertEqual(self.names(plan), ["P1", "b", "a", "P2", "c", "d"])
        self.assertEqual(reorder.order(plan)[1], ["b"])               # one task moved; the rest kept their order

    def test_moves_a_whole_phase_when_it_must(self):
        plan = [T(1, 1, "P1"), T(2, 2, "a", "5SS+2d"), T(3, 1, "P2"), T(4, 2, "b"), T(5, 2, "c")]
        self.assertEqual(self.names(plan), ["P2", "b", "c", "P1", "a"])   # a needs c: P1 goes after P2, intact

    def test_in_order_already(self):
        plan = [T(1, 1, "P1"), T(2, 2, "a"), T(3, 2, "b", "2FS-1d, 99, 1"), T(4, 2, "c", "abc")]
        self.assertEqual(reorder.order(plan), (plan, []))            # missing UIDs, summaries, text: Link Check's

    def test_refuses_when_the_outline_makes_it_impossible(self):
        plan = [T(1, 1, "P1"), T(2, 2, "a", "5"), T(3, 2, "x"), T(4, 1, "P2"), T(5, 2, "c", "3")]
        with self.assertRaisesRegex(reorder.ReorderError, r"P1, P2 each need another of them first .*UID 5 \(c\)"):
            reorder.order(plan)


class WorkbookTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.dir = Path(cls.tmp.name)
        cls.path = cls.dir / "Project_Planner.xlsx"
        build(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def edited(self, task_id, predecessors):
        wb = openpyxl.load_workbook(self.path)
        ws = wb["Tasks"]
        ws.cell(task_id + FIRST - 1, columns(ws)["Predecessors"]).value = predecessors
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    def test_reorders_and_keeps_everything(self):
        data = self.edited(22, "21, 23")                              # UAT now also waits for Training, listed below it
        xlsx, notes = reorder.convert(data)
        self.assertIn("Moved 1 tasks", notes[0])
        self.assertIn("Training & user guides", notes[0])
        wb = openpyxl.load_workbook(io.BytesIO(xlsx))
        rows = [r for r in table_rows(wb, "Tasks", "Tasks") if r["Task Name"]]
        uids = [r["UID"] for r in rows]
        self.assertEqual(uids, list(range(1, 22)) + [23, 22] + list(range(24, 31)))
        by_uid = {r["UID"]: r for r in rows}
        self.assertEqual((by_uid[22]["Predecessors"], by_uid[22]["Task Name"]), ("21, 23", "User acceptance testing"))
        self.assertEqual(by_uid[23]["Level"], 2)                      # still in Test & Release

    @unittest.skipUnless(SOFFICE, "LibreOffice not installed")
    def test_link_check_clears_after_reordering(self):
        data = self.edited(22, "21, 23")
        before = self.dir / "before.xlsx"
        before.write_bytes(data)
        link = lambda wb: {r["UID"]: r["Link Check"] for r in table_rows(wb, "Tasks", "Tasks") if r["Task Name"]}
        self.assertIn("UID 23 is below this task", link(recalc(before, self.dir))[22])
        after = self.dir / "after.xlsx"
        after.write_bytes(reorder.convert(data)[0])
        wb = recalc(after, self.dir)
        self.assertEqual(errors_in(wb), [])
        self.assertTrue(all(v in (None, "") for v in link(wb).values()), link(wb))

    def test_refuses_a_cycle_between_phases(self):
        data = self.edited(3, "2, 6")                                 # Initiation waits for Requirements, which waits for it
        with self.assertRaisesRegex(reorder.ReorderError, "Initiation, Requirements & Design each need"):
            reorder.convert(data)


if __name__ == "__main__":
    unittest.main()


@unittest.skipUnless(SOFFICE, "LibreOffice not installed")
class ProjectTypeTest(unittest.TestCase):
    """B-1018: Start > Project type sets the stages, sprint length and RAG slip thresholds; each stays editable."""

    def test_presets_follow_the_type_until_typed_over(self):
        from planner.settings import PRESETS
        from test_planner import named_cell
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "blank.xlsx"
            build(path, "--blank", "--project-type", "Construction")
            wb = recalc(path, tmp, status_date=False)
            self.assertEqual(errors_in(wb), [])
            stages = [(r["Stage"], r["% Complete"]) for r in table_rows(wb, "Settings", "StageTable")]
            self.assertEqual(stages, PRESETS["Construction"]["stages"])
            self.assertEqual([named_cell(wb, n).value for n in ("SprintLength", "RAGAmberSlipDays", "RAGRedSlipDays")],
                             [28, 5, 15])

            def switch(wb):
                named_cell(wb, "ProjectType").value = "IT infrastructure"
                named_cell(wb, "RAGRedSlipDays").value = 7                      # typed: stays
            wb = recalc(path, tmp, switch, status_date=False)
            self.assertEqual([r["Stage"] for r in table_rows(wb, "Settings", "StageTable")],
                             [n for n, _ in PRESETS["IT infrastructure"]["stages"]])
            self.assertEqual([named_cell(wb, n).value for n in ("SprintLength", "RAGAmberSlipDays", "RAGRedSlipDays")],
                             [14, 2, 7])
            self.assertEqual(named_cell(openpyxl.load_workbook(path), "ProjectType").parent.title, "Start")


class ImportProjectTypeTest(unittest.TestCase):
    def test_ms_project_import_takes_a_type(self):
        from planner import mspdi, mspdi_import
        from test_planner import named_cell
        xml = mspdi.convert((ROOT / "Project_Planner.xlsx").read_bytes()).encode()
        wb = openpyxl.load_workbook(io.BytesIO(mspdi_import.convert(xml, "Construction")[0]))
        self.assertEqual(named_cell(wb, "ProjectType").value, "Construction")
        stage = {r["Task Name"]: r["Stage"] for r in table_rows(wb, "Tasks", "Tasks") if r["Task Name"]}
        self.assertEqual(stage["Core API endpoints"], "Under Way")       # 50% done: the Construction stage at 50%
